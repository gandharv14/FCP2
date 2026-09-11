"""Certify shared source-link exchanges against original formula and period evidence."""
from collections import OrderedDict
from datetime import date, datetime
import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string, get_column_letter
from spoken_formula import clean_sheet


def _reference(value):
    sheet, coordinate = value.rsplit('!', 1)
    coordinate = coordinate.replace('$', '')
    column, row = coordinate_from_string(coordinate)
    return clean_sheet(sheet), column.upper(), row


def _tokens(formula):
    if not isinstance(formula, str) or not formula.startswith('='):
        return []
    return [token for token in Tokenizer(formula).items if token.type != 'WHITE-SPACE']


def _operand_reference(value, sheet):
    return _reference(value if '!' in value else sheet + '!' + value)


def _year_value(reference, values, visiting=None):
    """Derive calendar headers only from typed integers, bare links, and +1.

    No cached formula result or general financial evaluator is consulted. The
    caller separately proves the original header's explicit year label.
    """
    visiting = set() if visiting is None else visiting
    if reference in visiting or len(visiting) >= 64 or reference not in values:
        raise ValueError('year header dependency is missing or cyclic')
    value = values[reference]
    if type(value) in (int, float) and value == int(value):
        year = int(value)
    else:
        tokens = _tokens(value)
        if len(tokens) == 1 and tokens[0].subtype == 'RANGE':
            year = _year_value(_operand_reference(tokens[0].value, reference[0]),
                               values, visiting | {reference})
        elif (len(tokens) == 3 and tokens[0].subtype == 'RANGE'
              and tokens[1].type == 'OPERATOR-INFIX' and tokens[1].value == '+'
              and tokens[2].subtype == 'NUMBER' and tokens[2].value == '1'):
            year = _year_value(_operand_reference(tokens[0].value, reference[0]),
                               values, visiting | {reference}) + 1
        else:
            raise ValueError('unsupported year header derivation')
    if not 1900 <= year <= 2200:
        raise ValueError('numeric header is outside calendar-year bounds')
    return year


def _labelled_year(header, label, axis, values):
    if (label[0] != header[0] or label[2] != header[2]
            or axis[0] != header[0] or axis[2] != header[2]
            or not column_index_from_string(label[1]) < column_index_from_string(axis[1])
            <= column_index_from_string(header[1])):
        raise ValueError('year label or axis is not on the original header row')
    text = values.get(label)
    if not isinstance(text, str) or not re.fullmatch(
            r'(?:fiscal |calendar )?year', text.strip(), re.I):
        raise ValueError('numeric header lacks an explicit original year label')
    for column in range(column_index_from_string(label[1]) + 1,
                        column_index_from_string(header[1])):
        value = values.get((header[0], get_column_letter(column), header[2]))
        if isinstance(value, str) and value.strip() and not value.startswith('='):
            raise ValueError('intervening header label changes the year scope')
    # A certified axis has at least three consecutive calendar years. Controls
    # before that axis are not used as year evidence; every axis cell is derived
    # from original typed inputs, never a cached value or integer equality alone.
    start = column_index_from_string(axis[1])
    first = _year_value(axis, values)
    for column in range(start, max(start + 2, column_index_from_string(header[1])) + 1):
        actual = _year_value((header[0], get_column_letter(column), header[2]), values)
        if actual != first + column - start:
            raise ValueError('year axis is not consecutive calendar years')
    return _year_value(header, values)



def _relative_year(header, axis, values):
    """Prove an explicitly labelled project-year axis from original text only."""
    if (axis[0] != header[0] or axis[2] != header[2]
            or column_index_from_string(axis[1]) > column_index_from_string(header[1])):
        raise ValueError('relative-year axis does not contain its header')
    start = column_index_from_string(axis[1])
    end = max(start + 2, column_index_from_string(header[1]))
    periods = []
    for column in range(start, end + 1):
        value = values.get((header[0], get_column_letter(column), header[2]))
        match = re.fullmatch(r'Year (0|[1-9][0-9]{0,2})', value, re.I) if isinstance(value, str) else None
        if match is None:
            raise ValueError('relative-year header is not original explicit Year text')
        periods.append(int(match.group(1)))
    if periods != list(range(periods[0], periods[0] + len(periods))):
        raise ValueError('relative-year axis is not consecutive')
    return periods[column_index_from_string(header[1]) - start]


def prove_groups(claims, disclosure, source, evidence_path):
    source, evidence_path = Path(source), Path(evidence_path)
    evidence = json.loads(evidence_path.read_text())
    sha = lambda path: hashlib.sha256(Path(path).read_bytes()).hexdigest()
    if evidence.get('source_sha256') != sha(source) or evidence.get('disclosure_sha256') != sha(disclosure):
        raise ValueError('source-link scope evidence hashes changed')
    rows = evidence.get('rows', [])
    claim_ids = [claim['record_id'] for claim in claims]
    if len(claim_ids) != len(set(claim_ids)):
        raise ValueError('original claim IDs are not unique')
    if evidence.get('schema_version') not in (None, 'dialogue-source-link-scope/v1', 'dialogue-source-link-scope/v2', 'dialogue-source-link-scope/v3'):
        raise ValueError('unsupported source-link scope evidence schema')
    relative = evidence.get('schema_version') == 'dialogue-source-link-scope/v3'
    partial = relative or evidence.get('schema_version') == 'dialogue-source-link-scope/v2'
    if partial:
        if evidence.get('claim_ids') != claim_ids:
            raise ValueError('partial scope must bind every original claim in order')
        mapping = {row['claim_id']: row for row in rows}
        if len(mapping) != len(rows) or not set(mapping) <= set(claim_ids):
            raise ValueError('partial scope contains duplicate or unknown claims')
    else:
        mapping = {row['target_reference']: row for row in rows}
        targets = [claim['reviewer_only']['cells'][0] for claim in claims
                   if len(claim['reviewer_only']['cells']) == 1]
        if (len(mapping) != len(rows) or len(mapping) != len(claims)
                or len(targets) != len(claims) or set(mapping) != set(targets)):
            raise ValueError('source-link scope must cover every original claim exactly')
    groups = OrderedDict()
    workbook = load_workbook(source, read_only=True, data_only=False)
    proven = []
    try:
        requested = set()
        for row in rows:
            if partial and row.get('period_kind') not in (('date', 'year', 'relative_year') if relative else ('date', 'year')):
                raise ValueError('partial scope needs an explicit period kind')
            requested.update(_reference(row[key]) for key in
                             ('target_reference', 'source_reference', 'target_period_header', 'source_period_header'))
            if relative and row.get('period_kind') == 'relative_year':
                for side in ('target', 'source'):
                    header = _reference(row[side + '_period_header'])
                    axis = _reference(row[side + '_period_axis_start'])
                    if axis[0] != header[0] or axis[2] != header[2]:
                        raise ValueError('relative-year axis has foreign sheet or row')
                    start = column_index_from_string(axis[1])
                    end = max(start + 2, column_index_from_string(header[1]))
                    requested.update((header[0], get_column_letter(column), header[2])
                                     for column in range(start, end + 1))
            if partial and row.get('period_kind') == 'year':
                for side in ('target', 'source'):
                    label = _reference(row[side + '_period_label'])
                    header = _reference(row[side + '_period_header'])
                    requested.add(label)
                    axis = _reference(row[side + '_year_axis_start'])
                    requested.add(axis)
                    if axis[0] == header[0] and axis[2] == header[2]:
                        start = column_index_from_string(axis[1])
                        requested.update((header[0], get_column_letter(column), header[2])
                                         for column in range(start, max(start + 2, column_index_from_string(header[1])) + 1))
                    if label[0] == header[0] and label[2] == header[2]:
                        requested.update((header[0], get_column_letter(column), header[2])
                                         for column in range(column_index_from_string(label[1]),
                                                             column_index_from_string(header[1]) + 1))
        if partial:
            requested.update(_reference(ref) for ref in evidence.get('year_evidence_cells', []))
        values = {}
        for sheet in {key[0] for key in requested}:
            positions = [(column_index_from_string(col), row)
                         for name, col, row in requested if name == sheet]
            wanted = {f'{col}{row}' for name, col, row in requested if name == sheet}
            for row in workbook[sheet].iter_rows(
                    min_col=min(col for col, _ in positions), max_col=max(col for col, _ in positions),
                    min_row=min(row for _, row in positions), max_row=max(row for _, row in positions)):
                for cell in row:
                    if getattr(cell, 'coordinate', None) in wanted:
                        value = cell.value
                        # A displayed string beginning '=' is not a formula.
                        # Preserve that distinction before any primitive proof.
                        if isinstance(value, str) and value.startswith('=') and cell.data_type != 'f':
                            value = ('literal_text', value)
                        values[(sheet, *coordinate_from_string(cell.coordinate))] = value
        for claim in claims:
            private = claim['reviewer_only']
            row = mapping.get(claim['record_id']) if partial else mapping[private['cells'][0]]
            if row is None:
                groups[('unshared', claim['record_id'])] = [claim['record_id']]
                continue
            if private.get('entry') != 'source_selection' or len(private.get('cells', [])) != 1:
                raise ValueError('shared scope supports only direct source selections')
            target, origin = _reference(row['target_reference']), _reference(row['source_reference'])
            if target != _reference(private['cells'][0]):
                raise ValueError('scope target differs from its original claim')
            target_header = _reference(row['target_period_header'])
            source_header = _reference(row['source_period_header'])
            formula = values.get(target)
            tokens = _tokens(formula)
            if (formula != row.get('original_formula') or len(tokens) != 1
                    or tokens[0].subtype != 'RANGE' or _operand_reference(tokens[0].value, target[0]) != origin):
                raise ValueError('scope formula is not the exact original direct reference')
            if (target_header[:2] != target[:2] or source_header[:2] != origin[:2]
                    or target_header[2] >= target[2] or source_header[2] >= origin[2]):
                raise ValueError('source and target period headers do not prove matching periods')
            if relative and row.get('period_kind') == 'relative_year':
                periods = [_relative_year(header, _reference(row[side + '_period_axis_start']), values)
                           for side, header in (('target', target_header), ('source', source_header))]
                if type(row.get('period')) is not int or periods[0] != periods[1] or periods[0] != row['period']:
                    raise ValueError('source and target relative-year headers differ')
            elif partial and row.get('period_kind') == 'year':
                periods = [_labelled_year(header, _reference(row[side + '_period_label']),
                                           _reference(row[side + '_year_axis_start']), values)
                           for side, header in (('target', target_header), ('source', source_header))]
                if type(row.get('period')) is not int or periods[0] != periods[1] or periods[0] != row['period']:
                    raise ValueError('source and target year headers do not prove matching periods')
            else:
                periods = [values.get(header) for header in (target_header, source_header)]
                if (not all(isinstance(period, (date, datetime)) for period in periods)
                        or periods[0] != periods[1]
                        or periods[0].strftime('%Y-%m-%d') != row.get('period')):
                    raise ValueError('source and target period headers do not prove matching periods')
            key = ('source', target[0], target[2], origin[0], origin[2], claim['row_label'], claim['spoken'],
                   tuple(claim['must_say']))
            if relative:
                key += (row['period_kind'],)
            groups.setdefault(key, []).append(claim['record_id'])
            proven.append(claim['record_id'])
    finally:
        workbook.close()
    for claim in claims:
        if claim['record_id'] in proven:
            claim['spoken'] += ' Use the matching period for this link.'
            claim['must_say'].append('matching period')
    flattened = [identity for group in groups.values() for identity in group]
    if sorted(flattened) != sorted(claim_ids):
        raise ValueError('shared exchanges lost original claim coverage')
    return {'schema_version': 'dialogue-shared-source-links/v1',
            'source_sha256': evidence['source_sha256'], 'disclosure_sha256': evidence['disclosure_sha256'],
            'evidence_sha256': sha(evidence_path), 'groups': list(groups.values()),
            'proven_claim_ids': proven}
