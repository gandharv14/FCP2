from __future__ import annotations

import contextlib
import io
import copy
import hashlib
import json
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import openpyxl

from xl_task_build import read_env_key
import xl_variable_source_audit
from xl_variable_source_audit import (
    build_inventory,
    resolve_segmentation_directory,
    resolve_table,
    validate_row_ids,
    validate_table,
)

try:
    # The real downstream consumer; requires fastmcp at import time.
    import xl_variable_mcp

    def run_downstream_import(markdown_path, output_path):
        with contextlib.redirect_stdout(io.StringIO()):
            xl_variable_mcp.main(
                ["import", str(markdown_path), str(output_path)])
        return json.loads(Path(output_path).read_text(encoding="utf-8"))
except ModuleNotFoundError:
    # Same parser xl_variable_mcp.py's import subcommand delegates to.
    from mcp_env.import_table import import_table

    def run_downstream_import(markdown_path, output_path):
        result = import_table(Path(markdown_path))
        Path(output_path).write_text(
            json.dumps(result, indent=2, ensure_ascii=False) + "\n",
            encoding="utf-8")
        return result


SEGMENTS = {
    "inputs": [
        {
            "band": "Inputs!A1",
            "sheet": "Inputs",
            "label": "Tax rate",
            "kind": "label",
            "vtype": "text",
            "cells": ["Inputs!A1"],
        },
        {
            "band": "Inputs!B1",
            "sheet": "Inputs",
            "label": "Tax rate / %",
            "kind": "input",
            "vtype": "numeric",
            "cells": ["Inputs!B1"],
        },
    ],
    "embedded_literals": [],
}


def save_workbook(path):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Inputs"
    sheet["A1"] = "Tax rate"
    sheet["B1"] = 0.25
    book.save(path)


MODEL_HEADER = "| Variable | Inventory row IDs | Plausible external source(s) |"
FINAL_HEADER = (
    "| Variable | Workbook cells and values | Plausible external source(s) |"
)

# Sorted inventory order for the two-variable fixture below:
#   R0001 = Inputs!B2:D2 "Revenue / 000$"  (values 100, 110, 121)
#   R0002 = Inputs!B1    "Tax rate / %"    (value 0.25)
GOOD_MODEL_TABLE = "\n".join([
    MODEL_HEADER,
    "|---|---|---|",
    "| Tax assumptions | R0002 | Tax authority (https://example.gov/rates) |",
    "| Historical revenue | R0001 | Audited financial statements |",
])

EXPECTED_FINAL_TABLE = "\n".join([
    FINAL_HEADER,
    "|---|---|---|",
    "| Tax assumptions | `Inputs!B1` Tax rate / %: `0.25` "
    "| Tax authority (https://example.gov/rates) |",
    "| Historical revenue | `Inputs!B2:D2` Revenue / 000$: `[100, 110, 121]` "
    "| Audited financial statements |",
])


def write_two_variable_fixture(root):
    """Workbook + segments with one label, one scalar, and one series band."""
    artifact = root / "0001-inputs.xlsx"
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Inputs"
    sheet["A1"] = "Tax rate"
    sheet["B1"] = 0.25
    sheet["A2"] = "Revenue"
    sheet["B2"] = 100
    sheet["C2"] = 110
    sheet["D2"] = 121
    book.save(artifact)
    seg_dir = root / "seg"
    seg_dir.mkdir(exist_ok=True)
    (seg_dir / "segments.json").write_text(json.dumps({
        "inputs": [
            {"band": "Inputs!A1", "sheet": "Inputs", "label": "Tax rate",
             "kind": "label", "vtype": "text", "cells": ["Inputs!A1"]},
            {"band": "Inputs!B1", "sheet": "Inputs", "label": "Tax rate / %",
             "kind": "input", "vtype": "numeric", "cells": ["Inputs!B1"]},
            {"band": "Inputs!B2:D2", "sheet": "Inputs",
             "label": "Revenue / 000$", "kind": "input", "vtype": "numeric",
             "cells": ["Inputs!B2", "Inputs!C2", "Inputs!D2"]},
        ],
        "embedded_literals": [],
    }), encoding="utf-8")
    return artifact, seg_dir


def make_model_stub(replies):
    """Sequential _call_model replacement; records every call's messages."""
    calls = []

    def stub(endpoint, api_key, model, project_id, messages):
        calls.append(messages)
        return replies[min(len(calls) - 1, len(replies) - 1)], "stop"

    stub.calls = calls
    return stub


class VariableSourceAuditTests(unittest.TestCase):
    def test_rl_key_takes_precedence_over_other_audit_keys(self):
        with tempfile.TemporaryDirectory() as temp:
            env_file = Path(temp) / ".env"
            env_file.write_text(
                "lbx_api_key=stale-key\n"
                "aligerr_org_key=org-key\n"
                "rl_api_key=current-key\n",
                encoding="utf-8",
            )
            self.assertEqual(read_env_key(env_file), "current-key")

    def test_inventory_is_deterministic_and_skips_labels(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            artifact = root / "0001-inputs.xlsx"
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "Inputs"
            sheet["A1"] = "Tax rate"
            sheet["B1"] = 0.25
            book.save(artifact)

            seg_dir = root / "seg"
            seg_dir.mkdir()
            (seg_dir / "segments.json").write_text(json.dumps({
                "inputs": [
                    {
                        "band": "Inputs!A1",
                        "sheet": "Inputs",
                        "label": "Tax rate",
                        "kind": "label",
                        "vtype": "text",
                        "cells": ["Inputs!A1"],
                    },
                    {
                        "band": "Inputs!B1",
                        "sheet": "Inputs",
                        "label": "Tax rate / %",
                        "kind": "input",
                        "vtype": "numeric",
                        "cells": ["Inputs!B1"],
                    },
                ],
                "embedded_literals": [],
            }), encoding="utf-8")

            first = build_inventory("0001", artifact, seg_dir)
            second = build_inventory("0001", artifact, seg_dir)
            self.assertEqual(first, second)
            self.assertEqual(len(first["variables"]), 1)
            self.assertEqual(first["variables"][0]["row_id"], "R0001")
            self.assertEqual(first["variables"][0]["value_summary"],
                             {"count": 1, "value": 0.25})

    def test_inventory_rows_carry_deterministic_row_ids(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            artifact, seg_dir = write_two_variable_fixture(root)
            first = build_inventory("0001", artifact, seg_dir)
            second = build_inventory("0001", artifact, seg_dir)
            self.assertEqual(first, second)
            self.assertEqual(
                [(row["row_id"], row["band"]) for row in first["variables"]],
                [("R0001", "Inputs!B2:D2"), ("R0002", "Inputs!B1")],
            )

    def _run_inventory_main(self, inputs, seg_root, audit_root):
        with contextlib.redirect_stdout(io.StringIO()):
            return xl_variable_source_audit.main([
                "0001",
                "--inputs-root", str(inputs),
                "--seg-root", str(seg_root),
                "--audit-root", str(audit_root),
                "--inventory-only",
            ])

    def test_main_resolves_xlsm_inputs_artifact(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            inputs = root / "inputs"
            inputs.mkdir()
            save_workbook(inputs / "0001-inputs.xlsm")
            seg_root = root / "seg"
            (seg_root / "0001").mkdir(parents=True)
            (seg_root / "0001" / "segments.json").write_text(
                json.dumps(SEGMENTS), encoding="utf-8")
            audit_root = root / "runs"

            rc = self._run_inventory_main(inputs, seg_root, audit_root)
            self.assertEqual(rc, 0)
            inventory = json.loads(
                (audit_root / "0001-variable-sources"
                 / "0001-inputs-variable-sources.inventory.json"
                 ).read_text(encoding="utf-8"))
            self.assertEqual(len(inventory["variables"]), 1)

    def test_main_prefers_xlsx_when_both_suffixes_exist(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            inputs = root / "inputs"
            inputs.mkdir()
            save_workbook(inputs / "0001-inputs.xlsx")
            # A stale sibling with a different value must lose to the .xlsx.
            book = openpyxl.Workbook()
            sheet = book.active
            sheet.title = "Inputs"
            sheet["A1"] = "Tax rate"
            sheet["B1"] = 0.99
            book.save(inputs / "0001-inputs.xlsm")
            seg_root = root / "seg"
            (seg_root / "0001").mkdir(parents=True)
            (seg_root / "0001" / "segments.json").write_text(
                json.dumps(SEGMENTS), encoding="utf-8")
            audit_root = root / "runs"

            rc = self._run_inventory_main(inputs, seg_root, audit_root)
            self.assertEqual(rc, 0)
            inventory = json.loads(
                (audit_root / "0001-variable-sources"
                 / "0001-inputs-variable-sources.inventory.json"
                 ).read_text(encoding="utf-8"))
            self.assertEqual(inventory["variables"][0]["value_summary"],
                             {"count": 1, "value": 0.25})

    def test_main_missing_artifact_raises_same_error(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            inputs = root / "inputs"
            inputs.mkdir()
            seg_root = root / "seg"
            (seg_root / "0001").mkdir(parents=True)
            (seg_root / "0001" / "segments.json").write_text(
                json.dumps(SEGMENTS), encoding="utf-8")
            with self.assertRaises(FileNotFoundError):
                self._run_inventory_main(inputs, seg_root, root / "runs")

    def test_table_validation_rejects_unknown_refs_and_values(self):
        rows = [json.dumps({
            "band": "Inputs!B1",
            "sheet": "Inputs",
            "label": "Tax rate / %",
            "kind": "input",
            "value_type": "numeric",
            "value_summary": {"count": 1, "value": 0.25},
        }, separators=(",", ":"))]
        metadata = {"sheet_names": ["Inputs"]}
        valid = (
            "| Variable | Workbook cells and values | "
            "Plausible external source(s) |\n"
            "|---|---|---|\n"
            "| Tax rate | `Inputs!B1` = 0.25 | Tax authority |"
        )
        self.assertEqual(validate_table(valid, rows, metadata), [])

        unknown_ref = valid.replace("Inputs!B1", "Inputs!C1")
        self.assertIn(
            "reference not in inventory: Inputs!C1",
            validate_table(unknown_ref, rows, metadata),
        )

        unknown_value = valid.replace("0.25", "0.3")
        violations = validate_table(unknown_value, rows, metadata)
        self.assertTrue(any("value not in inventory: 0.3" in item
                            for item in violations))

    def test_pinned_inactive_segmentation_generation_is_resolved(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source_dir = root / "source-generation"
            generation_dir = root / "segmentation-generation"
            generation_dir.mkdir()
            manifest = generation_dir / "generation-manifest.json"
            manifest.write_text('{"generation_id":"seg-id"}', encoding="utf-8")

            with (
                mock.patch(
                    "xl_source_publication.resolve_source_generation_by_id",
                    return_value=(source_dir, {"generation_id": "source-id"}),
                ) as source_resolver,
                mock.patch(
                    "xl_seg.publication.resolve_generation_by_id",
                    return_value=(generation_dir, {"generation_id": "seg-id"}),
                ) as segmentation_resolver,
            ):
                resolved, provenance = resolve_segmentation_directory(
                    root / "seg-root",
                    "0001",
                    segmentation_generation_id="seg-id",
                    source_generation_root=root / "source-root",
                    source_generation_id="source-id",
                )

            self.assertEqual(resolved, generation_dir)
            self.assertEqual(provenance["generation_id"], "seg-id")
            self.assertEqual(
                provenance["source_generation_id"], "source-id"
            )
            source_resolver.assert_called_once_with(
                root / "source-root" / "0001", "source-id"
            )
            segmentation_resolver.assert_called_once_with(
                root / "seg-root" / "0001",
                "seg-id",
                source_generation_dir=source_dir,
                require_pass=True,
            )

    def test_partial_generation_binding_is_rejected(self):
        with self.assertRaisesRegex(RuntimeError, "pinned audit requires"):
            resolve_segmentation_directory(
                "seg_out",
                "0001",
                segmentation_generation_id="seg-id",
            )

    def _generate(self, root, stub):
        artifact, seg_dir = write_two_variable_fixture(root)
        out_dir = root / "audit"
        with mock.patch.object(xl_variable_source_audit, "_call_model", stub):
            metadata = xl_variable_source_audit.generate_audit(
                "0001", artifact, seg_dir, out_dir, "test-key",
                log=lambda *args: None)
        return metadata, out_dir / "0001-inputs-variable-sources.md"

    def test_stubbed_reply_resolves_to_downstream_compatible_table(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            stub = make_model_stub([GOOD_MODEL_TABLE])
            metadata, markdown_path = self._generate(root, stub)

            self.assertEqual(len(stub.calls), 1)
            self.assertEqual(metadata["prompt_version"],
                             "variable-source-audit-v3")
            self.assertTrue(metadata["row_id_resolution"])
            self.assertEqual(metadata["validation"]["status"], "passed")
            self.assertIn("row_id_citations",
                          metadata["validation"]["checks"])

            markdown = markdown_path.read_text(encoding="utf-8")
            # Byte-identical final table in the historical downstream format.
            self.assertIn(EXPECTED_FINAL_TABLE, markdown)
            # No row_id ever leaks into the downstream artifact.
            self.assertIsNone(
                xl_variable_source_audit.ROW_ID_RE.search(markdown))
            # The import-parser header contract still matches: the header row
            # names a variable, a value, and a source column.
            header = next(line for line in markdown.splitlines()
                          if line.startswith("|"))
            self.assertEqual(header, FINAL_HEADER)
            lowered = header.lower()
            self.assertTrue("variable" in lowered and "value" in lowered
                            and "source" in lowered)

            draft = run_downstream_import(markdown_path,
                                          root / "draft.json")
            self.assertEqual(draft["row_count"], 2)
            self.assertEqual(
                [(row["draft_id"], row["value_text"]) for row in draft["rows"]],
                [("tax-assumptions", "Inputs!B1 Tax rate / %: 0.25"),
                 ("historical-revenue",
                  "Inputs!B2:D2 Revenue / 000$: [100, 110, 121]")],
            )
            self.assertEqual(draft["rows"][0]["source_urls"],
                             ["https://example.gov/rates"])

    def test_unknown_row_id_fails_validation_with_precise_message(self):
        bad = GOOD_MODEL_TABLE.replace("R0002", "R9999")
        violations = validate_row_ids(
            bad, {"R0001", "R0002"}, {"R0001", "R0002"})
        self.assertIn("unknown row_id: R9999 (row Tax assumptions)",
                      violations)

        out_of_batch = validate_row_ids(
            GOOD_MODEL_TABLE, {"R0001"}, {"R0001", "R0002"})
        self.assertIn("row_id not in this batch: R0002 (row Tax assumptions)",
                      out_of_batch)

        duplicate = validate_row_ids(
            GOOD_MODEL_TABLE.replace("R0001", "R0002"),
            {"R0001", "R0002"}, {"R0001", "R0002"})
        self.assertIn("duplicate row_id citation: R0002", duplicate)

        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            stub = make_model_stub([bad])
            with self.assertRaises(RuntimeError) as ctx:
                self._generate(root, stub)
            self.assertIn("unknown row_id: R9999", str(ctx.exception))
            # One retry through the correction loop before failing.
            self.assertEqual(len(stub.calls), 2)
            failure = json.loads(
                (root / "audit" / "0001-inputs-variable-sources.metadata.json"
                 ).read_text(encoding="utf-8"))
            self.assertEqual(failure["status"], "failed")

    def test_model_written_reference_is_never_trusted(self):
        # A hand-written reference in the row-id column is rejected outright.
        invented = "\n".join([
            MODEL_HEADER,
            "|---|---|---|",
            "| Staff expenses | Staff Exp!G66 | HR system |",
        ])
        violations = validate_row_ids(
            invented, {"R0001", "R0002"}, {"R0001", "R0002"})
        self.assertTrue(any(
            "row_id column may cite only supplied row_ids" in item
            for item in violations))

        # In a table where other rows cite ids, a row without any row_id
        # is flagged too.
        two_rows = invented + "\n| Tax assumptions | R0002 | Tax authority |"
        self.assertIn(
            "row cites no row_id: Staff expenses",
            validate_row_ids(two_rows, {"R0001", "R0002"},
                             {"R0001", "R0002"}))

        mixed = invented.replace("Staff Exp!G66", "R0002, Staff Exp!G66")
        self.assertTrue(any(
            "row_id column may cite only supplied row_ids" in item
            for item in validate_row_ids(
                mixed, {"R0001", "R0002"}, {"R0001", "R0002"})))

        # End to end: the correction loop recovers, and the invented
        # reference never reaches the downstream Markdown.
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            stub = make_model_stub([mixed, GOOD_MODEL_TABLE])
            metadata, markdown_path = self._generate(root, stub)
            self.assertEqual(len(stub.calls), 2)
            self.assertEqual(metadata["validation"]["status"], "passed")
            markdown = markdown_path.read_text(encoding="utf-8")
            self.assertNotIn("Staff Exp!G66", markdown)
            self.assertIn("`Inputs!B1` Tax rate / %: `0.25`", markdown)

    def test_no_candidate_sentinel_row_is_allowed(self):
        table = "\n".join([
            MODEL_HEADER,
            "|---|---|---|",
            "| No externally sourced candidates identified | — | — |",
        ])
        self.assertEqual(validate_row_ids(table, set(), set()), [])
        final = resolve_table(table, {})
        self.assertIn(
            "| No externally sourced candidates identified | — | — |", final)
        self.assertTrue(final.startswith(FINAL_HEADER))

    def test_cache_reuse_and_prompt_version_invalidation(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            metadata, _ = self._generate(root, make_model_stub([GOOD_MODEL_TABLE]))
            self.assertFalse(metadata["cache_hit"])

            # A current complete-coverage cache reuses without a model call.
            second_stub = make_model_stub([GOOD_MODEL_TABLE])
            reused, _ = self._generate(root, second_stub)
            self.assertTrue(reused["cache_hit"])
            self.assertEqual(len(second_stub.calls), 0)

            # A cache written by the older prompt version is invalidated.
            metadata_path = (root / "audit"
                             / "0001-inputs-variable-sources.metadata.json")
            stale = json.loads(metadata_path.read_text(encoding="utf-8"))
            stale["prompt_version"] = "variable-source-audit-v1"
            metadata_path.write_text(json.dumps(stale, indent=2) + "\n",
                                     encoding="utf-8")
            third_stub = make_model_stub([GOOD_MODEL_TABLE])
            refreshed, _ = self._generate(root, third_stub)
            self.assertFalse(refreshed["cache_hit"])
            self.assertEqual(len(third_stub.calls), 1)

    def test_cached_markdown_tampering_requires_fresh_audit(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            _, markdown = self._generate(root, make_model_stub([GOOD_MODEL_TABLE]))
            markdown.write_text(markdown.read_text() + "\nInvented extra source claim.\n")
            stub = make_model_stub([GOOD_MODEL_TABLE])
            result, _ = self._generate(root, stub)
            self.assertFalse(result["cache_hit"])
            self.assertEqual(len(stub.calls), 1)

    def test_cache_preserves_inventory_bytes_and_rejects_changed_inventory(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            self._generate(root, make_model_stub([GOOD_MODEL_TABLE]))
            inventory_path = root / 'audit/0001-inputs-variable-sources.inventory.json'
            inventory = json.loads(inventory_path.read_text())
            # Simulate a valid historical serialization with a different key order.
            historical = json.dumps(dict(reversed(list(inventory.items()))), indent=4) + '\n'
            inventory_path.write_text(historical)
            stub = make_model_stub([GOOD_MODEL_TABLE])
            result, _ = self._generate(root, stub)
            self.assertTrue(result['cache_hit'])
            self.assertEqual(inventory_path.read_text(), historical)
            self.assertEqual(len(stub.calls), 0)
            inventory_path.write_text('{}')
            result, _ = self._generate(root, stub)
            self.assertFalse(result['cache_hit'])
            self.assertEqual(len(stub.calls), 1)

    def test_interrupted_audit_does_not_leave_complete_metadata(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            self._generate(root, make_model_stub([GOOD_MODEL_TABLE]))
            metadata_path = root / 'audit/0001-inputs-variable-sources.metadata.json'
            previous = json.loads(metadata_path.read_text())
            previous['prompt_version'] = 'obsolete'
            metadata_path.write_text(json.dumps(previous))
            def interrupted(*args, **kwargs):
                raise KeyboardInterrupt()
            with self.assertRaises(KeyboardInterrupt):
                self._generate(root, interrupted)
            self.assertEqual(json.loads(metadata_path.read_text())['status'], 'in_progress')

    def test_inventory_only_change_invalidates_audit_before_publication(self):
        from unittest.mock import patch
        from xl_variable_source_audit import publish_inventory_only
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            self._generate(root, make_model_stub([GOOD_MODEL_TABLE]))
            output = root / 'audit'
            metadata = output / '0001-inputs-variable-sources.metadata.json'
            inventory = output / '0001-inputs-variable-sources.inventory.json'
            previous_metadata, previous_inventory = metadata.read_bytes(), inventory.read_bytes()
            publish_inventory_only('0001', json.loads(previous_inventory), output)
            self.assertEqual(metadata.read_bytes(), previous_metadata)
            self.assertEqual(inventory.read_bytes(), previous_inventory)
            with patch('xl_variable_source_audit._publish_inventory', side_effect=KeyboardInterrupt):
                with self.assertRaises(KeyboardInterrupt):
                    publish_inventory_only('0001', {}, output)
            self.assertEqual(json.loads(metadata.read_text())['status'], 'incomplete')
            self.assertEqual(inventory.read_bytes(), previous_inventory)

    def test_cache_request_configuration_and_generator_are_bound(self):
        for field in ("request_config_sha256", "generator_version"):
            with tempfile.TemporaryDirectory() as temp:
                root = Path(temp)
                self._generate(root, make_model_stub([GOOD_MODEL_TABLE]))
                path = root / "audit/0001-inputs-variable-sources.metadata.json"
                value = json.loads(path.read_text())
                value[field] = "obsolete"
                path.write_text(json.dumps(value))
                stub = make_model_stub([GOOD_MODEL_TABLE])
                result, _ = self._generate(root, stub)
                self.assertFalse(result["cache_hit"])
                self.assertEqual(len(stub.calls), 1)


    def test_incomplete_rows_and_nonempty_sentinel_fail_closed(self):
        subset = "\n".join(GOOD_MODEL_TABLE.splitlines()[:3])
        self.assertIn("missing inventory row dispositions: R0001",
                      validate_row_ids(subset, {"R0001", "R0002"}, {"R0001", "R0002"}))
        sentinel = MODEL_HEADER + "\n|---|---|---|\n| No externally sourced candidates identified | — | — |"
        self.assertTrue(validate_row_ids(sentinel, {"R0001"}, {"R0001"}))
        with tempfile.TemporaryDirectory() as temp:
            stub = make_model_stub([subset])
            with self.assertRaisesRegex(RuntimeError, "missing inventory row dispositions"):
                self._generate(Path(temp), stub)
            self.assertEqual(len(stub.calls), 2)

    def _coverage_files(self, root, table=GOOD_MODEL_TABLE):
        metadata, markdown = self._generate(root, make_model_stub([table]))
        metadata_path = markdown.with_suffix(".metadata.json")
        return metadata_path, markdown.with_suffix(".inventory.json"), markdown, Path(metadata["coverage_path"])

    def test_exclusions_preserve_values_and_separate_dispositions(self):
        table = GOOD_MODEL_TABLE.replace("Tax assumptions", "Excluded: Internal model tax scenario")
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            paths = self._coverage_files(root, table)
            report = xl_variable_source_audit.validate_saved_audit_coverage(*paths[:3])
            self.assertEqual(report["excluded_row_ids"], ["R0002"])
            self.assertEqual(report["included_row_ids"], ["R0001"])
            draft = run_downstream_import(paths[2], root / "draft.json")
            self.assertEqual(draft["row_count"], 2)
            self.assertTrue(draft["rows"][0]["variable_text"].startswith("Excluded:"))
            self.assertIn("0.25", draft["rows"][0]["value_text"])
        for reason, source in [("Excluded", "Workbook evidence"),
                               ("Excluded: not externally sourced", "Workbook evidence"),
                               ("Excluded: Internal assumption", "—")]:
            bad = MODEL_HEADER + "\n|---|---|---|\n| " + reason + " | R0001 | " + source + " |"
            self.assertTrue(validate_row_ids(bad, {"R0001"}, {"R0001"}))

    def test_saved_coverage_rejects_tampering_and_missing_evidence(self):
        for damage in ("missing", "bytes", "path", "subset_rehashed", "final_rehashed"):
            with self.subTest(damage=damage), tempfile.TemporaryDirectory() as temp:
                paths = self._coverage_files(Path(temp))
                metadata_path, inventory_path, markdown, coverage_path = paths
                metadata = json.loads(metadata_path.read_text())
                if damage == "missing":
                    coverage_path.unlink()
                elif damage == "bytes":
                    coverage_path.write_text(coverage_path.read_text() + " ")
                elif damage == "path":
                    metadata["coverage_path"] = str(coverage_path.with_name("other.json"))
                elif damage == "subset_rehashed":
                    coverage = json.loads(coverage_path.read_text())
                    coverage["batches"][0]["model_table"] = "\n".join(GOOD_MODEL_TABLE.splitlines()[:3])
                    coverage_path.write_text(json.dumps(coverage))
                    metadata["coverage_sha256"] = hashlib.sha256(coverage_path.read_bytes()).hexdigest()
                else:
                    markdown.write_text(markdown.read_text().replace("Historical revenue", "Other revenue"))
                    metadata["markdown_sha256"] = hashlib.sha256(markdown.read_bytes()).hexdigest()
                metadata_path.write_text(json.dumps(metadata))
                before = {str(p): p.read_bytes() for p in paths if p.exists()}
                with self.assertRaises(ValueError):
                    xl_variable_source_audit.validate_saved_audit_coverage(*paths[:3])
                self.assertEqual(before, {str(p): p.read_bytes() for p in paths if p.exists()})
                stub = make_model_stub([GOOD_MODEL_TABLE])
                refreshed, _ = self._generate(Path(temp), stub)
                self.assertFalse(refreshed["cache_hit"])
                self.assertEqual(len(stub.calls), 1)

    def test_inventory_duplicate_identity_and_malformed_schema_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            paths = self._coverage_files(Path(temp))
            inventory = json.loads(paths[1].read_text())
            coverage = json.loads(paths[3].read_text())
            for damage in ("duplicate", "missing_id", "bad_count", "missing_scope", "not_list"):
                altered, proof = copy.deepcopy(inventory), copy.deepcopy(coverage)
                if damage == "duplicate":
                    altered["variables"][1]["row_id"] = altered["variables"][0]["row_id"]
                elif damage == "missing_id":
                    altered["variables"][0].pop("row_id")
                elif damage == "bad_count":
                    altered["variables"][0]["value_summary"]["count"] = "three"
                elif damage == "missing_scope":
                    altered["metadata"] = {}
                else:
                    altered["variables"] = {}
                proof["inventory_sha256"] = xl_variable_source_audit._canonical_hash(altered)
                self.assertTrue(xl_variable_source_audit.validate_audit_coverage(paths[2].read_text(), altered, proof), damage)

    def test_genuine_empty_inventory_requires_explicit_sentinel(self):
        inventory = {"metadata": {"workbook": "0001", "artifact": "0001.xlsx", "sheet_names": ["Inputs"]}, "variables": []}
        table = MODEL_HEADER + "\n|---|---|---|\n| No externally sourced candidates identified | — | — |"
        coverage = {"schema_version": "variable-source-audit-coverage/v1", "inventory_sha256": xl_variable_source_audit._canonical_hash(inventory),
                    "batches": [{"row_ids": [], "model_table": table}]}
        self.assertEqual(xl_variable_source_audit.validate_audit_coverage(resolve_table(table, {}), inventory, coverage), [])
        coverage["batches"][0]["model_table"] = MODEL_HEADER + "\n|---|---|---|"
        self.assertTrue(xl_variable_source_audit.validate_audit_coverage("", inventory, coverage))

    def test_complete_multibatch_partition_and_cross_batch_rejection(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            artifact, segments = write_two_variable_fixture(root)
            inventory = build_inventory("0001", artifact, segments)
            # Enough source context to exceed the actual production chunk size.
            for row in inventory["variables"]:
                row["context"] = "x" * 60000
            batches, final = [], []
            for row in inventory["variables"]:
                table = MODEL_HEADER + "\n|---|---|---|\n| Historical evidence | " + row["row_id"] + " | Workbook-supplied record |"
                batches.append({"row_ids": [row["row_id"]], "model_table": table})
                final.append(resolve_table(table, {row["row_id"]: row}))
            coverage = {"schema_version": "variable-source-audit-coverage/v1", "inventory_sha256": xl_variable_source_audit._canonical_hash(inventory), "batches": batches}
            text = "# Audit\n" + "\n".join(final)
            self.assertEqual(xl_variable_source_audit.validate_audit_coverage(text, inventory, coverage), [])
            coverage["batches"][1]["model_table"] = coverage["batches"][0]["model_table"]
            self.assertTrue(xl_variable_source_audit.validate_audit_coverage(text, inventory, coverage))

    def test_interrupted_coverage_publication_cannot_reuse_completed_audit(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            paths = self._coverage_files(root)
            metadata = json.loads(paths[0].read_text())
            metadata["prompt_version"] = "obsolete"
            paths[0].write_text(json.dumps(metadata))
            original = xl_variable_source_audit._atomic_text
            def interrupted(path, text):
                if str(path).endswith(".md"):
                    raise KeyboardInterrupt()
                return original(path, text)
            with mock.patch.object(xl_variable_source_audit, "_atomic_text", interrupted):
                with self.assertRaises(KeyboardInterrupt):
                    self._generate(root, make_model_stub([GOOD_MODEL_TABLE]))
            self.assertEqual(json.loads(paths[0].read_text())["status"], "in_progress")
            with self.assertRaisesRegex(ValueError, "not complete"):
                xl_variable_source_audit.validate_saved_audit_coverage(*paths[:3])

    def test_raw_and_final_table_structure_cannot_hide_extra_columns(self):
        extra = GOOD_MODEL_TABLE.replace("Historical revenue |", "Historical revenue | Unbound content |")
        self.assertTrue(validate_row_ids(extra, {"R0001", "R0002"}, {"R0001", "R0002"}))
        with tempfile.TemporaryDirectory() as temp:
            paths = self._coverage_files(Path(temp))
            inventory, coverage = json.loads(paths[1].read_text()), json.loads(paths[3].read_text())
            changed = paths[2].read_text().replace("Historical revenue |", "Historical revenue | Unbound content |")
            self.assertTrue(xl_variable_source_audit.validate_audit_coverage(changed, inventory, coverage))

    def test_downstream_rejects_stale_versions_model_routing_and_actual_prompt(self):
        for field in ("generator_version", "prompt_version", "model", "endpoint", "project_id", "request_config_sha256", "system_prompt_sha256"):
            with self.subTest(field=field), tempfile.TemporaryDirectory() as temp:
                paths = self._coverage_files(Path(temp))
                metadata = json.loads(paths[0].read_text())
                metadata[field] = "stale"
                paths[0].write_text(json.dumps(metadata))
                with self.assertRaises(xl_variable_source_audit.AuditRevalidationRequired):
                    xl_variable_source_audit.validate_saved_audit_coverage(*paths[:3])
        with tempfile.TemporaryDirectory() as temp:
            paths = self._coverage_files(Path(temp))
            with mock.patch.object(xl_variable_source_audit, "SYSTEM_PROMPT", "New instructions without a version bump"):
                with self.assertRaises(xl_variable_source_audit.AuditRevalidationRequired):
                    xl_variable_source_audit.validate_saved_audit_coverage(*paths[:3])

    def test_explicit_custom_configuration_validates_only_with_expected_configuration(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp); artifact, seg_dir = write_two_variable_fixture(root)
            config = {"model": "approved-custom-model", "endpoint": "https://example.test/proxy", "project_id": "custom-project"}
            with mock.patch.object(xl_variable_source_audit, "_call_model", make_model_stub([GOOD_MODEL_TABLE])):
                result = xl_variable_source_audit.generate_audit("0001", artifact, seg_dir, root / "audit", "key", log=lambda *args: None, **config)
            markdown = Path(result["markdown"]); paths = (markdown.with_suffix(".metadata.json"), Path(result["inventory"]), markdown)
            with self.assertRaises(xl_variable_source_audit.AuditRevalidationRequired):
                xl_variable_source_audit.validate_saved_audit_coverage(*paths)
            self.assertTrue(xl_variable_source_audit.validate_saved_audit_coverage(*paths, **config)["passed"])

    def test_saved_request_and_response_cannot_be_rehashed_away(self):
        for kind in ("request", "response", "history"):
            with self.subTest(kind=kind), tempfile.TemporaryDirectory() as temp:
                paths = self._coverage_files(Path(temp))
                coverage = json.loads(paths[3].read_text()); batch = coverage["batches"][0]
                if kind == "history":
                    batch["history"] = []
                else:
                    binding = batch[kind]; path = Path(binding["path"]); value = json.loads(path.read_text())
                    if kind == "request":
                        value["messages"][0]["content"] = "Different instructions"
                    else:
                        value["content"] = value["content"].replace("Historical revenue", "Different selection")
                    path.write_text(json.dumps(value)); binding["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
                paths[3].write_text(json.dumps(coverage)); metadata = json.loads(paths[0].read_text())
                metadata["coverage_sha256"] = hashlib.sha256(paths[3].read_bytes()).hexdigest(); paths[0].write_text(json.dumps(metadata))
                with self.assertRaises(ValueError):
                    xl_variable_source_audit.validate_saved_audit_coverage(*paths[:3])

    def test_complete_looking_table_with_incomplete_model_finish_is_rejected(self):
        for finish in ('length', 'content_filter', ''):
            with self.subTest(finish=finish), tempfile.TemporaryDirectory() as temp:
                root = Path(temp)
                with self.assertRaisesRegex(RuntimeError, 'model completion is incomplete'):
                    self._generate(root, lambda *args: (GOOD_MODEL_TABLE, finish))
                metadata = json.loads((root / 'audit/0001-inputs-variable-sources.metadata.json').read_text())
                self.assertEqual(metadata['status'], 'failed')
                # Retain the actual received response and its consumed attempt.
                responses = list((root / 'audit/.audit-attempts').rglob('*.response.json'))
                self.assertEqual(len(responses), 1)
                self.assertEqual(json.loads(responses[0].read_text())['finish_reason'], finish)

    def test_saved_audit_rejects_rehashed_incomplete_model_finish(self):
        with tempfile.TemporaryDirectory() as temp:
            paths = self._coverage_files(Path(temp))
            metadata_path, inventory_path, markdown, coverage_path = paths
            coverage = json.loads(coverage_path.read_text())
            batch = coverage['batches'][0]
            response_path = Path(batch['response']['path'])
            response = json.loads(response_path.read_text())
            response['finish_reason'] = 'length'
            response_path.write_text(json.dumps(response))
            binding = xl_variable_source_audit._file_binding(response_path)
            batch['response'] = binding
            batch['response_history'][-1] = binding
            coverage_path.write_text(json.dumps(coverage))
            metadata = json.loads(metadata_path.read_text())
            metadata['coverage_sha256'] = hashlib.sha256(coverage_path.read_bytes()).hexdigest()
            metadata['finish_reasons'] = ['length']
            metadata_path.write_text(json.dumps(metadata))
            with self.assertRaisesRegex(ValueError, 'model completion is incomplete'):
                xl_variable_source_audit.validate_saved_audit_coverage(metadata_path, inventory_path, markdown)

    def test_restart_after_transport_interruption_consumes_persisted_attempt(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp); calls = []
            def interrupted(*args):
                calls.append(1); raise KeyboardInterrupt()
            for _ in range(2):
                with self.assertRaises(KeyboardInterrupt):
                    self._generate(root, interrupted)
            self.assertEqual(len(calls), 2)
            stub = make_model_stub([GOOD_MODEL_TABLE])
            with self.assertRaisesRegex(RuntimeError, "semantic budget exhausted"):
                self._generate(root, stub)
            self.assertEqual(len(stub.calls), 0)
            self.assertEqual(len(list((root / "audit/.audit-attempts").rglob("*.request.json"))), 2)

    def test_interrupted_publication_resumes_accepted_response_without_new_attempt(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp); atomic = xl_variable_source_audit._atomic_text
            def interrupted(path, text):
                if str(path).endswith(".coverage.json"):
                    raise KeyboardInterrupt()
                return atomic(path, text)
            stub = make_model_stub([GOOD_MODEL_TABLE])
            with mock.patch.object(xl_variable_source_audit, "_atomic_text", interrupted):
                with self.assertRaises(KeyboardInterrupt):
                    self._generate(root, stub)
            result, markdown = self._generate(root, stub)
            self.assertEqual(len(stub.calls), 1)
            self.assertTrue(xl_variable_source_audit.validate_saved_audit_coverage(markdown.with_suffix(".metadata.json"), Path(result["inventory"]), markdown)["passed"])

    def _inherited_receipt(self, root):
        prior_root = root / "prior"; prior_root.mkdir()
        paths = self._coverage_files(prior_root)
        coverage = json.loads(paths[3].read_text()); attempt = coverage["batches"][0]
        # Legacy evidence may lack the new journal schema but retains actual
        # request/response payloads, bound by a new receipt without modifying them.
        receipt = root / "inherited.json"
        receipt.write_text(json.dumps({"schema_version": "audit-inherited-budget/v1", "workbook": "0001",
            "attempts": [{"request": attempt["request"], "response": attempt["response"]}]}))
        return {"path": str(receipt), "sha256": hashlib.sha256(receipt.read_bytes()).hexdigest()}

    def test_hash_bound_inherited_call_leaves_exactly_one_attempt(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp); inherited = self._inherited_receipt(root)
            new = root / "fresh"; new.mkdir(); artifact, seg_dir = write_two_variable_fixture(new)
            stub = make_model_stub(["\n".join(GOOD_MODEL_TABLE.splitlines()[:3])])
            for _ in range(2):
                with mock.patch.object(xl_variable_source_audit, "_call_model", stub):
                    with self.assertRaises(RuntimeError):
                        xl_variable_source_audit.generate_audit("0001", artifact, seg_dir, new / "audit", "key", inherited_budget=inherited, log=lambda *args: None)
            self.assertEqual(len(stub.calls), 1)

    def test_altered_inherited_evidence_and_wrong_workbook_fail_before_call(self):
        for damage in ("bytes", "workbook", "duplicate"):
            with self.subTest(damage=damage), tempfile.TemporaryDirectory() as temp:
                root = Path(temp); inherited = self._inherited_receipt(root)
                path = Path(inherited["path"]); receipt = json.loads(path.read_text())
                if damage == "bytes":
                    evidence = Path(receipt["attempts"][0]["response"]["path"]); evidence.write_text(evidence.read_text() + " ")
                else:
                    if damage == "workbook": receipt["workbook"] = "other"
                    else: receipt["attempts"] *= 2
                    path.write_text(json.dumps(receipt)); inherited["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
                new = root / "fresh"; new.mkdir(); artifact, seg_dir = write_two_variable_fixture(new)
                stub = make_model_stub([GOOD_MODEL_TABLE])
                with mock.patch.object(xl_variable_source_audit, "_call_model", stub):
                    with self.assertRaises(ValueError):
                        xl_variable_source_audit.generate_audit("0001", artifact, seg_dir, new / "audit", "key", inherited_budget=inherited, log=lambda *args: None)
                self.assertEqual(len(stub.calls), 0)

    def test_inherited_budget_debits_only_the_matching_batch(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp); artifact, seg_dir = write_two_variable_fixture(root)
            inventory = build_inventory("0001", artifact, seg_dir)
            for row in inventory["variables"]: row["context"] = "x" * 60000
            chunks = xl_variable_source_audit._split_chunks(inventory)
            self.assertEqual(len(chunks), 2)
            def table(row_id):
                return MODEL_HEADER + "\n|---|---|---|\n| Supplied historical fact | " + row_id + " | Internal record |"
            request = root / "original-request.json"; response = root / "original-response.json"
            request.write_text(json.dumps({"endpoint": "https://old.example.test", "model": "previous-model",
                "messages": xl_variable_source_audit._base_messages(*chunks[0], 1, 2)}))
            response.write_text(json.dumps({"content": table("R0001"), "finish_reason": "stop"}))
            receipt = root / "inherited.json"; receipt.write_text(json.dumps({"schema_version": "audit-inherited-budget/v1", "workbook": "0001",
                "attempts": [{"request": xl_variable_source_audit._file_binding(request), "response": xl_variable_source_audit._file_binding(response)}]}))
            inherited = xl_variable_source_audit._file_binding(receipt)
            self.assertEqual(xl_variable_source_audit._inherited_batch_counts(inherited, inventory), {1: 1, 2: 0})
            sentinel = MODEL_HEADER + "\n|---|---|---|\n| No externally sourced candidates identified | — | — |"
            stub = make_model_stub([table("R0001"), sentinel, table("R0002")])
            with mock.patch.object(xl_variable_source_audit, "build_inventory", return_value=inventory), mock.patch.object(xl_variable_source_audit, "_call_model", stub):
                result = xl_variable_source_audit.generate_audit("0001", artifact, seg_dir, root / "audit", "key", inherited_budget=inherited, log=lambda *args: None)
            self.assertEqual(len(stub.calls), 3)
            markdown = Path(result["markdown"])
            self.assertTrue(xl_variable_source_audit.validate_saved_audit_coverage(markdown.with_suffix(".metadata.json"), Path(result["inventory"]), markdown)["passed"])
            # The same prior call cannot be distributed across a changed partition.
            for row in inventory["variables"]: row.pop("context")
            with self.assertRaisesRegex(ValueError, "cannot map exactly"):
                xl_variable_source_audit._inherited_batch_counts(inherited, inventory)

    def test_deleted_request_history_does_not_restore_an_attempt(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            def interrupted(*args): raise KeyboardInterrupt()
            with self.assertRaises(KeyboardInterrupt): self._generate(root, interrupted)
            next((root / "audit/.audit-attempts").rglob("*.request.json")).unlink()
            stub = make_model_stub([GOOD_MODEL_TABLE])
            with self.assertRaisesRegex(RuntimeError, "history is incomplete"):
                self._generate(root, stub)
            self.assertEqual(len(stub.calls), 0)

    def test_remaining_single_attempt_does_not_reset_semantic_budget(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            artifact, seg_dir = write_two_variable_fixture(root)
            stub = make_model_stub(["\n".join(GOOD_MODEL_TABLE.splitlines()[:3])])
            with mock.patch.object(xl_variable_source_audit, "_call_model", stub):
                with self.assertRaisesRegex(RuntimeError, "missing inventory row dispositions"):
                    xl_variable_source_audit.generate_audit("0001", artifact, seg_dir, root / "audit", "key", semantic_attempt_limit=1, log=lambda *args: None)
            self.assertEqual(len(stub.calls), 1)
            with self.assertRaisesRegex(ValueError, "semantic_attempt_limit"):
                xl_variable_source_audit.generate_audit("0001", artifact, seg_dir, root / "audit", "key", semantic_attempt_limit=3)


if __name__ == "__main__":
    unittest.main()
