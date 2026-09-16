"""Source-proved operand occurrences for independent dialogue review only.

No answer values are projected. The reviewer receives exact source operand
identity and whether it collides with an answer; the writer receives neither.
This is evidence for judgment, never a semantic PASS override.
"""
from collections import Counter
import hashlib
import json
import math
import re
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from scripts.batch.harbor_admission import read,record,require
from scripts.batch.harbor_tracker_lifecycle import checked_record
from scripts.batch.harbor_readiness import skill_module,ROOT
from xl_ast_graph import parse_formula

SCHEMA='dialogue-source-operands/v1'
GUIDANCE='''\nAdditional independently checked source-operand evidence is supplied only to you, the reviewer. It is not an answer key. An operand is not a graded answer merely because it is numeric. Assess the exact per-claim provenance, occurrence budget and target_collision field; do not infer a target collision from appearance. A collision is permitted only for its independently licensed occurrence, never for an extra answer disclosure. This evidence does not override any accuracy, completeness or naturalness requirement. Return your own full review.\n'''


class SourceFormulaTranslator(Translator):
    """Preserve qualified error operands while translating ordinary references."""
    @classmethod
    def translate_range(cls, range_str, rdelta, cdelta):
        # Translator splits at the final exclamation mark in #REF! and drops
        # the entire token. An error reference has no coordinates to shift.
        # Keep its exact sheet spelling and error, including escaped quotes.
        if re.fullmatch(r"(?:'(?:[^']|'')+'|[A-Za-z_][A-Za-z0-9_.]*)!#REF!", range_str):
            return range_str
        return super().translate_range(range_str, rdelta, cdelta)


def literal_operands(ast):
    """Signed numeric leaves, consuming a unary sign exactly once."""
    if ast.kind=='const' and ast.name=='number':
        number=float(ast.value)
        return Counter({number:1}) if math.isfinite(number) else Counter()
    if ast.kind=='prefix' and ast.name in ('+','-') and len(ast.args)==1:
        child=ast.args[0]
        if child.kind=='const' and child.name=='number':
            number=float(child.value)*(-1 if ast.name=='-' else 1)
            return Counter({number:1}) if math.isfinite(number) else Counter()
    result=Counter()
    for child in ast.args:result.update(literal_operands(child))
    return result


def canonical_operand_phrases(number):
    """Bounded display aliases for already proven integer source operands.

    This is not general NLP parsing and does not claim to detect arbitrary
    verbal answers. Ordinary semantic review remains mandatory for all prose.
    """
    if not math.isfinite(number) or int(number)!=number or abs(number)>999:return []
    units=('zero','one','two','three','four','five','six','seven','eight','nine','ten','eleven','twelve','thirteen','fourteen','fifteen','sixteen','seventeen','eighteen','nineteen')
    tens=('','','twenty','thirty','forty','fifty','sixty','seventy','eighty','ninety')
    def words(n):
        if n<20:return units[n]
        if n<100:return tens[n//10]+(' '+units[n%10] if n%10 else '')
        return units[n//100]+' hundred'+(' '+words(n%100) if n%100 else '')
    plain=words(abs(int(number)))
    return ['negative '+plain,'minus '+plain] if number<0 else [plain]



def bound_restatement(ast,prose,value,occurrences):
    """One source MAX zero bound may be restated as the same floor, not an answer."""
    while ast.kind=='prefix' and ast.name=='+' and len(ast.args)==1:ast=ast.args[0]
    if value!=0 or len(occurrences)!=2 or ast.kind!='func' or ast.name.upper()!='MAX' or len(ast.args)!=2:return None
    zero=[a for a in ast.args if a.kind=='const' and a.name=='number' and float(a.value)==0]
    if len(zero)!=1 or literal_operands(ast)[0.0]!=1:return None
    pattern=r'\bthe greater of [^.;\n]+? and (?P<bound>zero|0(?:\.0+)?)\s*,\s*so (?P<restatement>zero|0(?:\.0+)?) is the floor\b'
    matches=list(re.finditer(pattern,prose,re.I))
    if len(matches)!=1:return None
    match=matches[0];spans={(o['start'],o['end']) for o in occurrences}
    if spans!={match.span('bound'),match.span('restatement')}:return None
    return {'code':'redundant_zero_floor_restatement','source_operation':'MAX(expr,0)',
        'logical_operand_occurrences':1,'observed_spans':sorted(occurrences,key=lambda x:x['start']),
        'meaning':'Both spans name the same source-proved lower bound; independent semantic review remains required.'}


def column_increment_evidence(claim,payload,*,ordinal,source,task_dir,root,disclose):
    """Re-derive a legacy empty formula field from its complete source record."""
    reviewer=claim.get('reviewer_only') or {}
    if reviewer.get('source')!='convention_detector' or reviewer.get('entry')!='column_literal_increment':return None
    require(reviewer.get('evidence')=='','only the original empty formula field is supported')
    matches=[r for r in payload['records']if r.get('entry')=='column_literal_increment' and r.get('cell_keys')==reviewer.get('cells')]
    require(len(matches)==1,'unique complete increment record required')
    rec=matches[0];extract=skill_module(root,'additional-assumptions-dialogue','extract_claims')
    expected=extract.build_claim(rec,ordinal)
    formula=expected['reviewer_only']['evidence']
    require(isinstance(formula,str) and formula.startswith('='),'source increment formula missing')
    expected['reviewer_only']['evidence']=''
    require(expected==claim,'legacy increment claim differs beyond empty formula evidence')
    gold=disclose.Book(source['path']);delivered=disclose.Book(disclose.find_environment(Path(task_dir)))
    targets,_=disclose.load_key(Path(task_dir))
    closure=disclose.graded_closure(gold,{disclose.parse_ref(x,gold.sheets[0])for x in targets})
    require(disclose.ship_column_literal_increment(dict(rec),dict(gold=gold,delivered=delivered,targets=closure)),
            'original source increment full-family or graded closure proof failed')
    return dict(reviewer,evidence=formula,source='custom_method_detector')


def project(claims,draft,*,source,answer_key,task_dir,root=ROOT):
    checked_record(source);checked_record(answer_key)
    require(Path(answer_key['path'])==Path(task_dir)/'tests/answer_key.json','root-selected current answer key required')
    validator=skill_module(root,'additional-assumptions-dialogue','validate_dialogue');d=validator.disclose
    payload=read(record(Path(task_dir)/'tests/disclosure.json'))
    require(payload.get('golden')==source['path'] and payload.get('golden_sha256')==source['sha256'],'disclosure golden source differs')
    # Applied stages remove disclosure prose, so exact current claims are
    # independently checked by the composing readiness gate. Here we prove each
    # literal's source cells/formula, rather than trusting the extracted text.
    targets=d.numeric_targets(Path(task_dir));key_cells,_=d.load_key(Path(task_dir))
    target_cells={d.parse_ref(ref,'') for ref in key_cells if '!' in ref}
    turns=validator.parse_turns(draft);mapped,faults=validator.map_turns_to_claims(turns,claims['claims'])
    require(not faults,'operand projection requires exact claim-to-turn mapping')
    book=load_workbook(source['path'],read_only=True,data_only=False,keep_links=False)
    rows=[]
    try:
        for ordinal,claim in enumerate(claims['claims'],1):
            reviewer=column_increment_evidence(claim,payload,ordinal=ordinal,source=source,task_dir=task_dir,root=root,disclose=d) or claim.get('reviewer_only') or {}
            evidence=reviewer.get('evidence') or ''
            row={'record_id':claim['record_id'],'operands':[]}
            if reviewer.get('source')!='custom_method_detector':rows.append(row);continue
            refs=[d.parse_ref(ref,'') for ref in reviewer.get('cells') or []]
            cells=[tuple(ref.rsplit('!',1)) for ref in refs]
            require(cells and len(set(cells))==len(cells),'method operand cells missing or repeated')
            require(not set(refs)&target_cells,'target cells cannot license source operands')
            require(not set(d.refs_in(evidence,claim.get('sheet') or ''))&target_cells,'target-derived method cannot license operands')
            require(isinstance(evidence,str) and evidence.startswith('='),'method formula evidence missing')
            anchor=cells[0]
            require(anchor[0] in book.sheetnames,'source method sheet missing')
            original_cell=book[anchor[0]][anchor[1]]
            original=original_cell.value
            require(original_cell.data_type=='f' and isinstance(original,str) and original==evidence,'method evidence differs from source formula')
            proof=[]
            for sheet,coordinate in cells:
                require(sheet==anchor[0] and sheet in book.sheetnames,'method copied scope changed sheet')
                cell=book[sheet][coordinate]
                formula=cell.value
                expected=SourceFormulaTranslator(original,origin=anchor[1]).translate_formula(coordinate)
                require(cell.data_type=='f' and isinstance(formula,str) and formula==expected,'source method copied formula changed')
                proof.append({'cell':sheet+'!'+coordinate,'formula_sha256':hashlib.sha256(formula.encode()).hexdigest()})
            budget=literal_operands(parse_formula(original))
            prose='\n'.join(t['text'] for t in mapped[claim['record_id']] if validator.is_senior(t['speaker']))
            for value,count in sorted(budget.items()):
                occurrences=[]
                for match in d.NUMBER_RE.finditer(prose):
                    try:observed=float(match.group().replace(',',''))
                    except ValueError:continue
                    if d.same_number(observed,value):occurrences.append({'text':match.group(),'start':match.start(),'end':match.end()})
                for phrase in canonical_operand_phrases(value):
                    for match in re.finditer(r'(?<![\w-])'+re.escape(phrase)+r'(?![\w-])',prose,re.I):
                        # Do not count a smaller word inside a larger number,
                        # or a positive alias inside an explicitly signed one.
                        number_words=set('zero one two three four five six seven eight nine ten eleven twelve thirteen fourteen fifteen sixteen seventeen eighteen nineteen twenty thirty forty fifty sixty seventy eighty ninety hundred thousand million billion'.split())
                        left=re.search(r'([A-Za-z]+)[ \t-]+$',prose[:match.start()])
                        right=re.match(r'^[ \t-]+([A-Za-z]+)',prose[match.end():])
                        if left and left.group(1).lower() in number_words|{'negative','minus'}:continue
                        if right and right.group(1).lower() in number_words:continue
                        occurrences.append({'text':match.group(),'start':match.start(),'end':match.end()})
                collision=any(d.same_number(value,float(target)) for target in targets)
                alias=bound_restatement(parse_formula(original),prose,value,occurrences)
                effective=alias['logical_operand_occurrences'] if alias is not None else len(occurrences)
                require(not collision or effective<=count,'extra target occurrence exceeds the exact source operand budget')
                item={'value':value,'licensed_occurrences':count,'observed_occurrences':sorted(occurrences,key=lambda x:x['start']),
                    'target_collision':collision,'within_occurrence_budget':effective<=count,'source_evidence':proof}
                if alias is not None:item['restatement_warning']=alias
                row['operands'].append(item)
            rows.append(row)
    finally:book.close()
    checked_record(source);checked_record(answer_key)
    return {'schema_version':SCHEMA,'source_sha256':source['sha256'],'answer_key_sha256':answer_key['sha256'],
            'claims_sha256':hashlib.sha256(json.dumps(claims,sort_keys=True,separators=(',',':'),allow_nan=False).encode()).hexdigest(),
            'draft_sha256':hashlib.sha256(draft.encode()).hexdigest(),'claims':rows,
            'target_values_disclosed':False,'semantic_review_required':True}
