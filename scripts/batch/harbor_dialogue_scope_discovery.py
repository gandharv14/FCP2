"""Conservative first-round discovery of explicit relative-Year source-link scope.

Only original literal Year N axes and exact typed direct formulas participate.
No cache, financial evaluator, column-offset period inference or model call.
The installed owning scope proof and all later dialogue gates remain mandatory.
"""
from collections import defaultdict
from datetime import date,datetime
from pathlib import Path
import copy,re
from openpyxl import load_workbook
from openpyxl.formula.tokenizer import TokenizerError
from openpyxl.utils.exceptions import CellCoordinatesException
from openpyxl.utils.cell import get_column_letter,column_index_from_string,quote_sheetname
from scripts.batch.harbor_admission import read,record,require
from scripts.batch.harbor_model_executor import canonical,sha,retain,checked_bytes
from scripts.batch.harbor_readiness import skill_module
ROOT=Path(__file__).resolve().parents[2]
SCHEMA='dialogue-relative-year-discovery/v1'
MAX_CLAIMS=10000
MAX_SCANNED_CELLS=2000000
YEAR=re.compile(r'Year (0|[1-9][0-9]{0,2})',re.I)


class ScanLimit(ValueError):pass


def _ref(key):return quote_sheetname(key[0])+'!'+get_column_letter(key[1])+str(key[2])


def _key(scope,text):
    s,c,r=scope._reference(text);c=column_index_from_string(c)
    require(1<=c<=16384 and 1<=r<=1048576,'invalid source coordinate')
    return s,c,r


def _scan(book,bounds,remaining):
    for sheet,(lastrow,lastcol,firstrow,firstcol) in bounds.items():
        if sheet not in book.sheetnames:continue
        for row in book[sheet].iter_rows(min_row=firstrow,max_row=lastrow,min_col=firstcol,max_col=lastcol):
            for cell in row:
                remaining[0]-=1
                if remaining[0]<0:raise ScanLimit('bounded original-cell scan exhausted')
                if cell.value is not None:yield sheet,cell


def _derive(claims,source,disclosure,*,root):
    checked_bytes(source);checked_bytes(disclosure)
    scope=skill_module(root,'additional-assumptions-dialogue','claim_scope')
    ids=[c['record_id'] for c in claims];require(len(ids)==len(set(ids)),'duplicate original claim IDs')
    payload=dict(schema_version='dialogue-source-link-scope/v3',source_sha256=source['sha256'],disclosure_sha256=disclosure['sha256'],claim_ids=ids,rows=[])
    decisions={i:'individual: not a direct source selection' for i in ids}
    if len(claims)>MAX_CLAIMS:return payload,dict(decisions={i:'individual: bounded claim scan exceeded' for i in ids},observed=[])
    targets={};bounds={};remaining=[MAX_SCANNED_CELLS]
    for c in claims:
        private=c.get('reviewer_only',{})
        if private.get('entry')!='source_selection' or len(private.get('cells',[]))!=1:continue
        try:key=_key(scope,private['cells'][0])
        except (ValueError,KeyError,IndexError,TokenizerError,CellCoordinatesException):continue
        targets[c['record_id']]=key;s,col,row=key
        if s not in bounds:bounds[s]=(row,col,row,col)
        a,b,d,e=bounds[s];bounds[s]=(max(a,row),max(b,col),min(d,row),min(e,col))
    book=load_workbook(source['path'],read_only=True,data_only=False);observed={};links={}
    try:
        wanted=set(targets.values());formulas={}
        for sheet,cell in _scan(book,bounds,remaining):
            key=(sheet,cell.column,cell.row)
            if key in wanted and cell.data_type=='f':formulas[key]=cell.value
        for identity,target in targets.items():
            formula=formulas.get(target);decisions[identity]='individual: unsupported original formula'
            try:
                tokens=scope._tokens(formula)
                if len(tokens)!=1 or tokens[0].subtype!='RANGE':continue
                text=tokens[0].value
                origin=_key(scope,text if '!' in text else quote_sheetname(target[0])+'!'+text)
                if '[' in origin[0] or origin[0] not in book.sheetnames:continue
            except (ValueError,KeyError,IndexError,TokenizerError,CellCoordinatesException):continue
            links[identity]=(target,origin,formula)
        bounds={}
        for target,origin,_ in links.values():
            for sheet,col,row in (target,origin):
                a,b,_,_=bounds.get(sheet,(1,1,1,1));bounds[sheet]=(max(a,row-1),max(b,min(16384,col+2)),1,1)
        header_cells={};date_cells=defaultdict(list)
        for sheet,cell in _scan(book,bounds,remaining):
            if isinstance(cell.value,(date,datetime)):
                date_cells[(sheet,cell.row)].append(cell.column)
            if cell.data_type not in ('s','inlineStr') or not isinstance(cell.value,str):continue
            match=YEAR.fullmatch(cell.value)
            if match:header_cells[(sheet,cell.column,cell.row)]=(int(match[1]),cell.value,cell.data_type)
        byrow=defaultdict(list)
        for (sheet,col,row),(year,text,kind) in header_cells.items():byrow[(sheet,row)].append((col,year))
        axes=[]
        for (sheet,row),items in sorted(byrow.items()):
            run=[]
            for col,year in sorted(items):
                if run and (col!=run[-1][0]+1 or year!=run[-1][1]+1):
                    if len(run)>=3:axes.append((sheet,row,tuple(run)))
                    run=[]
                run.append((col,year))
            if len(run)>=3:axes.append((sheet,row,tuple(run)))
        date_axes=[]
        for (sheet,row),columns in sorted(date_cells.items()):
            run=[]
            for col in sorted(columns):
                if run and col!=run[-1]+1:
                    if len(run)>=3:date_axes.append((sheet,row,run[0],run[-1]))
                    run=[]
                run.append(col)
            if len(run)>=3:date_axes.append((sheet,row,run[0],run[-1]))
        def candidates(key):
            sheet,col,row=key
            return [a for a in axes if a[0]==sheet and a[1]<row and a[2][0][0]<=col<=a[2][-1][0]]
        families=defaultdict(list);claims_by_id={c['record_id']:c for c in claims}
        for identity,(target,origin,formula) in links.items():
            ta,sa=candidates(target),candidates(origin)
            if any(a[0]==key[0] and a[1]<key[2] and a[2]<=key[1]<=a[3] for key in (target,origin) for a in date_axes):
                decisions[identity]='individual: competing unsupported date axis';continue
            if len(ta)!=1 or len(sa)!=1:
                decisions[identity]='individual: ambiguous or absent complete Year axis';continue
            target_axis,source_axis=ta[0],sa[0]
            ty=dict(target_axis[2])[target[1]];sy=dict(source_axis[2])[origin[1]]
            if ty!=sy:decisions[identity]='individual: unequal explicit Year labels';continue
            row=dict(claim_id=identity,target_reference=_ref(target),source_reference=_ref(origin),original_formula=formula,
                target_period_header=_ref((target[0],target[1],target_axis[1])),source_period_header=_ref((origin[0],origin[1],source_axis[1])),
                target_period_axis_start=_ref((target[0],target_axis[2][0][0],target_axis[1])),source_period_axis_start=_ref((origin[0],source_axis[2][0][0],source_axis[1])),period_kind='relative_year',period=ty)
            claim=claims_by_id[identity]
            family=(target[0],target[2],origin[0],origin[2],claim['row_label'],claim['spoken'],tuple(claim['must_say']))
            families[family].append((identity,row,target_axis,source_axis))
        selected={}
        for family,members in families.items():
            if len(members)<2:
                for identity,*_ in members:decisions[identity]='individual: no repeated identical source rule'
                continue
            for identity,row,ta,sa in members:
                selected[identity]=row;decisions[identity]='proposed: unique paired original Year axes'
                target,origin,formula=links[identity];observed[_ref(target)]=dict(data_type='f',value=formula)
                for axis in (ta,sa):
                    for col,year in axis[2]:
                        key=(axis[0],col,axis[1]);_,text,kind=header_cells[key]
                        observed[_ref(key)]=dict(data_type=kind,value=text)
        payload['rows']=[selected[i] for i in ids if i in selected]
    except ScanLimit:
        payload['rows']=[];observed={};decisions={i:'individual: bounded original-cell scan exhausted' for i in ids}
    finally:book.close()
    checked_bytes(source);checked_bytes(disclosure)
    return payload,dict(decisions=decisions,observed=[dict(reference=k,**v) for k,v in sorted(observed.items())])


def _inputs(stage,source,root):
    stage=Path(stage);checked_bytes(source)
    disclosure=record(stage/'tests/disclosure.json');instruction=record(stage/'instruction.md')
    payload=read(disclosure)
    require(payload.get('golden_sha256')==source['sha256'] and payload.get('golden')==source['path'],'disclosure golden source differs from selected original')
    extractor=skill_module(root,'additional-assumptions-dialogue','extract_claims')
    claims=extractor.extract(stage)
    return disclosure,instruction,extractor,claims


def prepare(stage,source,directory,*,root=ROOT):
    """Fresh private pre-call evidence. Caller retains returned scope in Selection."""
    stage,directory=Path(stage).resolve(),Path(directory).resolve()
    require(not directory.exists(),'existing scope discovery requires exact validation, not regeneration')
    require(directory!=stage and stage not in directory.parents and directory not in stage.parents,'scope evidence must be outside shipped stage')
    disclosure,instruction,extractor,base=_inputs(stage,source,root)
    payload,details=_derive(base['claims'],source,disclosure,root=root)
    directory.mkdir(parents=True,mode=0o700)
    def put(name,value):retain(directory/name,canonical(value));return record(directory/name)
    scope=put('scope.json',payload);diagnostic=put('discovery.json',details)
    # Owning proof, not discovery heuristics, determines exact accepted grouping.
    actual=extractor.extract(stage,source_workbook=source['path'],scope_evidence=scope['path'])
    grouped=any(len(group)>1 for group in actual['shared_source_links']['groups'])
    if not grouped:actual=base
    proof=put('prepared.json',dict(schema_version=SCHEMA,producer=record(__file__),scope_validator=record(Path(skill_module(root,'additional-assumptions-dialogue','claim_scope').__file__)),extractor=record(Path(extractor.__file__)),source=source,disclosure=disclosure,instruction=instruction,
        original_claims_sha256=sha(canonical(base['claims'])),claim_ids=[c['record_id'] for c in base['claims']],scope=scope,discovery=diagnostic,used=grouped,final_claims_sha256=sha(canonical(actual)),semantic_attempts_added=0,semantic_budget_reset=False))
    return proof


def validate(stage,binding,*,root=ROOT):
    """Recompute exact first-call discovery without modifying any source or claim."""
    return _validate(stage,binding,root=root)


def _validate(stage,binding,*,root=ROOT,equivalence=False):
    value=read(binding);require(value.get('schema_version')==SCHEMA,'unknown discovery schema')
    directory=Path(binding['path']).parent
    require(binding==record(directory/'prepared.json') and value['scope']==record(directory/'scope.json') and value['discovery']==record(directory/'discovery.json'),'scope evidence namespace or bytes changed')
    disclosure,instruction,extractor,base=_inputs(stage,value['source'],root)
    payload,details=_derive(base['claims'],value['source'],disclosure,root=root)
    require(payload==read(value['scope']) and details==read(value['discovery']),'source formula/header discovery changed')
    actual=extractor.extract(Path(stage),source_workbook=value['source']['path'],scope_evidence=value['scope']['path'])
    grouped=any(len(g)>1 for g in actual['shared_source_links']['groups'])
    if not grouped:actual=base
    expected=dict(schema_version=SCHEMA,producer=record(__file__),scope_validator=record(Path(skill_module(root,'additional-assumptions-dialogue','claim_scope').__file__)),extractor=record(Path(extractor.__file__)),source=value['source'],disclosure=disclosure,instruction=instruction,
        original_claims_sha256=sha(canonical(base['claims'])),claim_ids=[c['record_id'] for c in base['claims']],scope=value['scope'],discovery=value['discovery'],used=grouped,final_claims_sha256=sha(canonical(actual)),semantic_attempts_added=0,semantic_budget_reset=False)
    if equivalence:
        from scripts.batch.harbor_tracker_lifecycle import checked_record
        for key in ('producer','scope_validator','extractor'):checked_record(value[key])
        value={k:v for k,v in value.items() if k not in ('producer','scope_validator','extractor')}
        expected={k:v for k,v in expected.items() if k not in ('producer','scope_validator','extractor')}
    require(value==expected,'scope discovery selection changed')
    return actual,({'source_workbook':value['source']['path'],'scope_evidence':value['scope']['path']} if grouped else {})


def initialize_dialogue_claims(stage,source,dialogue_directory,*,root=ROOT):
    """Integration hook, called once before claims/scaffold/model preparation.

    Later _dialogue checks must use validate() against this same receipt.
    Root retains receipt and returned scope in current typed content Selection.
    """
    directory=Path(dialogue_directory)
    require(not directory.exists(),'existing dialogue requires an explicit root recovery, not first-call discovery')
    proof=prepare(stage,source,directory/'source-scope',root=root)
    claims,options=validate(stage,proof,root=root)
    return dict(claims=claims,receipt=proof,scope_evidence=read(proof)['scope'] if options else None,extraction_options=options)
