"""Deterministic bounded disclosure shards with explicit unknown/blank semantics.

This module makes no model requests and never changes source workbooks. Full
source snapshots remain root evidence; each shard declares its exact selected
references and recursively rederived dependencies in both workbook versions.
"""
from pathlib import Path
import copy
import hashlib
import json
import re
from collections import Counter
from openpyxl.utils import range_boundaries
from scripts.batch import harbor_disclosure_payload_repair as tables
from scripts.batch import harbor_method_assessments as assessments
from scripts.batch import harbor_semantic_readiness as semantic
from scripts.batch.harbor_admission import record

SCHEMA='disclosure-partition-evidence/v1'
BOUNDED_SCHEMA='disclosure-partition-evidence/v2'
BOUNDS_SCHEMA='harbor-disclosure-partition-bounds/v1'
WIRE_BOUNDS_SCHEMA='harbor-disclosure-partition-bounds/v2'
WIRE_SCHEMA='disclosure-partition-evidence/v3'
EMPTY_SCHEMA='disclosure-partition-evidence/v4'
AUTHORIZATION_SHA256='acffe75c8ed401f66484a5bb6c39a15a0a3e209fc976967da0bddd6125126fd0'
WORKBOOK_SCHEMA='explicit-workbook-evidence-shard/v1'
SEMANTICS='selected_refs and blank_refs use cell-reference-runs/v1: [sheet,column_number,start_row,count], all addresses inclusive; expanding and lexically sorting selected_refs defines zero-based dependency_edges indexes. Only selected_refs are described. present_rows use row_columns; number_format_index selects number_formats. Four entries mean cached_value is absent; five mean it is present, including null. Exact source values, formulas, formats and types are retained. blank_refs are explicitly absent from the complete source snapshot. Every other address is unknown in this shard, never implicitly blank. Full snapshot hashes bind all omitted evidence.'
MAX_BYTES=750000
MAX_SENTENCES=100

class ScopeDiagnosisRequired(ValueError):pass

def canonical(value):return json.dumps(value,ensure_ascii=False,separators=(',',':'),sort_keys=True,allow_nan=False).encode()
def digest(value):return hashlib.sha256(canonical(value)).hexdigest()
def require(ok,message):
    if not ok:raise ValueError(message)

def selected_bounds(binding,body,golden_path,delivered_path):
    """Source-scoped representation limits; never a model/lease authorization."""
    from scripts.batch.harbor_tracker_lifecycle import checked_record
    from scripts.batch.harbor_admission import read
    checked_record(binding);value=read(binding)
    wire=value.get('schema_version')==WIRE_BOUNDS_SCHEMA
    require(set(value)==({'schema_version','authorization','full_body_sha256','sources','max_bytes','max_sentences','max_batches'}|({'wire_encoding'}if wire else set()))
            and value['schema_version']in(BOUNDS_SCHEMA,WIRE_BOUNDS_SCHEMA),'closed partition bounds required')
    if wire:
        from scripts.batch import harbor_disclosure_partition_wire as W
        require(isinstance(value['wire_encoding'],str) and value['wire_encoding']in (W.SCHEMA,getattr(W,'COLUMNAR_SCHEMA',None),getattr(W,'RUN_SCHEMA',None)),'unknown lossless shard encoding')
    checked_record(value['authorization'])
    require(value['authorization']['sha256']==AUTHORIZATION_SHA256,'partition bounds authorization differs')
    require(value['full_body_sha256']==digest(body) and value['sources']=={'golden':record(golden_path),'delivered':record(delivered_path)},'partition bounds source or body differs')
    require(type(value['max_bytes'])is int and (0 if wire else MAX_BYTES)<value['max_bytes']<=2000000
            and type(value['max_sentences'])is int and 1<=value['max_sentences']<=MAX_SENTENCES
            and type(value['max_batches'])is int and 1<=value['max_batches']<=64,'partition bounds exceed supported finite limits')
    return value

def _refs(row):
    """Uncertain scopes cannot disappear through an empty filter."""
    out=set();unknown=False
    for key in ('cell_keys','stated_cell_keys'):
        value=row.get(key)
        if value is None and key=='stated_cell_keys':continue
        if not isinstance(value,list) or not value or not all(tables.canonical_cell(x) for x in value):unknown=True
        else:out.update(value)
    source=row.get('source_cell')
    if source is not None:
        if tables.canonical_cell(source):out.add(source)
        elif source!='':unknown=True
    return out,unknown

def encode_refs(refs):
    from openpyxl.utils.cell import coordinate_from_string,column_index_from_string
    groups={}
    for ref in sorted(set(refs)):
        require(tables.canonical_cell(ref),'invalid reference')
        sheet,coord=ref.rsplit('!',1);column,row=coordinate_from_string(coord);groups.setdefault((sheet,column_index_from_string(column)),[]).append(row)
    runs=[]
    for (sheet,col),rows in sorted(groups.items()):
        rows=sorted(rows);start=last=rows[0]
        for row in rows[1:]:
            if row==last+1:last=row
            else:runs.append([sheet,col,start,last-start+1]);start=last=row
        runs.append([sheet,col,start,last-start+1])
    return dict(schema_version='cell-reference-runs/v1',runs=runs)

def decode_refs(value):
    from openpyxl.utils import get_column_letter
    require(isinstance(value,dict) and set(value)=={'schema_version','runs'} and value['schema_version']=='cell-reference-runs/v1' and isinstance(value['runs'],list),'unknown reference runs')
    refs=[]
    for run in value['runs']:
        require(isinstance(run,list) and len(run)==4,'reference run shape')
        sheet,col,start,count=run
        require(isinstance(sheet,str) and all(type(x) is int for x in (col,start,count)) and 1<=col<=16384 and 1<=start<=1048576 and 1<=count<=1048576-start+1,'reference run bounds')
        require(len(refs)+count<=2000000,'reference expansion requires diagnosis')
        refs.extend(sheet+'!'+get_column_letter(col)+str(row) for row in range(start,start+count))
    require(len(refs)==len(set(refs)) and encode_refs(refs)==value,'noncanonical/duplicate reference runs')
    return sorted(refs)

class _Workbook:
    def __init__(self,path,evidence):
        from xl_ast_graph import AstGraph
        self.path=Path(path);self.binding=record(self.path)
        require(evidence==semantic.workbook_evidence(self.binding),'full workbook evidence differs from current source')
        self.evidence=evidence;self.cells={};self.sheets={};self.cache={};self.named_dependencies={}
        for sheet in evidence['sheets']:
            self.sheets[sheet['sheet']]=sheet
            for cell in sheet['cells']:
                key=sheet['sheet']+'!'+cell['address'];require(key not in self.cells,'duplicate full snapshot cell');self.cells[key]=cell
        self.graph=AstGraph(self.path,max_range_expand=50000,verbose=False)
        try:self.graph._load();self.graph._collect_names()
        finally:self.graph._release_workbook()
    def close(self):self.graph.close()
    def static_name(self,token,formula_cell):
        """Resolve only literal absolute named areas, retaining their source definition."""
        sheet,name,external=self.graph._split_ref(token)
        if external is not None:return None
        lookup=sheet or formula_cell.rsplit('!',1)[0];key=name.upper();scope=None
        if (lookup,key) in self.graph.names_local:
            definition=self.graph.names_local[(lookup,key)];scope=lookup
        elif key in self.graph.names_global:definition=self.graph.names_global[key]
        else:return None
        refs=set();bad=False
        try:
            from xl_ast_graph import split_areas
            text=str(definition).strip();text=text[1:].strip()if text.startswith('=')else text
            areas=split_areas(text);require(bool(areas),'empty named area')
            for area in areas:
                target,span,ext=self.graph._split_ref(area)
                require(ext is None and target in self.sheets and
                        re.fullmatch(r'\$[A-Za-z]{1,3}\$[1-9][0-9]*(?::\$[A-Za-z]{1,3}\$[1-9][0-9]*)?',span),
                        'named dependency is not a literal absolute area')
                lo,row,hi,last=range_boundaries(span.replace('$',''))
                require(1<=lo<=hi<=16384 and 1<=row<=last<=1048576
                        and (hi-lo+1)*(last-row+1)<=50000,'named dependency exceeds bounded area')
                refs.update(self.graph.cell_id(target,r,c)for r in range(row,last+1)for c in range(lo,hi+1))
                require(len(refs)<=50000,'named dependency union exceeds bound')
        except Exception:bad=True
        if not bad:
            item=dict(name=key,scope_sheet=scope,definition=definition,targets=sorted(refs))
            self.named_dependencies.setdefault(formula_cell,{})[(scope,key)]=item
        return refs,bad
    def static_index(self,node,key):
        """Resolve a fixed snapshot's named integer case selector, without caches.

        Keep the name definition, literal selector, ROW reference and selected
        value. Formula-valued, missing or ambiguous selectors remain unknown.
        """
        if node.kind!='func' or node.name.upper()!='INDEX' or len(node.args)!=2:return None
        array,offset=node.args
        if array.kind!='ref' or offset.kind!='infix' or offset.name!='+' or len(offset.args)!=2:return None
        row,case=offset.args
        if row.kind!='func' or row.name.upper()!='ROW' or len(row.args)!=1 or row.args[0].kind!='ref' or case.kind!='ref':return None
        sheet,span,external=self.graph._split_ref(array.name)
        match=re.fullmatch(r'\$?([A-Za-z]{1,3}):\$?\1',span)
        if external is not None or match is None:return None
        sheet=sheet or key.rsplit('!',1)[0]
        refsheet,ref,external=self.graph._split_ref(row.args[0].name)
        refsheet=refsheet or key.rsplit('!',1)[0]
        owner=re.fullmatch(r'\$?([A-Za-z]{1,3})\$?([1-9][0-9]*)',ref)
        if external is not None or sheet not in self.sheets or refsheet!=sheet or owner is None or owner[1].upper()!=match[1].upper():return None
        named=self.static_name(case.name,key)
        if named is None:return None
        targets,bad=named
        if bad or len(targets)!=1:return None
        selector=next(iter(targets));cell=self.cells.get(selector)
        if cell is None or cell['data_type']!='n' or type(cell['value'])is not int:return None
        number=int(owner[2])+cell['value']
        from openpyxl.utils.cell import column_index_from_string
        if not 1<=int(owner[2])<=1048576 or not 1<=number<=1048576 or column_index_from_string(match[1])>16384:return None
        return {selector,sheet+'!'+owner[1].upper()+owner[2],sheet+'!'+match[1].upper()+str(number)}

    def dependencies(self,key):
        from xl_ast_graph import parse_formula
        if key in self.cache:return self.cache[key]
        cell=self.cells.get(key);refs=set();unknown=False
        if cell and cell['data_type']=='f':
            try:
                formula=cell['value']
                if isinstance(formula,dict):
                    # Workbook evidence preserves array metadata as well as text.
                    # Only a single-cell owner has the same dependency semantics
                    # as the original text expression; retain the complete value
                    # in the shard and leave every other formula object unknown.
                    require(set(formula)=={'formula_type','text','attributes'}
                            and formula['formula_type']=='array'
                            and isinstance(formula['text'],str)
                            and formula['text'].startswith('=')
                            and formula['attributes']=={'t':'array','ref':cell['address']}
                            and key.rsplit('!',1)[-1]==cell['address'],
                            'unsupported array formula ownership')
                    formula=formula['text']
                stack=[parse_formula(formula)]
                while stack:
                    node=stack.pop()
                    selected_index=self.static_index(node,key)
                    if selected_index is not None:
                        refs.update(selected_index);continue
                    stack.extend(node.args)
                    if node.kind=='func' and node.name.upper().replace('_XLFN.','') in {'INDIRECT','OFFSET'}:unknown=True
                    if node.kind=='ref':
                        named=self.static_name(node.name,key)
                        if named is not None:
                            found,bad=named;refs.update(found);unknown|=bad;continue
                        sheet_part,span,external=self.graph._split_ref(node.name)
                        if ':' in span and external is None:
                            try:
                                bounds=range_boundaries(span.replace('$',''))
                                if all(x is not None for x in bounds) and (bounds[2]-bounds[0]+1)*(bounds[3]-bounds[1]+1)<=50000 and (sheet_part is None or sheet_part in self.sheets):
                                    refsheet=sheet_part or key.rsplit('!',1)[0]
                                    refs.update(self.graph.cell_id(refsheet,r,c) for r in range(bounds[1],bounds[3]+1) for c in range(bounds[0],bounds[2]+1))
                                else:unknown=True
                            except Exception:unknown=True
                        if unknown:continue  # Never expand an unresolved or oversized scope.
                        for item in self.graph.resolve(node.name,key.rsplit('!',1)[0]):
                            if item[0]=='cell':refs.add(self.graph.cell_id(*item[1:]))
                            elif item[0]=='range':
                                sheet,span=item[1:3]
                                refs.update(self.graph.cell_id(sheet,r,c) for r,c in self.graph.index[sheet].cells(*range_boundaries(span)))
                                # Bounded literal ranges above include their blank members.
                                if (sheet,span) in self.graph._truncated_ranges:unknown=True
                            else:unknown=True
            except Exception:unknown=True
        self.cache[key]=(refs,unknown);return refs,unknown
    def shard(self,seeds,unknown=False):
        selected=set(seeds);todo=list(seeds);edges=set();fallback=unknown
        if any(k.rsplit('!',1)[0] not in self.sheets for k in selected):raise ScopeDiagnosisRequired('reference targets an absent sheet; cannot certify it as a blank cell')
        while True:
            while todo and not fallback:
                key=todo.pop();refs,bad=self.dependencies(key)
                if bad:fallback=True;break
                edges.update((key,ref) for ref in refs)
                for ref in refs-selected:selected.add(ref);todo.append(ref)
            if fallback:
                selected=set(self.cells)|set(seeds);edges=set()
                # Unknown/dynamic formula dependencies are not falsely certified.
                break
            sheets={key.rsplit('!',1)[0] for key in selected}
            labels={key for key,cell in self.cells.items() if key.rsplit('!',1)[0] in sheets and
                    ((isinstance(cell['value'],str) and cell['data_type']!='f') or isinstance(cell.get('cached_value'),str))}
            added=labels-selected
            if not added:break
            selected.update(added);todo.extend(added)
        rows=[];formats=[];ref_index={ref:i for i,ref in enumerate(sorted(selected))}
        for sheet in self.evidence['sheets']:
            name=sheet['sheet'];present=[]
            for cell in sheet['cells']:
                if name+'!'+cell['address'] not in selected:continue
                fmt=cell['number_format']
                if fmt not in formats:formats.append(fmt)
                item=[cell['address'],cell['value'],formats.index(fmt),cell['data_type']]
                if 'cached_value' in cell:item.append(cell['cached_value'])
                present.append(item)
            rows.append({**{k:copy.deepcopy(v) for k,v in sheet.items() if k!='cells'},'present_rows':present})
        metadata={k:copy.deepcopy(v) for k,v in self.evidence.items() if k not in {'sheets','blank_cell_semantics'}}
        named=[dict(formula_cell=cell,**item)for cell in sorted(selected)
               for _,item in sorted(self.named_dependencies.get(cell,{}).items(),key=lambda pair:(pair[0][0]or'',pair[0][1]))]
        if named:metadata['resolved_defined_names']=copy.deepcopy(named)
        return dict(schema_version=WORKBOOK_SCHEMA,semantics=SEMANTICS,source=self.binding,
                    snapshot_metadata=metadata,
                    selected_refs=encode_refs(selected),blank_refs=encode_refs(selected-set(self.cells)),sheets=rows,number_formats=formats,row_columns=list(tables.TABLE_COLUMNS),
                    dependency_edges=[[ref_index[a],ref_index[b]] for a,b in sorted(edges)],closure='full_snapshot_fallback' if fallback else 'exact_ast_closure')


def decode_workbook(shard):
    require(isinstance(shard,dict) and set(shard)=={'schema_version','semantics','source','snapshot_metadata','selected_refs','blank_refs','sheets','dependency_edges','closure','number_formats','row_columns'},'unknown shard fields')
    require(shard['schema_version']==WORKBOOK_SCHEMA and shard['semantics']==SEMANTICS,'unknown shard semantics')
    require(shard['row_columns']==tables.TABLE_COLUMNS and isinstance(shard['number_formats'],list) and all(isinstance(x,str) for x in shard['number_formats']) and len(shard['number_formats'])==len(set(shard['number_formats'])),'unknown cell table declaration');selected=decode_refs(shard['selected_refs']);blank=decode_refs(shard['blank_refs']);require(selected==sorted(set(selected)) and blank==sorted(set(blank)) and all(tables.canonical_cell(x) for x in selected),'invalid selected references')
    require(set(blank)<=set(selected),'blank references not selected');present={}
    for sheet in shard['sheets']:
        require(isinstance(sheet,dict) and 'present_rows' in sheet and 'cells' not in sheet,'invalid sheet shard')
        for row in sheet['present_rows']:
            require(isinstance(row,list) and len(row) in (4,5) and type(row[2]) is int and 0<=row[2]<len(shard['number_formats']),'cell presence shape changed')
            value=dict(address=row[0],value=row[1],number_format=shard['number_formats'][row[2]],data_type=row[3])
            if len(row)==5:value['cached_value']=row[4]
            key=sheet['sheet']+'!'+row[0];require(key in selected and key not in present,'duplicate/unselected cell');present[key]=value
    require(set(present).isdisjoint(blank) and set(present)|set(blank)==set(selected),'selected cell coverage differs')
    edges=shard['dependency_edges'];require(isinstance(edges,list) and len({tuple(x) for x in edges})==len(edges),'dependency edge duplicates/order')
    require(all(isinstance(e,list) and len(e)==2 and all(type(i) is int and 0<=i<len(selected) for i in e) for e in edges),'dependency outside selected closure')
    decoded_edges=[[selected[a],selected[b]] for a,b in edges];require(decoded_edges==sorted(decoded_edges),'dependency edge order differs')
    require(shard['closure'] in {'exact_ast_closure','full_snapshot_fallback'},'unknown closure mode')
    return dict(present=present,blank_refs=blank,selected_refs=selected,dependency_edges=decoded_edges,unknown_elsewhere=True)


def paired_shards(workbooks,seeds,unknown=False):
    """Use the same reference closure, reading each workbook's own cell bytes.

    A delivered target can be blank because its formula was intentionally
    removed. Its dependencies must still be described on that side. Only
    addresses cross the pair; values, formulas, blanks and dependency edges
    always come from the selected workbook itself.
    """
    require(set(workbooks)=={'golden','delivered'},'exact authenticated workbook pair required')
    selected=set(seeds)
    for _ in range(64):
        shards={name:book.shard(selected,unknown)for name,book in workbooks.items()}
        scopes={name:set(decode_refs(shard['selected_refs']))for name,shard in shards.items()}
        union=set().union(*scopes.values())
        require(selected<=union,'paired reference closure lost selected evidence')
        if all(refs==union for refs in scopes.values()):return shards
        selected=union
    raise ScopeDiagnosisRequired('paired dependency closure exceeded bounded fixed-point iterations')


def _manifest(body,root):
    require(set(body)=={'manifest','manifest_sha256','workbooks','instruction','disclosure'},'full disclosure body fields differ')
    manifest=body['manifest'];disclosure=body['disclosure'];require(set(manifest)=={'schema_version','input_hashes','agent_record_ids','sentences'} and manifest['schema_version']=='harbor-disclosure-sentence-manifest/v1','unknown full manifest schema');renderer=semantic.skill_module(root,'task-disclosure','disclose')
    selected=renderer.agent_records(disclosure['records']);require(selected==disclosure['agent_records'],'agent records differ from owning renderer')
    ids=[];bands={};examples={}
    for row in selected:
        existing=row.get('record_id') or row.get('id');rid=existing if isinstance(existing,str) and existing.strip() else 'record-'+semantic._hash(semantic._canonical(row))[:24]
        band=row.get('band') or ''
        ids.append(rid);bands.setdefault(band,[]).append(rid);examples.setdefault(band,row)
    require(len(ids)==len(set(ids)),'duplicate record identity')
    section=renderer.render_section(selected)
    prefixes=['- '+renderer.compact_cells(examples[band].get('cells',[]))+': ' for band in sorted(bands)]
    lines=semantic._disclosure_bullets(section,prefixes)
    require(len(lines)==len(bands) and (body['instruction'].count(section)==1 if section else not selected),'instruction grouping differs')
    sentences=[]
    for order,(band,line) in enumerate(zip(sorted(bands),lines)):
        require(body['instruction'].count(line)==1,'ambiguous sentence')
        row=dict(order=order,text=line,record_ids=bands[band]);sentences.append({'sentence_id':'sentence-'+semantic._hash(semantic._canonical(row))[:24],**row})
    require(manifest['agent_record_ids']==ids and manifest['sentences']==sentences,'full manifest coverage/identity differs')
    return dict(zip(ids,selected))


def _instruction(body,root,sentences):
    renderer=semantic.skill_module(root,'task-disclosure','disclose');section=renderer.render_section(body['disclosure']['agent_records']);allowed={x['text'] for x in sentences}
    projected=section
    for row in body['manifest']['sentences']:
        if row['text'] not in allowed:
            require(projected.count(row['text'])==1,'unassigned instruction span differs')
            projected=projected.replace(row['text'],'',1)
    require(all(projected.count(line)==1 for line in allowed),'assigned instruction line differs')
    return body['instruction'].replace(section,projected,1)


def plan(body,golden_path,delivered_path,*,root,max_sentences=MAX_SENTENCES,max_bytes=MAX_BYTES,bounds_policy=None,empty_review_contract=None):
    bounds=selected_bounds(bounds_policy,body,golden_path,delivered_path) if bounds_policy is not None else None
    wire=bool(bounds and bounds.get('wire_encoding'))
    if wire:from scripts.batch import harbor_disclosure_partition_wire as W
    require(type(max_sentences) is int and 1<=max_sentences<=(bounds['max_sentences']if bounds else MAX_SENTENCES) and type(max_bytes) is int and 1<=max_bytes<=(bounds['max_bytes']if bounds else MAX_BYTES),'unapproved shard bounds')
    root=Path(root);agents=_manifest(body,root);workbooks={};result=[]
    require(set(body['workbooks'])=={'golden','delivered'},'complete workbook pair required')
    try:
        for name,path in [('golden',golden_path),('delivered',delivered_path)]:workbooks[name]=_Workbook(path,body['workbooks'][name])
        full_manifest=body['manifest'];root_digest=digest(body)
        review_scope=None
        if not full_manifest['sentences']:
            require(empty_review_contract is not None and not wire,'empty disclosure requires an explicit full-source review contract')
            from scripts.batch import harbor_disclosure_review_round as rounds
            selected=rounds.read(empty_review_contract)
            helper=rounds.bound_module(selected['implementation'])
            review_scope=helper.review_scope(empty_review_contract,body,workbooks)
        else:require(empty_review_contract is None,'empty review contract cannot change nonempty partitions')
        hashes=full_manifest['input_hashes']
        require(set(hashes)=={'renderer','instruction','disclosure','golden','delivered'},'manifest input hash fields differ')
        require(hashes['golden']==workbooks['golden'].binding['sha256'] and hashes['delivered']==workbooks['delivered'].binding['sha256'] and hashes['renderer']==record(root/'.cursor/skills/task-disclosure/scripts/disclose.py')['sha256'] and hashes['instruction']==hashlib.sha256(body['instruction'].encode()).hexdigest(),'manifest inputs differ from exact current source')
        def batch(sentences):
            ids=[rid for s in sentences for rid in s['record_ids']];selected=[agents[rid] for rid in ids];bands={x.get('band') or '' for x in selected}
            raw=[x for x in body['disclosure']['records'] if (x.get('band') or '') in bands]
            seeds=set();unknown=False
            for row in selected+raw:
                refs,bad=_refs(row);seeds.update(refs);unknown|=bad
            if review_scope is not None:
                seeds=set(review_scope['graded_targets']);unknown=True
            shards=paired_shards(workbooks,seeds,unknown)
            scope=set(seeds)
            def relevant(row):
                refs,bad=_refs(row);return bad or bool(refs&scope)
            disclosure={**{k:copy.deepcopy(v) for k,v in body['disclosure'].items() if k not in {'records','agent_records','method_assessments'}},
                        'records':copy.deepcopy(raw),'agent_records':copy.deepcopy(selected),'method_assessments':assessments.filtered(body['disclosure']['method_assessments'],relevant)}
            if review_scope is not None:disclosure=copy.deepcopy(body['disclosure'])
            subset={**copy.deepcopy(full_manifest),'agent_record_ids':ids,'sentences':copy.deepcopy(sentences)}
            identity={'root_body_sha256':root_digest,'sentence_ids':[s['sentence_id'] for s in sentences]}
            if review_scope is not None:identity['review_contract']=empty_review_contract
            payload=dict(schema_version=SCHEMA,batch_id='disclosure-batch-'+digest(identity)[:24],root_body_sha256=root_digest,
                         full_manifest_sha256=body['manifest_sha256'],full_manifest_digest=digest(full_manifest),manifest=subset,manifest_sha256=digest(subset),
                         instruction=_instruction(body,root,sentences),instruction_projection={'schema_version':'assigned-disclosure-lines/v1','full_instruction_sha256':hashlib.sha256(body['instruction'].encode()).hexdigest(),'semantics':'All non-disclosure instruction text is unchanged. Only assigned exact disclosure sentences appear; the immutable full instruction remains root evidence.'},disclosure=disclosure,workbooks=shards)
            if review_scope is not None:
                payload.update(schema_version=EMPTY_SCHEMA,review_scope=review_scope,instruction=body['instruction'])
            for shard in shards.values():decode_workbook(shard)
            if wire:
                payload['schema_version']=WIRE_SCHEMA
                payload['workbooks']=W.encode(shards) if bounds['wire_encoding']==W.SCHEMA else W.encode(shards,encoding=bounds['wire_encoding'])
            if len(canonical(payload))>max_bytes:
                if not sentences:raise ScopeDiagnosisRequired('full empty-manifest review exceeds explicit byte bound')
                if len(sentences)==1:raise ScopeDiagnosisRequired('single sentence exceeds bound: '+json.dumps({'bytes':len(canonical(payload)),'components':{k:len(canonical(v)) for k,v in payload.items()},'disclosure_components':{k:len(canonical(v)) for k,v in disclosure.items()},'closures':{k:v['closure'] for k,v in shards.items()}},sort_keys=True))
                middle=len(sentences)//2;return batch(sentences[:middle])+batch(sentences[middle:])
            return [payload]
        if review_scope is not None:result.extend(batch([]))
        for offset in range(0,len(full_manifest['sentences']),max_sentences):
            result.extend(batch(full_manifest['sentences'][offset:offset+max_sentences]))
            require(bounds is None or len(result)<=bounds['max_batches'],'partition count exceeds explicit finite bounds')
        require(bounds is None or len(result)<=bounds['max_batches'],'partition count exceeds explicit finite bounds')
        sentence_ids=[s['sentence_id'] for b in result for s in b['manifest']['sentences']]
        record_ids=[r for b in result for r in b['manifest']['agent_record_ids']]
        require(sentence_ids==[s['sentence_id'] for s in full_manifest['sentences']] and Counter(record_ids)==Counter(full_manifest['agent_record_ids']) and len(record_ids)==len(set(record_ids)),'global coverage differs')
        ownership={}
        for b in result:
            pair=W.decode(b['workbooks'])if wire else b['workbooks']
            for name,shard in pair.items():
                for ref in decode_refs(shard['selected_refs']):ownership.setdefault((name,ref),b['batch_id'])
        cross=sorted({(b['batch_id'],name,a,z,ownership[(name,z)]) for b in result for name,shard in (W.decode(b['workbooks'])if wire else b['workbooks']).items() for a,z in decode_workbook(shard)['dependency_edges'] if ownership[(name,z)]!=b['batch_id']})
        return dict(schema_version=EMPTY_SCHEMA if review_scope is not None else WIRE_SCHEMA if wire else BOUNDED_SCHEMA if bounds else SCHEMA,**({'bounds_policy':bounds_policy}if bounds else {}),**({'empty_review_contract':empty_review_contract}if review_scope is not None else {}),root_body_sha256=root_digest,full_manifest=copy.deepcopy(full_manifest),full_manifest_sha256=body['manifest_sha256'],
                    source_bindings={name:book.binding for name,book in workbooks.items()},max_sentences=max_sentences,max_bytes=max_bytes,
                    batches=result,coverage={'sentence_ids':sentence_ids,'record_ids':record_ids,'cross_batch_dependencies':[list(x) for x in cross]})
    finally:
        for book in workbooks.values():book.close()


def validate_plan(value,body,golden_path,delivered_path,*,root):
    require(isinstance(value,dict),'partition plan required')
    empty=value.get('schema_version')==EMPTY_SCHEMA
    bounded=value.get('schema_version')in(BOUNDED_SCHEMA,WIRE_SCHEMA)or(empty and 'bounds_policy'in value)
    require(value.get('schema_version')in(SCHEMA,BOUNDED_SCHEMA,WIRE_SCHEMA,EMPTY_SCHEMA) and set(value)==({'schema_version','root_body_sha256','full_manifest','full_manifest_sha256','source_bindings','max_sentences','max_bytes','batches','coverage'}|({'bounds_policy'}if bounded else set())|({'empty_review_contract'}if empty else set())),'unknown partition plan fields')
    expected=plan(body,golden_path,delivered_path,root=root,max_sentences=value['max_sentences'],max_bytes=value['max_bytes'],bounds_policy=value.get('bounds_policy'),empty_review_contract=value.get('empty_review_contract'))
    require(canonical(value)==canonical(expected),'partition projection/coverage differs from exact current sources')
    return expected
