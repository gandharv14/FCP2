"""Fresh atomic normalization production, separate from tracker/quality approval.

Root supplies already admitted bindings and immutable original lane-budget
snapshots. All current audit, workbook, review and generation evidence is
independently rechecked. This module never approves a lane or changes trackers.
"""
from __future__ import annotations

import argparse
import copy
from dataclasses import asdict, dataclass
from contextlib import redirect_stdout
import json
from pathlib import Path
import re
import sys
import uuid

import openpyxl
import gen_normalizer as gen
import plain_eligibility
from mcp_env.build import validate_spec
from xl_variable_mcp import import_table, verify_against_golden
from xl_artifact_paths import resolve_workbook_artifact
from xl_release_publication import resolve_build_context
from scripts.qualified_publication import validate_audit_bindings
from scripts.batch.harbor_workbook_evidence import typed_cell_value
from scripts.batch.harbor_admission import read, record, require
from scripts.batch.harbor_tracker_lifecycle import checked_record
from scripts.batch.harbor_review_context import normalization_body, normalization_review
from scripts.batch.harbor_normalization_compact import json_value
from scripts.batch import harbor_normalization_compact as K
from scripts.batch.harbor_normalization_equivalence import validate_emission, expected_outputs
from scripts.batch.harbor_model_review_provenance import Policy, Call
from scripts.batch import harbor_model_executor as executor

SCHEMA = 'harbor-normalization-production/v1'
LANE = 'harbor-normalization-engineer'
MAX_CONTEXT_BYTES = 4 * 1024 * 1024
PROMPT = '''You are a fresh independent financial normalization reviewer. Review every
candidate reference independently using the exact supplied workbook evidence.
A mechanically maskable cell is not automatically an external research input.
Distinguish genuine historical or externally supplied primitive assumptions from
internal projections, management-selected scenarios, calculated answers and
uncertain provenance. Do not invent a company, currency, period, scenario,
verified real-world source, URL, or financial convention. Supplied-workbook
provenance is not evidence of independent external verification. Audit source
recommendations (including Excluded labels) are context, not binding eligibility.
Do not add references or change any supplied value, formula, target or tolerance.
If evidence does not establish eligibility and all four dimensions, reject it
with a specific reason. Preserve audit source notes and URLs; do not claim they
were retrieved. Return one JSON object with the exact top-level field
"schema_version": "normalization-review/v1" (not "schema"), and exact
draft_sha256, source_sha256, baseline_sha256, a nonempty reviewer, and a cells
mapping containing every enumerated candidate reference exactly once. Each
entry requires eligible:boolean, reason:nonempty string and evidence_refs:a
nonempty list of specific workbook references or audit draft IDs. Eligible
entries additionally require grounded entity, period, scenario and basis
strings. An optional variable_name must preserve the supplied financial meaning.
No overall PASS, omitted cells, inferred hidden values or extra references.'''


@dataclass(frozen=True)
class Budget:
    ordinal: int
    limit: int
    attempt_id: str
    evidence: tuple


@dataclass(frozen=True)
class Selection:
    workbook: str
    bindings: dict
    run: Path
    policy: Policy
    budget: Budget
    first_normalization: dict | None = None


def raw(record_):
    checked_record(record_)
    value = Path(record_['path']).read_bytes()
    require(executor.sha(value) == record_['sha256'], 'normalization evidence changed while reading')
    return value


def keep(path, value):
    """Stream canonical JSON into an immutable, atomically linked artifact."""
    import os
    import uuid
    path=executor.safe(path)
    temporary=path.with_name('.'+path.name+'.'+str(uuid.uuid4())+'.tmp')
    try:
        fd=os.open(temporary,os.O_WRONLY|os.O_CREAT|os.O_EXCL|os.O_NOFOLLOW,0o600)
        with os.fdopen(fd,'wb') as handle:
            if isinstance(value,bytes):handle.write(value)
            else:
                for chunk in json.JSONEncoder(indent=2,sort_keys=True,allow_nan=False).iterencode(value):
                    handle.write(chunk.encode())
                handle.write(b'\n')
            handle.flush();os.fsync(handle.fileno())
        if path.exists():
            require(not path.is_symlink() and record(path)['sha256']==record(temporary)['sha256'],
                    'immutable normalization artifact changed: '+path.name)
        else:
            os.link(temporary,path);executor.sync_dir(path.parent)
    finally:temporary.unlink(missing_ok=True)
    return record(path)


def store_evidence(path,value):
    projected=json_value(value)
    if K.encoded_size(projected)>MAX_CONTEXT_BYTES:
        body=projected if isinstance(projected,dict) else dict(candidates=projected,candidate_primitive_checks=[])
        projected=dict(schema_version='normalization-indexed-artifact/v1',
                       kind='body' if isinstance(value,dict) else 'analysis',evidence=K.compact(body))
    return keep(path,projected)


def expand_artifact(value):
    if isinstance(value,dict) and value.get('schema_version')=='normalization-indexed-artifact/v1':
        require(set(value)=={'schema_version','kind','evidence'} and value['kind'] in {'body','analysis'},'invalid indexed artifact')
        body=K.expand(value['evidence'])
        return body if value['kind']=='body' else body['candidates']
    return value


def code_bindings():
    import scripts.batch.harbor_review_context as context
    import scripts.qualified_publication as qualified
    import scripts.batch.harbor_workbook_evidence as workbook_evidence
    return {str(Path(module.__file__).resolve()): executor.sha(Path(module.__file__).read_bytes())
            for module in (sys.modules[__name__], gen, context, qualified, workbook_evidence, K)}


def selection_data(s):
    require(isinstance(s, Selection) and isinstance(s.policy, Policy) and isinstance(s.budget, Budget), 'root-selected typed normalization recipe required')
    require(re.fullmatch(r'[0-9]{4}',s.workbook) is not None, 'invalid workbook identity')
    run = executor.safe(s.run)
    require(run.name == s.workbook+'-variable-sources', 'normalization run must match workbook')
    require(type(s.budget.ordinal) is int and type(s.budget.limit) is int and 1 <= s.budget.ordinal <= s.budget.limit <= 2, 'original bounded review budget required')
    require(str(uuid.UUID(s.budget.attempt_id)) == s.budget.attempt_id, 'canonical review attempt identity required')
    require(isinstance(s.budget.evidence,tuple) and bool(s.budget.evidence), 'bound original review budget evidence required')
    for item in s.budget.evidence:
        data = raw(item)
        try: evidence = json.loads(data)
        except (ValueError,UnicodeDecodeError): continue
        if isinstance(evidence,dict) and evidence.get('schema_version') == 'normalization-format-recovery/v1':
            from scripts.batch.harbor_normalization_recovery import validate_receipt
            validate_receipt(item, selection=s)
    require(s.policy == Policy(s.policy.project_id) and bool(s.policy.project_id.strip()), 'normalization requires configured independent Sol High endpoint')
    b=s.bindings
    for key in ('source_file','original_raw_source','baseline_inputs','baseline_segmentation','curation',
                'variable_source_audit','variable_source_inventory','variable_source_metadata'):
        raw(b[key])
    require(Path(b['original_raw_source']['path']).stem == s.workbook, 'original workbook identity differs')
    require(Path(b['source_file']['path']).name == s.workbook+'.xlsx', 'current source identity differs')
    actual = resolve_workbook_artifact(Path(b['baseline_inputs']['path']).parent,s.workbook,'%s-inputs')
    require(actual.resolve() == Path(b['baseline_inputs']['path']).resolve(), 'normalizer baseline resolution differs')
    audit = validate_audit_bindings(b, expected_generations={k:b[k+'_generation']['id'] for k in ('source','segmentation')})
    context = resolve_build_context(s.workbook,
        source_root=Path(b['source_generation']['path']).parent.parent, source_generation_id=b['source_generation']['id'],
        segmentation_root=Path(b['segmentation_generation']['path']).parent.parent, segmentation_generation_id=b['segmentation_generation']['id'])
    require(Path(context['source_path']).resolve() == Path(b['source_file']['path']).resolve(), 'resolved source path differs')
    curation = Path(b['segmentation_generation']['path'])/'curation.toml'
    require(curation.read_bytes() == raw(b['curation']), 'retained curated targets/tolerances differ')
    first = None
    if s.first_normalization:
        first = json.loads(raw(s.first_normalization))
        require(first.get('workbook') == s.workbook and type(first.get('atomic_variables')) is int and first['atomic_variables'] >= 0,
                'invalid inherited first-normalization snapshot')
    plan = dict(schema_version=SCHEMA, workbook=s.workbook, bindings=b, run=str(run), policy=asdict(s.policy),
                budget=asdict(s.budget), first_normalization=s.first_normalization,
                audit_validation_files=audit, pipeline_bindings=context['bindings'], code=code_bindings(),
                semantic_budget_reset=False)
    # Convert tuples only in the saved JSON representation, never lose types in analysis.
    return json_value(plan), first


def workbook_context(source, baseline, rows):
    """Deterministic sparse local evidence, with no cached-formula values."""
    gold=openpyxl.load_workbook(source,data_only=False)
    base=openpyxl.load_workbook(baseline,data_only=False)
    try:
        wanted={};checks=[]
        from collections import defaultdict
        from bisect import bisect_left
        row_index={};header_index={}
        for sheet in gold.sheetnames:
            by_row=defaultdict(list);by_col=defaultdict(list)
            for observed in sorted(gold[sheet]._cells.values(),key=lambda c:(c.row,c.column)):
                if observed.value is None:continue
                by_row[observed.row].append(observed)
                if observed.data_type!='f' and isinstance(observed.value,str):by_col[observed.column].append(observed)
            row_index[sheet]=by_row
            header_index[sheet]={c:([v.row for v in values],values) for c,values in by_col.items()}
        for row in rows:
            if not row['clean']:continue
            require(len(row['cells'])==1, 'normalization must expose atomic candidate cells')
            sheet,address,value=row['cells'][0];cell=gold[sheet][address];bc=base[sheet][address]
            require(cell.data_type!='f' and bc.data_type!='f' and type(cell.value) is type(bc.value) and cell.value==bc.value==value,
                    'candidate is not a preserved primitive')
            selected=wanted.setdefault(sheet,set());selected.add(address)
            # Keep the full populated local row and nearby column headings. No
            # invented labels; uncertain meaning remains the reviewer's rejection.
            for observed in row_index[sheet].get(cell.row,[]):selected.add(observed.coordinate)
            for row_number in range(max(1,cell.row-4),cell.row):
                selected.update(c.coordinate for c in row_index[sheet].get(row_number,[]) if c.column<=cell.column+2)
            positions,headers=header_index[sheet].get(cell.column,([],[]))
            stop=bisect_left(positions,cell.row)
            selected.update(c.coordinate for c in headers[max(0,stop-4):stop])
            checks.append(dict(reference=row['reference'],source_data_type=cell.data_type,baseline_data_type=bc.data_type,
                source_equals_baseline=True,number_format=cell.number_format,duplicate_refs=row['duplicate_refs']))
        context={}
        for sheet,refs in wanted.items():
            cells=sorted((gold[sheet][r] for r in refs),key=lambda c:(c.row,c.column))
            context[sheet]=[dict(reference=sheet+'!'+c.coordinate,value=typed_cell_value(c),data_type=c.data_type,number_format=c.number_format) for c in cells]
        return json_value(context),json_value(checks)
    finally:gold.close();base.close()


def analyse(s):
    b=s.bindings
    return gen.analyse(s.workbook,b['source_file']['path'],b['source_file']['sha256'],
                       runs_root=s.run.parent,inputs_root=Path(b['baseline_inputs']['path']).parent)


def prepare(s):
    plan,first=selection_data(s);run=executor.safe(s.run)
    if run.exists():
        require((run/'production-selection.json').exists(), 'fresh run required; existing unbound normalization artifacts')
        require(read(record(run/'production-selection.json')) == plan, 'normalization selection/budget changed')
    else:run.mkdir(parents=True,mode=0o700)
    keep(run/'production-selection.json',plan)
    for key,suffix in [('variable_source_metadata','.metadata.json'),('variable_source_inventory','.inventory.json')]:
        keep(run/(s.workbook+'-inputs-variable-sources'+suffix),raw(s.bindings[key]))
    if first is not None:keep(run/'first_normalization.json',raw(s.first_normalization))
    # Do not load or execute an existing generated script. Import exact audit bytes.
    draft=import_table(Path(s.bindings['variable_source_audit']['path']))
    require(len(draft['rows']) == draft['row_count'] and all(r.get('normalization_status')=='needs_review' for r in draft['rows']), 'lossless audit import incomplete')
    keep(run/'draft.json',json.dumps(draft,indent=2,ensure_ascii=False).encode()+b'\n')
    actual,rows=analyse(s);require(actual==draft,'normalization draft differs from lossless import')
    context,checks=workbook_context(s.bindings['source_file']['path'],s.bindings['baseline_inputs']['path'],rows)
    b=s.bindings;draft_record=record(run/'draft.json')
    body=dict(source=b['source_file'],baseline=b['baseline_inputs'],draft=draft_record,
        source_sha256=b['source_file']['sha256'],baseline_sha256=b['baseline_inputs']['sha256'],draft_sha256=draft_record['sha256'],
        candidates=json_value([r for r in rows if r['clean']]),context=context,candidate_primitive_checks=checks,
        audit=b['variable_source_audit'],current_source_generation=b['source_generation'],current_segmentation_generation=b['segmentation_generation'])
    if b.get('audit_equivalence') is not None:body['audit_equivalence']=b['audit_equivalence']
    normalization_body(body,source=b['source_file'],baseline=b['baseline_inputs'],draft=draft_record,rows=rows,bindings=b)
    oversized=K.encoded_size(body)>MAX_CONTEXT_BYTES
    evidence=store_evidence(run/'review-evidence.json',body);analysis=store_evidence(run/'analysis.json',rows)
    if oversized or len(raw(evidence)) > MAX_CONTEXT_BYTES:
        raise K.RequestTooLarge('complete normalization context exceeds bounded request size; no candidates may be dropped')
    clean=[r for r in rows if r['clean']]
    request=None
    if clean:
        payload=dict(model=s.policy.model,reasoning_effort=s.policy.reasoning_effort,messages=[
            dict(role='system',content=PROMPT),dict(role='user',content=raw(evidence).decode())])
        request=keep(run/'review-request.json',payload)
    manifest=dict(schema_version='prepared-normalization-review/v1',status='review_required' if clean else 'no_maskable_candidates',
        request=request,evidence=evidence,analysis=analysis,draft=draft_record,selection=record(run/'production-selection.json'),
        candidate_count=len(clean),candidate_references=sorted({r['reference'] for r in clean}),
        policy=asdict(s.policy),budget=asdict(s.budget),semantic_budget_reset=False)
    prepared=keep(run/'prepared.json',json_value(manifest))
    return prepared


def prepared_call(s,prepared):
    expected=prepare(s);require(prepared==expected,'prepared normalization evidence changed')
    manifest=read(prepared);require(manifest['request'] is not None,'no maskable candidates; no model call is authorized')
    return executor.PreparedCall(manifest['request'],s.policy,
        executor.SemanticScope(s.workbook,s.bindings['original_raw_source']['sha256'],LANE,'normalization-review'),
        s.budget.ordinal,s.budget.limit,s.budget.attempt_id,s.budget.evidence,response_mode='json_object')


def dispatch(s,prepared,campaign,*,credential):
    """The only network-capable entry; owning root authorizes this exact request."""
    return executor.execute(prepared_call(s,prepared),campaign,credential=credential)


def emit_phase1(s,prepared,call=None):
    require(prepare(s)==prepared,'prepared normalization evidence changed')
    manifest=read(prepared);draft,rows=analyse(s);run=s.run;b=s.bindings
    for name in ('source_profiles.json','source-profile-mapping.json','source-display-review.json','oracle_forced_exclusions.json'):
        require(not (run/name).exists(),'fresh phase1 cannot consume unreviewed later-stage inputs: '+name)
    if manifest['request'] is not None:
        require(isinstance(call,Call) and call.parsed is not None,'fresh independent parsed model call required')
        result=normalization_review(call,prepared,source=b['source_file'],baseline=b['baseline_inputs'],draft=record(run/'draft.json'),rows=rows,bindings=b,policy=s.policy)
        submitted=read(call.submitted)
        require(submitted.get('attempt_id')==s.budget.attempt_id and submitted.get('semantic_ordinal')==s.budget.ordinal
                and submitted.get('semantic_limit')==s.budget.limit and submitted.get('scope')==asdict(prepared_call(s,prepared).scope),
                'normalization review does not use the original root budget identity')
        keep(run/'normalization-review.json',raw(call.parsed))
        require(read(record(run/'normalization-review.json'))==result['review'],'review is not the exact parsed model response')
        reviewed=copy.deepcopy(rows);gen.apply_semantic_review(reviewed,run)
    else:
        require(call is None,'empty candidate set cannot consume or fabricate a semantic call')
        require(not (run/'normalization-review.json').exists(),'unexpected review for zero maskable candidates')
    plan=read(record(run/'production-selection.json'))
    projection=gen.emit(s.workbook,copy.deepcopy(draft),copy.deepcopy(rows),runs_root=run.parent,pipeline_bindings=plan['pipeline_bindings'],dry_run=True)
    first=run/'first_normalization.json'
    expected_first=raw(s.first_normalization) if s.first_normalization else executor.canonical(projection['snapshot'])
    keep(first,expected_first)
    script=projection['script'];keep(run/f'normalize_{s.workbook}.py',script.encode())
    names=('normalized.json','exclusions.json','normalization_report.json')
    # Compile only freshly generated current code; retained scripts are evidence.
    _,expected=expected_outputs(s.workbook,draft,rows,run,plan['pipeline_bindings'])
    for path,value in expected.items():
        if path.exists():
            require(path.read_bytes()==(json.dumps(value,indent=2,ensure_ascii=False,sort_keys=True)+'\n').encode(),
                    'retained interrupted normalization differs; refusing to overwrite evidence')
    prior=None
    for index in range(2):
        namespace={'__name__':'harbor_current_normalizer','__file__':str(run/f'normalize_{s.workbook}.py')}
        exec(compile(script,'<current-normalizer>','exec'),namespace);namespace['main']()
        current={name:record(run/name) for name in names}
        if prior is not None:require(current==prior,'normalizer is not byte deterministic')
        prior=current
    validate_emission(s.workbook,draft,rows,run,plan['pipeline_bindings'])
    spec=read(record(run/'normalized.json'))
    if spec['variables']:validate_spec(spec)
    require(not verify_against_golden(spec,Path(b['source_file']['path'])),'normalized primitive values differ from source')
    current_plan,_=selection_data(s)
    require(current_plan==plan,'normalization inputs/code changed during emission')
    phase1=keep(run/'normalized.phase1.json',(run/'normalized.json').read_bytes())
    mode=plain_eligibility.evaluate(run);mode_record=keep(run/'plain-eligibility.json',mode)
    status='phase1_evidence_ready' if mode['mode'] in {'plain','mcp'} and (mode['mode']!='plain' or mode['eligible'] is True) else 'normalization_rejected'
    result=dict(schema_version=SCHEMA,status=status,workbook=s.workbook,prepared=prepared,
        phase1=phase1,outputs=prior,first_normalization=record(first),plain_eligibility=mode_record,
        review_call=asdict(call) if call else None,lane_approval=False,semantic_budget_reset=False,
        next_gate='source_profiler' if status=='phase1_evidence_ready' else 'orchestrator_failure_review')
    return keep(run/'phase1-result.json',result)


def load_selection(binding):
    data=read(binding)
    require(set(data)=={'workbook','bindings','run','policy','budget','first_normalization'},'exact root selection fields required')
    budget=data['budget'];require(set(budget)=={'ordinal','limit','attempt_id','evidence'},'invalid original review budget')
    return Selection(data['workbook'],data['bindings'],Path(data['run']),Policy(**data['policy']),
                     Budget(budget['ordinal'],budget['limit'],budget['attempt_id'],tuple(budget['evidence'])),data['first_normalization'])


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('operation',choices=('prepare','emit-phase1'))
    parser.add_argument('--selection',required=True);parser.add_argument('--selection-sha256',required=True)
    parser.add_argument('--prepared');parser.add_argument('--prepared-sha256')
    parser.add_argument('--call');parser.add_argument('--call-sha256')
    args=parser.parse_args(argv)
    selection=load_selection(dict(path=str(Path(args.selection).absolute()),sha256=args.selection_sha256))
    with redirect_stdout(sys.stderr):
        if args.operation=='prepare':result=prepare(selection)
        else:
            require(args.prepared and args.prepared_sha256,'exact prepared binding required')
            require(bool(args.call)==bool(args.call_sha256),'call path/hash must be supplied together')
            call=Call(**read(dict(path=str(Path(args.call).absolute()),sha256=args.call_sha256))) if args.call else None
            result=emit_phase1(selection,dict(path=str(Path(args.prepared).absolute()),sha256=args.prepared_sha256),call)
    print(json.dumps(result,sort_keys=True));return 0


if __name__=='__main__':raise SystemExit(main())
