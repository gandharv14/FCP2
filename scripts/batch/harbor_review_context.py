"""Concrete current-data adapters for retained independent model reviews."""
import copy
import json
from pathlib import Path
from openpyxl import load_workbook
from scripts.batch.harbor_workbook_evidence import typed_cell_value
from scripts.batch.harbor_admission import read,record,require
from scripts.batch.harbor_tracker_lifecycle import checked_record
from scripts.batch.harbor_model_review_provenance import validate_call,require_independent


def json_value(value):return json.loads(json.dumps(value,default=str,allow_nan=False))


class NormalizationReviewRejected(ValueError):
    """Completed model content failed the owning normalization contract."""
    failure_class = 'invalid_model_content'
    semantic_attempt_consumed = True
    automatic_retry = False
    source_invalidity_claimed = False

    def __init__(self, reason_code, message):
        self.reason_code = reason_code
        super().__init__(message)


def validate_normalization_review(review, rows):
    """Pure preflight of unchanged gen.apply_semantic_review's contract.

    Exact candidate coverage is also mandatory in this transport-context gate.
    This validates decisions, never applies them or changes any reviewer text.
    """
    def check(ok, code, message):
        if not ok: raise NormalizationReviewRejected(code, message)
    candidates=[row for row in rows if row['clean']]
    check(isinstance(review,dict) and review.get('schema_version')=='normalization-review/v1'
        and bool(str(review.get('reviewer') or '').strip()) and isinstance(review.get('cells'),dict),
        'invalid_schema','invalid normalization semantic review: schema_version/reviewer/cells required')
    check(set(review['cells'])=={row['reference'] for row in candidates},
        'incomplete_coverage','model review is not exactly complete')
    for row in candidates:
        check(all(review.get(key)==value for key,value in row.get('evidence',{}).items()),
            'stale_evidence','stale normalization semantic review')
        decision=review['cells'][row['reference']]
        check(isinstance(decision,dict) and type(decision.get('eligible')) is bool
            and bool(str(decision.get('reason') or '').strip())
            and isinstance(decision.get('evidence_refs'),list) and bool(decision['evidence_refs'])
            and all(isinstance(value,str) and value.strip() for value in decision['evidence_refs']),
            'invalid_decision','missing or incomplete cell review: '+row['reference'])
        if decision['eligible']:
            for key in ('entity','period','scenario','basis'):
                check(isinstance(decision.get(key),str) and bool(decision[key].strip()),
                    'missing_dimension','missing reviewed dimension: '+key)
            check('variable_name' not in decision or (isinstance(decision['variable_name'],str)
                  and bool(decision['variable_name'].strip())),
                  'invalid_variable_name','invalid reviewed variable_name: '+row['reference'])
    return True


def normalization_body(body,*,source,baseline,draft,rows,bindings):
    """Reconstruct every claimed primitive/context observation from workbook bytes."""
    for key,expected in [('source',source),('baseline',baseline),('draft',draft)]:
        checked_record(body[key]);checked_record(expected)
        require(body[key]['sha256']==expected['sha256'] and body.get(key+'_sha256')==expected['sha256'],'normalization context input changed')
    for key in ('current_source_generation','current_segmentation_generation'):
        require(body[key]==bindings[key.removeprefix('current_')],'review generation pair changed')
    expected=bindings['variable_source_audit']
    require(body['audit']==expected,'review audit applicability changed');checked_record(expected)
    equivalence=bindings.get('audit_equivalence')
    require(body.get('audit_equivalence')==equivalence,'review audit applicability changed')
    if equivalence is not None:
        checked_record(equivalence)
    from scripts.batch.harbor_normalization_compact import json_value, equal
    clean=[row for row in rows if row['clean']]
    require(equal(body['candidates'],json_value(clean)),'review candidate coverage/data changed')
    gold=load_workbook(source['path'],data_only=False);base=load_workbook(baseline['path'],data_only=False)
    try:
        context={}
        for sheet,entries in body['context'].items():
            actual=[];seen=set()
            for item in entries:
                ref=item['reference'];require(ref.startswith(sheet+'!') and ref not in seen,'ambiguous context cell');seen.add(ref)
                cell=gold[sheet][ref[len(sheet)+1:]]
                actual.append({'reference':ref,'value':typed_cell_value(cell),'data_type':cell.data_type,'number_format':cell.number_format})
            context[sheet]=actual
        require(json_value(context)==body['context'],'review context workbook observations changed')
        checks=[]
        for row in clean:
            require(len(row['cells'])==1,'primitive review needs explicit atomic candidate context')
            sheet,address,value=row['cells'][0];a=gold[sheet][address];b=base[sheet][address]
            require(a.data_type!='f' and b.data_type!='f' and type(a.value) is type(b.value) and a.value==b.value==value,'review candidate is not a preserved literal assumption')
            checks.append({'reference':row['reference'],'source_data_type':a.data_type,'baseline_data_type':b.data_type,'source_equals_baseline':True,'number_format':a.number_format,'duplicate_refs':row['duplicate_refs']})
        require(equal(json_value(checks),body['candidate_primitive_checks']),'candidate primitive evidence changed')
    finally:gold.close();base.close()
    return True


def normalization_review(call,prepared,*,source,baseline,draft,rows,bindings,policy,other_calls=()):
    manifest=read(prepared)
    for key in ('request','evidence','analysis'):checked_record(manifest[key])
    require(manifest['request']['sha256']==call.request['sha256'],'call did not use prepared request')
    from scripts.batch.harbor_normalization_producer import expand_artifact
    from scripts.batch.harbor_normalization_compact import json_value as shared_json_value, equal
    require(equal(expand_artifact(read(manifest['analysis'])),shared_json_value(rows)),'prepared analysis differs from current reconstruction')
    body=expand_artifact(read(manifest['evidence']))
    normalization_body(body,source=source,baseline=baseline,draft=draft,rows=rows,bindings=bindings)
    request=read(manifest['request'])
    require(request['messages'][1]['content']==Path(manifest['evidence']['path']).read_text(),'review request omitted or changed evidence')
    result=validate_call(call,request,policy=policy);require_independent(result,other_calls)
    validate_normalization_review(result['review'],rows)
    for item in (manifest.get('budget') or {}).get('evidence',[]):
        checked_record(item)
        try:budget_evidence=read(item)
        except (ValueError,UnicodeDecodeError):continue
        if isinstance(budget_evidence,dict) and budget_evidence.get('schema_version')=='normalization-format-recovery/v1':
            from scripts.batch.harbor_normalization_recovery import validate_prepared_budget
            validate_prepared_budget(item,manifest,source=source,baseline=baseline,draft=draft,bindings=bindings,policy=policy)
    return result


def naturalization_review(call,request_binding,*,source,candidate,policy,writer_call,writer_request,writer_bodies,prior_reviews=(),root):
    """Use the fixed supported semantic validator, never a caller predicate."""
    from scripts.batch.harbor_readiness import skill_module
    semantic_validator=skill_module(root,'naturalize-finance-task-instruction','span_restoration')
    checked_record(request_binding);request=read(request_binding)
    require(request_binding['sha256']==call.request['sha256'],'review call request changed')
    require(json.loads(request['messages'][1]['content'])=={'source':source,'candidate':candidate},'review source/candidate differs from current exact context')
    result=validate_call(call,request,policy=policy)
    require(writer_call is not None,'writer provenance is required for independent review')
    checked_record(writer_request)
    require(writer_request['sha256']==writer_call.request['sha256'],'writer request changed')
    writer=validate_call(writer_call,read(writer_request),policy=policy)
    require(set(writer_bodies)=={'preamble_body','input_body'},'exact writer body evidence required')
    require(set(writer['review'])==set(writer_bodies),'writer response shape changed')
    for key,binding in writer_bodies.items():
        checked_record(binding);require(Path(binding['path']).read_text()==writer['review'][key],'submitted writer body is not the model output')
    require_independent(result,[writer,*prior_reviews])
    semantic_validator.semantic_evidence(Path(call.request['path']).read_bytes(),Path(call.response['path']).read_bytes(),source,candidate)
    return result
