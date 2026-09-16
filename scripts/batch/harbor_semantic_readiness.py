"""Current disclosure and dialogue review context, independent of handoff flags."""
from collections import Counter
from dataclasses import dataclass
from pathlib import Path
import copy,hashlib,json
from openpyxl import load_workbook
from scripts.batch.harbor_workbook_evidence import typed_cell_value
from scripts.batch.harbor_admission import record,read,require
from scripts.batch.harbor_tracker_lifecycle import checked_record
from scripts.batch.harbor_readiness import skill_module
from scripts.batch.harbor_model_review_provenance import validate_call,require_independent

ROOT=Path(__file__).resolve().parents[2]


def _canonical(v):return json.dumps(v,sort_keys=True,ensure_ascii=False,separators=(',',':')).encode()
def _hash(v):return hashlib.sha256(v).hexdigest()
def _json(v):return json.loads(json.dumps(v,default=str,allow_nan=False))


def workbook_evidence(binding):
    """Same complete populated-cell snapshot sent by the supported review lane."""
    checked_record(binding);path=Path(binding['path']);book=load_workbook(path,data_only=False);cached=load_workbook(path,data_only=True)
    try:
        sheets=[]
        for ws in book:
            cells=[]
            # iter_rows() creates every omitted blank cell in the bounding
            # rectangle. A single distant formatted cell can allocate billions
            # of empty Cell objects although the evidence lists populated cells.
            from openpyxl.utils import get_column_letter
            dimensions=f'A1:{get_column_letter(ws.max_column)}{ws.max_row}'
            for _,cell in sorted(ws._cells.items()):
                if cell.value is None:continue
                value=typed_cell_value(cell)
                item={'address':cell.coordinate,'value':value,'number_format':cell.number_format,'data_type':cell.data_type}
                if cell.data_type=='f':item['cached_value']=cached[ws.title][cell.coordinate].value
                cells.append(item)
            sheets.append({'sheet':ws.title,'state':ws.sheet_state,'dimensions':dimensions,'merged_ranges':[str(x) for x in ws.merged_cells.ranges],'cells':cells})
        return _json({'path':str(path),'sha256':binding['sha256'],'blank_cell_semantics':'Every populated cell is listed; omitted addresses have no value.','sheets':sheets})
    finally:book.close();cached.close()


def _disclosure_bullets(section,prefixes):
    """Extract entire renderer-owned bullets, retaining embedded label newlines."""
    if not prefixes:
        require(not section,'unexpected unbound disclosure section')
        return []
    require(section.endswith('\n'),'disclosure renderer terminator changed')
    starts=[]
    for prefix in prefixes:
        marker='\n'+prefix
        require(section.count(marker)==1,'ambiguous or missing disclosure bullet anchor')
        starts.append(section.index(marker)+1)
    require(starts==sorted(set(starts)),'disclosure bullet order or identity changed')
    ends=[position-1 for position in starts[1:]]+[len(section)-1]
    return [section[start:end] for start,end in zip(starts,ends)]


def disclosure_manifest(instruction,disclosure,golden,delivered,*,root=ROOT):
    renderer=skill_module(root,'task-disclosure','disclose')
    payload=read(disclosure);text=Path(instruction['path']).read_text();checked_record(instruction)
    selected=renderer.agent_records(payload.get('records') or [])
    require(payload.get('agent_records')==selected,'agent records differ from current renderer')
    ids=[];by_band={};bands={}
    for item in selected:
        existing=item.get('record_id') or item.get('id')
        rid=existing if isinstance(existing,str) and existing.strip() else 'record-'+_hash(_canonical(item))[:24]
        band=item.get('band') or ''
        ids.append(rid);by_band.setdefault(band,[]).append(rid);bands.setdefault(band,item)
    require(len(ids)==len(set(ids)),'duplicate disclosure record identity')
    section=renderer.render_section(selected)
    prefixes=['- '+renderer.compact_cells(bands[band].get('cells',[]))+': ' for band in sorted(by_band)]
    lines=_disclosure_bullets(section,prefixes)
    require(len(lines)==len(by_band),'disclosure sentence grouping changed')
    require((text.count(section)==1) if section else not selected and renderer.SECTION_START not in text,'reviewed instruction does not contain exact disclosure')
    sentences=[]
    for order,(band,line) in enumerate(zip(sorted(by_band),lines)):
        require(text.count(line)==1,'ambiguous sentence occurrence')
        item={'order':order,'text':line,'record_ids':by_band[band]}
        sentences.append({'sentence_id':'sentence-'+_hash(_canonical(item))[:24],**item})
    return {'schema_version':'harbor-disclosure-sentence-manifest/v1','input_hashes':{'renderer':record(root/'.cursor/skills/task-disclosure/scripts/disclose.py')['sha256'],'instruction':instruction['sha256'],'disclosure':disclosure['sha256'],'golden':golden['sha256'],'delivered':delivered['sha256']},'agent_record_ids':ids,'sentences':sentences}


def validate_disclosure_review(review,manifest,manifest_sha):
    require(set(review)=={'schema_version','model','input_hashes','manifest_sha256','sentences','blocking_findings','non_blocking_flags','passed'},'review schema keys changed')
    require(review['schema_version']=='harbor-disclosure-faithfulness/v1' and review['model']=='gpt-5.6-sol-high','review schema/model mismatch')
    require(review['input_hashes']==manifest['input_hashes'] and review['manifest_sha256']==manifest_sha,'review input or manifest changed')
    actual=review['sentences'];expected=manifest['sentences'];require(isinstance(actual,list) and len(actual)==len(expected),'review sentence coverage incomplete')
    ids=[];records=[]
    for got,want in zip(actual,expected):
        require(set(got)=={'sentence_id','text','record_ids','golden_evidence','delivered_evidence','verdict','findings'},'sentence review shape changed')
        require(all(got[k]==want[k] for k in ('sentence_id','text','record_ids')),'review sentence identity/order/text changed')
        ids.append(got['sentence_id']);records+=got['record_ids']
        require(got['verdict'] in ('pass','non_blocking'),'blocking or unknown sentence verdict')
        for key in ('golden_evidence','delivered_evidence','findings'):require(isinstance(got[key],list),'review evidence/findings must be arrays')
        require(got['golden_evidence'] and got['delivered_evidence'],'independent workbook observations missing')
    require(len(ids)==len(set(ids)) and len(records)==len(set(records)) and Counter(records)==Counter(manifest['agent_record_ids']),'review record coverage incomplete/duplicated')
    require(review['blocking_findings']==[] and isinstance(review['non_blocking_flags'],list) and review['passed'] is True,'review has blocking findings')
    return {'sentences':len(actual),'records':len(records)}


def disclosure_review(call,manifest_binding,*,instruction,disclosure,golden,delivered,policy,other_calls=(),root=ROOT,revision=None):
    """Instruction is the owning validated NAT source snapshot, before rewriting."""
    for b in (manifest_binding,instruction,disclosure,golden,delivered):checked_record(b)
    manifest=disclosure_manifest(instruction,disclosure,golden,delivered,root=root)
    composition=None
    if revision is not None and read(revision).get('schema_version')=='harbor-disclosure-manifest-revision/v1':
        from scripts.batch import harbor_disclosure_manifest_revision as composition_module
        composition=read(revision);composition_module.owned(revision,composition)
        require(manifest==read(composition['full_manifest']),'full current disclosure differs from manifest correction')
        manifest=read(composition['manifest'])
    require(read(manifest_binding)==manifest,'manifest is not independently regenerated current context')
    request=read(call.request);body=json.loads(request['messages'][1]['content'])
    if read(call.submitted).get('physical_repair') is not None:
        from scripts.batch.harbor_disclosure_review_resume import original_review_request
        original=original_review_request(call)
        body=json.loads(original['messages'][1]['content'])
    expected_disclosure=read(disclosure)
    if 'evidence_projection' in body:
        from scripts.batch.harbor_disclosure_evidence import decode_body,scoped_disclosure
        body=decode_body(body)
        expected_disclosure=scoped_disclosure(expected_disclosure)
    require(set(body)=={'manifest','manifest_sha256','workbooks','instruction','disclosure'},'review request fields changed')
    require(body['manifest']==manifest and body['manifest_sha256']==manifest_binding['sha256'],'request manifest differs')
    require(body['instruction']==Path(instruction['path']).read_text() and body['disclosure']==expected_disclosure,'request instruction/disclosure differs')
    require(set(body['workbooks'])=={'golden','delivered'},'complete workbook evidence required')
    for role,binding in [('golden',golden),('delivered',delivered)]:
        fresh=workbook_evidence(binding);old=copy.deepcopy(body['workbooks'][role])
        # Immutable source generation paths can change without changing bytes.
        require(isinstance(old.get('path'),str) and old['path'],'review workbook locator missing')
        from scripts.batch.harbor_model_executor import canonical
        old['path']=fresh['path'];require(canonical(old)==canonical(fresh),'review workbook snapshot differs from current bytes')
    from scripts.batch.harbor_disclosure_revision import validate_review_history
    history=validate_review_history(revision,call,manifest_binding,instruction,disclosure)
    result=validate_call(call,request,policy=policy);require_independent(result,other_calls)
    if composition is not None:
        _,old_result,_=composition_module.original(composition)
        require_independent(result,[old_result,*other_calls])
        counts=composition_module.compose(read(composition['old_manifest']),read(composition['full_manifest']),
            old_result['review'],composition['old_manifest']['sha256'],manifest,result['review'],manifest_binding['sha256'])
        return {**history,**counts,'provenance':result,'retained_provenance':old_result}
    return {**history,**validate_disclosure_review(result['review'],manifest,manifest_binding['sha256']),'provenance':result}


def canonical_claims(value):
    """The sole historical schema extension: absent and explicit null mean no grouping."""
    result=copy.deepcopy(value)
    for obj in (result,result.get('writer_pack')):
        if isinstance(obj,dict) and obj.get('shared_source_links') is None:obj.pop('shared_source_links',None)
    return result


def dialogue_review(call,*,claims,notes,reviewed_claims,policy,writer_calls,workbook,current_source_sha256,historical=None,root=ROOT,original_sha256=None,source=None,answer_key=None,task_dir=None,pre_review_repair=None):
    """Exact source claims/draft and owning semantic coverage, never a passed flag."""
    require(1<=len(writer_calls)<=2,'one or two independent writer-call identities required')
    require(len({(w.request['sha256'],w.response['sha256']) for w in writer_calls})==len(writer_calls),'duplicate writer-call evidence')
    current=read(claims);prior=read(reviewed_claims);checked_record(notes)
    require(canonical_claims(current)==canonical_claims(prior),'dialogue claims are not exactly current')
    request=read(call.request);body=json.loads(request['messages'][1]['content'])
    extended=set(body)=={'claims.json','draft.md','round','operand_evidence'}
    repaired=set(body)=={'claims.json','draft.md','round','source_proved_correction'}
    require(repaired==(pre_review_repair is not None),'dialogue repair receipt/context differs')
    require(extended or repaired or set(body)=={'claims.json','draft.md','round'},'dialogue request context shape changed')
    require(not repaired or (historical is None and body['round']==1),'pre-review correction requires first independent review')
    require(not extended or (historical is None and body['round']==2),'source operand extension requires a new second-round review')
    validator=skill_module(root,'additional-assumptions-dialogue','validate_dialogue')
    require(body['claims.json']==prior and validator.strip_claim_comments(body['draft.md'])==Path(notes['path']).read_text(),'review claims/draft differs from exact current content')
    require(type(body['round']) is int and body['round'] in (1,2),'dialogue review round invalid')
    if historical is not None:
        require(historical.current_source_sha256==current_source_sha256,'historical/current source binding changed')
        history=_historical_evidence(historical,workbook=workbook,notes=notes)
        require(reviewed_claims['sha256']==history['reviewed_claims_sha256'],'historical claims are not the published reuse context')
        require(_hash(body['draft.md'].encode())==history['reviewed_draft_sha256'],'historical reviewed draft differs from retained original')
        result=_historical_call(call,request,policy=policy,evidence=history)
        result['evidence']=history['evidence']
        writers=[_historical_call(w,read(w.request),policy=policy,evidence=history,json_response=False) for w in writer_calls]
        require(all(result['response_id']!=w['response_id'] for w in writers),'historical reviewer reused writer response')
    else:
        from scripts.batch.harbor_instruction_producer import skill_prompt
        prompt=skill_prompt(root,'additional-assumptions-dialogue','You are reviewing notes from a deal-team conversation')
        if repaired:
            from scripts.batch import harbor_dialogue_pre_review_repair as Q
            require(body['source_proved_correction']==Q.review_projection(pre_review_repair,prior,body['draft.md'],root=root),'source-proved correction context differs')
            prompt+=Q.REVIEW_GUIDANCE
        if extended:
            from scripts.batch import harbor_dialogue_operand_evidence as operands
            require(source is not None and answer_key is not None and task_dir is not None,'root-selected operand proof inputs required')
            require(source['sha256']==current_source_sha256,'operand source differs from current source')
            require(body['operand_evidence']==operands.project(prior,body['draft.md'],source=source,answer_key=answer_key,task_dir=task_dir,root=root),'source operand projection changed')
            prompt+=operands.GUIDANCE
        require(request['messages'][0]['content']==prompt,'dialogue reviewer prompt changed')
        submitted=read(call.submitted)
        require(submitted.get('scope')=={'workbook_id':workbook,'original_sha256':original_sha256,'lane':'harbor-dialogue-author','budget_id':'dialogue-review'} and submitted.get('semantic_ordinal')==body['round'] and submitted.get('semantic_limit')==2,'dialogue reviewer original/round scope changed')
        result=validate_call(call,request,policy=policy)
        from scripts.batch.harbor_dialogue_authoring import verify_pair
        if repaired:
            writers=Q.verify_pair(pre_review_repair,writer_calls,policy=policy,claims=current,draft=body['draft.md'],round=body['round'],root=root,workbook=workbook,original_sha256=original_sha256);Q.validate_review_call(pre_review_repair,call)
        else:
            writers=verify_pair(writer_calls,policy=policy,claims=current,draft=body['draft.md'],round=body['round'],root=root,workbook=workbook,original_sha256=original_sha256)
        require_independent(result,writers)
    require(len({w['response_id'] for w in writers})==len(writers),'duplicate writer response identity')
    review=result['review'];validator=skill_module(root,'additional-assumptions-dialogue','validate_dialogue')
    require(set(review)=={'agent_model','round','accuracy','naturalness','passed'},'dialogue review schema changed')
    require(review['agent_model']=='gpt-5.6-sol-high' and review['round']==body['round'],'dialogue review model/round differs')
    source_warning_proof=None
    if not validator.review_passed(review,current,draft_text=body['draft.md']) and historical is None and not repaired:
        from scripts.batch import harbor_dialogue_source_warnings as source_warnings
        require(source is not None and task_dir is not None and source['sha256']==current_source_sha256,
                'source-proved warning requires current workbook authority')
        directory=Path(reviewed_claims['path']).parent
        round_dir=directory/('round-%02d'%body['round'])
        require(body['round']==1 and len(writer_calls)==2,'source warning requires exact first-round author pair')
        composer=skill_module(root,'additional-assumptions-dialogue','compose_draft')
        template,slots=composer.compose(current,current['writer_pack'])
        require((directory/'draft_template.md').read_text()==template,'warning template differs from current claims')
        from scripts.batch.harbor_dialogue_authoring import verify_author
        _,incoming,faults,_=verify_author(writer_calls[0],policy=policy,kind='writer',round=1,
            template=template,slots=slots,pack=current['writer_pack'],root=root)
        require(not faults and (round_dir/'writer.filled.md').read_text()==incoming,'warning writer bytes differ from actual Call')
        _,filled,faults,_=verify_author(writer_calls[1],policy=policy,kind='paraphraser',round=1,
            template=template,slots=slots,pack=current['writer_pack'],incoming=incoming,root=root)
        require(not faults and (round_dir/'paraphraser.filled.md').read_text()==filled,'warning paraphraser bytes differ from actual Call')
        source_warning_proof=source_warnings.validate(call.parsed,reviewed_claims,record(round_dir/'draft.md'),
            record(directory/'draft_template.md'),record(round_dir/'writer.filled.md'),
            record(round_dir/'paraphraser.filled.md'),source,task_dir=task_dir,root=root)
        source_warnings.revalidate(source_warning_proof,review,current,body['draft.md'],root=root)
    require(validator.review_passed(review,current,draft_text=body['draft.md'],source_warning_proof=source_warning_proof),
            'dialogue required-claim/semantic coverage failed')
    verified={'claims':len(current['claims']),'round':body['round'],'provenance':result}
    if source_warning_proof is not None:verified['source_warning_proof']=source_warning_proof
    return verified


@dataclass(frozen=True)
class HistoricalApplicability:
    workbook: str
    publication_root: Path
    source_root: Path
    segmentation_root: Path
    task_root: Path
    release_id: str
    parent_review: dict
    handoff: dict
    publication_evidence: dict
    reuse: dict
    current_source_sha256: str


@dataclass(frozen=True)
class HistoricalCall:
    request: dict
    response: dict
    parsed: dict | None = None


def _historical_evidence(h,*,workbook,notes):
    """Read the exact published historical tuple; this grants no current PASS."""
    import xl_release_publication as publication
    require(isinstance(h,HistoricalApplicability) and h.workbook==workbook,'historical evidence belongs to another task')
    for binding in (h.parent_review,h.handoff,h.publication_evidence,h.reuse):checked_record(binding)
    directory,manifest=publication.resolve_current_release_integrity(h.publication_root,source_root=h.source_root,segmentation_root=h.segmentation_root,task_root=h.task_root,expected_release_id=h.release_id)
    identity=manifest['identity'];require(identity['workbook_id']==workbook,'published workbook differs')
    require(identity['bindings']['source_sha256']==h.current_source_sha256,'historical source differs from current source')
    parent=read(h.parent_review);handoff=read(h.handoff);evidence=read(h.publication_evidence);reuse=read(h.reuse)
    require(parent['handoff_sha256']==h.handoff['sha256'],'root review does not bind this handoff')
    require(handoff['workbook_id']==workbook and evidence.get(h.handoff['path'])==h.handoff['sha256'],'handoff is not in exact published evidence')
    taskdir=h.task_root/'task-generations'/identity['task_generation']['generation_id']
    task=read(record(taskdir/'generation-manifest.json'))
    task_hash=publication.canonical_hash(task['artifacts'])
    require(handoff['staged_tree_sha256']==parent['stage_sha256']==task_hash,'root-reviewed stage is not the published task')
    require(parent['current_source_generation']==identity['source_generation']['generation_id'] and parent['current_segmentation_generation']==identity['segmentation_generation']['generation_id'],'root-reviewed source/segmentation pair differs')
    require(record(taskdir/'environment/additional-assumptions.md')['sha256']==notes['sha256'],'current dialogue differs from published dialogue')
    require(evidence.get(h.reuse['path'])==h.reuse['sha256'],'reuse evidence is not publication-bound')
    require(reuse['budget_reset'] is False and reuse['new_writer_rounds']==0 and reuse['new_semantic_rounds']==0,'historical applicability must not spend/reset/transfer budgets')
    require(reuse['retained_writer_rounds_used'] in (1,2),'historical round evidence missing')
    bound={}
    for section in ('original_round2_evidence','original_writer_evidence'):
        for binding in reuse[section].values():
            checked_record(binding)
            require(binding['path'] not in evidence or evidence[binding['path']]==binding['sha256'],'historical derived context conflicts with publication evidence')
            bound[binding['path']]=binding['sha256']
    for path,sha in handoff['model_request_response_hashes'].items():
        require(evidence.get(path)==sha,'handoff model record not in publication evidence')
    claims_binding=reuse['comparisons']['claims.json']['old'];checked_record(claims_binding)
    require(claims_binding['path'] not in evidence or evidence[claims_binding['path']]==claims_binding['sha256'],'historical reviewed claims conflict with publication evidence')
    retained=[{'path':path,'sha256':sha} for path,sha in bound.items()]
    retained += [claims_binding,record(directory/'release-manifest.json'),record(h.publication_root/'current-release.json'),record(taskdir/'generation-manifest.json')]
    retained += [record(Path(identity[k]['manifest_path'])) for k in ('source_generation','segmentation_generation')]
    return {'evidence':retained,'reviewed_draft_sha256':reuse['original_round2_evidence']['draft.md']['sha256'],'reviewed_claims_sha256':claims_binding['sha256'],'bindings':bound,'model_bindings':handoff['model_request_response_hashes'],'release_id':h.release_id,'transport_journals_available':False,'provenance':'root-selected published historical raw responses; no fabricated invocation or submission fields'}


def _historical_call(call,expected_request,*,policy,evidence,json_response=True):
    require(isinstance(call,HistoricalCall),'typed historical call required')
    for binding in (call.request,call.response):
        checked_record(binding);require(evidence['bindings'].get(binding['path'])==binding['sha256'] and evidence['model_bindings'].get(binding['path'])==binding['sha256'],'call not bound by historical publication/reuse')
    request=read(call.request);response=read(call.response)
    require(request==expected_request and set(request)=={'model','reasoning_effort','messages'},'historical request/context changed')
    require(request['model']==policy.model and request['reasoning_effort']==policy.reasoning_effort,'historical model/configuration differs')
    messages=request['messages'];require(isinstance(messages,list) and len(messages)==2 and [m['role'] for m in messages]==['system','user'],'historical stateless messages required')
    require(all(set(m)=={'role','content'} and isinstance(m['content'],str) and m['content'].strip() for m in messages),'historical messages malformed')
    require(response.get('model')==policy.model and response.get('object')=='chat.completion' and isinstance(response.get('id'),str) and response['id'] and type(response.get('created')) is int,'historical response identity/model invalid')
    choices=response.get('choices');require(isinstance(choices,list) and len(choices)==1,'historical response choices incomplete')
    choice=choices[0];message=choice.get('message',{})
    require(choice.get('finish_reason')=='stop' and message.get('role')=='assistant' and not message.get('refusal') and not message.get('tool_calls') and isinstance(message.get('content'),str),'historical response incomplete/refused')
    value=json.loads(message['content']) if json_response else message['content']
    if call.parsed:
        checked_record(call.parsed);require(evidence['bindings'].get(call.parsed['path'])==call.parsed['sha256'],'historical parsed review not bound')
        require(read(call.parsed)==value,'historical parsed review differs from raw response')
    return {'response_id':response['id'],'review':value,'request_sha256':call.request['sha256'],'response_sha256':call.response['sha256'],'historical_release_id':evidence['release_id'],'transport_journals_available':False}
