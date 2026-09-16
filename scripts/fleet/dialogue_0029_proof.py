"""Closed, source-backed 0029 correction proof. Never author or review a task.

The exact private proposal is deterministic postprocessing, not an original
model response. All final clauses still require an ordinary independent review.
"""
from __future__ import annotations
import hashlib,json,re
from pathlib import Path

SOURCE_SHA='55029f5834ed624c303b4d7f00adb9bf0ad88c1ff10856e174e44cc7ae1f9509'
CLAIMS_SHA='25693e71f2a6589d7de87397c977784d0f46652ebaeb518e1f1a722251bc1371'
ORIGINAL_SHA='a7908b73e2cdffb7865491243cf8cf004a658d7b2b18cf8e189527068a73aa40'
PROPOSAL_SHA='d41833b0ef89d38f140e33dc7edcdc8085f0235fbd1f2b8b870d2e750a009183'
CARDINALITY_CLAIM='method_debt_movement::Illustrative Economics::Redeemable Preferred Stock::001'
SCHEMA='harbor-0029-source-proved-pre-review/v1'
# Replaced only in the hash-bound immutable successor, never at import runtime.
PROFILE=None

def require(ok,message):
    if not ok:raise ValueError(message)
def sha(raw):return hashlib.sha256(raw).hexdigest()
def canonical(v):return json.dumps(v,sort_keys=True,separators=(',',':'),ensure_ascii=False).encode()
def record(p):return {'path':str(Path(p).absolute()),'sha256':sha(Path(p).read_bytes())}
def read(b):
    require(record(b['path'])==b,'0029 bound evidence changed');return json.loads(Path(b['path']).read_bytes())
def _text(b):
    require(record(b['path'])==b,'0029 bound text changed');return Path(b['path']).read_text()
def blocks(text):
    matches=list(re.finditer(r'<!-- claim:([^\n]+) -->\n',text));result=[]
    require(matches and matches[0].start()==0,'0029 leading scaffold changed')
    for i,m in enumerate(matches):
        block=text[m.end():matches[i+1].start() if i+1<len(matches) else len(text)]
        turns=list(re.finditer(r'^\*\*([^*]+):\*\* ([^\n]+)',block,re.M))
        require(len(turns)==2,'0029 exact two-turn claim scaffold required')
        require(turns[0][1] in ('Analyst','Associate') and turns[1][1] in ('VP','Director','Managing Director'),'0029 cast changed')
        senior=turns[1]
        result.append({'id':m[1],'junior':turns[0][0],'senior_role':senior[1],'senior':senior[2],
            'scaffold':block[:senior.start(2)]+ '<SOURCE_PROVED_SENIOR>'+block[senior.end(2):]})
    return result

def build_proof(bindings):
    from openpyxl import load_workbook
    from openpyxl.formula.translate import Translator
    required={'source','claims','original_draft','proposal','v2_proof','v3_proof','diagnosis','sections'}
    require(set(bindings)==required,'0029 proof inputs differ')
    for name,digest in [('source',SOURCE_SHA),('claims',CLAIMS_SHA),('original_draft',ORIGINAL_SHA),('proposal',PROPOSAL_SHA)]:
        require(bindings[name]['sha256']==digest and record(bindings[name]['path'])==bindings[name],'0029 exact source/claim/draft identity differs')
    claims=read(bindings['claims']);old=_text(bindings['original_draft']);new=_text(bindings['proposal'])
    a,b=blocks(old),blocks(new);ids=[c['record_id'] for c in claims['claims']]
    require(len(ids)==len(set(ids))==43 and [x['id'] for x in a]==[x['id'] for x in b]==ids,'0029 complete43claim identity/order changed')
    require(all({k:x[k] for k in ('id','junior','senior_role','scaffold')}=={k:y[k] for k in ('id','junior','senior_role','scaffold')} for x,y in zip(a,b)),'0029 non-senior bytes changed')
    changed=[x['id'] for x,y in zip(a,b) if x['senior']!=y['senior']]
    require(len(changed)==36,'0029 exact36senior correction scope changed')
    v2=read(bindings['v2_proof']);v3=read(bindings['v3_proof']);d=read(bindings['diagnosis']);s=read(bindings['sections'])
    require(v2['source']==bindings['source'] and v2['claims']==bindings['claims'] and v2['original_draft']==bindings['original_draft'],'0029 original proposal provenance changed')
    require(v3['prior']==bindings['v2_proof'] and v3['proposal']==bindings['proposal'],'0029 proposal successor changed')
    require(v2['original_diagnosis']==bindings['diagnosis'] and v2['section_proof']==bindings['sections'],'0029 source proof linkage changed')
    require({c['record_id'] for c in v2['changes']}==set(changed),'0029 changed-clause proof coverage differs')
    workbook=load_workbook(bindings['source']['path'],read_only=False,data_only=False)
    formulas=[];source_cells={}
    def cell(sheet,ref):
        v=workbook[sheet][ref].value;source_cells[f"'{sheet}'!{ref}"]=v;return v
    try:
        for c in claims['claims']:
            refs=[x.rsplit('!',1)[1] for x in c['junior_scope']];sheet=c['sheet'];review=c['reviewer_only'];expected=review['evidence']
            require(refs and review['cells']==c['junior_scope'],'0029 source scope differs')
            instances=[]
            for ref in refs:
                formula=cell(sheet,ref);require(isinstance(formula,str) and formula.startswith('='),'0029 source formula missing')
                if review['source']=='custom_method_detector':require(formula==Translator(expected,origin=refs[0]).translate_formula(ref),'0029 source family formula differs')
                else:
                    targets={'041':"='Historical P&L Schedule'!$I$30",'042':"='Historical P&L Schedule'!$I$8",'043':"='Historical P&L Schedule'!$I$8"}
                    require(c['record_id'].split('::')[-1] in targets and formula==targets[c['record_id'].split('::')[-1]],'0029 convention source differs')
                instances.append({'cell':f"'{sheet}'!{ref}",'formula':formula})
            formulas.append({'record_id':c['record_id'],'instances':instances})
        require(len(d['changed_claims'])==24,'0029 historical-period proof cardinality differs')
        cached=load_workbook(bindings['source']['path'],read_only=False,data_only=True)
        try:
            require(cached['Historical P&L Schedule']['G6'].value==2019 and cached['Historical P&L Schedule']['F6'].value==2018,'0029 source period header values differ')
        finally:cached.close()
        for row in d['changed_claims']:
            c=next(c for c in claims['claims'] if c['record_id']==row['record_id']);sheet=c['sheet']
            require(cell(sheet,row['source_target'])==row['formula'],'0029 historical selector formula differs')
            require(row['true_header']==2019 and row['false_header']==2018,'0029 historical period order differs')
            require(cell(sheet,'G6')=='=F6+1' and cell(sheet,'F6')=='=E6+1','0029 source period header formula order differs')
        require(len(s['additional_distinct_branch_collisions'])==5,'0029 section correction cardinality differs')
        for ref,label in s['section_headers'].items():require(cell('Illustrative Economics',ref)==label,'0029 source section label differs')
        for row in s['additional_distinct_branch_collisions']:
            for scoped in row['scoped_source_formulas']:
                require(cell('Illustrative Economics',scoped['target'])==scoped['formula'],'0029 branch source formula differs')
                require(scoped['true_operand']!=scoped['false_operand'],'0029 source branches collapsed')
        first=claims['claims'][0];require(first['record_id']==CARDINALITY_CLAIM and first['reviewer_only']['evidence']=='=(C36-SUM(C32:$C$33))*$D$14' and '2-row' in first['must_say'],'0029 cardinality source differs')
        cardinality={'record_id':CARDINALITY_CLAIM,'source_cell':"'Illustrative Economics'!C34",'source_formula':first['reviewer_only']['evidence'],'range':"'Illustrative Economics'!C32:C33",'row_count':2,'column_count':1,'phrase':'locked 2-row input block','occurrences':1}
        require(b[0]['senior'].count(cardinality['phrase'])==1,'0029 cardinality phrase differs')
        require(a[0]==b[0],'0029 cardinality clause must remain unchanged')
    finally:workbook.close()
    return {'schema_version':SCHEMA,'bindings':bindings,'claims_count':43,'changed_claim_ids':changed,
        'unchanged_claim_ids':[x for x in ids if x not in changed],'source_formulas':formulas,'source_cells':source_cells,
        'cardinality':cardinality,'model_calls_added':0,'semantic_budget_reset':False,'independent_review_required':True,'qualified':False}

def profile():
    require(PROFILE is not None,'0029 immutable profile is not configured');return read(PROFILE)
def runtime_producer(selection):
    """Read the exact retained producer identity, never execute historical tools."""
    from scripts.batch.harbor_runtime_readiness import RuntimeProducerIdentity
    p=profile()['historical_runtime']
    if selection.observation.execution!=p['execution']:return None
    require(selection.content.workbook=='0029','0029 historical runtime workbook differs')
    report=read(p['execution']);command=read(p['main_build']);intent=read(p['producer_intent'])
    require(report['valid'] is True and report['run_id']==selection.observation.run_id,'0029 historical runtime identity differs')
    require(command['argv'][0]==p['docker_executable'] and intent['executable']['path']==p['python_executable'],'0029 historical producer locators changed')
    return RuntimeProducerIdentity(p['python_executable'],p['docker_executable'])
def verified_code(root):
    p=profile();old=read(p['previous']);expected={r['relative']:r['sha256'] for r in old['files']}
    for name,h in expected.items():require(sha((Path(old['root'])/name).read_bytes())==h,'0029 predecessor code changed')
    require(p['semantic_budget_reset'] is False and len(p['legacy_calls'])==8,'0029 successor scope differs')
    for name,change in p['patches'].items():
        if name in expected:require(change['before']==record(Path(old['root'])/name),'0029 patch predecessor differs')
        else:require(change['before'] is None,'0029 unexpected added-file predecessor')
        raw=Path(change['after']['path']).read_bytes();require(sha(raw)==change['after']['sha256'],'0029 exact patch witness changed');expected[name]=sha(raw)
    template=_text(p['proof_template']);needle='\nPROFILE=None\n';require(template.count(needle)==1,'0029 proof template ambiguous')
    expected['scripts/fleet/dialogue_0029_proof.py']=sha(template.replace(needle,'\nPROFILE='+repr(PROFILE)+'\n',1).encode())
    expected.pop('.pipeline-code.json',None)
    actual={str(f.relative_to(root)):sha(f.read_bytes()) for f in Path(root).rglob('*') if f.is_file() and '__pycache__' not in f.parts}
    actual.pop('.pipeline-code.json',None)
    require(read(record(Path(root)/'.pipeline-code.json'))==actual,'0029 successor manifest changed')
    require(actual==expected,'0029 unreviewed successor code delta')
    return sorted(name for name in actual if actual[name]!={r['relative']:r['sha256'] for r in old['files']}.get(name))
def correction(draft,claims,source,*,root):
    verified_code(root);p=profile();proof=read(p['source_proof']);actual=build_proof(proof['bindings'])
    require(actual==proof,'0029 full source proof changed')
    require(source==proof['bindings']['source'] and claims==read(proof['bindings']['claims']) and draft==_text(proof['bindings']['original_draft']),'0029 correction detached original source/authors')
    return _text(proof['bindings']['proposal']),proof

def cardinality_budget(text,claim,source):
    """License only one exact source-derived structural count, never a value."""
    from collections import Counter
    if claim.get('record_id')!=CARDINALITY_CLAIM:return Counter()
    p=profile();proof=read(p['source_proof']);b=proof['bindings']
    require(source==b['source'] and record(source['path'])==source,'0029 cardinality source binding differs')
    selected=read(b['claims'])['claims'][0];require(claim==selected,'0029 cardinality claim changed')
    senior=blocks(_text(b['proposal']))[0]['senior']
    # Any reworded second-round senior must prove its own count; this profile
    # does not automatically license a token in arbitrary future prose.
    if text!=senior:return Counter()
    require(text.count('locked 2-row input block')==1 and proof['cardinality']['row_count']==2,'0029 cardinality span changed')
    return Counter({2.0:1})

def fleet_legacy_plan(root,current):
    verified_code(Path(__file__).resolve().parents[2]);p=profile();key=str(Path(root));require(key in p['legacy_calls'],'0029 legacy Call not explicitly pinned')
    row=p['legacy_calls'][key]
    for b in row.values():require(record(b['path'])==b,'0029 legacy Call bytes changed')
    require(row['plan']==record(Path(root)/'plan.json') and row['outcome']==record(Path(root)/'outcome.json'),'0029 legacy Call location differs')
    old=read(row['plan']);require(read(row['outcome'])['status']=='completed' and 'fleet_transport_policy' not in old,'0029 incomplete or prospective Call cannot use historical inspection')
    expected=dict(old,fleet_transport_policy='single-dispatch-durable-global-gate/v1',max_transport_retries=0)
    require(canonical(current)==canonical(expected),'0029 historical Call identity/budget/configuration differs')
    return old

def require_fleet_custody(root):
    """Check the selected live central claim before owning permit mutation."""
    from scripts.fleet import model_gate
    require(model_gate.configured(),'0029 pre-review adoption requires live fleet custody')
    client,owner,_=model_gate.context();require(owner['workbook']=='0029','0029 foreign fleet workbook')
    verified_code(root);status=client.call('status')
    require(status['model_limit']==status['publication_limit']==1,'0029 fleet limits changed')
    rows=[r for r in status['workbooks'] if r['workbook']=='0029'];require(len(rows)==1,'0029 central row missing')
    row=rows[0]
    require(row['state']=='active' and row['owner']==owner['owner'] and row['fence']==owner['fence']
            and row['source_sha256']==SOURCE_SHA and row['version']==Path(root).name,'0029 central source/version/fence/custody differs')
    require(not any(o['state']=='hold' for o in status['operations']),'0029 global operation needs reconciliation')
    # The ordinary heartbeat API checks the secret fencing token on the server;
    # no token is returned or embedded in the successor's public proof.
    client.call('heartbeat',**owner)
    return {'workbook':'0029','owner':owner['owner'],'fence':owner['fence'],'version':row['version']}

def validate_custody(failed,binding,before,lease):
    """Keep the three original renewals and one owning bounded takeover."""
    from scripts.batch.harbor_nat_retry import _renewal_chain
    from scripts.batch.harbor_workbook_bootstrap import Lease
    from scripts.batch import harbor_lease_takeover as T
    from dataclasses import asdict
    p=profile();rows=[read(b) for b in p['custody_chain']]
    require(len(rows)==3 and rows[0]['old_lease']==failed,'0029 original failed custody differs')
    previous=failed
    for row in rows:
        require(row['old_lease']==previous and row['before']['sha256']==previous['sha256'],'0029 historical custody link differs')
        _renewal_chain(row['before'],row['old_lease'],read(row['after']),Lease(**row['lease']))
        previous=row['lease']
    first=rows[0]['before'];tracker=Path(p['legacy_calls'][next(iter(p['legacy_calls']))]['plan']['path']).parents[3]/'trackers/workbooks/0029.json'
    if binding is None:
        _renewal_chain(first,failed,before,lease)
    else:
        require(read(binding)['schema_version']==T.SCHEMA,'0029 requires owning expired-lease takeover completion')
        T.continuation_chain(first,failed,before,lease,takeover=binding,tracker_path=tracker)
    require(before['workbook_id']=='0029' and before['bindings']['original_raw_source']['sha256']==SOURCE_SHA,'0029 current custody source differs')
    return {'original_revision':failed['revision'],'current_revision':lease.revision,'historical_renewals':3,'takeover':binding,'semantic_budget_reset':False}
