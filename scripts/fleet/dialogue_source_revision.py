"""Explicit deterministic source corrections; original authored text stays intact.

This module produces a separately named derivative and an auditable byte delta.
It cannot authorize a model call, alter a task, or confer review qualification.
"""
from __future__ import annotations
import hashlib
import re
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from xl_ast_graph import parse_formula

SCHEMA='harbor-dialogue-source-revision/v1'


class SourceFormulaTranslator(Translator):
    @classmethod
    def translate_range(cls,range_str,rdelta,cdelta):
        if re.fullmatch(r"(?:'(?:[^']|'')+'|[A-Za-z_][A-Za-z0-9_.]*)!#REF!",range_str):
            return range_str
        return super().translate_range(range_str,rdelta,cdelta)



def require(ok,message):
    if not ok:raise ValueError(message)


def node(n):
    return (n.kind,n.name,getattr(n,'value',None),tuple(node(a) for a in n.args))


def reference(value,sheet):
    if '!' in value:
        s,c=value.rsplit('!',1);s=s.strip("'").replace("''", "'")
    else:s,c=sheet,value
    require(s==sheet and re.fullmatch(r'\$?[A-Z]+\$?[1-9][0-9]*',c),'foreign or non-cell source reference')
    return c.replace('$','')


def blocks(draft,claims):
    ids=[c['record_id']for c in claims['claims']]
    markers=list(re.finditer(r'<!-- claim:([^\n]+) -->',draft))
    require([m.group(1)for m in markers]==ids and len(set(ids))==len(ids),'exact original claim order required')
    return {m.group(1):(m.start(),markers[i+1].start() if i+1<len(markers)else len(draft))for i,m in enumerate(markers)}


def family(book,claim):
    r=claim['reviewer_only'];sheet=claim['sheet'];cells=r['cells']
    require(r['source']=='custom_method_detector' and cells==claim['junior_scope'] and cells,'source method family required')
    addresses=[reference(c,sheet)for c in cells];formula=r['evidence']
    require(book[sheet][addresses[0]].value==formula,'source formula differs from retained card')
    rows=[]
    for address in addresses:
        actual=book[sheet][address].value
        require(actual==SourceFormulaTranslator(formula,origin=addresses[0]).translate_formula(address),'copied source family differs')
        rows.append(dict(cell=address,formula=actual))
    return parse_formula(formula),rows


def label(book,sheet,ref):
    row=int(re.search(r'[0-9]+$',ref).group())
    # These methods explicitly use visible row labels in column B. Reject
    # missing/ambiguous labels rather than inferring names from reviewer prose.
    text=book[sheet].cell(row,2).value
    require(isinstance(text,str) and text.strip()==text and text and not any(c in text for c in '\n"'),'visible source row label required')
    return text


def unary_sum(book,claim,tree,block):
    require(tree.kind=='func' and tree.name=='IFERROR' and len(tree.args)==2,'IFERROR source shape required')
    minimum,zero=tree.args
    require(zero.kind=='const' and zero.value=='0' and minimum.kind=='func' and minimum.name=='MIN' and len(minimum.args)==2,'source outer bounds changed')
    z,maximum=minimum.args
    require(z.kind=='const' and z.value=='0' and maximum.kind=='func' and maximum.name=='MAX' and len(maximum.args)==2,'source candidate comparison changed')
    negative,other=maximum.args
    require(negative.kind=='prefix' and negative.name=='-' and len(negative.args)==1,'whole-expression source negation required')
    total=negative.args[0]
    require(total.kind=='infix' and total.name=='+' and len(total.args)==2,'negated sum required')
    balance,product=total.args
    require(balance.kind=='ref' and product.kind=='infix' and product.name=='*' and len(product.args)==2 and node(product.args[0])==node(balance) and product.args[1].kind=='ref','balance plus its hurdle product required')
    hurdle=product.args[1]
    require(re.fullmatch(r'\$[A-Z]+\$[1-9][0-9]*',hurdle.name),'fixed hurdle source required')
    require(other.kind=='infix' and other.name=='*' and len(other.args)==2 and other.args[0].kind=='prefix' and other.args[0].name=='-' and len(other.args[0].args)==1 and other.args[0].args[0].kind=='ref' and other.args[1].kind=='ref','other candidate changed')
    balance_label=label(book,claim['sheet'],balance.name);hurdle_label=label(book,claim['sheet'],hurdle.name)
    # Only the explicitly mistaken clause is replaceable, in a senior turn.
    pattern=r'(?:One candidate is|Start with) the row labelled "'+re.escape(balance_label)+r'" plus the sign-flipped product of that (?:balance|principal) and the locked input on the row labelled "'+re.escape(hurdle_label)+r'"\.'
    matches=list(re.finditer(pattern,block));require(len(matches)==1,'exact mistaken sum clause required')
    old=matches[0].group();new=('First add the row labelled "'+balance_label+'" plus the product of that same row and the locked input on the row labelled "'+hurdle_label+'". Then flip the sign of that entire sum to form the first candidate.')
    return old,new,dict(operator='negate_entire_sum',balance=balance.name,hurdle=hurdle.name,balance_label=balance_label,hurdle_label=hurdle_label)


def cumulative_sum(book,claim,tree,block):
    require(tree.kind=='func' and tree.name=='IF' and len(tree.args)==3,'conditional cumulative source required')
    test=tree.args[0]
    require(test.kind=='infix' and test.name=='<=' and len(test.args)==2 and test.args[1].kind=='const' and test.args[1].value=='0','cumulative zero test changed')
    summed=test.args[0]
    require(summed.kind=='func' and summed.name=='SUM' and len(summed.args)==1,'single cumulative range required')
    ran=summed.args[0]
    require(ran.kind=='ref' and ran.name.count(':')==1 and not ran.args,'actual source range required')
    start,end=ran.name.split(':')
    require(re.fullmatch(r'\$[A-Z]+\$[1-9][0-9]*',start) and re.fullmatch(r'[A-Z]+[1-9][0-9]*',end) and start.replace('$','')==end,'fixed start and moving first-period endpoint required')
    row_label=label(book,claim['sheet'],start)
    old='If the locked input on the row labelled "'+row_label+'" is at most zero,'
    require(block.count(old)==1,'exact mistaken single-input clause required')
    first=reference(claim['reviewer_only']['cells'][0],claim['sheet'])
    from openpyxl.utils.cell import coordinate_from_string,column_index_from_string
    require(column_index_from_string(coordinate_from_string(first)[0])==column_index_from_string(coordinate_from_string(end)[0])+1,'range endpoint must immediately precede this period')
    new=('Add the row labelled "'+row_label+'" from its fixed starting column through the column immediately before this period, extending the range as each period advances. If that cumulative total is at most zero,')
    return old,new,dict(operator='cumulative_sum',start=start,end=end,row_label=row_label)


def decimal_rounding(book,claim,tree,block):
    """Correct a precision operand misrendered as a rounding multiple."""
    require(tree.kind=='func' and tree.name=='ROUND' and len(tree.args)==2,
            'source must round to decimal precision')
    precision=tree.args[1]
    require(precision.kind=='const' and precision.name=='number'
            and re.fullmatch(r'\d+',str(precision.value)) is not None,
            'literal nonnegative decimal precision required')
    digits=int(precision.value)
    require(0<=digits<=15,'bounded decimal precision required')
    def rounding_nodes(n):
        return int(n.kind=='func' and n.name in ('ROUND','MROUND'))+sum(rounding_nodes(a)for a in n.args)
    require(rounding_nodes(tree)==1,'nested rounding needs separate source diagnosis')
    pattern=r'\bround(?:ed)?(?: it| the greater (?:value|amount))? to the nearest multiple of '+str(digits)+r'(?!\w|\.\d)'
    matches=list(re.finditer(pattern,block))
    require(len(matches)==1,'exact mistaken decimal-rounding clause required')
    old=matches[0].group()
    require(not any(old in text for text in claim.get('must_say',[])),
            'source label or required clause cannot be rewritten as rounding')
    new=old.replace('to the nearest multiple of '+str(digits),'to '+str(digits)+' decimal places')
    return old,new,dict(operator='decimal_rounding',precision=digits,
                        formula=claim['reviewer_only']['evidence'])


def visible_row_label(book,claim,tree,block):
    """Remove one surplus closing parenthesis from an exact visible row label."""
    cells=claim['reviewer_only']['cells'];labels={label(book,claim['sheet'],reference(c,claim['sheet']))for c in cells}
    require(labels=={claim['row_label']},'source row label is missing or inconsistent')
    value=claim['row_label'];require(value.count('(')==value.count(')'),'balanced visible source label required')
    old='the row labelled "'+value+')"';new='the row labelled "'+value+'"'
    require(block.count(old)==1 and new in claim['must_say'],'exact required label typo must be unique')
    return old,new,dict(operator='visible_row_label',label=value,removed_character=')')


def broken_reference_wording(book,claim,tree,block):
    """Express a proved broken CHOOSE option without printing its reference token."""
    errors=[]
    def visit(n,parent=None,index=None):
        if n.kind=='ref'and '#REF!'in n.name:
            require(re.fullmatch(r"(?:'(?:[^']|'')+'|[A-Za-z_][A-Za-z0-9_.]*)!#REF!",n.name) is not None
                and parent is not None and parent.kind=='func'and parent.name=='CHOOSE'
                and len(parent.args)==4 and index==3,'exact third CHOOSE option must be a qualified broken reference')
            errors.append(n.name)
        for i,a in enumerate(n.args):visit(a,n,i)
    visit(tree)
    require(errors and len(errors)<=2 and len(set(errors))==1,'one consistent source error operand required')
    token=errors[0];old='uses the broken reference "'+token+'" and yields a reference error'
    require(block.count(old)==1 and all(token not in clause for clause in claim['must_say']),
            'source error wording is ambiguous or an explicit required label')
    return old,'yields a reference error because that source option is a broken reference',dict(operator='broken_choose_option',option=3,source_operand=token,occurrences=len(errors),error_preserved=True)


def repair(draft,claims,source,selections):
    require(isinstance(selections,list) and 0<=len(selections)<=32,'bounded explicit clause selection required')
    # An explicit empty selection proves an unchanged candidate for a separately
    # authorized review. It adds no source claim or model allowance.
    ids=[x['record_id']for x in selections]
    require(len(set(ids))==len(ids) and all(set(x)=={'record_id','kind'}for x in selections),'unique closed correction selection required')
    require(sum(x['kind']!='decimal_rounding'for x in selections)<=8,
            'existing non-rounding correction allowance remains eight')
    locations=blocks(draft,claims);cards={c['record_id']:c for c in claims['claims']}
    require(set(ids)<=set(cards),'selected correction not in original cards')
    book=load_workbook(source,data_only=False,read_only=True);edits=[]
    try:
        for selected in selections:
            c=cards[selected['record_id']];lo,hi=locations[selected['record_id']];block=draft[lo:hi]
            ast,rows=family(book,c)
            fn={'negated_sum':unary_sum,'cumulative_range':cumulative_sum,'decimal_rounding':decimal_rounding,'visible_row_label':visible_row_label,'broken_reference_wording':broken_reference_wording}.get(selected['kind']);require(fn is not None,'unsupported source correction')
            old,new,meaning=fn(book,c,ast,block)
            require(block.count(old)==1,'ambiguous source correction')
            pos=block.index(old);prefix=block[:pos]
            roles=list(re.finditer(r'\*\*([^\n*]+):\*\*',prefix));require(roles and roles[-1].group(1)in('VP','Director','Managing Director'),'correction must be within a senior turn')
            edits.append(dict(record_id=selected['record_id'],kind=selected['kind'],start=lo+pos,end=lo+pos+len(old),old=old,new=new,source_cells=rows,meaning=meaning))
    finally:book.close()
    fixed=draft
    for e in sorted(edits,key=lambda e:e['start'],reverse=True):fixed=fixed[:e['start']]+e['new']+fixed[e['end']:]
    after=blocks(fixed,claims)
    for rid,(lo,hi)in locations.items():
        if rid not in ids:require(draft[lo:hi]==fixed[slice(*after[rid])],'unselected claim block changed')
    proof=dict(schema_version=SCHEMA,source_sha256=hashlib.sha256(Path(source).read_bytes()).hexdigest(),old_draft_sha256=hashlib.sha256(draft.encode()).hexdigest(),new_draft_sha256=hashlib.sha256(fixed.encode()).hexdigest(),edits=edits,claims=len(cards),unchanged_claims=len(cards)-len(edits),model_calls_added=0,task_mutations=0,qualified=False)
    return fixed,proof
