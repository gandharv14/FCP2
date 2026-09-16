"""Build explicit, reproducible operating variants of two macro-dependent tasks.

The gold oracle reads only original input cells, never cached model outputs.
A separate floating-point oracle uses openpyxl and annual closed-form sums.
The workbook change only appends a blank output sheet and its registrations.
"""
from __future__ import annotations

import copy
import hashlib
import json
import math
from decimal import Decimal, getcontext
from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import ZipFile
import xml.etree.ElementTree as ET

import openpyxl

from scripts.recursion_hardening_workbooks import NS, sha256, sheets

ROOT = Path(__file__).resolve().parents[1]
AUDIT = ROOT / 'reports/recursion/gemini-high-score-audit-20260915'
OUT = AUDIT / 'refactor/depth-fixes'
D = Decimal
getcontext().prec = 40
SOLAR_COST_ROWS = [199, 200, *range(205, 213)]
WIND_COST_ROWS = [219, 220, *range(225, 229)]


class Inputs:
    def __init__(self, path: Path, sheet: str):
        self.sheet = sheet
        with ZipFile(path) as z:
            self.cells = {c.attrib['r']: c for c in ET.fromstring(z.read(sheets(z)[sheet])).findall('.//s:sheetData/s:row/s:c', NS)}
        self.used = {}

    def __call__(self, cell: str, *, blank_zero: bool = False) -> Decimal:
        node = self.cells.get(cell)
        value = node.findtext('s:v', namespaces=NS) if node is not None else None
        if value is None and blank_zero:
            self.used[f"'{self.sheet}'!{cell}"] = {'value': None, 'interpretation': 'empty fixed-cost input contributes zero'}
            return D(0)
        if node is None or node.attrib.get('t', 'n') != 'n' or value is None or node.find('s:f', NS) is not None:
            raise ValueError(f'Missing or nonnumeric primitive input: {self.sheet}!{cell}')
        self.used[f"'{self.sheet}'!{cell}"] = {'value': value}
        return D(value)


def decimal_solar(v: Inputs):
    capacity, hours, price = v('K30'), v('K31'), v('K22')
    years, season = int(v('K18')), v('K38')
    fixed = sum(v('K'+str(r)) for r in SOLAR_COST_ROWS)
    degeneration, rows = D(1), []
    for year in range(1, years+1):
        availability = v('K'+str(120+year))
        degradation = v('K'+str(158+year))
        annual_cost = fixed + v('K'+str(44+year)) + v('K'+str(82+year))
        for semester, share in [(1, season), (2, 1-season)]:
            energy = capacity*hours*availability*degeneration*share
            revenue = energy*price/1000
            rows.append({'year':year, 'semester':semester, 'generation':energy, 'revenue':revenue, 'opex':annual_cost/2, 'contribution':revenue-annual_cost/2})
            degeneration /= 1+degradation/2
    total = {k:sum(row[k] for row in rows) for k in ['generation','revenue','opex','contribution']}
    return [next(x['generation'] for x in rows if x['year']==7 and x['semester']==2), sum(x['generation'] for x in rows if x['year']==20), total['revenue'],total['opex'],total['contribution'],total['opex']*1000/total['generation']], rows


def decimal_wind(v: Inputs):
    portfolio = v('F38')
    rows = []
    for col in 'KLMNOPQ':
        turbines, capacity = v(col+'34'), v(col+'33')*v(col+'34')
        hours = {'P90':v(col+'37')*(1+portfolio), 'P50':v(col+'41')}
        price, season = v(col+'25'), v(col+'56')
        fixed = sum(v(col+str(r), blank_zero=True) for r in WIND_COST_ROWS)
        for year in range(1,11):
            availability = v(col+str(183+year))
            variable = v(col+str(84+year))
            annual_fixed = turbines*v(col+str(117+year))+v(col+str(150+year))+fixed
            for scenario in ['P90','P50']:
                for semester,share in [(1,season),(2,1-season)]:
                    energy = capacity*hours[scenario]*availability*share
                    revenue = energy*price/1000
                    costs = energy*variable/1000+annual_fixed/2
                    rows.append({'project_column':col,'year':year,'semester':semester,'scenario':scenario,'generation':energy,'revenue':revenue,'opex':costs,'contribution':revenue-costs})
    p90 = [x for x in rows if x['scenario']=='P90']
    total = {k:sum(x[k] for x in p90) for k in ['generation','revenue','opex','contribution']}
    return [sum(x['revenue'] for x in p90 if x['year']==1),sum(x['contribution'] for x in p90 if x['year']==10 and x['semester']==2),total['contribution'],sum(x['contribution'] for x in rows if x['scenario']=='P50')-total['contribution'],total['opex'],1000*total['opex']/total['generation']], rows


def independent_float(path: Path, task: str):
    """Independent annual aggregation; no Decimal oracle or its intermediates."""
    w = openpyxl.load_workbook(path, data_only=True, read_only=False)
    if task=='0239':
        s=w['Inputs - Solar Parks']
        get=lambda r:float(s.cell(r,11).value)
        n=int(get(18)); base=get(30)*get(31); price=get(22); share=get(38)
        rates=[get(159+y)/2 for y in range(n)]
        first=[math.prod((1+r)**-2 for r in rates[:y]) for y in range(n)]
        second=[first[y]/(1+rates[y]) for y in range(n)]
        energy=[base*get(121+y)*(share*first[y]+(1-share)*second[y]) for y in range(n)]
        costs=[get(45+y)+get(83+y)+math.fsum(get(r) for r in SOLAR_COST_ROWS) for y in range(n)]
        revenue=math.fsum(energy)*price/1000; total_cost=math.fsum(costs)
        result=[base*get(127)*(1-share)*second[6],energy[19],revenue,total_cost,revenue-total_cost,total_cost*1000/math.fsum(energy)]
    else:
        s=w['Inputs - Wind Parks']; portfolio=float(s['F38'].value)
        annual=[]; increment=[]
        for c in range(11,18):
            get=lambda r:float(s.cell(r,c).value or 0)
            mw=get(33)*get(34); h90=get(37)*(1+portfolio); h50=get(41)
            fixed_other=math.fsum(get(r) for r in WIND_COST_ROWS)
            for y in range(10):
                gen=mw*h90*get(184+y)
                fixed=get(34)*get(118+y)+get(151+y)+fixed_other
                variable=get(85+y); revenue=gen*get(25)/1000; costs=gen*variable/1000+fixed
                annual.append({'y':y,'gen':gen,'rev':revenue,'costs':costs,'s2margin':gen*(1-get(56))*(get(25)-variable)/1000-fixed/2})
                increment.append(mw*(h50-h90)*get(184+y)*(get(25)-variable)/1000)
        generation=math.fsum(x['gen'] for x in annual); costs=math.fsum(x['costs'] for x in annual)
        result=[math.fsum(x['rev'] for x in annual if x['y']==0),math.fsum(x['s2margin'] for x in annual if x['y']==9),math.fsum(x['rev']-x['costs'] for x in annual),math.fsum(increment),costs,costs*1000/generation]
    w.close()
    return result


LABELS = {
 '0239': ['Year 7 second-half net generation (MWh)','Year 20 annual net generation (MWh)','Years 1-20 cumulative base-FiT revenue (kEUR)','Years 1-20 cumulative defined cash operating costs (kEUR)','Years 1-20 cumulative operating contribution (kEUR)','Years 1-20 uniform operating breakeven tariff (EUR/MWh)'],
 '0719': ['Year 1 P90 portfolio base-FiT revenue (kEUR)','Year 10 second-half P90 portfolio operating contribution (kEUR)','Years 1-10 cumulative P90 portfolio operating contribution (kEUR)','Years 1-10 P50 less P90 operating contribution (kEUR)','Years 1-10 cumulative P90 defined cash operating costs (kEUR)','Years 1-10 P90 uniform operating breakeven tariff (EUR/MWh)']}


def append_output_sheet(source: Path, output: Path, labels: list[str]):
    structural = ['xl/workbook.xml','xl/_rels/workbook.xml.rels','[Content_Types].xml']
    newpath = 'xl/worksheets/operatingReview.xml'
    with ZipFile(source) as old:
        if 'OperatingReview' in sheets(old) or newpath in old.namelist():
            raise ValueError('Output sheet already exists')
        book = old.read(structural[0]); rels=old.read(structural[1]); types=old.read(structural[2])
        tree=ET.fromstring(book); sheet_id=max(int(x.attrib['sheetId']) for x in tree.findall('s:sheets/s:sheet',NS))+1
        existing_ids={r.attrib['Id'] for r in ET.fromstring(rels)};rid='rIdOperatingReview'
        if rid in existing_ids:raise ValueError('Relationship id already exists')
        replacements={
          structural[0]:book.replace(b'</sheets>',f'<sheet xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" name="OperatingReview" sheetId="{sheet_id}" r:id="{rid}"/></sheets>'.encode()),
          structural[1]:rels.replace(b'</Relationships>',f'<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="/xl/worksheets/operatingReview.xml" Id="{rid}"/></Relationships>'.encode()),
          structural[2]:types.replace(b'</Types>',b'<Override PartName="/xl/worksheets/operatingReview.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>')}
        for name,after in replacements.items():
            if after==old.read(name):raise ValueError('Unexpected package namespace/registration shape')
            ET.fromstring(after)
        rows=['<row r="1"><c r="A1" t="inlineStr"><is><t>Operating review — explicit task conventions in instruction.md</t></is></c></row>', '<row r="3"><c r="A3" t="inlineStr"><is><t>Figure</t></is></c><c r="B3" t="inlineStr"><is><t>Computed output</t></is></c></row>']
        for row,label in enumerate(labels,5):rows.append(f'<row r="{row}"><c r="A{row}" t="inlineStr"><is><t>{escape(label)}</t></is></c><c r="B{row}"/></row>')
        worksheet=('<worksheet xmlns="'+NS['s']+'"><cols><col min="1" max="1" width="85" customWidth="1"/><col min="2" max="2" width="24" customWidth="1"/></cols><sheetData>'+''.join(rows)+'</sheetData></worksheet>').encode()
        output.parent.mkdir(parents=True,exist_ok=True)
        with ZipFile(output,'w') as new:
            new.comment=old.comment
            for info in old.infolist():new.writestr(copy.copy(info),replacements.get(info.filename,old.read(info.filename)))
            new.writestr(newpath,worksheet)
        with ZipFile(output) as new:
            assert new.namelist()==old.namelist()+[newpath]
            for name in old.namelist():assert new.read(name)==replacements.get(name,old.read(name))
            assert new.testzip() is None
        w=openpyxl.load_workbook(output,data_only=False,read_only=True)
        assert all(w['OperatingReview'][f'B{r}'].value is None for r in range(5,11));w.close()
    return {'original_worksheet_entries_byte_identical':True,'other_original_entries_byte_identical':True,'structural_registration_entries_appended':structural,'new_entry':newpath,'no_formulas_added':True,'all_new_target_cells_blank':True,'package_crc_valid':True}


def instruction(task: str, original: str):
    service=original.split('## Research data service',1)[1].split('## What to compute',1)[0]
    common=f'''# Rebuild the renewable operating review

Use `{task}-inputs.xlsx` to reconstruct the operating review defined below. The `OperatingReview` sheet contains blank output cells. This is a constant-price operating analysis of the same projects and supplied contracts; use its explicit conventions for these new outputs. The original debt-sizing model and its covenant remain available as context, but the required results are the operating figures below.

## Input and scope

Use the existing input assumptions and schedules in the workbook. All monetary outputs are in thousands of euros (kEUR), except the breakeven tariff in EUR/MWh; generation is MWh. A MW times an equivalent full-load hour yields one MWh. Read blank annual fixed-cost inputs as no charge; do not substitute any value into a nonblank input. Keep full precision.

Use complete operating years counted separately from each project's commissioning: year 1 corresponds to the first row of each numbered annual input schedule, year 2 to its next row, and so on. Each year has first and second halves. Allocate annual fixed costs equally between halves; apply variable charges to the actual production in each half. Use the supplied first-half production share and its complement for the second half. This review uses the initial FiT as a constant price on all modeled MWh, before tariff indexation, production caps, aggregation premiums, merchant prices and tariff step-downs. It does not include taxes, interest, debt payments, capex, depreciation, working-capital movements or terminal value. "Operating contribution" means the defined FiT revenue less the specifically included cash operating costs below; it is not the debt model's CFADS.
'''
    if task=='0239':
        conventions='''
## Solar conventions

Model Roc Du Doun (column K, `Inputs - Solar Parks`) for the useful life in K18: 20 full operating years. The zero-capacity Le Mouruen column contributes nothing. Use the supplied capacity, P90 equivalent hours **before CPV sensitivity**, initial FiT and first-half production share. This operating review has no CPV sensitivity adjustment.

For each year, use that year's availability and degradation rows. Availability and degradation are separate effects. Begin year 1 first half with accumulated degradation factor 1. After each half-year's generation, divide the accumulated factor by `(1 + that year's annual degradation rate / 2)` for the next half. Carry it forward across year boundaries. This preserves the native model's reciprocal, lagged half-year degradation convention; do not replace it with an annual `(1 - degradation)` haircut. Half-year generation equals capacity times P90 hours times that year's availability times the accumulated factor at the start of the half times its seasonal share.

Included annual cash operating costs are the numbered **LTSA – Solar Park** and **LTSA – Inverters** schedules, plus AFSA, OMSA, Site Maintenance & Environmental Measures, Panels Cleaning, Enedis, Electric Consumption & Telecom, Insurance, Auditor Costs, LTSA Trackers CPV and Rent. Use the row for the relevant operating year for the two LTSA schedules; the other listed amounts repeat each year at their stated constant prices. Do not include the combined AFSA/OMSA subtotal a second time.

The uniform operating breakeven tariff equals cumulative included cash operating costs in euros divided by cumulative net generation over all 20 years. It covers these operating costs only, without a capital return or discounting.
'''
    else:
        conventions='''
## Wind conventions

Model all seven projects in columns K:Q of `Inputs - Wind Parks` for their first 10 full operating years. Reconstruct each project's total capacity from turbine capacity and turbine count. For the P90 case, multiply the supplied **EOH – P90 before portfolio effect** by `(1 + portfolio effect)`; for the P50 case, use the supplied P50 hours with no portfolio uplift. In both cases apply the year-specific technical availability schedule and the project-specific seasonal split to production. There is no additional production or availability sensitivity.

Included cash operating costs are: the year-specific Variable Costs rate per MWh; the year-specific Fixed Costs amount **per turbine** multiplied by turbine count; the year-specific Other Fixed Costs amount per project; and the annual AFSA, OMSA, Leases, Insurance, Auditor Costs and Others Costs inputs. The latter named annual amounts repeat at their stated constant prices. Do not include the combined AFSA/OMSA subtotal again. An empty AFSA or OMSA input means no separate charge for that project; preserve all supplied nonblank amounts. Use the row for the relevant operating year in each of the three numbered O&M cost schedules and in the availability schedule.

Aggregate the projects after computing each project's production, revenue and costs. "P50 less P90 operating contribution" is the difference between the two complete 10-year portfolio operating contributions under identical tariffs and fixed costs. The P90 uniform operating breakeven tariff is cumulative P90 included cash operating costs in euros divided by cumulative P90 portfolio production; it is one common tariff that exactly covers those costs, not an average of individual project breakeven tariffs. It includes no capital return or discounting.
'''
    lines=common+conventions+'\n## Research data service\n'+service+'\nThe covenant DSCR remains a legitimate research input and is not a scored output in this operating review. The operating assumptions required above are supplied on the named input sheet.\n\n## What to compute\n\n| Figure | Output cell |\n|---|---|\n'
    for row,label in enumerate(LABELS[task],5):lines+=f'| {label} | `OperatingReview!B{row}` |\n'
    return lines+'\n## Output\n\nWrite `/workspace/answers.json` as a JSON object mapping the six exact cell references above to numbers. Do not round or add other figures.\n'


def build():
    selected={x['taskId']:x for x in json.loads((AUDIT/'refactor/selected-assets.json').read_text())}
    manifest=[]
    for task in ['0239','0719']:
        a=selected[task];source=Path(a['workbookPath']);before=sha256(source);assert before==a['workbookSha256']
        sheet='Inputs - Solar Parks' if task=='0239' else 'Inputs - Wind Parks';v=Inputs(source,sheet)
        primary,rows=(decimal_solar(v) if task=='0239' else decimal_wind(v));secondary=independent_float(source,task)
        checks=[]
        for row,(answer,other) in enumerate(zip(primary,secondary),5):
            assert math.isclose(float(answer),other,rel_tol=1e-12,abs_tol=1e-8)
            checks.append({'cell':f'OperatingReview!B{row}','decimal_oracle':str(answer),'independent_float_oracle':other,'absolute_difference':abs(float(answer)-other),'matches':True})
        dest=OUT/task;dest.mkdir(parents=True,exist_ok=True);output=dest/source.name
        package=append_output_sheet(source,output,LABELS[task])
        targets={f'OperatingReview!B{r}':float(value) for r,value in enumerate(primary,5)}
        key={'kind':'cell_value','tolerance':{'numeric_abs':1e-6,'numeric_rel':1e-6},'targets':targets,'groups':{k:[k] for k in targets}}
        (dest/'answer_key.json').write_text(json.dumps(key,indent=2)+'\n')
        (dest/'instruction.md').write_text(instruction(task,Path(a['promptPath']).read_text()))
        (dest/'oracle-periods.json').write_text(json.dumps(rows,indent=2,default=str)+'\n')
        (dest/'oracle-inputs.json').write_text(json.dumps(v.used,indent=2)+'\n')
        numeric_matches=[]
        w=openpyxl.load_workbook(source,data_only=True,read_only=False)
        for s in w:
            for row in s:
                for cell in row:
                    if isinstance(cell.value,(float,int)) and not isinstance(cell.value,bool):
                        for ref,value in targets.items():
                            if math.isclose(cell.value,value,rel_tol=1e-6,abs_tol=1e-6):numeric_matches.append({'target':ref,'source_cell':f'{s.title}!{cell.coordinate}'})
        w.close();assert not numeric_matches,'Investigate retained target numeric matches'
        validation={'taskId':task,'status':'fixed','oracle_checks':checks,'oracle_methods':['Decimal 40-digit half-year recursion from source XML primitive cells','Independent floating-point annual closed-form aggregation from openpyxl input reads'],'primitive_inputs':len(v.used),'input_availability':'Every used primitive is present; specifically identified blank annual fixed costs mean no charge. No model-derived answer cache is used.','package_validation':package,'original_workbook_sha256_before':before,'original_workbook_sha256_after':sha256(source),'retained_new_target_numeric_matches':numeric_matches,'economic_inputs_changed':False,'DSCR_input_changed':False,'old_key_unchanged':sha256(Path(a['answerKeyPath']))==a['answerKeySha256'],'scope_change':'Replaces macro-dependent project debt-model lookup with explicitly scoped constant-price operating review using the same source contracts and operating schedules. New targets are not claimed to be the original macro-model outputs.','unresolved_solvability_defects':[]}
        assert validation['original_workbook_sha256_after']==before and validation['old_key_unchanged']
        (dest/'validation.json').write_text(json.dumps(validation,indent=2)+'\n')
        record={'taskId':task,'status':'fixed','workbookSourceSha256':before,'workbookPath':str(output),'workbookSha256':sha256(output),'promptPath':str(dest/'instruction.md'),'answerKeyPath':str(dest/'answer_key.json'),'additionalMasks':[],'validationPath':str(dest/'validation.json'),'beforeTargets':json.loads(Path(a['answerKeyPath']).read_text())['targets'],'afterTargets':targets,'blocker':None,'repair_kind':'explicit_operating_submodel_replacement'}
        manifest.append(record)
    (OUT/'renewable-ready-manifest.json').write_text(json.dumps(manifest,indent=2)+'\n')
    print(json.dumps({'tasks':len(manifest),'targets':sum(len(x['afterTargets']) for x in manifest),'manifest':str(OUT/'renewable-ready-manifest.json')}))


if __name__=='__main__':
    build()
