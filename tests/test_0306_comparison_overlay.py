from pathlib import Path
import importlib.util,sys
import pytest
ROOT=Path(__file__).resolve().parents[1]
PATH=ROOT/'diagnostics/authoring-queue-extension-20260911T185027Z/four-resume-diagnostics/0306-comparison-repair-v1/overlay/xl_seg/evaluate.py'
@pytest.fixture(scope='module')
def evaluator():
 spec=importlib.util.spec_from_file_location('xl_seg.comparison_candidate',PATH)
 module=importlib.util.module_from_spec(spec);sys.modules[spec.name]=module;spec.loader.exec_module(module);return module
@pytest.mark.parametrize('number,text',[(37500,'N/A'),(-10,'N/A'),(0,''),(1e200,'0'),(1,'1'),(-1,'-999'),(12,' 12 ')])
def test_ordered_number_text_operators_keep_types(evaluator,number,text):
 for op,want in [('>',False),('>=',False),('<',True),('<=',True)]:
  assert evaluator._binary(op,number,text) is want
  assert evaluator._binary(op,text,number) is not want

def test_arithmetic_errors_and_array_shape_preserved(evaluator):
 e=evaluator
 assert e._binary('+',2,'3')==5
 assert repr(e._binary('+',2,'N/A'))=='#VALUE!'
 error=e.ExcelError('#N/A');unknown=e.Unresolved('unproved')
 assert e._binary('>',error,'text') is error
 assert e._binary('>',4,unknown) is unknown
 a=e.RangeValues([37500,'N/A',-1],1,3)
 assert list(e._binary('>',a,0))==[True,True,False]
 assert repr(e._binary('>',a,e.RangeValues([1,2],1,2)))=='#VALUE!'
 assert e._binary('>',None,-1) is True
 assert e._binary('>',True,0) is True

def test_actual_staffing_formula_from_only_supplied_primitives(tmp_path,evaluator):
 import openpyxl
 from xl_ast_graph import AstGraph,write_csv
 from xl_seg.model import load
 from xl_seg.project import build
 book=openpyxl.Workbook();ws=book.active;ws.title='Staffing Matrix';monthly=book.create_sheet('Monthly Operating Build')
 for addr,val in {'G12':'N/A','H12':2,'I12':44562,'J12':'N/A','W12':0,'X8':43890}.items():ws[addr]=val
 monthly['N12']=37500
 ws['X12']="=IF($G12=0,0,IF(AND($H12=1,X$8>$I12),1,IF(W12>=1,1,IF('Monthly Operating Build'!N$12>$J12,1,0))))"
 path=tmp_path/'staffing.xlsx';book.save(path);ast=AstGraph(path,verbose=False).build();out=tmp_path/'ast';out.mkdir();write_csv(out,ast.nodes,ast.edges)
 graph=load(out,'staffing');cg=build(graph);inputs={n.id for n in graph.cells() if n.kind in ('input','label')}
 result=evaluator.Evaluator(graph,cg,strict_proof=True).run(inputs)
 assert result.values['Staffing Matrix!X12']==0
 assert not result.oracle_reads
