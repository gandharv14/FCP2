"""Explicit empty CSE follower tags retain the unique array anchor's formula."""
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as E
import pytest
from openpyxl import Workbook
from openpyxl.worksheet.formula import ArrayFormula
from xl_ast_graph import AstGraph,FormulaError,write_csv
from xl_seg.model import load
from xl_seg.project import build
from xl_seg.evaluate import Evaluator,CalculationMetadata

NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'

def fixture(tmp_path,change='empty_follower',ref='D1:E1'):
    path=tmp_path/'source.xlsx';book=Workbook();ws=book.active;ws.title="1 O'Brien"
    ws['A1']=2;ws['A2']=3;ws['D1']=ArrayFormula(ref=ref,text='=TRANSPOSE(A1:A2)');ws['E1']=3
    book.save(path)
    with ZipFile(path) as z:parts={p:z.read(p) for p in z.namelist()}
    part='xl/worksheets/sheet1.xml';root=E.fromstring(parts[part]);row=root.find(f'.//{{{NS}}}row')
    anchor=root.find(f'.//{{{NS}}}c[@r="D1"]');follower=root.find(f'.//{{{NS}}}c[@r="E1"]')
    anchor.find(f'{{{NS}}}v').text='2'
    if change!='value_only':E.SubElement(follower,f'{{{NS}}}f',{'ca':'1'})
    if change=='nonempty_follower':follower.find(f'{{{NS}}}f').text='99'
    if change=='matching_nonempty_follower':follower.find(f'{{{NS}}}f').text='TRANSPOSE(A1:A2)'
    if change=='shared_follower':follower.find(f'{{{NS}}}f').set('t','shared')
    if change=='whitespace_follower':follower.find(f'{{{NS}}}f').text=' '
    if change=='duplicate_formula':E.SubElement(follower,f'{{{NS}}}f').text='99'
    if change=='duplicate_anchor_formula':E.SubElement(anchor,f'{{{NS}}}f').text='99'
    if change=='duplicate_cell':
        from copy import deepcopy
        row.append(deepcopy(follower))
    if change=='overlap':
        f=follower.find(f'{{{NS}}}f');f.attrib.update(t='array',ref='E1:F1');f.text='TRANSPOSE(A1:A2)'
    if change=='ordinary_empty':anchor.remove(anchor.find(f'{{{NS}}}f'))
    conflicts={'follower_ref':('ref','D1:E1'),'follower_si':('si','0'),
               'follower_dt2D':('dt2D','1'),'follower_dtr':('dtr','1'),
               'follower_r1':('r1','A1'),'follower_bx':('bx','1'),
               'follower_aca':('aca','1'),'invalid_ca':('ca','yes')}
    if change in conflicts:
        key,value=conflicts[change];follower.find(f'{{{NS}}}f').set(key,value)
    if change.startswith('ca_'):
        follower.find(f'{{{NS}}}f').attrib.update(t='normal',ca=change[3:])
    if change=='missing_cache':
        anchor.remove(anchor.find(f'{{{NS}}}v'));follower.remove(follower.find(f'{{{NS}}}v'))
    if change=='altered_caches':
        anchor.find(f'{{{NS}}}v').text='900';follower.find(f'{{{NS}}}v').text='901'
    parts[part]=E.tostring(root)
    with ZipFile(path,'w') as z:
        for p,v in parts.items():z.writestr(p,v)
        if change=='duplicate_workbook_part':z.writestr('xl/workbook.xml',parts['xl/workbook.xml'])
        if change=='duplicate_relationship_part':z.writestr('xl/_rels/workbook.xml.rels',parts['xl/_rels/workbook.xml.rels'])
    return path

@pytest.mark.parametrize('change',['empty_follower','value_only','missing_cache','altered_caches','ca_0','ca_1','ca_true','ca_false'])
def test_valid_array_members_rebuild_from_formula_only(tmp_path,change):
    ast=AstGraph(fixture(tmp_path,change),verbose=False).build()
    assert ast.failed==0
    for coordinate,col in [('D1',0),('E1',1)]:
        node=ast.nodes["1 O'Brien!"+coordinate]
        assert node['kind']=='formula' and node['array_formula'] is True
        assert node['array_anchor']=="1 O'Brien!D1"
        assert node['array_ref']=='D1:E1' and node['array_row']==0 and node['array_col']==col
        assert node['formula']=='=TRANSPOSE(A1:A2)'
    out=tmp_path/'ast';out.mkdir();write_csv(out,ast.nodes,ast.edges)
    graph=load(out,'source');cg=build(graph)
    result=Evaluator(graph,cg,strict_proof=True,calculation=CalculationMetadata(available=True,iterate=False),proof_outputs={"1 O'Brien!D1","1 O'Brien!E1"}).run({"1 O'Brien!A1","1 O'Brien!A2"})
    assert result.values["1 O'Brien!D1"]==2
    assert result.values["1 O'Brien!E1"]==3

@pytest.mark.parametrize('change',['nonempty_follower','matching_nonempty_follower','shared_follower','whitespace_follower','duplicate_formula','duplicate_anchor_formula','duplicate_cell','overlap','follower_ref','follower_si','follower_dt2D','follower_dtr','follower_r1','follower_bx','follower_aca','invalid_ca','duplicate_workbook_part','duplicate_relationship_part'])
def test_conflicting_formula_evidence_rejected(tmp_path,change):
    with pytest.raises(FormulaError):AstGraph(fixture(tmp_path,change),verbose=False).build()

@pytest.mark.parametrize('ref',['C1:E1','D2:E2','E1:D1','D1:XFE1','D0:E1','D1:E1048577','D:E','bad'])
def test_invalid_or_unowned_anchor_bounds_rejected(tmp_path,ref):
    with pytest.raises((FormulaError,ValueError)):AstGraph(fixture(tmp_path,ref=ref),verbose=False).build()

def test_ordinary_empty_formula_still_parse_error(tmp_path):
    ast=AstGraph(fixture(tmp_path,'ordinary_empty'),verbose=False).build()
    assert ast.failed==1
    assert ast.nodes["1 O'Brien!E1"]['parse_status']=='error'
    assert ast.nodes["1 O'Brien!E1"]['array_formula'] is False

def test_absolute_anchor_span_preserves_member_coordinates(tmp_path):
    ast=AstGraph(fixture(tmp_path,ref='$D$1:$E$1'),verbose=False).build()
    assert ast.failed==0
    assert ast.nodes["1 O'Brien!E1"]['array_col']==1
    assert ast.nodes["1 O'Brien!E1"]['array_ref']=='$D$1:$E$1'
