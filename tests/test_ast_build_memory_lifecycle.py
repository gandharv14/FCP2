"""Bound temporary workbook/dedup retention without changing graph semantics."""
import gc
from pathlib import Path
import sys
import weakref

import openpyxl
from openpyxl.workbook.defined_name import DefinedName
import pytest

from xl_ast_graph import AstGraph,write_csv


def fixture_workbook(path):
    wb=openpyxl.Workbook();ws=wb.active;ws.title='3 Statements'
    ws['A1']='Revenue';ws['B1']=100;ws['B2']=.2
    wb.defined_names.add(DefinedName('Growth',attr_text="'3 Statements'!$B$2"))
    ws.defined_names.add(DefinedName('Base',attr_text="'3 Statements'!$B$1"))
    ws['C1']='=Base*(1+Growth)'
    ws['C2']='=SUM(B1:C1)'
    ws['D1']='=D2+1';ws['D2']='=D1-1'
    other=wb.create_sheet("Valuation's Results")
    other['A1']="='3 Statements'!C2*2"
    wb.save(path);wb.close()


def test_workbooks_are_reclaimed_before_formula_walk(tmp_path,monkeypatch):
    path=tmp_path/'source.xlsx';fixture_workbook(path)
    load=openpyxl.load_workbook;references=[]
    def tracked(*args,**kwargs):
        workbook=load(*args,**kwargs)
        references.append(weakref.ref(workbook))
        return workbook
    monkeypatch.setattr(openpyxl,'load_workbook',tracked)
    walk=AstGraph._walk
    def inspect_walk(self):
        assert '_wb' not in self.__dict__
        assert len(references)==2 and all(ref() is None for ref in references)
        assert self.names_global['GROWTH']=="'3 Statements'!$B$2"
        assert self.names_local[('3 Statements','BASE')]=="'3 Statements'!$B$1"
        return walk(self)
    monkeypatch.setattr(AstGraph,'_walk',inspect_walk)
    graph=AstGraph(path,verbose=False).build()
    assert graph.parsed==5 and graph.failed==0
    assert graph._edge_seen is None


@pytest.mark.parametrize('read_values',[False,True])
def test_csv_rows_cycles_and_names_identical_to_unreleased_build(tmp_path,read_values):
    path=tmp_path/'source.xlsx';fixture_workbook(path)
    old=AstGraph(path,verbose=False,read_values=read_values)
    # Former phase order: retain both temporary indexes through postprocessing.
    old._load();old._collect_names();old._walk();old._adopt_static_cells();old._postprocess()
    current=AstGraph(path,verbose=False,read_values=read_values).build()
    assert current.nodes==old.nodes and current.edges==old.edges
    assert current.cycles==old.cycles and current.cyclic_nodes==old.cyclic_nodes
    assert current.names_global==old.names_global and current.names_local==old.names_local
    assert current.stats(current.nodes,current.edges)==old.stats(old.nodes,old.edges)
    for name,graph in [('old',old),('current',current)]:
        directory=tmp_path/name;directory.mkdir()
        write_csv(directory,graph.nodes,graph.edges)
    for name in ('nodes.csv','edges.csv'):
        assert (tmp_path/'old'/name).read_bytes()==(tmp_path/'current'/name).read_bytes()
    old._release_workbook()


@pytest.mark.parametrize('phase',['load','collect_names'])
def test_temporary_workbook_closed_on_extraction_failure(monkeypatch,phase):
    closed=[];references=[]
    class Workbook:
        def __init__(self):self.cycle=self
        def close(self):closed.append(True)
    def load(self):
        self._wb=Workbook();references.append(weakref.ref(self._wb))
        if phase=='load':raise ValueError('invalid array evidence')
    def collect(self):raise ValueError('invalid name evidence')
    monkeypatch.setattr(AstGraph,'_load',load)
    monkeypatch.setattr(AstGraph,'_collect_names',collect)
    graph=AstGraph('unused.xlsx',verbose=False)
    with pytest.raises(ValueError):graph.build()
    gc.collect()  # Exception traceback frames may temporarily retain locals.
    assert closed==[True] and '_wb' not in graph.__dict__
    assert references[0]() is None


def test_edge_dedup_index_released_before_scc(monkeypatch):
    graph=AstGraph('synthetic.xlsx',verbose=False)
    retained_bytes=[]
    monkeypatch.setattr(graph,'_load',lambda:None)
    monkeypatch.setattr(graph,'_collect_names',lambda:None)
    def walk():
        for i in range(20_000):
            # A large construction-only tuple index is avoidable at SCC peak.
            graph._edge_seen.add(('source'+str(i),'target'+str(i),'value',i,''))
        retained_bytes.append(sys.getsizeof(graph._edge_seen)+sum(map(sys.getsizeof,graph._edge_seen)))
    monkeypatch.setattr(graph,'_walk',walk)
    monkeypatch.setattr(graph,'_adopt_static_cells',lambda:None)
    def postprocess():
        assert graph._edge_seen is None
        assert retained_bytes[0]>1_000_000
    monkeypatch.setattr(graph,'_postprocess',postprocess)
    graph.build()


def test_late_edge_addition_rebuilds_exact_dedup_guard(tmp_path):
    path=tmp_path/'source.xlsx';fixture_workbook(path)
    graph=AstGraph(path,verbose=False).build()
    old_count=len(graph.edges);edge=graph.edges[0]
    assert graph._edge_seen is None
    graph._add_edge(edge['source'],edge['target'],edge['role'],edge['arg_index'],edge['op'],
                    edge['ref'],edge['via_range'],edge['cell'],graph.nodes[edge['target']]['sheet'])
    assert len(graph.edges)==old_count
    assert graph._edge_seen is not None
    graph._add_edge(edge['source'],edge['target'],edge['role'],999,edge['op'],
                    edge['ref'],edge['via_range'],edge['cell'],graph.nodes[edge['target']]['sheet'])
    assert len(graph.edges)==old_count+1


def test_failed_walk_does_not_retain_construction_index(monkeypatch):
    graph=AstGraph('unused.xlsx',verbose=False)
    monkeypatch.setattr(graph,'_load',lambda:None)
    monkeypatch.setattr(graph,'_collect_names',lambda:None)
    def fail():
        graph._edge_seen.add(('a','b','value',1,''))
        raise ValueError('invalid source')
    monkeypatch.setattr(graph,'_walk',fail)
    with pytest.raises(ValueError):graph.build()
    assert graph._edge_seen is None
