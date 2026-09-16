"""Exact source-backed RiskMeasures-style same-row references; no broad table inference."""
from pathlib import Path
import openpyxl
from openpyxl.worksheet.table import Table
import pytest
from xl_ast_graph import AstGraph


def workbook(path):
    book=openpyxl.Workbook();sheet=book.active;sheet.title='Industry Averages'
    headers=['Industry name','Number of firms','Beta','Unlevered beta','Unlevered beta corrected for cash','D/E','E/V','D/V']
    for col,label in enumerate(headers,2):sheet.cell(5,col,label)
    for row,de in ((6,0.5),(7,2.0)):
        sheet.cell(row,2,'Industry '+str(row));sheet.cell(row,7,de)
        sheet.cell(row,8,'=1-RiskMeasures[[#This Row],[D/V]]')
        sheet.cell(row,9,'=RiskMeasures[[#This Row],[D/E]]/(1+RiskMeasures[[#This Row],[D/E]])')
    sheet.add_table(Table(displayName='RiskMeasures',ref='B5:I7'));book.save(path);book.close()


@pytest.mark.parametrize('storage',['memory','disk'])
def test_real_formula_forms_resolve_each_owner_row_and_preserve_source(tmp_path,storage):
    path=tmp_path/'source.xlsx';workbook(path);original=path.read_bytes()
    with AstGraph(path,verbose=False,graph_storage=storage,storage_dir=tmp_path)as graph:
        graph.build()
        assert graph.parsed==4 and graph.failed==0
        assert not any(n['kind']=='name'for n in graph.nodes.values())
        assert not graph.cycles
        for row in (6,7):
            targets={'Industry Averages!H'+str(row):'Industry Averages!I'+str(row),
                     'Industry Averages!I'+str(row):'Industry Averages!G'+str(row)}
            for owner,dependency in targets.items():
                refs={edge['source']for edge in graph.edges if edge['cell']==owner and graph.nodes[edge['source']]['kind'] in ('input','formula')}
                assert refs=={dependency}
    assert path.read_bytes()==original


@pytest.mark.parametrize('token,owner',[
    ('RiskMeasures[[#This Row],[D/V]]','Industry Averages!H5'),
    ('RiskMeasures[[#This Row],[D/V]]','Industry Averages!H8'),
    ('RiskMeasures[[#This Row],[D/V]]','Industry Averages!A6'),
    ('RiskMeasures[[#This Row],[D/V]]','Other!H6'),
    ('RiskMeasures[[#This Row],[D/V]]',None),
    ('RiskMeasures[[#This Row],[Missing]]','Industry Averages!H6'),
    ('OtherTable[[#This Row],[D/V]]','Industry Averages!H6'),
    ('RiskMeasures[[#All],[D/V]]','Industry Averages!H6'),
    ('RiskMeasures[@[D/V]]','Industry Averages!H6'),
    ('RiskMeasures[[#This Row],[D/E]:[D/V]]','Industry Averages!H6'),
])
def test_unsupported_or_foreign_table_context_stays_unresolved(tmp_path,token,owner):
    path=tmp_path/'source.xlsx';workbook(path)
    with AstGraph(path,verbose=False)as graph:
        graph.build();assert graph.resolve(token,'Industry Averages',current_cell=owner)==[('name',token)]


def test_duplicate_column_name_and_totals_row_do_not_infer_cells(tmp_path):
    path=tmp_path/'source.xlsx';workbook(path)
    with AstGraph(path,verbose=False)as graph:
        graph._load()
        table=graph._wb.active.tables['RiskMeasures'];table.totalsRowCount=1
        graph.tables.clear();graph._collect_tables(graph._wb.active)
        assert graph.resolve('RiskMeasures[[#This Row],[D/V]]','Industry Averages',current_cell='Industry Averages!H7')==[('name','RiskMeasures[[#This Row],[D/V]]')]
        table.tableColumns[-1].name=table.tableColumns[-2].name
        graph.tables.clear();graph._collect_tables(graph._wb.active)
        assert graph.resolve('RiskMeasures[[#This Row],[E/V]]','Industry Averages',current_cell='Industry Averages!H6')==[('name','RiskMeasures[[#This Row],[E/V]]')]


def test_source_health_uses_same_metadata_and_keeps_real_external_references(tmp_path):
    from xl_source_health import inspect_workbook, validate_report
    path=tmp_path/'source.xlsx';workbook(path);original=path.read_bytes()
    report=inspect_workbook(path)
    assert validate_report(report,path)==report
    events=report['restriction_events'];rows=[e for e in events if e['event']=='structured_reference']
    assert len(rows)==6 and all(e['allowed'] for e in rows)
    assert not any(e['event']=='external_workbook_reference' for e in events)
    assert path.read_bytes()==original
    book=openpyxl.load_workbook(path);book.active['J6']="='[Other.xlsx]Sheet1'!A1";book.save(path);book.close()
    report=inspect_workbook(path)
    assert 'external_links_present' in report['reason_codes']
    assert any(e['event']=='external_workbook_reference' and not e['allowed'] for e in report['restriction_events'])


@pytest.mark.parametrize('formula,cell',[
    ('=RiskMeasures[[#This Row],[D/V]]','A6'),
    ('=RiskMeasures[[#This Row],[D/V]]','H8'),
    ('=RiskMeasures[[#This Row],[Missing]]','H6'),
    ('=OtherTable[[#This Row],[D/V]]','H6'),
    ('=RiskMeasures[[#All],[D/V]]','H6'),
    ("='Foreign'!RiskMeasures[[#This Row],[D/V]]",'H6'),
])
def test_health_cannot_admit_unsupported_table_context(tmp_path,formula,cell):
    from xl_source_health import inspect_workbook
    path=tmp_path/'source.xlsx';workbook(path)
    book=openpyxl.load_workbook(path);book.active[cell]=formula;book.save(path);book.close()
    report=inspect_workbook(path)
    assert any(e['location']=='Industry Averages!'+cell and not e['allowed'] for e in report['restriction_events'])


@pytest.mark.parametrize('fault',[None,'source','offset','target','token','edge','formula','event_kind'])
def test_restriction_cone_requires_source_metadata_owner_and_actual_edge(tmp_path,fault):
    import hashlib,copy
    from xl_ast_graph import write_csv
    from xl_source_health import inspect_workbook
    from xl_seg import model,restriction_cone as R
    path=tmp_path/'source.xlsx';workbook(path);digest=hashlib.sha256(path.read_bytes()).hexdigest()
    with AstGraph(path,verbose=False)as built:
        built.build();write_csv(tmp_path,built.nodes,built.edges)
    with model.load(tmp_path,'source')as graph:
        events=inspect_workbook(path)['restriction_events'];event=copy.deepcopy(events[0])
        proof=R._SameRowTableProof(path,'0'*64 if fault=='source'else digest)
        if fault=='offset':event['token_offset']+=1
        if fault=='target':event['resolved_cell'][-1]+=1
        if fault=='token':event['token']='RiskMeasures[[#This Row],[D/E]]'
        if fault=='event_kind':event['reason']='unproven'
        if fault=='formula':graph.nodes[event['location']].formula='=1+1'
        if fault=='edge':graph.out_edges.clear()
        if fault:
            with pytest.raises(R.RestrictionConeError):R._match_event_nodes(graph,[event],proof)
        else:
            matches=R._match_event_nodes(graph,events,proof)
            assert len(matches)==6 and all(nodes[0].split('!')[0]=='Industry Averages' for nodes in matches)
            with pytest.raises(R.RestrictionConeError):R._match_event_nodes(graph,events)
