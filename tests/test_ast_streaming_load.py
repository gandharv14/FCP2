"""The streamed (read_only) workbook load builds the same AST as full mode.

Byte identity of nodes.csv and edges.csv is the contract: the builder's code
hash is part of every source generation's identity, so a load-path change is
only acceptable when the artifact bytes it produces are unchanged.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.formula import ArrayFormula

import xl_ast_graph
from xl_ast_graph import AstGraph, write_csv


def _rewrite_member(path: Path, member: str, transform) -> None:
    with zipfile.ZipFile(path) as src:
        parts = {item.filename: src.read(item.filename) for item in src.infolist()}
    parts[member] = transform(parts[member])
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for name, data in parts.items():
            out.writestr(name, data)


def _rich_workbook(path: Path) -> None:
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "Model"
    ws["A1"] = "Revenue"
    ws["B1"] = 100
    ws["B2"] = 0.1
    ws["B3"] = "=B1*(1+B2)"
    ws["B4"] = "=SUM(B1:B3)"
    ws["C1"] = ArrayFormula(ref="C1:D2", text="=A1:B2*10")
    ws["E1"] = True
    ws["E2"] = "=IF(E1,1,0)"
    ws["F1"] = "=NA()"
    ws["G1"] = "=Rates!A1*Named"
    ws["H1"] = "=Local"
    ws["A10"] = "merged anchor"
    ws["B10"] = "stale text in a merged member"
    ws.merge_cells("A10:B10")
    rates = book.create_sheet("Rates")
    rates["A1"] = 0.05
    rates["A2"] = "=A1*2"
    rates["B1"] = "=Model!B3+A2"
    book.defined_names["Named"] = DefinedName("Named", attr_text="Rates!$A$1")
    local = DefinedName("Local", attr_text="Model!$B$1")
    ws.defined_names["Local"] = local
    book.save(path)


def _rows_out_of_order_and_wrong_dimension(xml: bytes) -> bytes:
    # Shrink the declared dimension so a trusting reader stops early, and
    # swap two rows so a trusting reader drops the one that arrives late.
    xml, count = re.subn(
        rb'<dimension ref="[^"]*"\s*/>', b'<dimension ref="A1:B2"/>', xml
    )
    assert count == 1
    rows = re.findall(rb"<row .*?</row>", xml, flags=re.S)
    assert len(rows) >= 3
    swapped = list(rows)
    swapped[1], swapped[2] = swapped[2], swapped[1]
    order = iter(swapped)
    return re.sub(rb"<row .*?</row>", lambda _m: next(order), xml, flags=re.S)


def _shared_formula(xml: bytes) -> bytes:
    # Excel writes runs of identical relative formulas as one master with
    # t="shared" and children that only carry si=. openpyxl never writes
    # these itself, so inject one by hand.
    master = (
        b'<c r="A20" t="str"><f t="shared" ref="A20:A22" si="0">B1*2</f><v>200</v></c>'
        b'<c r="A21"><f t="shared" si="0"/><v>0.2</v></c>'
        b'<c r="A22"><f t="shared" si="0"/><v>220</v></c>'
    )
    return xml.replace(
        b"</sheetData>",
        b'<row r="20">' + master[:master.find(b'<c r="A21"')] + b"</row>"
        b'<row r="21">' + master[master.find(b'<c r="A21"'):master.find(b'<c r="A22"')] + b"</row>"
        b'<row r="22">' + master[master.find(b'<c r="A22"'):] + b"</row>"
        b"</sheetData>",
        1,
    )


@pytest.fixture
def full_mode(monkeypatch: pytest.MonkeyPatch):
    """Force the builder back onto openpyxl's full load for a reference run."""

    real = openpyxl.load_workbook

    def full(path, **kwargs):
        kwargs.pop("read_only", None)
        return real(path, **kwargs)

    monkeypatch.setattr(xl_ast_graph.openpyxl, "load_workbook", full)
    yield
    monkeypatch.setattr(xl_ast_graph.openpyxl, "load_workbook", real)


def _csv_bytes(path: Path, out: Path) -> tuple[bytes, bytes, AstGraph]:
    built = AstGraph(path, verbose=False).build()
    out.mkdir(parents=True, exist_ok=True)
    write_csv(out, built.nodes, built.edges)
    return (out / "nodes.csv").read_bytes(), (out / "edges.csv").read_bytes(), built


def _build_both(tmp_path: Path, monkeypatch: pytest.MonkeyPatch, workbook: Path):
    streamed = _csv_bytes(workbook, tmp_path / "streamed")
    real = openpyxl.load_workbook

    def full(path, **kwargs):
        kwargs.pop("read_only", None)
        return real(path, **kwargs)

    monkeypatch.setattr(xl_ast_graph.openpyxl, "load_workbook", full)
    try:
        reference = _csv_bytes(workbook, tmp_path / "full")
    finally:
        monkeypatch.setattr(xl_ast_graph.openpyxl, "load_workbook", real)
    return streamed, reference


def test_streamed_load_is_byte_identical_to_full_mode(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    workbook = tmp_path / "rich.xlsx"
    _rich_workbook(workbook)
    _rewrite_member(workbook, "xl/worksheets/sheet1.xml", _shared_formula)

    (nodes, edges, built), (ref_nodes, ref_edges, ref_built) = _build_both(
        tmp_path, monkeypatch, workbook
    )

    assert nodes == ref_nodes
    assert edges == ref_edges
    assert built.sheet_names == ref_built.sheet_names
    assert built.names_global == ref_built.names_global
    assert built.names_local == ref_built.names_local
    assert built.formulas == ref_built.formulas
    assert built.values == ref_built.values
    # The shared-formula children were translated, not left as bare si=.
    assert built.formulas["Model"][(21, 1)]["text"] == "=B2*2"
    assert built.formulas["Model"][(22, 1)]["text"] == "=B3*2"
    # The array formula covers its whole ref on both paths.
    assert built.formulas["Model"][(2, 4)]["array_formula"] is True
    # The stale value in the merged member is blanked on both paths.
    assert (10, 2) not in built.values["Model"]
    assert (10, 2) not in ref_built.values["Model"]
    assert built.values["Model"][(10, 1)] == "merged anchor"
    # Nothing is left holding the archive open.
    assert built._wb is None


def test_streamed_load_ignores_a_wrong_dimension_and_row_order(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # openpyxl's read-only iter_rows() trusts <dimension> and row order; with
    # this fixture it silently drops most of the sheet. The builder must not.
    workbook = tmp_path / "wrong_dimension.xlsx"
    _rich_workbook(workbook)
    _rewrite_member(
        workbook, "xl/worksheets/sheet1.xml", _rows_out_of_order_and_wrong_dimension
    )

    trusting = openpyxl.load_workbook(workbook, read_only=True, data_only=False)
    trusting_cells = sum(
        1
        for row in trusting["Model"].iter_rows()
        for cell in row
        if cell.value is not None
    )
    streamed_cells = sum(
        1 for c in xl_ast_graph.iter_cells(trusting["Model"]) if c.value is not None
    )
    trusting.close()

    (nodes, edges, built), (ref_nodes, ref_edges, ref_built) = _build_both(
        tmp_path, monkeypatch, workbook
    )

    assert nodes == ref_nodes
    assert edges == ref_edges
    assert built.formulas == ref_built.formulas
    assert built.values == ref_built.values
    # The trusting iterator lost cells to the bad dimension; the builder's
    # streamer saw them all.
    assert trusting_cells < streamed_cells
    assert (4, 2) in built.formulas["Model"]
    assert (10, 1) in built.values["Model"]


def test_iter_cells_streams_read_only_sheets_and_keeps_full_mode_path(
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "plain.xlsx"
    _rich_workbook(workbook)

    full = openpyxl.load_workbook(workbook)
    streamed = openpyxl.load_workbook(workbook, read_only=True)
    try:
        full_cells = {
            (c.row, c.column, c.value, c.data_type)
            for c in xl_ast_graph.iter_cells(full["Model"])
            if c.value is not None
        }
        streamed_cells = {
            (c.row, c.column, c.value, c.data_type)
            for c in xl_ast_graph.iter_cells(streamed["Model"])
            if c.value is not None
        }
    finally:
        streamed.close()
    # ArrayFormula compares by identity, so compare its text.
    normalise = lambda cells: {  # noqa: E731
        (r, c, getattr(v, "text", v), t) for r, c, v, t in cells
    }
    assert normalise(streamed_cells) == normalise(full_cells)
