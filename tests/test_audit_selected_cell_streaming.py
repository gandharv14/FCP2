"""Small typed fixtures; no private corpus, service calls or historical imports."""
from collections import Counter
from contextlib import contextmanager
import datetime as dt
import json
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
import xml.etree.ElementTree as ET
import pytest
import openpyxl
import xl_variable_source_audit as A

NS='http://schemas.openxmlformats.org/spreadsheetml/2006/main'

def rewrite(path, transform):
    with ZipFile(path) as z:files={name:z.read(name) for name in z.namelist()}
    transform(files)
    with ZipFile(path,'w',ZIP_DEFLATED) as z:
        for name,data in files.items():z.writestr(name,data)


@pytest.fixture
def workbook(tmp_path):
    path=tmp_path/'9472-inputs.xlsx';b=openpyxl.Workbook();s=b.active;s.title='2024 Inputs'
    s.append([1,1.5,True,'label',None,'#DIV/0!'])
    s['A2']=dt.datetime(2024,2,29,12,30);s['B2']=dt.time(12,30);s['C2']=dt.timedelta(days=2,hours=3)
    s['D2']='=1+1';s['E2']='001';s['F2']='line1\nline2';s['XFD10']=42
    s['A3']='=2+2';s['B3']='shared fixture'
    for name in ["O'Brien", "'quoted'", '3! Cases', 'Plain', 'Unused']:
        b.create_sheet(name)['A1']=name
    b['Plain']['B1']=False;b['Plain'].sheet_state='hidden';b.properties.title='Synthetic typed inventory'
    b.save(path);b.close()
    def cached(files):
        xml=ET.fromstring(files['xl/worksheets/sheet1.xml'])
        for c in xml.findall('.//{%s}c'%NS):
            if c.get('r')=='D2':c.find('{%s}v'%NS).text='2'
            if c.get('r')=='A3':c.find('{%s}v'%NS).text=None
            if c.get('r')=='B3':
                for ch in list(c):c.remove(ch)
                c.set('t','s');ET.SubElement(c,'{%s}v'%NS).text='0'
        files['xl/worksheets/sheet1.xml']=ET.tostring(xml)
        files['xl/sharedStrings.xml']=('<sst xmlns="'+NS+'" count="1" uniqueCount="1"><si><t>shared fixture</t></si></sst>').encode()
        ct=ET.fromstring(files['[Content_Types].xml']);ET.SubElement(ct,'{http://schemas.openxmlformats.org/package/2006/content-types}Override',PartName='/xl/sharedStrings.xml',ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml')
        files['[Content_Types].xml']=ET.tostring(ct)
    rewrite(path,cached);return path


def naive(book,refs):
    # Independent retained old lookup for unaffected canonical sheet identities.
    result={}
    for ref in refs:
        sheet,_,coord=ref.rpartition('!');sheet=sheet.strip("'")
        result[ref]=None if sheet not in book.sheetnames else book[sheet][coord].value
    return result


def test_exact_native_types_missing_cells_caches_and_sparse_column(workbook):
    refs=["'2024 Inputs'!"+x for x in ('A1','B1','C1','D1','E1','F1','A2','B2','C2','D2','E2','F2','A3','B3','XFD10','ZZ9','A1048576')]
    refs+=['Plain!$B$1','Missing!A1','not-a-reference',refs[0]]
    b=openpyxl.load_workbook(workbook,data_only=True,read_only=True)
    try:
        expected=naive(b,refs);actual=A._selected_cell_values(b,refs)
        assert {k:(type(v),v) for k,v in actual.items()}=={k:(type(v),v) for k,v in expected.items()}
        assert actual["'2024 Inputs'!D2"]==2 and actual["'2024 Inputs'!A3"] is None
        assert actual["'2024 Inputs'!B3"]=='shared fixture'
        assert actual["'2024 Inputs'!XFD10"]==42
    finally:b.close()


@pytest.mark.parametrize('sheet,ref',[("O'Brien","'O''Brien'!$A$1"),("'quoted'","'''quoted'''!a1"),('3! Cases',"'3! Cases'!A1"),('2024 Inputs',"'2024 Inputs'!$A1"),("O'Brien","O'Brien!A1")])
def test_machine_sheet_identity_and_absolute_references(workbook,sheet,ref):
    b=openpyxl.load_workbook(workbook,data_only=True,read_only=True)
    try:
        expected=b[sheet]['A1'].value
        assert A._cell_value(b,ref)==expected
        assert A._selected_cell_values(b,[ref])[ref]==expected
    finally:b.close()


@pytest.mark.parametrize('coord',['A0','A1048577','XFE1','A1:B2','A:A','1:2','A-1','A1!B1','$$A1','A$0'])
def test_invalid_coordinates_do_not_expand_sheet(workbook,monkeypatch,coord):
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
    b=openpyxl.load_workbook(workbook,data_only=True,read_only=True)
    try:
        monkeypatch.setattr(ReadOnlyWorksheet,'_get_source',lambda *a:pytest.fail('invalid references must not open XML'))
        ref='Plain!'+coord
        if '!' in coord:assert A._selected_cell_values(b,[ref])[ref] is None
        else:
            with pytest.raises(ValueError):A._selected_cell_values(b,[ref])
    finally:b.close()


def test_malformed_quoted_reference_is_not_silently_stripped(workbook):
    b=openpyxl.load_workbook(workbook,data_only=True,read_only=True)
    try:
        for ref in ("'O'Brien'!A1","'Plain!A1","''Plain''!A1"):
            with pytest.raises(ValueError):A._selected_cell_values(b,[ref])
    finally:b.close()


def test_one_xml_pass_per_requested_sheet_and_no_dense_expansion(workbook,monkeypatch):
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
    b=openpyxl.load_workbook(workbook,data_only=True,read_only=True);opened=Counter();real=ReadOnlyWorksheet._get_source
    @contextmanager
    def source(ws):
        opened[ws.title]+=1
        with real(ws) as stream:yield stream
    try:
        monkeypatch.setattr(ReadOnlyWorksheet,'_get_source',source)
        monkeypatch.setattr(ReadOnlyWorksheet,'_get_row',lambda *a,**k:pytest.fail('no dense rows'))
        refs=["'2024 Inputs'!A%d"%i for i in range(1,2001)]+["'2024 Inputs'!XFD10"]*100+['Plain!A1']*100
        actual=A._selected_cell_values(b,refs)
        assert opened=={'2024 Inputs':1,'Plain':1}
        assert actual["'2024 Inputs'!XFD10"]==42 and actual["'2024 Inputs'!A2000"] is None
    finally:b.close()


def inventory_fixture(tmp_path):
    seg=tmp_path/'segments';seg.mkdir();rows=[]
    for n in range(1,21):
        rows.append(dict(band="'2024 Inputs'!A%d:F%d"%(n,n),sheet='2024 Inputs',label='Input %02d'%n,kind='input',vtype='mixed',cells=["'2024 Inputs'!%s%d"%(c,n) for c in 'ABCDEF']))
    rows+=[dict(band='skip-label',kind='label',cells=['Plain!A1']),dict(band=rows[0]['band'],kind='input',cells=['Plain!A1']),dict(band='empty-input',kind='input'),dict(band='empty-literal',kind='literal',label='Hardcoded parameter')]
    (seg/'segments.json').write_text(json.dumps(dict(inputs=rows,embedded_literals=[dict(band='Plain!B1',kind='input',cells=['Plain!$B$1'])])))
    return seg


def test_complete_inventory_json_and_chunk_parity(workbook,tmp_path,monkeypatch):
    seg=inventory_fixture(tmp_path);selection={'generation_id':'a'*64,'manifest_sha256':'b'*64}
    current=A.build_inventory('9472',workbook,seg,segmentation=selection)
    monkeypatch.setattr(A,'_selected_cell_values',naive)
    expected=A.build_inventory('9472',workbook,seg,segmentation=selection)
    assert A._canonical_hash(current)==A._canonical_hash(expected)
    assert json.dumps(current,ensure_ascii=False,sort_keys=True)==json.dumps(expected,ensure_ascii=False,sort_keys=True)
    assert A._split_chunks(current)==A._split_chunks(expected)
    assert len(current['variables'])==22
    assert current['metadata']['sheet_names']==['2024 Inputs',"O'Brien","'quoted'",'3! Cases','Plain','Unused']


def test_wrong_declared_dimensions_do_not_hide_selected_values(workbook):
    def changed(files):
        xml=ET.fromstring(files['xl/worksheets/sheet1.xml']);xml.find('{%s}dimension'%NS).set('ref','A1:A1')
        files['xl/worksheets/sheet1.xml']=ET.tostring(xml)
    rewrite(workbook,changed);b=openpyxl.load_workbook(workbook,data_only=True,read_only=True)
    try:assert A._selected_cell_values(b,["'2024 Inputs'!XFD10"])["'2024 Inputs'!XFD10"]==42
    finally:b.close()


def test_malformed_tail_and_duplicate_rows_reject(workbook):
    def changed(files):
        xml=ET.fromstring(files['xl/worksheets/sheet1.xml']);data=xml.find('{%s}sheetData'%NS)
        data.append(ET.fromstring(ET.tostring(data[0])));files['xl/worksheets/sheet1.xml']=ET.tostring(xml)
    rewrite(workbook,changed);b=openpyxl.load_workbook(workbook,data_only=True,read_only=True)
    try:
        with pytest.raises(ValueError,match='unordered'):A._selected_cell_values(b,["'2024 Inputs'!A1"])
    finally:b.close()


def test_escaped_sheet_correction_changes_inventory_and_forbids_old_equivalence(workbook,tmp_path,monkeypatch):
    seg=tmp_path/'quoted-seg';seg.mkdir()
    (seg/'segments.json').write_text(json.dumps({'inputs':[dict(band="'O''Brien'!A1",sheet="O'Brien",label='Exact source label',kind='input',cells=["'O''Brien'!A1"])]}))
    actual=A.build_inventory('9472',workbook,seg)
    monkeypatch.setattr(A,'_selected_cell_values',naive)
    old=A.build_inventory('9472',workbook,seg)
    assert actual['variables'][0]['value_summary']=={'count':1,'value':"O'Brien"}
    assert old['variables'][0]['value_summary']=={'count':0}
    assert A._canonical_hash(actual)!=A._canonical_hash(old)


def test_malformed_unselected_tail_rejects_and_closes_stream(workbook,monkeypatch):
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
    def changed(files):files['xl/worksheets/sheet1.xml']+=b'<broken'
    rewrite(workbook,changed);b=openpyxl.load_workbook(workbook,data_only=True,read_only=True)
    streams=[];real=ReadOnlyWorksheet._get_source
    @contextmanager
    def source(ws):
        with real(ws) as stream:
            streams.append(stream);yield stream
    try:
        monkeypatch.setattr(ReadOnlyWorksheet,'_get_source',source)
        with pytest.raises(ET.ParseError):A._selected_cell_values(b,["'2024 Inputs'!A1"])
        assert streams and all(s.closed for s in streams)
    finally:b.close()
