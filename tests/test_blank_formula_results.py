"""Blank scalar formula results retain Excel reference semantics."""
import pytest
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula
from xl_ast_graph import AstGraph, write_csv
from xl_seg.model import load
from xl_seg import evaluate as candidate, project

def graph_for(tmp_path, formulas, inputs=None):
    book=openpyxl.Workbook();ws=book.active;ws.title='S'
    for coord,value in (inputs or {}).items():ws[coord]=value
    for coord,value in formulas.items():ws[coord]=value
    path=tmp_path/'synthetic.xlsx';book.save(path)
    built=AstGraph(path,verbose=False).build();out=tmp_path/'ast';out.mkdir()
    write_csv(out,built.nodes,built.edges)
    graph=load(out,'synthetic')
    # Poison every saved formula answer; strict reconstruction must ignore them.
    for n in graph.cells():
        if n.kind=='formula':n.value='999999';n.value_type='int'
    return graph


def run(module,graph,omit=()):
    ev=module.Evaluator(graph,project.build(graph),strict_proof=True,
        calculation=module.CalculationMetadata(available=True,iterate=False))
    result=ev.run({n.id for n in graph.cells() if n.kind in {'input','label'}}-set(omit))
    return ev,result




@pytest.mark.parametrize('formula',['=A1','=INDEX(A1:A2,1)'])
def test_blank_materializes_as_zero_without_target_cache(tmp_path,formula):
    graph=graph_for(tmp_path,{'B1':formula})
    ev,result=run(candidate,graph)
    assert result.values['S!B1']==0
    assert ev._read('S!A1') is None
    assert not result.unresolved


@pytest.mark.parametrize('formula',['=A1','=INDEX(A1:A2,1)'])
def test_missing_known_primitive_frontier_remains_unresolved(tmp_path,formula):
    graph=graph_for(tmp_path,{'B1':formula},{'A1':7})
    ev,result=run(candidate,graph,omit={'S!A1'})
    assert isinstance(result.values['S!B1'],candidate.Unresolved)


def test_blank_reference_context_and_range_aggregations_are_preserved(tmp_path):
    graph=graph_for(tmp_path,{'B1':'=ISBLANK(INDEX(A1:A3,1))',
        'B2':'=SUM(A1:A3)','B3':'=COUNT(A1:A3)','B4':'=AVERAGE(A1:A3)'}, {'A2':8})
    ev,result=run(candidate,graph)
    assert [result.values[f'S!B{i}'] for i in range(1,5)]==[True,8,1,8]
    assert ev._read('S!A1') is None
    assert ev._read('S!A3') is None


def test_empty_text_and_errors_do_not_become_blank_zero(tmp_path):
    graph=graph_for(tmp_path,{'B1':'=""','B2':'=B1','B3':'=1/0','B4':'=INDEX(A1:A2,8)'})
    _,result=run(candidate,graph)
    assert result.values['S!B1']==result.values['S!B2']==''
    assert isinstance(result.values['S!B3'],candidate.ExcelError)
    assert result.values['S!B3'].code=='#DIV/0!'
    assert result.values['S!B4'].code=='#REF!'


def test_materialized_link_is_numeric_to_downstream_count(tmp_path):
    graph=graph_for(tmp_path,{'B1':'=A1','B2':'=COUNT(B1)','B3':'=ISBLANK(B1)'})
    _,result=run(candidate,graph)
    assert [result.values[f'S!B{i}'] for i in range(1,4)]==[0,1,False]


def test_array_and_intrinsic_array_evaluation_is_unchanged(tmp_path):
    graph=graph_for(tmp_path,{
        'D1':ArrayFormula(ref='D1:E2',text='=A1:B2*2'),
        'F1':'=SUMPRODUCT((A1:B2>2)*A1:B2)',
        'F2':ArrayFormula(ref='F2',text='=AVERAGE(A1:B2/2)'), 'F3':'=AVERAGE(A1:B2)'},
        {'A1':1,'B1':2,'A2':3,'B2':4})
    _,result=run(candidate,graph)
    assert [result.values[f'S!{c}'] for c in ('D1','E1','D2','E2')]==[2,4,6,8]
    assert result.values['S!F1']==7
    # Explicit single-cell array metadata retains array evaluation.
    assert result.values['S!F2']==1.25
    assert result.values['S!F3']==2.5


def test_index_helper_retains_scalar_and_array_reference_blanks():
    ev=candidate.Evaluator(None,None)
    vals=candidate.RangeValues([None,8],2,1,refs=['S!A1','S!A2'])
    assert ev._fn_index([vals,1]) is None
    column=ev._fn_index([vals,0,1])
    assert list(column)==[None,8]
    assert column.refs==('S!A1','S!A2')
