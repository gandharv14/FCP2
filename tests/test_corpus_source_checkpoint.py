"""Source checkpoints skip AST construction only after current production proofs."""
import copy
import json
from argparse import Namespace
from pathlib import Path
from unittest.mock import Mock
import zipfile

import pytest
from openpyxl import Workbook

from scripts import corpus_build as build
from scripts.corpus import digest, object_hash
from xl_source_recalc import prepare_source_generation
import xl_source_publication as publication


@pytest.fixture
def completed(tmp_path):
    raw = tmp_path / '0001.xlsx'
    wb = Workbook(); wb.calculation.fullCalcOnLoad = False; wb.calculation.calcMode = 'auto'; ws = wb.active; ws['A1'] = 3; ws['B1'] = '=A1*2'; wb.save(raw)
    with zipfile.ZipFile(raw) as z:
        parts = {name: z.read(name) for name in z.namelist()}
    parts['xl/worksheets/sheet1.xml'] = parts['xl/worksheets/sheet1.xml'].replace(b'<f>A1*2</f><v />', b'<f>A1*2</f><v>6</v>').replace(b'<f>A1*2</f><v></v>', b'<f>A1*2</f><v>6</v>')
    with zipfile.ZipFile(raw, 'w') as z:
        for name, value in parts.items(): z.writestr(name, value)
    args = Namespace(run_root=str(tmp_path / 'run'), source_inventory=None,
                     inventory_batch_id=None, trusted_public_key=str(tmp_path / 'key'))
    source = {'workbook_id': '0001', 'path': str(raw), 'sha256': digest(raw)}
    gen, manifest = prepare_source_generation(raw, '0001', Path(args.run_root) / 'sources/0001')
    previous = {'bindings': {'code_sha256': 'old-broad-fingerprint', 'source_sha256': source['sha256'],
                            'curation_sha256': None, 'inventory_sha256': None},
                'stage': 'segmentation', 'state': 'STOPPED', 'attempts': 1,
                'source_generation_id': manifest['generation_id'],
                'source_manifest_sha256': digest(gen / 'generation-manifest.json')}
    ledger = {'build_history': [previous], 'build': {'bindings': {**previous['bindings'], 'code_sha256': 'new-evaluator-or-docs'}, 'attempts': 2}}
    return args, source, ledger, gen


def reuse(fixture, **kwargs):
    args, source, ledger, _ = fixture
    return build.reuse_completed_source(args, source, ledger, source['path'], {}, kwargs)


def test_completed_native_generation_survives_unrelated_code_change(completed, monkeypatch):
    import xl_ast_graph
    ast = Mock(side_effect=AssertionError('AST must not be rebuilt'))
    monkeypatch.setattr(xl_ast_graph, 'main', ast)
    args, source, ledger, gen = completed
    original_history = copy.deepcopy(ledger['build_history'])
    original_files = {p.relative_to(gen): p.read_bytes() for p in gen.rglob('*') if p.is_file()}
    result = reuse(completed)
    assert result is not None and result[0] == gen
    assert result[2]['bindings']['original_sha256'] == source['sha256']
    assert result[2]['reused_from_attempt'] == 1
    assert ledger['build_history'] == original_history
    assert not ast.called
    assert {p.relative_to(gen): p.read_bytes() for p in gen.rglob('*') if p.is_file()} == original_files
    assert not (gen.parents[1] / 'current.json').exists()


def test_bound_source_checkpoint_revalidates_and_rejects_altered_binding(completed):
    first = reuse(completed)
    prev = completed[2]['build_history'][0]
    prev['source_checkpoint'] = first[2]
    prev['source_checkpoint_sha256'] = object_hash(first[2])
    assert reuse(completed) is not None
    prev['source_checkpoint']['bindings']['original_sha256'] = 'altered'
    assert reuse(completed) is None


@pytest.mark.parametrize('change', ['nodes', 'manifest', 'missing_manifest', 'prepared_source', 'inventory', 'curation', 'incomplete', 'unsafe_identity'])
def test_reuse_accepts_only_irrelevant_native_identity_changes(completed, tmp_path, change):
    args, source, ledger, gen = completed
    prev = ledger['build_history'][0]
    if change == 'nodes':
        node = gen / 'ast/0001/nodes.csv'; node.chmod(0o600); node.write_text('altered')
    elif change == 'manifest': prev['source_manifest_sha256'] = 'forged'
    elif change == 'missing_manifest': (gen / 'generation-manifest.json').unlink()
    elif change == 'prepared_source': source['path'] = str(tmp_path / 'changed.xlsx'); Path(source['path']).write_bytes(b'changed'); source['sha256'] = digest(source['path'])
    elif change == 'inventory':
        inv=tmp_path/'inventory.json';inv.write_text('{}');args.source_inventory=str(inv);args.inventory_batch_id='different'
    elif change == 'curation': ledger['build']['bindings']['curation_sha256'] = 'changed'
    elif change == 'incomplete': prev['stage'] = 'source'
    elif change == 'unsafe_identity': prev['source_generation_id'] = '../other'
    result = reuse(completed)
    if change in {'inventory', 'curation'}:
        assert result is not None
    else:
        assert result is None


def test_changed_current_builder_rejected_even_if_recorded_snapshot_is_intact(completed, monkeypatch, tmp_path):
    fake = tmp_path / 'xl_ast_graph.py'; fake.write_text('# different current builder\n')
    monkeypatch.setattr(publication, '_default_builder_paths', lambda: (fake,))
    assert reuse(completed) is None
    assert any('current production AST builder changed' in r['reason'] for r in completed[2]['build']['source_reuse_rejections'])


def test_current_health_change_rejected(completed, monkeypatch):
    import xl_source_health
    original = xl_source_health.inspect_workbook
    def current(*args, **kwargs):
        result = original(*args, **kwargs)
        if not kwargs: result = {**result, 'policy_version': 'changed-current-policy'}
        return result
    monkeypatch.setattr(xl_source_health, 'inspect_workbook', current)
    assert reuse(completed) is None


def test_changed_normalization_is_not_hidden_by_identical_candidate(completed, tmp_path):
    args, source, ledger, _ = completed
    receipt=tmp_path/'manifest.json';receipt.write_text('{}')
    ledger['build']['normalization']={'original_sha256':source['sha256'], 'candidate':source['path'],
        'candidate_sha256':source['sha256'], 'manifest':str(receipt), 'manifest_sha256':'old-proof'}
    with pytest.raises(ValueError, match='normalization evidence changed'):
        reuse(completed)


def test_current_trusted_documents_must_match_and_publication_gate_is_mandatory(completed, monkeypatch):
    # A superficially matching generation cannot waive production receipt checks.
    assert reuse(completed, request={'forged':'request'}, result={'forged':'result'}) is None
    gate = Mock(side_effect=publication.SourcePublicationError('forged trusted runner receipt'))
    monkeypatch.setattr(publication, 'publish_source_generation', gate)
    assert reuse(completed) is None
    assert gate.called
    assert 'forged trusted runner receipt' in completed[2]['build']['source_reuse_rejections'][-1]['reason']


def test_worker_reuses_source_but_always_invokes_new_segmentation(completed, monkeypatch):
    import xl_source_recalc, xl_segment, xl_preflight
    args, source, ledger, gen = completed
    args.recalc_root = None
    ledger.update(route='pass', state='READY')
    ledger['build'].update(curation_path=str(gen.parent/'absent-curation.toml'), state='RUNNING', stage='source')
    ledger['build']['bindings']['code_sha256']='current'
    path=Path(args.run_root)/'workbooks/0001/ledger.json';path.parent.mkdir(parents=True);path.write_text(json.dumps(ledger))
    monkeypatch.setattr(build,'code_hash',lambda:'current')
    constructor=Mock(side_effect=AssertionError('no fresh AST'))
    monkeypatch.setattr(xl_source_recalc,'prepare_source_generation',constructor)
    segment=Mock(side_effect=RuntimeError('fresh segmentation invoked'))
    monkeypatch.setattr(xl_segment,'segment',segment)
    with pytest.raises(RuntimeError,match='fresh segmentation invoked'): build.worker(args,source)
    assert segment.call_count==1 and constructor.call_count==0
    after=json.loads(path.read_text())
    assert after['build']['source_reused'] is True
    assert after['build']['source_checkpoint_sha256']==object_hash(after['build']['source_checkpoint'])
    assert after['build_history']==ledger['build_history']


def test_restricted_reuse_requires_current_approved_inventory_and_batch(completed, monkeypatch, tmp_path):
    import xl_inventory_approval as approval
    from xl_source_inventory import build_inventory_manifest
    args, source, ledger, _ = completed
    raw = tmp_path / 'restricted-original/0001.xlsx'; raw.parent.mkdir(); raw.write_bytes(Path(source['path']).read_bytes()); source['path']=str(raw)
    with zipfile.ZipFile(raw) as z: parts = {n:z.read(n) for n in z.namelist()}
    parts['xl/worksheets/sheet1.xml'] = parts['xl/worksheets/sheet1.xml'].replace(b'A1*2', b'OFFSET(A1,0,0)')
    with zipfile.ZipFile(raw,'w') as z:
        for n,v in parts.items(): z.writestr(n,v)
    source['sha256']=digest(raw)
    inventory=build_inventory_manifest(raw.parent,workbook_ids=['0001'])
    registry_core={'schema_version':'source-inventory-approval-registry/v1','approvals':[approval.approval_claims(inventory,batch_id='approved-batch')]}
    registry=tmp_path/'approval.json';registry.write_text(json.dumps({**registry_core,'registry_sha256':approval.object_hash(registry_core)}))
    monkeypatch.setattr(approval,'DEFAULT_REGISTRY',registry)
    inv=tmp_path/'inventory.json';inv.write_text(json.dumps(inventory));args.source_inventory=str(inv);args.inventory_batch_id='approved-batch'
    opts={'inventory':str(inv),'inventory_batch_id':'approved-batch'}
    gen,manifest=prepare_source_generation(raw,'0001',Path(args.run_root)/'sources/0001',**opts)
    prev=ledger['build_history'][0];prev.update(source_generation_id=manifest['generation_id'],source_manifest_sha256=digest(gen/'generation-manifest.json'))
    prev['bindings'].update(source_sha256=source['sha256'],inventory_sha256=digest(inv));ledger['build']['bindings'].update(prev['bindings'])
    assert build.reuse_completed_source(args,source,ledger,raw,opts,{}) is not None
    args.inventory_batch_id='unapproved-batch';opts['inventory_batch_id']='unapproved-batch'
    assert build.reuse_completed_source(args,source,ledger,raw,opts,{}) is None
    args.inventory_batch_id='approved-batch';opts['inventory_batch_id']='approved-batch'
    registry.write_text('{}')
    assert build.reuse_completed_source(args,source,ledger,raw,opts,{}) is None


def test_signed_recalculation_reuse_rechecks_signature_not_just_artifact_hashes(completed, monkeypatch, tmp_path):
    import base64
    import hashlib
    import shutil
    import subprocess
    import time
    import xl_source_recalc as recalc
    from xl_source_health import inspect_workbook
    args, source, ledger, _ = completed
    private=tmp_path/'private.pem';public=tmp_path/'public.pem'
    subprocess.run(['/usr/bin/openssl','genpkey','-algorithm','RSA','-pkeyopt','rsa_keygen_bits:2048','-out',str(private)],check=True,capture_output=True)
    public.write_bytes(subprocess.check_output(['/usr/bin/openssl','pkey','-in',str(private),'-pubout'],stderr=subprocess.DEVNULL));public.chmod(0o600)
    args.trusted_public_key=str(public)
    # Unit fixtures need no root ownership. Real cryptographic verification and
    # every receipt claim remain active; production does not receive this patch.
    verify=recalc.verify_signed_runner_receipt
    monkeypatch.setattr(recalc,'verify_signed_runner_receipt',lambda *a,**kw: verify(*a,require_root_owner=False,**kw))
    destination=tmp_path/'recalculated/0001.xlsx'
    request=recalc.create_recalc_request(source['path'],destination,allowed_root=destination.parent,engine_constraints={
        'required_engine':'excel-macos','allowed_engines':['excel-macos'],'permitted_versions':['fixture-v1'],
        'trusted_runner_public_key_sha256':digest(public)})
    canonical=lambda x:(json.dumps(x,sort_keys=True,separators=(',',':'))+'\n').encode()
    class Engine:
        name='excel-macos';version='fixture-v1';authoritative=True
        def execute(self, original, candidate, request_value):
            shutil.copyfile(original,candidate)
            payload={'request_sha256':request_value['request_sha256'],'source_sha256':request_value['source_sha256'],
                'output_sha256':digest(candidate),'engine':self.name,'engine_version':self.version,
                'calculation_complete':True,'isolation_enforced':True,'network_isolation_mechanism':'macos-pf-anchor',
                'isolation_controls':{x:True for x in ('dedicated_session','network_disabled','macros_disabled',
                    'add_ins_disabled','link_updates_disabled','prompts_suppressed')},
                'network_isolation_rules_sha256':recalc.PF_NETWORK_RULES_SHA256,'completed_at_ns':time.time_ns()}
            payload_path=tmp_path/'payload.json';payload_path.write_bytes(canonical(payload));sig=tmp_path/'signature.bin'
            subprocess.run(['/usr/bin/openssl','dgst','-sha256','-sign',str(private),'-out',str(sig),str(payload_path)],check=True,capture_output=True)
            self.evidence={'controls':{k:True for k in ('dedicated_session','network_disabled','macros_disabled','add_ins_disabled','link_updates_disabled','prompts_suppressed')},
                'runner_receipt':{'schema_version':'excel-runner-receipt/v1','signed_payload':payload,'signature_base64':base64.b64encode(sig.read_bytes()).decode()}}
    result=json.loads(json.dumps(recalc.execute_recalc(request,Engine(),allowed_root=destination.parent)))
    request=json.loads(json.dumps(request))
    options={'request':request,'result':result,'trusted_runner_public_key':public,'original_source_path':source['path']}
    gen,manifest=prepare_source_generation(destination,'0001',Path(args.run_root)/'sources/0001',**options)
    req=tmp_path/'request.json';res=tmp_path/'result.json';req.write_bytes(canonical(request));res.write_bytes(canonical(result))
    evidence={'request':str(req),'request_sha256':digest(req),'result':str(res),'result_sha256':digest(res)}
    prev=ledger['build_history'][0];prev.update(source_generation_id=manifest['generation_id'],source_manifest_sha256=digest(gen/'generation-manifest.json'),recalculation=evidence)
    ledger['build']['recalculation']=copy.deepcopy(evidence)
    assert build.reuse_completed_source(args,source,ledger,destination,{},options) is not None
    # Forge the receipt, then recompute ordinary document/manifest checksums.
    # Cryptographic revalidation must still reject it.
    forged=copy.deepcopy(result)
    forged['engine_evidence']['runner_receipt']['signature_base64']=base64.b64encode(b'forged signature').decode()
    forged['result_sha256']=hashlib.sha256(canonical({k:v for k,v in forged.items() if k!='result_sha256'})).hexdigest()
    result_file=gen/'result.json';result_file.chmod(0o600);result_file.write_bytes(canonical(forged))
    manifest['artifacts']['result.json'].update(sha256=digest(result_file),size_bytes=result_file.stat().st_size)
    manifest['bindings']['result_sha256']=digest(result_file)
    manifest_file=gen/'generation-manifest.json';manifest_file.chmod(0o600);manifest_file.write_bytes(canonical(manifest))
    prev['source_manifest_sha256']=digest(manifest_file);res.write_bytes(canonical(forged));evidence['result_sha256']=digest(res)
    ledger['build']['recalculation']=copy.deepcopy(evidence);options['result']=forged
    # This demonstrates why validating the generation's hashes alone is insufficient.
    publication.validate_source_generation(gen)
    assert build.reuse_completed_source(args,source,ledger,destination,{},options) is None
    assert 'signature' in ledger['build']['source_reuse_rejections'][-1]['reason']


@pytest.mark.parametrize('changed', ['inventory', 'request', 'result', 'trust_anchor'])
def test_worker_rejects_expected_evidence_drift_before_source_work(completed, monkeypatch, tmp_path, changed):
    import xl_source_recalc
    args, source, ledger, gen = completed
    args.recalc_root = None
    ledger.update(route='pass',state='READY')
    ledger['build'].update(curation_path=str(tmp_path/'absent-curation.toml'),state='RUNNING',stage='source')
    ledger['build']['bindings']['code_sha256']='current'
    if changed=='inventory':
        inv=tmp_path/'changed-inventory.json';inv.write_text('{}');args.source_inventory=str(inv);args.inventory_batch_id='batch'
    else:
        args.recalc_root=str(tmp_path/'recalc');job=Path(args.recalc_root)/'0001';job.mkdir(parents=True)
        (job/'request.json').write_text(json.dumps({'source_sha256':source['sha256']}));(job/'result.json').write_text('{}')
        key=Path(args.trusted_public_key);key.write_text('original key')
        ledger['failure']='trusted_recalculation_required'
        ledger['build']['bindings'].update(recalc_request_sha256=digest(job/'request.json'),recalc_result_sha256=digest(job/'result.json'),trusted_public_key_sha256=digest(key))
        target={'request':job/'request.json','result':job/'result.json','trust_anchor':key}[changed]
        if changed=='request':target.write_text(json.dumps({'source_sha256':source['sha256'],'altered':True}))
        else:target.write_text('{"altered":true}')
    path=Path(args.run_root)/'workbooks/0001/ledger.json';path.parent.mkdir(parents=True);path.write_text(json.dumps(ledger))
    monkeypatch.setattr(build,'code_hash',lambda:'current')
    constructor=Mock(side_effect=AssertionError('must reject before AST'))
    monkeypatch.setattr(xl_source_recalc,'prepare_source_generation',constructor)
    with pytest.raises(ValueError,match='changed after checkpoint binding'):build.worker(args,source)
    assert not constructor.called
    assert json.loads(path.read_text())['build_history']==ledger['build_history']


@pytest.mark.parametrize('case', ['same', 'other', 'missing', 'invalid'])
def test_required_source_generation_never_falls_back_to_ast(completed, monkeypatch, tmp_path, case):
    import xl_source_recalc
    import xl_segment
    args, source, ledger, gen = completed
    args.recalc_root = None
    args.require_source_generation = gen.name if case in {'same','missing'} else ('0'*64 if case=='other' else 'invalid')
    ledger.update(route='pass',state='READY')
    ledger['build'].update(curation_path=str(tmp_path/'absent.toml'),state='RUNNING',stage='source')
    ledger['build']['bindings']['code_sha256']='current'
    if case=='missing': ledger['build_history']=[]
    path=Path(args.run_root)/'workbooks/0001/ledger.json';path.parent.mkdir(parents=True);path.write_text(json.dumps(ledger))
    monkeypatch.setattr(build,'code_hash',lambda:'current')
    constructor=Mock(side_effect=AssertionError('required checkpoint must never rebuild AST'))
    monkeypatch.setattr(xl_source_recalc,'prepare_source_generation',constructor)
    segment=Mock(side_effect=RuntimeError('validated source reached segmentation'))
    monkeypatch.setattr(xl_segment,'segment',segment)
    expected=RuntimeError if case=='same' else ValueError
    message='validated source reached segmentation' if case=='same' else 'AST rebuild forbidden'
    with pytest.raises(expected,match=message):build.worker(args,source)
    assert not constructor.called
    after=json.loads(path.read_text())
    assert after['build_history']==ledger['build_history']
    assert after['build']['attempts']==2
    assert segment.call_count==(1 if case=='same' else 0)


def test_required_generation_reaches_guarded_child_and_consumes_next_attempt(completed, monkeypatch, tmp_path):
    args,source,ledger,gen=completed
    prior=copy.deepcopy(ledger['build_history'][0]);source['format']='.xlsx'
    old={'source':source,'state':'READY','route':'pass','build':prior,'build_history':[]}
    path=Path(args.run_root)/'workbooks/0001/ledger.json';path.parent.mkdir(parents=True);path.write_text(json.dumps(old))
    args.inventory=str(tmp_path/'inventory.json');args.curation_root=str(tmp_path/'curations');args.recalc_root=None
    args.code_root=str(tmp_path);args.timeout=900;args.memory_mb=8192;args.require_source_generation=gen.name
    observed=[]
    def stop(command,**kw):
        observed.append((command,kw));return 124
    monkeypatch.setattr(build,'_run_in_process_group',stop)
    result=build.execute_one(args,source,'new-code')
    assert result['attempts']==2 and result['state']=='STOPPED'
    after=json.loads(path.read_text());assert after['build_history']==[prior]
    command,kw=observed[0]
    assert command[command.index('--require-source-generation')+1]==gen.name
    assert kw['timeout_seconds']==900 and kw['memory_limit_bytes']==8*1024**3
