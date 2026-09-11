"""Real immutable source/segmentation fixtures and full dialogue checks."""
import copy
import json
import re
import shutil
import zipfile
from pathlib import Path
from types import SimpleNamespace
import pytest
import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'.cursor/skills/additional-assumptions-dialogue/scripts'))
from openpyxl import Workbook
import extract_claims as X
import validate_dialogue as V
from xl_source_recalc import prepare_source_generation
from xl_seg import bands, model, project
import xl_segment


def save(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value))


def make(tmp, literal='B3', formula=None, pointer=None):
    raw=tmp/'raw/0001.xlsx';raw.parent.mkdir(parents=True)
    wb=Workbook();ws=wb.active;ws.title='Ops'
    ws['A1']=literal;ws['A16']='Metric';ws['B16']=formula or f'=IF(A1="{literal}",2,3)'
    ws['A20']='Target';ws['B20']='=10+7'
    if pointer:ws['Z1']=pointer
    wb.calculation.calcMode='auto';wb.calculation.fullCalcOnLoad=False
    wb.save(raw);wb.close()
    with zipfile.ZipFile(raw) as z:parts={n:z.read(n) for n in z.namelist()}
    parts['xl/worksheets/sheet1.xml']=re.sub(rb'(<c[^>]* r="B16"[^>]*>.*?<f>.*?</f>)(?:<v></v>|<v />)',rb'\1<v>2</v>',parts['xl/worksheets/sheet1.xml'])
    parts['xl/worksheets/sheet1.xml']=re.sub(rb'(<c[^>]* r="B20"[^>]*>.*?<f>.*?</f>)(?:<v></v>|<v />)',rb'\1<v>17</v>',parts['xl/worksheets/sheet1.xml'])
    with zipfile.ZipFile(raw,'w') as z:
        for n,data in parts.items():z.writestr(n,data)
    gen,m=prepare_source_generation(raw,'0001',tmp/'sources/0001')
    ast=gen/m['layout']['ast_directory'];gold=gen/m['layout']['source_workbook']
    bg=bands.build(project.build(model.load(ast,'0001')))
    target=next(b for b in bg.bands.values() if 'Ops!B20' in b.cells)
    out=tmp/'seg';(out/'0001').mkdir(parents=True)
    (out/'0001/curation.toml').write_text('[[output]]\nband = '+json.dumps(target.id)+'\ninclude = true\n')
    xl_segment.segment('0001',SimpleNamespace(source_generation_path=str(gen),source_generation_id=None,ast_dir=str(ast.parent),source=str(raw.parent),out=str(out),threshold=6,top=40,sample=400,lineage_max=4000,recurate=False,llm=False,no_auto_adjudication=True,no_verify=False,allow_legacy_ast=False))
    sg,=list((out/'0001/generations').iterdir())
    stage=tmp/'0001-outputs';(stage/'tests').mkdir(parents=True)
    binding=dict(source_generation_id=gen.name,source_manifest_sha256=X._literal_digest(gen/'generation-manifest.json'),segmentation_generation_id=sg.name,segmentation_manifest_sha256=X._literal_digest(sg/'generation-manifest.json'))
    rec={'entry':'method_revenue','source':'custom_method_detector','value':'out_of_catalogue','cells':['Ops!B16'],'evidence':formula or f'=IF(A1="{literal}",2,3)',
         'fields':{'label':'"Metric"','band':'Ops!B16','representative':'Ops!B16','steps':'return the applicable category'}}
    d=dict(schema_version='2.0',task=stage.name,golden=str(gold),golden_sha256=X._literal_digest(gold),pipeline_bindings=binding,segmentation_mode='strict',segmentation_generation=dict(generation_id=sg.name,directory=str(sg),mode='strict'),selection={'closure_source':'ast'},ast_dir=str(ast.parent),seg_root=str(out),agent_records=[rec])
    save(stage/'tests/disclosure.json',d)
    save(stage/'tests/pipeline_bindings.json',binding)
    (stage/'task.toml').write_text('[metadata]\nworkbook="0001"\nsource_file="0001.xlsx"\npipeline_bindings="tests/pipeline_bindings.json"\nsource_generation_id="'+gen.name+'"\nsegmentation_generation_id="'+sg.name+'"\n')
    shutil.copyfile(sg/'generation-manifest.json',stage/'tests/segmentation_generation_manifest.json')
    save(stage/'tests/inputs_generation.json',dict(pipeline_bindings=binding))
    save(stage/'tests/answer_key.json',dict(kind='cell_value',tolerance={'numeric_abs':1e-6,'numeric_rel':1e-6},targets={'Ops!B20':17},groups={'target':['Ops!B20']}))
    (stage/'instruction.md').write_text('# Task\n## Workbook disclosure\nModel assumptions.\n')
    return stage,d,rec


@pytest.fixture
def task(tmp_path):
    stage,d,rec=make(tmp_path)
    payload=X.extract(stage)
    return stage,d,rec,payload


def prose(payload,tail=''):
    c=payload['claims'][0]
    return f'<!-- claim:{c["record_id"]} -->\n**Analyst:** For `Ops!B16`: Please explain the Ops tab.\n**VP:** For the row labelled "Metric", use the text literal "B3" as the comparison category, only for this cell. {tail}\n'


def test_exact_bound_comparison_label_passes_full_draft(task):
    stage,d,rec,payload=task
    result=V.check_draft(prose(payload),payload,stage)
    assert result['passed'], result
    proof=payload['claims'][0]['reviewer_only']['reference_literal_proof']
    assert proof['values']==['B3'] and proof['occurrence_limit']==1
    assert proof['source_sha256']==d['golden_sha256']
    assert 'reference_literal_proof' not in json.dumps(payload['writer_pack'])


@pytest.mark.parametrize('extra',[
    'The text literal "B3" also appears here.', 'Also use B3.',
    'Also use "B3".', 'Read Ops!B3.', 'Read B3:C3.',
    'Read $B$3.', 'Read \'Ops\'!B3.',
])
def test_extra_reference_occurrences_fail(task,extra):
    stage,d,rec,payload=task
    assert not V.check_draft(prose(payload,extra),payload,stage)['passed']


@pytest.mark.parametrize('change',['proof','claim','formula','source','curation','target','manifest','inputs','missing','wrong_source'])
def test_altered_owning_evidence_rejects(task,change,tmp_path):
    stage,d,rec,payload=task
    if change=='proof':payload['claims'][0]['reviewer_only']['reference_literal_proof']['occurrence_limit']=2
    elif change=='claim':payload['claims'][0]['row_label']='Other'
    elif change=='formula':payload['claims'][0]['reviewer_only']['evidence']='=IF(A1="B3",4,5)'
    elif change=='source':
        Path(d['golden']).chmod(0o644);Path(d['golden']).write_bytes(b'changed')
    elif change=='curation':(Path(d['seg_root'])/'0001/curation.toml').write_text('changed')
    elif change=='target':save(stage/'tests/answer_key.json',{'targets':{'Ops!B16':2}})
    elif change=='manifest':save(stage/'tests/segmentation_generation_manifest.json',{})
    elif change=='inputs':save(stage/'tests/inputs_generation.json',{})
    elif change=='missing':(stage/'tests/segmentation_generation_manifest.json').unlink()
    elif change=='wrong_source':
        alt=tmp_path/'foreign.xlsx';shutil.copyfile(d['golden'],alt);d['golden']=str(alt);save(stage/'tests/disclosure.json',d)
    assert not V.check_draft(prose(payload),payload,stage)['passed']


def test_other_claim_and_junior_cannot_borrow_occurrence(task):
    stage,d,rec,payload=task
    for added in ['<!-- claim:foreign -->\n**VP:** The text literal "B3".', '**Analyst:** The text literal "B3".']:
        assert not V.check_draft(prose(payload)+added,payload,stage)['passed']


def test_same_named_row_does_not_exempt_extra_occurrence(task):
    stage,d,rec,payload=task
    payload['claims'].append(dict(record_id='other',row_label='B3',sheet='Ops',must_say=[],spoken=''))
    assert not V.check_draft(prose(payload,'B3'),payload,stage)['passed']


def test_original_numeric_answer_scan_still_sees_prose(task):
    stage,d,rec,payload=task
    result=V.check_draft(prose(payload,'The final answer is 17.'),payload,stage)
    assert not result['passed'] and any('numeric' in x or 'target' in x for x in result['faults']), result


def test_missing_and_wrong_typed_literal_still_fail(task):
    stage,d,rec,payload=task
    for text in [prose(payload).replace('"B3"','"b3"'),prose(payload).replace('the text literal','the row labelled')]:
        assert not V.check_draft(text,payload,stage)['passed']


@pytest.mark.parametrize('literal',['B031','A1'])
def test_other_bare_comparison_labels_supported(tmp_path,literal):
    stage,d,rec=make(tmp_path,literal)
    payload=X.extract(stage)
    text=prose(payload).replace('"B3"',json.dumps(literal))
    result=V.check_draft(text,payload,stage)
    assert result['passed'], result


@pytest.mark.parametrize('literal',['Sheet!C5','$B$5','A:B','3:3','B3:C3'])
def test_ranges_sheets_absolute_and_axes_never_licensed(literal):
    with pytest.raises(ValueError,match='prohibited'):
        X.formula_text_atoms({'evidence':f'=IF(A1="{literal}",2,3)'})


@pytest.mark.parametrize('pointer', ['=INDIRECT("B3")','=INDIRECT(A2)','=ADDRESS(3,2)','=OFFSET(A1,1,1)','=HYPERLINK(A2)','=unknownUDF(A2)'])
def test_actual_source_pointer_and_cross_cell_pointer_rejected(tmp_path,pointer):
    from openpyxl import load_workbook
    path=tmp_path/'source.xlsx';book=Workbook();book.active['A1']='=IF(B1="B3",2,3)'
    book.active['A2']='B3';book.active['B2']=pointer;book.save(path);book.close()
    book=load_workbook(path,data_only=False)
    try:
        with pytest.raises(ValueError,match='reference-interpreting'):
            X.validate_reference_label_workbook(book)
    finally:book.close()


@pytest.mark.parametrize('local',[False,True])
def test_global_and_sheet_local_pointer_names_rejected(tmp_path,local):
    from openpyxl import load_workbook
    from openpyxl.workbook.defined_name import DefinedName
    path=tmp_path/'source.xlsx';book=Workbook();book.active['A1']='=IF(B1="B3",2,3)'
    owner=book.active if local else book
    owner.defined_names.add(DefinedName('Pointer',attr_text='INDIRECT("B3")',localSheetId=0 if local else None))
    book.save(path);book.close();book=load_workbook(path,data_only=False)
    try:
        with pytest.raises(ValueError,match='defined name'):
            X.validate_reference_label_workbook(book)
    finally:book.close()


def test_bare_literal_return_not_comparison_is_unproved(tmp_path):
    stage,d,rec=make(tmp_path,formula='=IF(A1="B3",2,3)&"B3"')
    # Source output selection remains independent and passes; literal recovery
    # must still refuse because the same string is also returned as a value.
    with pytest.raises(ValueError,match='solely an equality'):
        X.extract(stage)


def test_duplicate_encoded_literal_and_grouped_claim_cannot_borrow_proof(task):
    stage,d,rec,payload=task
    extra=r'The text literal "\u00423" is also used.'
    assert not V.check_draft(prose(payload,extra),payload,stage)['passed']
    second=dict(record_id='other',row_label='Metric',sheet='Ops',must_say=[],spoken='')
    payload['claims'].append(second)
    text=prose(payload).replace('**Analyst:**','<!-- claim:other -->\n**Analyst:**')
    assert not V.check_draft(text,payload,stage)['passed']


def test_no_unbound_proof_in_ordinary_card(task):
    stage,d,rec,payload=task
    c=payload['claims'][0]
    c['must_say']=[a for a in c['must_say'] if not a.startswith('the text literal')]
    assert not V.check_draft(prose(payload),payload,stage)['passed']


def test_unchanged_proof_can_revalidate_but_cannot_be_renominated(task):
    stage,d,rec,payload=task
    expected=X.reference_literal_proofs(stage,payload['claims'])
    assert expected[c:=payload['claims'][0]['record_id']]==payload['claims'][0]['reviewer_only']['reference_literal_proof']
    d['agent_records'][0]['evidence']='=IF(A1="B3",999,111)';save(stage/'tests/disclosure.json',d)
    changed=X.build_claim(d['agent_records'][0],1)
    with pytest.raises(ValueError,match='original source cell'):
        X.reference_literal_proofs(stage,[changed])


def test_complete_foreign_pair_substitution_cannot_replace_package_identity(task,tmp_path):
    stage,d,rec,payload=task
    foreign,fd,frec=make(tmp_path/'foreign','A1')
    # All claim-owned and disclosure-nominated pair locators replaced together;
    # the independent packaged task metadata/bindings remain the authoring root.
    fd['task']=stage.name
    save(stage/'tests/disclosure.json',fd)
    for name in ('segmentation_generation_manifest.json','inputs_generation.json'):
        shutil.copyfile(foreign/'tests'/name,stage/'tests'/name)
    card=X.build_claim(frec,1)
    with pytest.raises(ValueError,match='packaged task identity'):
        X.reference_literal_proofs(stage,[card])


def scoped_card(cells,index=1):
    return X.build_claim(dict(entry='method_revenue',value='out_of_catalogue',cells=cells,
        evidence='=IF(A1="No",2,3)',fields={'label':'"Metric"','band':cells[0] if cells else 'Ops!B16','calculation_kind':'single-cell calculation','steps':'return the applicable category','representative':cells[0] if cells else 'Ops!B16'}),index)


@pytest.mark.parametrize('cells',[
    ['Ops!B16'], ['Ops!C16','Ops!D16','Ops!E16'], ['Ops!C16','Ops!F16'],
    ["'O''Brien! $ Plan'!$B$16"],
])
def test_junior_scope_preserves_exact_disclosed_cells_without_fill(cells):
    from aa_lib import junior_scope_text
    c=scoped_card(cells)
    assert c['junior_scope']==cells
    assert X.writer_claim(c)['junior_scope']==cells
    skeleton=X.draft_skeleton([c],['Analyst'],['VP'])
    assert junior_scope_text(c) in skeleton.split('**Analyst:**')[1].split('**VP:**')[0]
    assert not any(ref in skeleton.split('**VP:**')[1] for ref in cells)
    if cells==['Ops!C16','Ops!F16']:
        assert 'D16' not in skeleton and 'E16' not in skeleton and 'C16:F16' not in skeleton


def test_same_row_claims_and_certified_group_keep_distinct_scope_order():
    from aa_lib import junior_scope_text
    cards=[scoped_card(['Ops!B16'],1),scoped_card(['Ops!C16','Ops!D16'],2)]
    shared={'groups':[[c['record_id'] for c in cards]]}
    skeleton=X.draft_skeleton(cards,['Analyst'],['VP'],shared)
    assert skeleton.index(cards[0]['record_id'])<skeleton.index(cards[1]['record_id'])
    assert skeleton.index(junior_scope_text(cards[0]))<skeleton.index(junior_scope_text(cards[1]))
    assert skeleton.count('**Analyst:**')==1


def test_representative_alone_does_not_invent_junior_scope():
    c=scoped_card([])
    assert c['junior_scope']==[]
    assert 'B16' not in X.draft_skeleton([c],['Analyst'],['VP'])


def test_full_draft_requires_junior_scope_and_still_rejects_senior_refs(task):
    from aa_lib import junior_scope_text
    stage,_,_,_=task;c=scoped_card(['Ops!B16']);payload={'claims':[c]}
    text=f'<!-- claim:{c["record_id"]} -->\n**Analyst:** {junior_scope_text(c)}: Please explain the Ops tab.\n**VP:** For the row labelled "Metric", use the text literal "No", only for this cell.\n'
    assert V.check_draft(text,payload,stage)['passed']
    for altered in [text.replace('For `Ops!B16`','For `Ops!C16`'), text.replace('For `Ops!B16`: ',''),text+'**VP:** Use Ops!B16.']:
        assert not V.check_draft(altered,payload,stage)['passed']


def test_literal_matching_real_junior_cell_scope_is_not_an_extra_literal(tmp_path):
    # The junior's exact immutable scope is a real allowed reference; only the
    # senior's typed comparison occurrence is licensed as a financial label.
    stage,d,rec=make(tmp_path,'B16');payload=X.extract(stage)
    result=V.check_draft(prose(payload).replace('"B3"','"B16"'),payload,stage)
    assert result['passed'], result


def test_actual_compose_assemble_and_full_check_preserve_frozen_junior_scope(task):
    from compose_draft import compose,assemble_slots,strip_slot_scaffold
    stage,_,_,payload=task
    template,slots=compose(payload,payload['writer_pack'])
    assert '**Analyst:** Ops: For `Ops!B16`:' in template
    answers={'schema_version':'dialogue-slot-prose/v1','slots':{
        'c001-junior':'How should the Metric row be built?',
        'c001-senior':'For the row labelled "Metric", use the text literal "B3" as the comparison category, only for this cell.'}}
    filled=assemble_slots(template,slots,answers)
    result=V.check_draft(strip_slot_scaffold(filled),payload,stage)
    assert result['passed'],result
    changed=filled.replace('For `Ops!B16`:','For `Ops!C16`:')
    assert V.fill_faults(template,changed)
    assert not V.check_draft(strip_slot_scaffold(changed),payload,stage)['passed']
    answers['slots']['c001-junior']='Use the text literal "B3".'
    filled=assemble_slots(template,slots,answers)
    assert not V.check_draft(strip_slot_scaffold(filled),payload,stage)['passed']


def test_actual_composer_rejects_pack_scope_drift(task):
    from compose_draft import compose
    _,_,_,payload=task
    wrong=copy.deepcopy(payload['writer_pack']);wrong['claims'][0]['junior_scope']=['Ops!C16']
    with pytest.raises(ValueError,match='scope'):compose(payload,wrong)


def test_actual_composer_preserves_certified_member_scopes_in_order():
    from compose_draft import compose
    from aa_lib import junior_scope_text
    cards=[scoped_card(['Ops!C16'],1),scoped_card(['Ops!F16'],2)]
    # Both single-cell source rules have identical semantics but distinct exact
    # disclosure coordinates, as allowed by an independently certified group.
    shared={'schema_version':'dialogue-shared-source-links/v1','groups':[[c['record_id'] for c in cards]]}
    payload={'claims':cards,'shared_source_links':shared}
    pack={'claims':[X.writer_claim(c) for c in cards],'shared_source_links':shared}
    template,slots=compose(payload,pack)
    assert template.index(junior_scope_text(cards[0]))<template.index(junior_scope_text(cards[1]))
    assert 'D16' not in template and 'E16' not in template
    assert slots['slots']['c001-junior']['claim_ids']==[c['record_id'] for c in cards]
