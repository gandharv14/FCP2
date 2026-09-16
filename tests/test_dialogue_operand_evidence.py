from collections import Counter
from pathlib import Path
import json
import pytest
from openpyxl import Workbook
from scripts.batch import harbor_dialogue_operand_evidence as operands
from scripts.batch.harbor_readiness import skill_module,ROOT
from scripts.batch.harbor_admission import record
from xl_ast_graph import parse_formula

@pytest.fixture
def v():return skill_module(ROOT,'additional-assumptions-dialogue','validate_dialogue')

@pytest.mark.parametrize('number',[0,1,-1,2,12,365,10000])
def test_every_target_collision_requires_its_own_occurrence_budget(v,tmp_path,monkeypatch,number):
    monkeypatch.setattr(v.disclose,'numeric_targets',lambda _: [float(number)])
    assert v.numeric_faults_with_budget(str(number),tmp_path,{})
    assert not v.numeric_faults_with_budget(str(number),tmp_path,Counter({float(number):1}))
    assert len(v.numeric_faults_with_budget(f'{number} then {number}',tmp_path,Counter({float(number):1})))==1
    assert v.numeric_faults_with_budget(str(number),tmp_path,None)

@pytest.mark.parametrize('number',[0,1,-1,2,12,365])
def test_exact_required_label_permits_only_one_occurrence(v,tmp_path,monkeypatch,number):
    monkeypatch.setattr(v.disclose,'numeric_targets',lambda _: [float(number)])
    budget=v.licensed_row_label_literals({'row_label':f'Period {number}'})
    assert not v.numeric_faults_with_budget(f'Period {number}',tmp_path,budget)
    assert v.numeric_faults_with_budget(f'Period {number}; the answer is {number}',tmp_path,budget)

@pytest.fixture
def context(tmp_path):
    source=tmp_path/'source.xlsx';book=Workbook();sheet=book.active;sheet.title='S'
    sheet['A1']='=CHOOSE(Circ,A3,-100)';sheet['B1']='=CHOOSE(Circ,B3,-100)';book.save(source)
    stage=tmp_path/'task';(stage/'tests').mkdir(parents=True)
    key=stage/'tests/answer_key.json';key.write_text(json.dumps({'targets':{'S!D4':42}}))
    (stage/'tests/disclosure.json').write_text(json.dumps({'golden':str(source),'golden_sha256':record(source)['sha256']}))
    claim={'record_id':'method::S::Interest::001','sheet':'S','reviewer_only':{'source':'custom_method_detector','evidence':'=CHOOSE(Circ,A3,-100)','cells':['S!A1','S!B1']}}
    claims={'claims':[claim]}
    draft='<!-- claim:method::S::Interest::001 -->\n**Analyst:** What is the fallback?\n\n**VP:** Use negative one hundred.\n'
    return source,stage,key,claims,draft

def project(context):
    source,stage,key,claims,draft=context
    return operands.project(claims,draft,source=record(source),answer_key=record(key),task_dir=stage)


@pytest.mark.parametrize('sheet',["'Topline Build - 2020 Revenue'", "'Director''s Sheet'", 'Sheet1'])
def test_qualified_ref_error_survives_actual_relative_translation(sheet):
    formula=f'=IF($D$16=1,CHOOSE($D$15,{sheet}!I13,{sheet}!#REF!),J$2)'
    t=operands.SourceFormulaTranslator(formula,origin='J107')
    assert t.translate_formula('J107')==formula
    assert t.translate_formula('K108')==f'=IF($D$16=1,CHOOSE($D$15,{sheet}!J14,{sheet}!#REF!),K$2)'


@pytest.mark.parametrize('change',[None,'different_error','different_sheet','missing_error','changed_reference','changed_literal'])
def test_qualified_error_projection_preserves_all_original_scope_checks(context,change):
    from openpyxl import load_workbook
    source,stage,key,claims,draft=context
    original="=IF(A3,-100,'Other Sheet'!#REF!)"
    copied="=IF(B3,-100,'Other Sheet'!#REF!)"
    if change=='different_error':copied=copied.replace('#REF!','#VALUE!')
    elif change=='different_sheet':copied=copied.replace('Other Sheet','Foreign Sheet')
    elif change=='missing_error':copied=copied.replace("'Other Sheet'!#REF!",'')
    elif change=='changed_reference':copied=copied.replace('B3','B4')
    elif change=='changed_literal':copied=copied.replace('-100','-99')
    book=load_workbook(source);book['S']['A1']=original;book['S']['B1']=copied;book.save(source)
    claims['claims'][0]['reviewer_only']['evidence']=original
    disclosure=stage/'tests/disclosure.json';value=json.loads(disclosure.read_text())
    value['golden_sha256']=record(source)['sha256'];disclosure.write_text(json.dumps(value))
    if change is None:
        result=project((source,stage,key,claims,draft))
        assert result['semantic_review_required'] and not result['target_values_disclosed']
        assert result['claims'][0]['operands'][0]['value']==-100
    else:
        with pytest.raises(ValueError,match='source method copied formula changed'):
            project((source,stage,key,claims,draft))

def test_signed_exact_source_operand_is_review_evidence_not_target_or_approval(context):
    r=project(context);item=r['claims'][0]['operands'][0]
    assert item['value']==-100 and item['licensed_occurrences']==1
    assert item['observed_occurrences'][0]['text']=='negative one hundred'
    assert item['within_occurrence_budget'] and not item['target_collision']
    assert r['semantic_review_required'] and not r['target_values_disclosed']
    assert 'passed' not in r and set(r)=={'schema_version','source_sha256','answer_key_sha256','claims_sha256','draft_sha256','claims','target_values_disclosed','semantic_review_required'}

def test_source_operand_target_collision_and_extra_occurrence_are_explicit(context):
    source,stage,key,claims,draft=context;key.write_text(json.dumps({'targets':{'S!D4':-100}}))
    item=project(context)['claims'][0]['operands'][0]
    assert item['target_collision'] and item['within_occurrence_budget']
    with pytest.raises(ValueError,match='extra target occurrence'):
        project((source,stage,key,claims,draft+'\nNegative one hundred.'))


@pytest.mark.parametrize('change',['missing_source','source_bytes','formula','copied_formula','literal_text','target_formula','target_ref','missing_key','key_bytes','disclosure_source','claim_map'])
def test_missing_altered_or_target_derived_inputs_reject(context,change):
    source,stage,key,claims,draft=context;src=record(source);kb=record(key)
    if change=='missing_source':source.unlink()
    elif change=='source_bytes':source.write_bytes(b'corrupt')
    elif change in ('formula','copied_formula','literal_text'):
        from openpyxl import load_workbook
        book=load_workbook(source)
        if change=='literal_text':book['S']['A1'].data_type='s'
        else:book['S']['A1' if change=='formula' else 'B1']='=CHOOSE(Circ,A3,-99)'
        book.save(source)
        src=record(source);d=json.loads((stage/'tests/disclosure.json').read_text());d['golden_sha256']=src['sha256'];(stage/'tests/disclosure.json').write_text(json.dumps(d))
    elif change=='target_formula':key.write_text(json.dumps({'targets':{'S!A1':42}}));kb=record(key)
    elif change=='target_ref':key.write_text(json.dumps({'targets':{'S!A3':42}}));kb=record(key)
    elif change=='missing_key':key.unlink()
    elif change=='key_bytes':key.write_text('{}')
    elif change=='disclosure_source':(stage/'tests/disclosure.json').write_text('{}')
    else:draft=draft.replace('method::S::Interest::001','unknown')
    with pytest.raises((ValueError,FileNotFoundError)):
        operands.project(claims,draft,source=src,answer_key=kb,task_dir=stage)

@pytest.mark.parametrize('formula,expected',[('=-1',{-1.0:1}),('=+1',{1.0:1}),('=CHOOSE(A1,B1,-100)',{-100.0:1}),('=1+1',{1.0:2})])
def test_exact_signed_operand_accounting(formula,expected):
    assert operands.literal_operands(parse_formula(formula))==Counter(expected)

@pytest.mark.parametrize('number',[0,1,-1,2,12,365,-100])
def test_actual_formula_literal_budget_matches_sign_and_never_extra_target(v,tmp_path,number):
    (tmp_path/'tests').mkdir();(tmp_path/'tests/answer_key.json').write_text(json.dumps({'targets':{'S!Z99':number}}))
    claim={'sheet':'S','reviewer_only':{'source':'custom_method_detector','evidence':f'=IF(A1,{number},A2)','cells':['S!B1']}}
    budget=v.licensed_formula_literals(claim,tmp_path)
    assert budget==Counter({float(number):1})
    assert not v.numeric_faults_with_budget(str(number),tmp_path,budget)
    assert v.numeric_faults_with_budget(f'{number} and {number}',tmp_path,budget)


def alias(formula,text):
    import re
    occurrences=[{'text':m.group(),'start':m.start(),'end':m.end()} for m in re.finditer(r'\bzero\b',text)]
    return operands.bound_restatement(parse_formula(formula),text,0.0,occurrences)


def test_source_max_zero_restatement_records_both_spans_and_one_bound():
    text='The row is the greater of that result and zero, so zero is the floor.'
    result=alias('=+MAX(A1-B1,0)',text)
    assert result['logical_operand_occurrences']==1 and len(result['observed_spans'])==2
    assert result['code']=='redundant_zero_floor_restatement'


@pytest.mark.parametrize('formula,text',[
    ('=MIN(A1,0)','The row is the greater of that result and zero, so zero is the floor.'),
    ('=-MAX(A1,0)','The row is the greater of that result and zero, so zero is the floor.'),
    ('=MAX(A1,0)+1','The row is the greater of that result and zero, so zero is the floor.'),
    ('=MAX(A1,1)','The row is the greater of that result and zero, so zero is the floor.'),
    ('=MAX(A1,0)','The row is the greater of that result and zero. The answer is zero.'),
    ('=MAX(A1,0)','The row is the greater of that result and zero, so zero is the floor. Output zero.'),
])
def test_unrelated_or_extra_zero_is_never_an_alias(formula,text):
    assert alias(formula,text) is None


def test_actual_project_accepts_only_proved_restatement(context):
    from openpyxl import load_workbook
    source,stage,key,claims,draft=context
    book=load_workbook(source);book['S']['A1']='=MAX(A3,0)';book['S']['B1']='=MAX(B3,0)';book.save(source)
    c=claims['claims'][0];c['reviewer_only']['evidence']='=MAX(A3,0)'
    key.write_text(json.dumps({'targets':{'S!D4':0}}))
    d=stage/'tests/disclosure.json';v=json.loads(d.read_text());v['golden_sha256']=record(source)['sha256'];d.write_text(json.dumps(v))
    draft=draft.replace('Use negative one hundred.','The row is the greater of that result and zero, so zero is the floor.')
    item=project((source,stage,key,claims,draft))['claims'][0]['operands'][0]
    assert item['within_occurrence_budget'] and item['target_collision'] and item['licensed_occurrences']==1
    assert len(item['observed_occurrences'])==2 and item['restatement_warning']
    with pytest.raises(ValueError,match='extra target occurrence'):
        project((source,stage,key,claims,draft+' The answer is zero.'))

@pytest.fixture
def legacy_increment(tmp_path):
    import copy
    d=skill_module(ROOT,'task-disclosure','disclose');extract=skill_module(ROOT,'additional-assumptions-dialogue','extract_claims')
    stage=tmp_path/'task';(stage/'environment').mkdir(parents=True);(stage/'tests').mkdir()
    source=tmp_path/'source.xlsx';book=Workbook();s=book.active;s.title='S';s['C28']='WACC';s['C31']=.10;s['C32']='=C31+0.01';s['C33']='=C32+0.01';book.save(source)
    s['C32']=None;s['C33']=None;delivered=stage/'environment/model-inputs.xlsx';book.save(delivered)
    G,V=d.Book(source),d.Book(delivered);rec=d.column_increment_record(G,V,['S!C32','S!C33']);assert rec is not None
    # Extraction normally follows render's status decision.
    rec['disposition']='disclose';rec['rendered']='Source increment'
    claim=extract.build_claim(rec,1);assert claim['reviewer_only']['evidence']=='=C31+0.01';claim['reviewer_only']['evidence']=''
    key=stage/'tests/answer_key.json';key.write_text(json.dumps({'targets':{'S!Z9':7}}))
    (stage/'tests/disclosure.json').write_text(json.dumps(dict(golden=str(source),golden_sha256=record(source)['sha256'],records=[rec])))
    draft='<!-- claim:'+claim['record_id']+' -->\n**Analyst:** What is the convention?\n\n**VP:** Use this period plus 0.01.\n'
    return source,stage,key,{'claims':[claim]},draft

@pytest.mark.parametrize('fault',[None,'source','claim','record','family','heading','graded_closure'])
def test_legacy_increment_rederived_from_real_formula_family(legacy_increment,fault):
    import copy
    from openpyxl import load_workbook
    source,stage,key,claims,draft=legacy_increment;before=copy.deepcopy(claims)
    if fault in ('source','family','heading'):
        book=load_workbook(source)
        if fault=='source':book['S']['C32']='=C31+0.02'
        elif fault=='family':book['S']['C34']='=C33+0.01'
        else:book['S']['C28']='Foreign'
        book.save(source)
        path=stage/'tests/disclosure.json';p=json.loads(path.read_text());p['golden_sha256']=record(source)['sha256'];path.write_text(json.dumps(p))
    elif fault=='claim':claims['claims'][0]['must_say'].append('invented condition')
    elif fault=='record':
        path=stage/'tests/disclosure.json';p=json.loads(path.read_text());p['records'].append(copy.deepcopy(p['records'][0]));path.write_text(json.dumps(p))
    elif fault=='graded_closure':key.write_text(json.dumps({'targets':{'S!C31':.10}}))
    if fault:
        with pytest.raises(ValueError):project((source,stage,key,claims,draft))
    else:
        result=project(legacy_increment);item=result['claims'][0]['operands'][0]
        assert item['value']==.01 and item['licensed_occurrences']==1 and len(item['source_evidence'])==2
        assert item['within_occurrence_budget'] and not item['target_collision'] and claims==before
