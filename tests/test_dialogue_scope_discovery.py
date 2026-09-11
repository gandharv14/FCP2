from pathlib import Path
import copy,json
from datetime import datetime
import pytest
from openpyxl import Workbook
from openpyxl.utils.cell import get_column_letter,quote_sheetname
from scripts.batch import harbor_dialogue_scope_discovery as D
from scripts.batch.harbor_admission import record,read
from scripts.batch import harbor_readiness as readiness
ROOT=Path(readiness.__file__).resolve().parents[2]


@pytest.fixture
def fixture(tmp_path):
    book=Workbook();book.active.title='3 Summary';book.create_sheet("O'Brien Inputs")
    source=tmp_path/'original.xlsx';stage=tmp_path/'0001-outputs';(stage/'tests').mkdir(parents=True)
    records=[]
    for year in range(11):
        col=get_column_letter(year+3);origin=get_column_letter(year+2)
        book['3 Summary'][col+'24']='='+quote_sheetname("O'Brien Inputs")+'!$'+origin+'$139'
        book['3 Summary'][col+'16']='Year '+str(year)
        book["O'Brien Inputs"][origin+'3']='Year '+str(year)
        book["O'Brien Inputs"][origin+'139']=year+10
        records.append(dict(entry='source_selection',value='source',cells=[quote_sheetname('3 Summary')+'!'+col+'24'],fields=dict(label='"Income Before Fixed Charges"',ingredient='cell '+quote_sheetname("O'Brien Inputs")+'!'+origin+'139 on the row labelled "Income Before Fixed Charges"')))
    records.insert(3,dict(entry='method_operating_expense',value='percent_of_revenue',cells=[quote_sheetname('3 Summary')+'!P40'],fields=dict(label='"Operating Expense"')))
    (stage/'instruction.md').write_text('# Financial task\n\n## Workbook disclosure\n\n- Original bound disclosure.\n')
    def save():
        book.save(source);b=record(source)
        (stage/'tests/disclosure.json').write_text(json.dumps(dict(golden=b['path'],golden_sha256=b['sha256'],agent_records=records)))
        return b
    return book,stage,save,tmp_path


def prepared(fixture):
    b,stage,save,tmp=fixture;source=save()
    proof=D.prepare(stage,source,tmp/'discovery',root=ROOT)
    return proof,D.validate(stage,proof,root=ROOT)


def test_eleven_links_use_literal_paired_axes_with_exact_mixed_coverage(fixture):
    book,stage,save,tmp=fixture
    proof,(claims,options)=prepared(fixture);v=read(proof);scope=read(v['scope']);details=read(v['discovery'])
    assert v['used'] and len(scope['rows'])==11 and len(scope['claim_ids'])==12
    assert [len(g) for g in claims['shared_source_links']['groups']]==[11,1]
    assert [c['record_id'] for c in claims['claims']]==scope['claim_ids']
    assert claims['claims'][3]['reviewer_only']['entry']=='method_operating_expense'
    assert 'matching period' not in claims['claims'][3]['must_say']
    assert options==dict(source_workbook=str((tmp/'original.xlsx').resolve()),scope_evidence=v['scope']['path'])
    assert all(x['data_type'] in ('f','s','inlineStr') for x in details['observed'])
    assert not any(x['reference'].endswith('139') for x in details['observed'])


@pytest.mark.parametrize('fault',['duplicate_axis','broken_axis','date_axis','competing_date_axis','numeric_axis','formula_axis','formula_literal','dynamic','named_range','unequal_periods'])
def test_unproved_or_ambiguous_links_stay_individual(fixture,fault):
    book,stage,save,tmp=fixture
    for year in range(11):
        col=get_column_letter(year+2)
        if fault=='duplicate_axis':book["O'Brien Inputs"][col+'4']='Year '+str(year)
        elif fault=='broken_axis':book["O'Brien Inputs"][col+'3']='Year '+str(year*2)
        elif fault=='date_axis':book["O'Brien Inputs"][col+'3']=datetime(2020+year,1,1)
        elif fault=='competing_date_axis':book["O'Brien Inputs"][col+'4']=datetime(2020+year,1,1)
        elif fault=='numeric_axis':book["O'Brien Inputs"][col+'3']=year
        elif fault=='formula_axis':book["O'Brien Inputs"][col+'3']='="Year '+str(year)+'"'
        elif fault=='unequal_periods':book["O'Brien Inputs"][col+'3']='Year '+str(year+20)
        else:
            target=book['3 Summary'][get_column_letter(year+3)+'24']
            if fault=='formula_literal':target.data_type='s'
            elif fault=='dynamic':target.value='=INDIRECT("B139")'
            else:target.value='=AnUnknownName'
    proof,(actual,options)=prepared(fixture)
    assert read(proof)['used'] is False and options=={} and actual['shared_source_links'] is None
    assert read(read(proof)['scope'])['rows']==[]
    assert not any('matching period' in c['must_say'] for c in actual['claims'])


@pytest.mark.parametrize('fault',['source','disclosure','claim_ids','formula','archive_header','producer','instruction'])
def test_retained_discovery_rejects_tampering_without_regeneration(fixture,fault):
    book,stage,save,tmp=fixture;proof,_=prepared(fixture);v=read(proof)
    if fault=='source':book['3 Summary']['Z99']='changed';book.save(tmp/'original.xlsx')
    elif fault=='disclosure':(stage/'tests/disclosure.json').write_text('{}')
    elif fault=='instruction':(stage/'instruction.md').write_text('changed')
    elif fault=='producer':v['producer']['sha256']='0'*64;Path(proof['path']).write_bytes(D.canonical(v));proof=record(proof['path'])
    else:
        key='discovery' if fault=='archive_header' else 'scope';p=Path(v[key]['path']);x=read(v[key])
        if fault=='claim_ids':x['claim_ids'].pop()
        elif fault=='formula':x['rows'][0]['original_formula']='=99'
        else:x['observed'][0]['value']='Year 99'
        p.write_bytes(D.canonical(x))
    with pytest.raises((ValueError,KeyError)):D.validate(stage,proof,root=ROOT)


def test_first_call_hook_cannot_reset_existing_dialogue_and_bounded_scan_is_noop(fixture,monkeypatch):
    book,stage,save,tmp=fixture;source=save();directory=tmp/'dialogue'
    result=D.initialize_dialogue_claims(stage,source,directory,root=ROOT)
    assert result['scope_evidence'] is not None
    with pytest.raises(ValueError):D.initialize_dialogue_claims(stage,source,directory,root=ROOT)
    monkeypatch.setattr(D,'MAX_SCANNED_CELLS',2)
    proof=D.prepare(stage,source,tmp/'bounded',root=ROOT);actual,options=D.validate(stage,proof,root=ROOT)
    assert not read(proof)['used'] and options=={} and len(actual['claims'])==12
    assert all('bounded' in reason for reason in read(read(proof)['discovery'])['decisions'].values())


def test_formula_cached_answers_are_never_period_or_source_evidence(fixture):
    import zipfile
    from xml.etree import ElementTree as ET
    book,stage,save,tmp=fixture;original=save()
    first=D.prepare(stage,original,tmp/'first',root=ROOT);one,_=D.validate(stage,first,root=ROOT)
    source=Path(original['path']);ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(source) as archive:files={n:archive.read(n) for n in archive.namelist()}
    for name,raw in list(files.items()):
        if not name.startswith('xl/worksheets/') or not name.endswith('.xml'):continue
        doc=ET.fromstring(raw)
        for cell in doc.findall('.//s:c',ns):
            if cell.find('s:f',ns) is None:continue
            value=cell.find('s:v',ns)
            if value is None:value=ET.SubElement(cell,'{'+ns['s']+'}v')
            value.text='999999999'
        files[name]=ET.tostring(doc)
    with zipfile.ZipFile(source,'w') as archive:
        for name,raw in files.items():archive.writestr(name,raw)
    revised=record(source);payload=json.loads((stage/'tests/disclosure.json').read_text());payload['golden_sha256']=revised['sha256'];(stage/'tests/disclosure.json').write_text(json.dumps(payload))
    second=D.prepare(stage,revised,tmp/'second',root=ROOT);two,_=D.validate(stage,second,root=ROOT)
    assert one['claims']==two['claims']
    assert read(read(first)['scope'])['rows']==read(read(second)['scope'])['rows']
    assert read(first)['source']['sha256']!=read(second)['source']['sha256']


@pytest.mark.parametrize('fault',['dropped_claim','forged_header','altered_decision'])
def test_rehashed_scope_receipt_still_requires_independent_exact_discovery(fixture,fault):
    book,stage,save,tmp=fixture;proof,_=prepared(fixture);v=read(proof)
    name='discovery' if fault=='altered_decision' else 'scope';binding=v[name];value=read(binding)
    if fault=='dropped_claim':value['rows'].pop();value['claim_ids'].pop()
    elif fault=='forged_header':value['rows'][0]['source_period_header']="'O''Brien Inputs'!C3"
    else:value['decisions'][next(iter(value['decisions']))]='proved regardless of evidence'
    path=Path(binding['path']);path.write_bytes(D.canonical(value));v[name]=record(path)
    Path(proof['path']).write_bytes(D.canonical(v));proof=record(proof['path'])
    with pytest.raises(ValueError):D.validate(stage,proof,root=ROOT)
