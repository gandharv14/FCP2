from pathlib import Path
from dataclasses import replace,asdict
import json
import pytest
from openpyxl import Workbook
from scripts.batch import harbor_instruction_producer as p,harbor_instruction_worker as w
from scripts.batch import harbor_dialogue_semantic_retry as retry,harbor_dialogue_scope_context as scope
from scripts.batch import harbor_tracker_lifecycle as life
from scripts.batch.harbor_admission import record
from scripts.batch.harbor_workbook_bootstrap import Lease
from scripts.batch.harbor_lane_notes import RootContext
from scripts.batch.harbor_readiness import Selection,PreDialogueWitness
from scripts.batch.harbor_complete_readiness import ReadinessSelection
from test_harbor_instruction_producer import setup


def prose(task,directory=None):
    directory=directory or task.run/'dialogue';slots=p.read(record(directory/'slots.json'))
    return {'schema_version':'dialogue-slot-prose/v1','slots':{sid:('Which source should I use?' if row['kind']=='junior' else 'Take the row labelled "Income Before Fixed Charges" from the row labelled "INCOME BEFORE FIXED CHARGES" on the Operations tab. Use the matching period for this link.') for sid,row in slots['slots'].items()}}

@pytest.fixture
def failed(setup,monkeypatch):
    task,campaign,respond,calls=setup
    book=Workbook();book.active.title='Summary';book.create_sheet('Operations');entries=[];rows=[]
    for i in range(3):
        col=chr(ord('C')+i);origin=chr(ord('B')+i)
        book['Summary'][col+'16']=f'Year {i}';book['Operations'][origin+'3']=f'Year {i}'
        book['Summary'][col+'24']='=Operations!'+origin+'$139';book['Operations'][origin+'139']=10+i
        entries.append({'entry':'source_selection','value':'source','cells':['Summary!'+col+'24'],'fields':{'label':'"Income Before Fixed Charges"','ingredient':'cell Operations!'+origin+'139 on the row labelled "INCOME BEFORE FIXED CHARGES"'}})
    source=task.run.parent/'dialogue-source.xlsx';book.save(source)
    (task.stage/'tests/disclosure.json').write_text(json.dumps({'golden':str(source),'golden_sha256':record(source)['sha256'],'agent_records':entries,'records':entries}))
    tracker=Path(task.tracker['path']);before=json.loads(tracker.read_text())
    for lane in (p.NAT,'harbor-independent-verifier'):before['lane_state'][lane]['state']='passed'
    for name in ('source_generation','segmentation_generation','curation'):before['bindings'][name]=before['raw_source']
    before['bindings']['source_file']=record(source);tracker.write_text(json.dumps(before))
    authority=life.RootAuthority('harbor-orchestrator',10**12,lambda *a:True)
    update=life.lease_operation(tracker,operation='acquire',authority=authority,lease_kind='task_agent',owner='fixture',scope='workbook:0245',expected_revision=before['revision'],expected_sha256=record(tracker)['sha256'],duration_ns=10**12)
    before=p.read(record(tracker));lease=Lease('fixture',before['leases']['task_agent']['token'],update['revision'],update['sha256'])
    task=replace(task,tracker=record(tracker),stage_tree=p.tree(task.stage));directory=task.run/'dialogue'
    # Explicit retained pre-discovery context. Real existing first calls use
    # this ordinary branch, without fabricating a discovery receipt afterward.
    extractor=p.skill_module(task.root,'additional-assumptions-dialogue','extract_claims');p.put(directory/'claims.json',extractor.extract(task.stage));p.initialize_dialogue(task)
    old_dispatch=p.dispatch
    monkeypatch.setattr(p,'dispatch',lambda _task,*a,**kw:old_dispatch(replace(task,tracker=record(tracker)),*a,**kw))
    worker=w.Worker(RootContext(tracker.parent),task.run.parent,campaign)
    norm=task.run/'normalization';norm.mkdir();(norm/'normalized.json').write_text('{}')
    selection=ReadinessSelection(Selection('0245',task.stage,w.static.publication._tree_hash(task.stage)[0],before['bindings'],norm,task.run/'nat',record(directory/'claims.json'),None,model_policy=task.policy),None,None,None,None)
    writer=p.prepare_dialogue_writer(task);respond(writer,prose(task));assert p.consume_dialogue_writer(task,writer,campaign)['passed']
    para=p.prepare_dialogue_paraphraser(task);respond(para,prose(task));assert p.consume_dialogue_paraphraser(task,para,writer,campaign)['draft_report']['passed']
    review=p.prepare_dialogue_review(task);claims=p.read(record(directory/'claims.json'))
    value={'agent_model':'gpt-5.6-sol-high','round':1,'accuracy':{'verdict':'pass','claims':[{'record_id':c['record_id'],'must_say':'entailed','findings':[]} for c in claims['claims']],'extras':[],'cell_refs_in_senior_turns':[]},'naturalness':{'verdict':'fail','findings':['Repetitive consecutive source exchanges']},'passed':False}
    respond(review,value);assert not p.consume_dialogue_review(task,review,writer,para,campaign)['passed']
    initial=w.keep(worker.run/'initial.json',{'selection':w.codec.encode(selection),'lease':asdict(lease),'code':{'old-code':'0'*64}});previous=None
    for name in w.STEPS[:w.STEPS.index('dialogue-review')]:
        intent=w.keep(w.step_directory(worker,name)/'intent.json',{'initial':initial,'previous':previous,'step':name})
        previous=w.keep(w.step_directory(worker,name)/'complete.json',{'step':name,'status':'completed','selection':w.codec.encode(selection),'lease':asdict(lease),'intent':intent})
    step=w.step_directory(worker,'dialogue-review');w.request_saved(step/'request.json',review)
    intent=w.keep(step/'intent.json',{'initial':initial,'previous':previous,'step':'dialogue-review','selection':w.codec.encode(selection),'lease':asdict(lease)})
    blocked=w.keep(step/'blocked.json',{'step':'dialogue-review','status':'DIALOGUE_REPAIR_REQUIRED','selection':w.codec.encode(selection),'lease':asdict(lease),'intent':intent})
    for claim in claims['claims']:
        target=claim['reviewer_only']['cells'][0];col=target.rsplit('!',1)[1][0];i=ord(col)-ord('C');origin=chr(ord('B')+i)
        rows.append(dict(claim_id=claim['record_id'],target_reference=target,source_reference='Operations!'+origin+'139',original_formula='=Operations!'+origin+'$139',target_period_header='Summary!'+col+'16',source_period_header='Operations!'+origin+'3',target_period_axis_start='Summary!C16',source_period_axis_start='Operations!B3',period_kind='relative_year',period=i))
    evidence=p.put(task.run.parent/'scope.json',dict(schema_version='dialogue-source-link-scope/v3',source_sha256=record(source)['sha256'],disclosure_sha256=record(task.stage/'tests/disclosure.json')['sha256'],claim_ids=[c['record_id'] for c in claims['claims']],rows=rows))
    # Current full financial/runtime gates are independently tested elsewhere;
    # the fixture stubs only that fixed owning boundary. No callback is accepted
    # by production. All scope, model journals, history, worker and CAS are real.
    monkeypatch.setattr(retry.structural,'_prefix',lambda *a:[record(task.stage/'instruction.md')])
    return worker,selection,lease,blocked,task,campaign,respond,calls,evidence


def test_root_grouping_preserves_original_calls_and_requires_new_full_review(failed):
    worker,s,lease,blocked,task,campaign,respond,calls,evidence=failed
    originals={str(p):record(p) for p in (task.run/'dialogue').rglob('*') if p.is_file()}
    worker,s,lease,permit=retry.authorize(worker,s,lease,blocked=blocked,scope_evidence=evidence)
    proof=s.content.dialogue_scope_receipt;old,new,value=scope.validate_transition(proof)
    assert len(old['claims'])==len(new['claims'])==3 and new['shared_source_links']['groups']==[[c['record_id'] for c in old['claims']]]
    assert len(calls)==3 and all(record(path)==b for path,b in originals.items())
    pending=w.step(worker,s,lease,checkpoint=permit);req=w.request_saved(w.step_directory(worker,'dialogue-writer')/'request.json')
    assert req.prepared.ordinal==2 and len(json.loads(p.read(req.prepared.request)['messages'][1]['content'])['slots.json']['slot_order'])==2
    directory=Path(value['directory']);respond(req,prose(task,directory));done=w.step(worker,s,lease,checkpoint=pending.evidence)
    pending=w.step(worker,done.selection,done.lease,checkpoint=done.evidence);para=w.request_saved(w.step_directory(worker,'dialogue-paraphraser')/'request.json');respond(para,prose(task,directory));done=w.step(worker,pending.selection,pending.lease,checkpoint=pending.evidence)
    pending=w.step(worker,done.selection,done.lease,checkpoint=done.evidence);review=w.request_saved(w.step_directory(worker,'dialogue-review')/'request.json')
    body=json.loads(p.read(review.prepared.request)['messages'][1]['content']);assert len(body['claims.json']['claims'])==3 and body['round']==2
    # Independent response covers every original ID; no prior PASS is reused.
    response=p.read(record(task.run/'dialogue/round-01/review.json'));response['round']=2;response['passed']=True;response['naturalness']={'verdict':'pass','findings':[]}
    respond(review,response);current=replace(task,tracker=record(task.tracker['path']))
    assert p.consume_dialogue_review(current,review,req,para,campaign)['passed'] and len(calls)==6
    assert all(record(path)==b for path,b in originals.items())
    with pytest.raises(ValueError):p.prepare_dialogue_writer(current,3)

@pytest.mark.parametrize('field', ['source','scope','original_scope','witness','history','original_inputs','artifacts'])
def test_mutated_bound_bytes_rejected(failed,field):
    worker,s,lease,blocked,task,campaign,respond,calls,evidence=failed
    proof=scope.prepare_transition(task,evidence);old,new,v=scope.validate_transition(proof)
    binding=v[field] if field in ('source','scope','original_scope') else next(iter(v[field].values()))
    path=Path(binding['path']);path.write_bytes(path.read_bytes()+b' ')
    with pytest.raises((ValueError,AssertionError)):scope.validate_transition(proof)
    assert len(calls)==3

@pytest.mark.parametrize('field,value', [('consumed_rounds',0),('consumed_rounds',True),('semantic_limit',2.0),('semantic_limit',3),('new_review_required',False),('semantic_budget_reset',True),('ordered_claim_ids',[]),('extra',True)])
def test_rehashed_receipt_cannot_change_contract(failed,field,value):
    *_,task,campaign,respond,calls,evidence=failed
    proof=scope.prepare_transition(task,evidence);v=p.read(proof);v[field]=value
    Path(proof['path']).write_text(json.dumps(v));changed=record(proof['path'])
    with pytest.raises(ValueError):scope.validate_transition(changed)


def test_prepared_receipt_is_not_dispatch_authority(failed):
    *_,task,campaign,respond,calls,evidence=failed
    proof=scope.prepare_transition(task,evidence)
    assert scope.selected_transition(task) is None and len(calls)==3
    with pytest.raises(ValueError):scope.prepare_transition(task,evidence)

@pytest.mark.parametrize('point',['prepared','committed'])
def test_interrupted_root_transaction_preserves_budget_and_reconciles_only_commit(failed,point):
    worker,s,lease,blocked,task,campaign,respond,calls,evidence=failed
    before=record(task.tracker['path'])
    def fault(where):
        if where==point:raise InterruptedError(where)
    with pytest.raises(InterruptedError):retry.authorize(worker,s,lease,blocked=blocked,scope_evidence=evidence,_fault=fault)
    permit=record(worker.run/'dialogue-semantic-retry-02/permit.json')
    assert len(calls)==3
    if point=='prepared':
        assert record(task.tracker['path'])==before
        with pytest.raises(ValueError):retry.recover(worker,permit,lease)
    else:
        recovered,selected,newlease,complete=retry.recover(worker,permit,lease)
        assert newlease.revision==lease.revision+1
        assert scope.validate_transition(selected.content.dialogue_scope_receipt)
        assert p.read(record(task.tracker['path']))['lane_state'][p.DIALOGUE]['attempt']==1


def test_legacy_selection_codec_adds_only_absent_optional_scope(failed):
    worker,s,*_=failed
    encoded=w.codec.encode(s.content);encoded['value'].pop('dialogue_scope_receipt')
    decoded=w.codec.decode(encoded);assert decoded.dialogue_scope_receipt is None
    encoded['value'].pop('stage')
    with pytest.raises(ValueError):w.codec.decode(encoded)


def test_current_claims_must_exactly_match_grouped_derivation(failed):
    *_,task,campaign,respond,calls,evidence=failed
    proof=scope.prepare_transition(task,evidence);old,new,v=scope.validate_transition(proof)
    new['claims'][0]['must_say'].append('invented rule')
    with pytest.raises(ValueError):scope.validate_transition(proof,current_claims=new)

from test_dialogue_scope_discovery import fixture as public_fixture

@pytest.mark.parametrize('grouped',[True,False])
def test_fresh_discovery_is_bound_before_first_writer_and_rederived_after_apply(setup,public_fixture,monkeypatch,grouped):
    task,campaign,respond,calls=setup
    book,stage,save,tmp=public_fixture
    if not grouped:
        for i in range(11):book["O'Brien Inputs"].cell(3,i+2).value=None
    source=save()
    import shutil
    shutil.copytree(stage,task.stage,dirs_exist_ok=True);stage=task.stage
    tracker=Path(task.tracker['path']);before=p.read(record(tracker))
    before['bindings']['source_file']=source
    for lane in (p.NAT,'harbor-independent-verifier'):before['lane_state'][lane]['state']='passed'
    tracker.write_text(json.dumps(before))
    task=replace(task,tracker=record(tracker),stage=stage,stage_tree=p.tree(stage))
    result=p.initialize_dialogue(task)
    assert len(calls)==0 and bool(result['scope_evidence'])==grouped
    request=p.prepare_dialogue_writer(task)
    assert request.prepared.ordinal==1 and len(calls)==0
    assert p.initialize_dialogue(task)==result
    witness=tmp/'retained'/stage.name;witness.mkdir(parents=True)
    import shutil
    shutil.copytree(stage,witness,dirs_exist_ok=True)
    selected=Selection('0245',stage,w.static.publication._tree_hash(stage)[0],before['bindings'],tmp/'norm',tmp/'nat',result['claims'],None,scope_evidence=result['scope_evidence'],dialogue_scope_receipt=result['scope_receipt'])
    # Owning NAT verification is independently covered; this test concentrates
    # on scope's binding to its returned exact witness, including after apply.
    monkeypatch.setattr(w.static,'disclosure_verification_stage',lambda *a,**k:witness)
    original=p.read(result['scope_receipt'])
    (stage/'instruction.md').write_text('Current applied dialogue, original instruction preserved privately.')
    assert scope.validate_selected(selected)
    assert p.read(result['scope_receipt'])==original
    (witness/'instruction.md').write_text('changed pre-dialogue source')
    with pytest.raises(ValueError):scope.validate_selected(selected)


def test_grouped_receipt_is_bound_to_selected_task_source_scope_and_nat_witness(failed,monkeypatch):
    worker,s,lease,blocked,task,*rest=failed;evidence=rest[-1]
    worker,s,lease,permit=retry.authorize(worker,s,lease,blocked=blocked,scope_evidence=evidence)
    v=p.read(s.content.dialogue_scope_receipt);witness=Path(v['directory'])/'witness'/task.stage.name
    monkeypatch.setattr(w.static,'disclosure_verification_stage',lambda *a,**k:witness)
    assert scope.validate_selected(s.content)
    for changed in (replace(s.content,scope_evidence=None),replace(s.content,stage=task.stage.parent/'other'),replace(s.content,bindings={**s.content.bindings,'source_file':s.content.bindings['curation']})):
        with pytest.raises(ValueError):scope.validate_selected(changed)
    (witness/'instruction.md').write_text('changed')
    with pytest.raises(ValueError):scope.validate_selected(s.content)
