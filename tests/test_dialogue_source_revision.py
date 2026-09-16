import copy
from pathlib import Path
import pytest
from openpyxl import Workbook,load_workbook
from scripts.fleet.dialogue_source_revision import repair

@pytest.fixture
def case(tmp_path):
    p=tmp_path/'source.xlsx';w=Workbook();s=w.active;s.title='Model'
    for cell,value in {'B35':'Beginning Balance','B11':'Hurdle','B54':'Loan Repayment','F38':'=IFERROR(MIN(0,MAX(-(F35+F35*$C$11),-F31*$C$9)),0)','G38':'=IFERROR(MIN(0,MAX(-(G35+G35*$C$11),-G31*$C$9)),0)','D53':'=IF(SUM($C$54:C54)<=0,1,0)','E53':'=IF(SUM($C$54:D54)<=0,1,0)'}.items():s[cell]=value
    w.save(p)
    def card(rid,cells):return dict(record_id=rid,sheet='Model',junior_scope=['Model!'+c for c in cells],reviewer_only=dict(source='custom_method_detector',cells=['Model!'+c for c in cells],evidence=s[cells[0]].value))
    claims=dict(claims=[card('interest',['F38','G38']),card('debt',['D53','E53']),dict(record_id='unchanged')])
    draft='<!-- claim:interest -->\n**Analyst:** What convention?\n\n**VP:** One candidate is the row labelled "Beginning Balance" plus the sign-flipped product of that balance and the locked input on the row labelled "Hurdle". Keep all other rules.\n\n<!-- claim:debt -->\n**Analyst:** What total?\n\n**Director:** If the locked input on the row labelled "Loan Repayment" is at most zero, follow the original branches.\n\n<!-- claim:unchanged -->\n**VP:** Keep this exact byte sequence.\n'
    selected=[dict(record_id='interest',kind='negated_sum'),dict(record_id='debt',kind='cumulative_range')]
    return p,claims,draft,selected

def test_exact_source_derivative_preserves_others(case):
    p,c,d,s=case;fixed,proof=repair(d,c,p,s)
    assert 'flip the sign of that entire sum' in fixed
    assert 'If that cumulative total is at most zero' in fixed
    assert fixed.split('<!-- claim:unchanged -->')[1]==d.split('<!-- claim:unchanged -->')[1]
    assert proof['unchanged_claims']==1 and proof['model_calls_added']==proof['task_mutations']==0
    assert len(proof['edits'])==2 and not proof['qualified']
    assert d==case[2] and c==case[1]

@pytest.mark.parametrize('formula',['=IFERROR(MIN(0,MAX(F35-F35*$C$11,-F31*$C$9)),0)','=IFERROR(MIN(1,MAX(-(F35+F35*$C$11),-F31*$C$9)),0)','=IFERROR(MIN(0,MAX(-(F35+G35*$C$11),-F31*$C$9)),0)'])
def test_rejects_different_source_ast(case,formula):
    p,c,d,s=case;w=load_workbook(p);w['Model']['F38']=formula;w.save(p)
    with pytest.raises(ValueError):repair(d,c,p,s)

def test_rejects_uncopied_family(case):
    p,c,d,s=case;w=load_workbook(p);w['Model']['G38']='=0';w.save(p)
    with pytest.raises(ValueError,match='copied'):repair(d,c,p,s)

def test_rejects_changed_claim_formula(case):
    p,c,d,s=case;c=copy.deepcopy(c);c['claims'][0]['reviewer_only']['evidence']='=0'
    with pytest.raises(ValueError):repair(d,c,p,s)

def test_rejects_junior_edit(case):
    p,c,d,s=case;d=d.replace('**VP:** One candidate','**Analyst:** One candidate')
    with pytest.raises(ValueError,match='senior'):repair(d,c,p,s)

def test_rejects_duplicate_or_unknown_selection(case):
    p,c,d,s=case
    with pytest.raises(ValueError):repair(d,c,p,s+s[:1])
    with pytest.raises(ValueError):repair(d,c,p,[dict(record_id='foreign',kind='negated_sum')])

def test_rejects_unscoped_cumulative_range(case):
    p,c,d,s=case;w=load_workbook(p);w['Model']['D53']='=IF(SUM(C54:C54)<=0,1,0)';w.save(p)
    c['claims'][1]['reviewer_only']['evidence']=w['Model']['D53'].value
    with pytest.raises(ValueError):repair(d,c,p,[s[1]])

def test_corrected_text_is_not_reprocessed(case):
    p,c,d,s=case;fixed,_=repair(d,c,p,s)
    with pytest.raises(ValueError):repair(fixed,c,p,s)

@pytest.fixture
def source_text_case(tmp_path):
    p=tmp_path/'source.xlsx';w=Workbook();s=w.active;s.title='Model';w.create_sheet('Revenue')
    s['B2']='(+) Other Services';s['C2']='=1+2'
    s['B3']='Services';s['C3']="=CHOOSE($A$1,'Revenue'!A1,'Revenue'!A2,'Revenue'!#REF!)";s['D3']="=CHOOSE($A$1,'Revenue'!B1,'Revenue'!B2,'Revenue'!#REF!)"
    w.save(p)
    def card(rid,cells,label):return dict(record_id=rid,sheet='Model',row_label=label,must_say=['the row labelled "'+label+'"'],junior_scope=['Model!'+c for c in cells],reviewer_only=dict(source='custom_method_detector',cells=['Model!'+c for c in cells],evidence=s[cells[0]].value))
    claims={'claims':[card('label',['C2'],'(+) Other Services'),card('error',['C3','D3'],'Services')]}
    draft='<!-- claim:label -->\n**VP:** Calculate the row labelled "(+) Other Services)" only here.\n<!-- claim:error -->\n**Director:** Option 3 uses the broken reference "\'Revenue\'!#REF!" and yields a reference error.\n'
    selected=[dict(record_id='label',kind='visible_row_label'),dict(record_id='error',kind='broken_reference_wording')]
    return p,claims,draft,selected

def test_source_label_and_broken_option_keep_semantics(source_text_case):
    p,c,d,s=source_text_case;before=p.read_bytes();fixed,proof=repair(d,c,p,s)
    assert '"(+) Other Services"'in fixed and 'source option is a broken reference'in fixed
    assert '#REF!'not in fixed and len(proof['edits'])==2
    assert proof['edits'][1]['meaning']['source_operand']=="'Revenue'!#REF!"and p.read_bytes()==before
    assert 'Option 3'in fixed and proof['model_calls_added']==0

@pytest.mark.parametrize('fault',['source_label','claim_label','required_error','other_error','other_cell','not_option3','extra_occurrence','junior'])
def test_text_derivative_fails_closed(source_text_case,fault):
    p,c,d,s=source_text_case;c=copy.deepcopy(c)
    if fault=='source_label':
        w=load_workbook(p);w['Model']['B2']='Different';w.save(p)
    elif fault=='claim_label':c['claims'][0]['row_label']='Different'
    elif fault=='required_error':c['claims'][1]['must_say'].append("'Revenue'!#REF!")
    elif fault in('other_error','other_cell','not_option3'):
        w=load_workbook(p)
        if fault=='other_error':w['Model']['D3']="=CHOOSE($A$1,'Revenue'!B1,'Revenue'!B2,'Other'!#REF!)"
        elif fault=='other_cell':w['Model']['D3']="=CHOOSE($A$1,'Revenue'!B1,'Revenue'!C2,'Revenue'!#REF!)"
        else:
            w['Model']['C3']="=CHOOSE($A$1,'Revenue'!#REF!,'Revenue'!A1,'Revenue'!A2)";w['Model']['D3']="=CHOOSE($A$1,'Revenue'!#REF!,'Revenue'!B1,'Revenue'!B2)";c['claims'][1]['reviewer_only']['evidence']=w['Model']['C3'].value
        w.save(p)
    elif fault=='extra_occurrence':d=d.replace('only here.','the row labelled "(+) Other Services)".')
    else:d=d.replace('**Director:**','**Analyst:**')
    with pytest.raises(ValueError):repair(d,c,p,s)

def test_explicit_empty_revision_preserves_draft_and_grants_no_qualification(case):
    p,c,d,_=case;fixed,proof=repair(d,c,p,[])
    assert fixed==d and proof['old_draft_sha256']==proof['new_draft_sha256']
    assert proof['edits']==[] and proof['unchanged_claims']==3
    assert proof['model_calls_added']==proof['task_mutations']==0 and proof['qualified']is False
