"""Hermetic 0257-derived copied geometry and label regressions."""
import copy
import importlib.util
import json
from pathlib import Path
import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
MODULE = ROOT / '.cursor/skills/task-disclosure/scripts/disclose.py'
spec = importlib.util.spec_from_file_location('scope_disclose', MODULE)
disclose = importlib.util.module_from_spec(spec); spec.loader.exec_module(disclose)
disclose.REPO_ROOT = ROOT
disclose.REGISTRY_PATH = ROOT/'.cursor/skills/task-disclosure/REGISTRY.md'
SPAN = ['S!K153','S!K154','S!K155']

def fixture(tmp_path, *, scenarios=False, uniform=False):
    path=tmp_path/'gold.xlsx'; wb=openpyxl.Workbook(); ws=wb.active; ws.title='S'
    labels=['Rent','Payroll','Insurance']
    for n,row in enumerate(range(153,156)):
        ws[f'B{row}']='Maintenance per unit' if uniform or scenarios else labels[n]
        if scenarios: ws[f'C{row}']=['Best Case','Base Case','Worst Case'][n]
        ws[f'D{row}']='£ / unit / year'; ws[f'K{row}']='=9000/12'
    ws['K152']='=INDEX(K153:K155,MATCH($C152,$C153:$C155,0))'; ws['K156']=800
    ws['C152']='Best Case'; ws['B152']='Selected maintenance cost'
    ws['Z99']='=5+2'; ws['B99']='Unrelated graded output'
    wb.save(path); gold=disclose.Book(path)
    task=tmp_path/'task'; (task/'environment').mkdir(parents=True); (task/'tests').mkdir()
    for sheet in wb:
        for cell in list(sheet._cells.values()):
            if cell.data_type=='f': cell.value=None
    wb.save(task/'environment/delivered.xlsx')
    (task/'tests/answer_key.json').write_text(json.dumps({'targets':{'S!Z99':7}}))
    return gold,path,task

def record(gold, cells):
    profile=disclose.formula_profile(gold, {'cell_keys':cells,'label':gold.row_label(cells[0]),'pattern':''})
    fields=disclose.custom_sentence_fields(profile,gold); fields.pop('_coverage_complete',None)
    return {'band':disclose.compact_cells([disclose.pretty(c) for c in cells]).strip('`'),
            'cells':[disclose.pretty(c) for c in cells], 'cell_keys':cells,
            'entry':'method_operating_expense','family':'method_operating_expense',
            'value':'out_of_catalogue','disposition':'disclosed','fields':fields}

def faithcheck(gold,path,task,records):
    (task/'tests/disclosure.json').write_text(json.dumps({'records':records}))
    return disclose.faithcheck_task(task,task.parent/'runs',golden=str(path))

def test_observed_scenario_labels_preserve_full_vertical_family(tmp_path):
    gold,path,task=fixture(tmp_path,scenarios=True)
    assert [gold.row_label(k) for k in SPAN]==['Maintenance per unit']*3
    assert disclose.full_copied_scope_vertical(gold,[SPAN[0]],set())==SPAN
    band=disclose.make_vertical_band(gold,'S',11,SPAN)
    assert disclose.method_subbands(gold,band)==[band]
    rec=record(gold,SPAN)
    assert 'copied down the rows' in disclose.render_sentence(rec)
    result=faithcheck(gold,path,task,[rec]); assert result['passed'],result

def test_qualified_financial_names_not_scenario_markers(tmp_path):
    gold,_,_=fixture(tmp_path)
    for text in ['Best Case Revenue','Worst Case Recovery','Worst case loss reserve','Base Case Sales']:
        assert gold.is_row_name('S',2,text)

def test_complete_heterogeneous_family_passes_all_faithchecks(tmp_path):
    gold,path,task=fixture(tmp_path); records=[record(gold,[k]) for k in SPAN]
    assert disclose._heterogeneous_scope_partitioned(gold,SPAN,records,'method_operating_expense')
    result=faithcheck(gold,path,task,records); assert result['passed'],result

@pytest.mark.parametrize('mutation',['missing','duplicate_record','duplicate_cell','overbroad','wrong_label','wrong_entry','contradictory_cell_keys','missing_machine_keys'])
def test_partition_rejects_incomplete_or_misleading_coverage(tmp_path,mutation):
    gold,_,_=fixture(tmp_path); records=[record(gold,[k]) for k in SPAN]
    if mutation=='missing': records.pop(1)
    elif mutation=='duplicate_record': records.append(copy.deepcopy(records[0]))
    elif mutation=='duplicate_cell': records[0]['cells']*=2
    elif mutation=='overbroad': records[0]['cells'].append('S!K156')
    elif mutation=='wrong_label': records[1]['fields']['label']='"Rent"'
    elif mutation=='wrong_entry': records[1]['entry']='method_returns'
    elif mutation=='missing_machine_keys': records[0]['cell_keys']=[]
    elif mutation=='contradictory_cell_keys': records[0]['cell_keys'].append('S!K156')
    assert not disclose._heterogeneous_scope_partitioned(gold,SPAN,records,'method_operating_expense')

def test_unrelated_record_does_not_change_exact_partition(tmp_path):
    gold,_,_=fixture(tmp_path); records=[record(gold,[k]) for k in SPAN]
    records.append(record(gold,['S!Z99']))
    assert disclose._heterogeneous_scope_partitioned(gold,SPAN,records,'method_operating_expense')

def test_uniform_label_fragments_remain_rejected(tmp_path):
    gold,path,task=fixture(tmp_path,uniform=True); records=[record(gold,[k]) for k in SPAN]
    result=faithcheck(gold,path,task,records)
    assert not result['passed'] and any(f['kind']=='scope' for f in result['faults'])

def test_graded_heterogeneous_partition_still_rejected(tmp_path):
    gold,path,task=fixture(tmp_path); records=[record(gold,[k]) for k in SPAN]
    (task/'tests/answer_key.json').write_text(json.dumps({'targets':{SPAN[1]:9000/12}}))
    result=faithcheck(gold,path,task,records)
    assert not result['passed'] and any(f['kind']=='closure' for f in result['faults'])

def test_changed_mechanics_cannot_join_band(tmp_path):
    gold,path,task=fixture(tmp_path,uniform=True)
    wb=openpyxl.load_workbook(path); wb['S']['K154']='=9000/13'; wb.save(path); gold=disclose.Book(path)
    result=faithcheck(gold,path,task,[record(gold,SPAN)])
    assert not result['passed'] and any(f['claim']=='stated band mixes different mechanics' for f in result['faults'])
