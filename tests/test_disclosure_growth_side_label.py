"""An isolated growth assumption must not relabel the adjacent EBITDA grid."""
import importlib.util
from pathlib import Path
import sys
import openpyxl
import pytest


@pytest.fixture
def disclose(monkeypatch):
    root=Path(__file__).resolve().parents[1]/'.cursor/skills/task-disclosure/scripts'
    monkeypatch.syspath_prepend(str(root))
    spec=importlib.util.spec_from_file_location('growth_label_disclose',root/'disclose.py')
    module=importlib.util.module_from_spec(spec);monkeypatch.setitem(sys.modules,spec.name,module);spec.loader.exec_module(module)
    return module


def workbook(tmp_path,*,header=True,parameter=.015,separator=True):
    wb=openpyxl.Workbook();s=wb.active;s.title='Project A_Val'
    s['B13']='EBITDA';s['E13']='Long-term growth';s['G13']=parameter
    if not separator:s['H13']=99
    for col in range(9,17):
        if header:s.cell(5,col,2011+col-9)
        s.cell(13,col,f"='Project A'!{openpyxl.utils.get_column_letter(col-4)}5")
    source=wb.create_sheet('Project A');source['B5']='GAAP EBITDA'
    for col in range(5,13):source.cell(5,col,100+col)
    path=tmp_path/'source.xlsx';wb.save(path);return path


def test_actual_j13_to_p13_uses_ebitda_without_relabeling_growth_input(disclose,tmp_path):
    book=disclose.Book(workbook(tmp_path))
    for col in 'JKLMNOP':assert book.row_label(f'Project A_Val!{col}13')=='EBITDA'
    assert book.row_label('Project A_Val!G13')=='Long-term growth'
    assert book.row_label('Project A_Val!E13')=='Long-term growth'
    assert book.row_label('Project A!F5')=='GAAP EBITDA'


@pytest.mark.parametrize('options',[{'header':False},{'parameter':'unknown'},{'separator':False}])
def test_without_geometry_proof_nearest_block_still_governs(disclose,tmp_path,options):
    book=disclose.Book(workbook(tmp_path,**options))
    assert book.row_label('Project A_Val!J13')==('unknown' if options.get('parameter')=='unknown' else 'Long-term growth')


def test_short_side_by_side_blocks_are_not_replaced_by_leftmost_label(disclose,tmp_path):
    wb=openpyxl.Workbook();s=wb.active;s.title='Summary'
    s['B9']='Parking Spaces';s['E9']='Long-term growth';s['G9']=.015;s['J9']=.02
    p=tmp_path/'blocks.xlsx';wb.save(p)
    assert disclose.Book(p).row_label('Summary!J9')=='Long-term growth'
