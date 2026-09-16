from pathlib import Path
import pytest,openpyxl
from openpyxl.workbook.defined_name import DefinedName
from scripts.batch import harbor_disclosure_partition_evidence as E,harbor_semantic_readiness as S
from scripts.batch.harbor_admission import record

def book(tmp_path,definition,formula='=QDSO+1',local=None):
 p=tmp_path/'names.xlsx';w=openpyxl.Workbook();m=w.active;m.title='Model';s=w.create_sheet('Input Tab');s['D15']=40;s['D17']=60;m['B1']=formula
 w.defined_names.add(DefinedName('QDSO',attr_text=definition))
 if local is not None:m.defined_names.add(DefinedName('QDSO',attr_text=local,localSheetId=0))
 w.save(p);w.close();return E._Workbook(p,S.workbook_evidence(record(p)))

def test_absolute_source_name_keeps_definition_edges_and_blank_members(tmp_path):
 w=book(tmp_path,"'Input Tab'!$D$15:$D$17")
 try:
  shard=w.shard({'Model!B1'});v=E.decode_workbook(shard)
  assert shard['closure']=='exact_ast_closure'
  assert 'Input Tab!D16'in v['blank_refs']
  assert ['Model!B1','Input Tab!D15']in v['dependency_edges']
  assert shard['snapshot_metadata']['resolved_defined_names']==[dict(formula_cell='Model!B1',name='QDSO',scope_sheet=None,definition="'Input Tab'!$D$15:$D$17",targets=['Input Tab!D15','Input Tab!D16','Input Tab!D17'])]
  # Other batches do not inherit names used by an earlier shard.
  assert 'resolved_defined_names'not in w.shard({'Input Tab!D15'})['snapshot_metadata']
 finally:w.close()

@pytest.mark.parametrize('definition',["'Input Tab'!D15","'[other.xlsx]Input Tab'!$D$15",'OFFSET(Model!$B$1,0,0)',"'Absent'!$D$15",'#REF!','OtherName',"'Input Tab'!$A$1:$XFD$1048576"])
def test_uncertain_names_keep_full_source_fallback(tmp_path,definition):
 w=book(tmp_path,definition)
 try:assert w.shard({'Model!B1'})['closure']=='full_snapshot_fallback'
 finally:w.close()

def test_local_name_shadows_global_without_using_global_target(tmp_path):
 w=book(tmp_path,"'Input Tab'!$D$15",local="'Input Tab'!$D$17")
 try:
  shard=w.shard({'Model!B1'});v=E.decode_workbook(shard)
  assert 'Input Tab!D17'in v['selected_refs'] and 'Input Tab!D15'not in v['selected_refs']
  assert shard['snapshot_metadata']['resolved_defined_names'][0]['scope_sheet']=='Model'
 finally:w.close()

def test_broken_local_name_does_not_fall_back_to_valid_global(tmp_path):
 w=book(tmp_path,"'Input Tab'!$D$15",local='#REF!')
 try:assert w.shard({'Model!B1'})['closure']=='full_snapshot_fallback'
 finally:w.close()
