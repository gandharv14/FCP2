import importlib.util,sys
from pathlib import Path
import openpyxl,pytest

@pytest.fixture
def d(monkeypatch):
    p=Path(__file__).resolve().parents[1]/'.cursor/skills/task-disclosure/scripts/disclose.py'
    monkeypatch.syspath_prepend(str(p.parent));spec=importlib.util.spec_from_file_location('locator_disclose',p)
    m=importlib.util.module_from_spec(spec);monkeypatch.setitem(sys.modules,spec.name,m);spec.loader.exec_module(m)
    return m

@pytest.fixture
def example(d,tmp_path):
    wb=openpyxl.Workbook();ws=wb.active;ws.title='Cut #1';ws['D100']="='Total Market'!D118"
    other=wb.create_sheet('Total Market');other['D118']='FLAT BRANCH MORTGAGE'
    p=tmp_path/'source.xlsx';wb.save(p)
    (tmp_path/'tests').mkdir();(tmp_path/'tests/answer_key.json').write_text('{"targets":{"Total Market!Z1":1.0}}')
    rec=dict(band="`'Cut #1'!D100`",cells=["'Cut #1'!D100"],cell_keys=['Cut #1!D100'],
        entry='source_selection',value='source',disposition='disclosed',fields=dict(label='"FLAT BRANCH MORTGAGE"',
        ingredient="cell 'Total Market'!D118 on the row labelled \"FLAT BRANCH MORTGAGE\""))
    return d.Book(p),rec

def test_actual_cut_sheet_number_is_metadata_with_retained_warning(d,example,tmp_path):
    gold,rec=example;section=d.render_section([rec])
    assert d.audit_text(section,tmp_path,records=[rec],gold=gold)==[]
    warnings=d.metadata_literal_warnings(section,tmp_path,[rec],gold)
    assert len(warnings)==1 and warnings[0]['literal']=='1' and warnings[0]['locator']=="`'Cut #1'!D100`"

@pytest.mark.parametrize('text',[' The answer is 1.',' Use 1 as the amount.'])
def test_same_number_in_body_stays_blocking(d,example,tmp_path,text):
    gold,rec=example;section=d.render_section([rec]).rstrip()+text+'\n'
    assert d.audit_text(section,tmp_path,records=[rec],gold=gold)==['numeric literal 1 matches target 1.0']

@pytest.mark.parametrize('change',['foreign-sheet','wrong-cell','detached-key','unknown-row'])
def test_unverified_locator_does_not_exempt_numbers(d,example,tmp_path,change):
    gold,rec=example
    if change=='foreign-sheet':rec.update(cells=["'Other #1'!D100"],cell_keys=['Other #1!D100'])
    elif change=='wrong-cell':rec.update(cells=["'Cut #1'!D999"],cell_keys=['Cut #1!D999'])
    elif change=='detached-key':rec['cell_keys']=['Cut #1!E100']
    else:rec.update(cells=["'Cut #1'!row 100"],cell_keys=['Cut #1!D100'])
    assert d.audit_text(d.render_section([rec]),tmp_path,records=[rec],gold=gold)

def test_locator_looking_text_inside_body_is_not_header(d,example,tmp_path):
    gold,rec=example;section=d.render_section([rec]).rstrip()+" Mention `'Cut #1'!D100`.\n"
    assert d.audit_text(section,tmp_path,records=[rec],gold=gold)==['numeric literal 1 matches target 1.0']

@pytest.fixture
def source_label(d,tmp_path):
    wb=openpyxl.Workbook();ws=wb.active;ws.title='Master Version - Cut #3'
    ws['F75']='Non-Bank/CU';ws['G75']="='Total Market'!D191"
    other=wb.create_sheet('Total Market');other['C191']='MORTGAGE 1 INC';other['D191']=1
    p=tmp_path/'source.xlsx';wb.save(p);gold=d.Book(p)
    (tmp_path/'tests').mkdir();(tmp_path/'tests/answer_key.json').write_text('{"targets":{"Total Market!Z1":1.0}}')
    target='Master Version - Cut #3!G75';source='Total Market!D191'
    rec=dict(band="`'Master Version - Cut #3'!G75`",cells=["'Master Version - Cut #3'!G75"],cell_keys=[target],
             source_cell=source,entry='source_selection',value='source',disposition='disclosed',
             fields=dict(label=d.q('Non-Bank/CU'),ingredient=d.describe_ref(gold,d.pretty(source),'Master Version - Cut #3','Non-Bank/CU')))
    sentence=d.render_sentence(rec)
    section='- '+d.compact_cells(rec['cells'])+': On this sheet, '+sentence[0].lower()+sentence[1:]+'\n'
    return gold,rec,section

def test_reanchored_source_label_digit_is_verified_metadata(d,source_label,tmp_path):
    gold,rec,section=source_label
    assert d.audit_text(section,tmp_path,records=[rec],gold=gold)==[]
    offsets=d.verified_source_label_numeric_offsets(section,tmp_path,[rec],gold)
    assert len(offsets)==1 and next(iter(offsets.values()))['source_cell']=='Total Market!D191'
    assert any(x.get('kind')=='verified_source_row_label' for x in d.metadata_literal_warnings(section,tmp_path,[rec],gold))

@pytest.mark.parametrize('change',['source','ingredient','formula','closure'])
def test_source_label_exemption_requires_exact_provenance(d,source_label,tmp_path,change):
    gold,rec,section=source_label;closure=None
    if change=='source':rec['source_cell']='Total Market!D192'
    elif change=='ingredient':rec['fields']['ingredient']+=' extra'
    elif change=='formula':gold.formula[rec['cell_keys'][0]]='=1'
    else:closure={rec['cell_keys'][0]}
    assert not d.verified_source_label_numeric_offsets(section,tmp_path,[rec],gold,closure)
    assert d.audit_text(section,tmp_path,records=[rec],gold=gold,closure=closure)

def test_verified_label_does_not_allow_extra_body_answer(d,source_label,tmp_path):
    gold,rec,section=source_label
    assert d.audit_text(section.rstrip()+' Use 1.\n',tmp_path,records=[rec],gold=gold)==['numeric literal 1 matches target 1.0']

@pytest.fixture
def verified_method(d,tmp_path):
    wb=openpyxl.Workbook();ws=wb.active;ws.title='Debt'
    ws['A1']='Note 1 Amount';ws['C1']=7;ws['A2']='Interest';ws['E2']=4
    ws['F2']='=IFERROR($C$1+E2+$C$1+1,0)';ws['G2']='=IFERROR($C$1+F2+$C$1+1,0)'
    path=tmp_path/'source.xlsx';wb.save(path);gold=d.Book(path)
    (tmp_path/'tests').mkdir();(tmp_path/'tests/answer_key.json').write_text('{"targets":{"Debt!Z9":1.0}}')
    cells=['Debt!F2','Debt!G2'];band=dict(band=d.compact_cells([d.pretty(c)for c in cells]),cell_keys=cells,label='Interest',pattern='copied')
    profile=d.formula_profile(gold,band);assert profile['complete']
    fields=d.custom_sentence_fields(profile,gold,lock_cells=cells);assert fields.pop('_coverage_complete')
    rec=d.record('method_interest','out_of_catalogue',cells,profile['formula'],[],'source method',fields=fields)
    rec.update(source='custom_method_detector',disposition='disclosed',coverage_complete=True,method_profile={k:v for k,v in profile.items()if k!='ast'})
    section=d.render_section([rec]);return gold,rec,section


def test_method_numeric_metadata_and_constant_are_exact_source_spans(d,verified_method,tmp_path):
    gold,rec,section=verified_method
    assert 'occurrence 1'in section and 'Note 1 Amount'in section
    assert d.verified_method_numeric_offsets(section,tmp_path,[rec],gold)
    assert d.audit_text(section,tmp_path,records=[rec],gold=gold)==[]
    assert d.audit_text(section.rstrip()+' The answer is 1.\n',tmp_path,records=[rec],gold=gold)==['numeric literal 1 matches target 1.0']

@pytest.mark.parametrize('defect',['formula','fields','profile','source','closure','target','foreign-bullet'])
def test_method_numeric_offsets_reject_detached_source_or_extra_prose(d,verified_method,tmp_path,defect):
    gold,rec,section=verified_method;closure=None
    if defect=='formula':gold.formula['Debt!F2']='=2'
    elif defect=='fields':rec['fields']['steps']+=' The answer is 1.';section=d.render_section([rec])
    elif defect=='profile':rec['method_profile']['numbers'].append(1)
    elif defect=='source':rec['source']='unknown'
    elif defect=='closure':closure={'Debt!C1'}
    elif defect=='target':(tmp_path/'tests/answer_key.json').write_text('{"targets":{"Debt!C1":1.0}}')
    else:section=section.replace('- `Debt!F2:G2`:','- `Other!F2:G2`:')
    assert not d.verified_method_numeric_offsets(section,tmp_path,[rec],gold,closure)


def test_plain_source_band_and_markdown_locator_are_exact_equivalents(d,verified_method,tmp_path):
    gold,rec,_=verified_method
    rec['method_profile']['band']=rec['method_profile']['band'][1:-1]
    rec['fields']['band']=rec['fields']['band'][1:-1]
    section=d.render_section([rec])
    assert d.verified_method_numeric_offsets(section,tmp_path,[rec],gold)
    rec['method_profile']['band']='Debt!Z9'
    assert not d.verified_method_numeric_offsets(section,tmp_path,[rec],gold)
