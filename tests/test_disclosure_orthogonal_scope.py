"""Complete intersecting source families keep their distinct row labels."""
import copy,json
import openpyxl,pytest
from test_disclosure_copied_scope import disclose,record,faithcheck


def fixture(tmp_path):
 p=tmp_path/'gold.xlsx';b=openpyxl.Workbook();s=b.active;s.title='S'
 s['B140']='Replacement Capex';s['B141']='Vehicle Wrap';s['B3']='Deal Year';s['B4']='Exit Year';s['C4']=5
 for c in ['N','O','P','Q']:
  s[c+'3']=3;s[c+'140']='=IF('+c+'$3>$C$4,0,5000)'
 s['N141']='=IF(N$3>$C$4,0,5000)';s['O141']='=IF(O$3>$C$4,0,0)';s['M140']='=IF(M$3>$C$4,0,0)';s['Z99']='=5+2';s['B99']='Graded output';b.save(p)
 gold=disclose.Book(p);task=tmp_path/'task';(task/'environment').mkdir(parents=True);(task/'tests').mkdir()
 for cell in list(s._cells.values()):
  if cell.data_type=='f':cell.value=None
 b.save(task/'environment/delivered.xlsx');(task/'tests/answer_key.json').write_text(json.dumps({'targets':{'S!Z99':7}}))
 rows=[]
 for cells in [['S!N140','S!O140','S!P140','S!Q140'],['S!N141']]:
  r=record(gold,cells);r.update(source='custom_method_detector',coverage_complete=True,method_profile={'complete':True});rows.append(r)
 return gold,p,task,rows


def test_original_full_faithcheck_accepts_complete_orthogonal_source_partition(tmp_path):
 gold,p,task,rows=fixture(tmp_path);span=['S!N140','S!N141']
 assert not disclose._heterogeneous_scope_partitioned(gold,span,rows,'method_operating_expense')
 assert disclose._orthogonal_scope_partitioned(gold,span,rows,'method_operating_expense',set())
 value=faithcheck(gold,p,task,rows);assert value['passed'],value


@pytest.mark.parametrize('change',['partial','gap','duplicate','wrong_label','wrong_entry','different_mechanics','unknown_source','unqualified','target','machine_keys','same_label'])
def test_no_incomplete_or_foreign_orthogonal_partition(tmp_path,change):
 gold,p,task,rows=fixture(tmp_path);span=['S!N140','S!N141'];closure=set()
 if change=='partial':rows[0]['cells'].pop();rows[0]['cell_keys'].pop()
 elif change=='gap':rows[0]['cells'].pop(1);rows[0]['cell_keys'].pop(1)
 elif change=='duplicate':rows.append(copy.deepcopy(rows[1]))
 elif change=='wrong_label':rows[0]['fields']['label']='"Vehicle Wrap"'
 elif change=='wrong_entry':rows[0]['entry']='method_returns'
 elif change=='different_mechanics':
  b=openpyxl.load_workbook(p);b['S']['O140']='=IF(O$3>$C$4,0,6000)';b.save(p);gold=disclose.Book(p)
 elif change=='unknown_source':rows[0]['source']='unverified'
 elif change=='unqualified':rows[0]['coverage_complete']=False
 elif change=='target':closure={'S!O140'}
 elif change=='machine_keys':rows[0]['cell_keys']=[]
 elif change=='same_label':
  b=openpyxl.load_workbook(p);b['S']['B141']='Replacement Capex';b.save(p);gold=disclose.Book(p);rows[1]['fields']['label']='"Replacement Capex"'
 assert not disclose._orthogonal_scope_partitioned(gold,span,rows,'method_operating_expense',closure)
