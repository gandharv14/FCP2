import copy
import hashlib
import openpyxl
import pytest
from openpyxl.worksheet.formula import ArrayFormula
from scripts.batch import harbor_disclosure_partition_evidence as E
from scripts.batch import harbor_semantic_readiness as S
from scripts.batch.harbor_admission import record


def workbook(tmp_path, text='=IFERROR(SUM((A1:A2)/A$3),"N/A")', ref='C1'):
    path=tmp_path/'array.xlsx';book=openpyxl.Workbook();sheet=book.active;sheet.title='Model'
    sheet['A1']=5;sheet['A2']=7;sheet['A3']=2;sheet['Z99']=901
    sheet['C1']=ArrayFormula(ref=ref,text=text);book.save(path);book.close()
    return path


def test_single_cell_array_retains_exact_source_and_dependency_closure(tmp_path):
    path=workbook(tmp_path);before=hashlib.sha256(path.read_bytes()).hexdigest()
    evidence=S.workbook_evidence(record(path));book=E._Workbook(path,evidence)
    try:
        shard=book.shard({'Model!C1'});decoded=E.decode_workbook(shard)
        assert shard['closure']=='exact_ast_closure'
        assert decoded['dependency_edges']==[['Model!C1','Model!A1'],['Model!C1','Model!A2'],['Model!C1','Model!A3']]
        original=next(c for c in evidence['sheets'][0]['cells']if c['address']=='C1')
        assert decoded['present']['Model!C1']==original
        assert 'Model!Z99' not in decoded['present'] and decoded['unknown_elsewhere']
    finally:book.close()
    assert hashlib.sha256(path.read_bytes()).hexdigest()==before


@pytest.mark.parametrize('text,ref',[('=INDIRECT("A1")','C1'),('=OFFSET(A1,0,0)','C1'),('=A1:A2','C1:C2')])
def test_dynamic_or_multiple_cell_arrays_retain_full_snapshot(tmp_path,text,ref):
    path=workbook(tmp_path,text,ref);book=E._Workbook(path,S.workbook_evidence(record(path)))
    try:
        shard=book.shard({'Model!C1'})
        assert shard['closure']=='full_snapshot_fallback'
        assert 'Model!Z99' in E.decode_workbook(shard)['present']
    finally:book.close()


@pytest.mark.parametrize('mutation',['other_owner','extra_attribute','extra_field','table','missing_text','non_text'])
def test_unrecognized_formula_objects_remain_unknown(tmp_path,mutation):
    path=workbook(tmp_path);book=E._Workbook(path,S.workbook_evidence(record(path)))
    value=book.cells['Model!C1']['value']
    if mutation=='other_owner':value['attributes']['ref']='D1'
    elif mutation=='extra_attribute':value['attributes']['aca']='1'
    elif mutation=='extra_field':value['other']='unrecognized'
    elif mutation=='table':value['formula_type']='dataTable'
    elif mutation=='missing_text':del value['text']
    else:value['text']=None
    try:assert book.dependencies('Model!C1')==(set(),True)
    finally:book.close()
