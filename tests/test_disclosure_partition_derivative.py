import ast,copy,json,shutil
from pathlib import Path
from types import SimpleNamespace as NS
import pytest
from openpyxl import Workbook
from scripts.batch import harbor_disclosure_partition_derivative as D
from scripts.batch import harbor_disclosure_review_round as Q
from scripts.batch import harbor_disclosure_round_adoption as A
from scripts.batch.harbor_admission import record
from scripts.fleet import disclosure_named_copy_repair as repair
from test_disclosure_named_copy import D as renderer


def save(path,value):
    path.parent.mkdir(parents=True,exist_ok=True);path.write_text(json.dumps(value))
    return record(path)


@pytest.fixture
def candidate(tmp_path):
    root=tmp_path/'code';skill=root/'.cursor/skills/task-disclosure';(skill/'scripts').mkdir(parents=True)
    shutil.copyfile(renderer.REGISTRY_PATH,skill/'REGISTRY.md')
    source=Path(renderer.__file__).read_text();nodes=ast.parse(source).body
    edits=[]
    for node in nodes:
        if isinstance(node,ast.FunctionDef)and node.name in ('copy_lock_note','_copied_named_operands'):
            edits.append((node.lineno-1,node.end_lineno,''))
    lines=source.splitlines(keepends=True)
    for a,b,text in sorted(edits,reverse=True):lines[a:b]=[text]
    oldtext=''.join(lines).replace('def _cell_copy_lock_note(', 'def copy_lock_note(')
    oldpath=skill/'scripts/disclose.py';oldpath.write_text(oldtext)
    old=Q.bound_module(record(oldpath));goldfile=tmp_path/'source.xlsx'
    wb=Workbook();ws=wb.active;ws.title='Debt Schedule';ws['F20']='=Circ+$D$14+F16';ws['G20']='=Circ+$D$14+G16';wb.save(goldfile)
    gold=old.Book(goldfile);cells=['Debt Schedule!F20','Debt Schedule!G20']
    rec=dict(entry='method_interest',value='out_of_catalogue',disposition='disclosed',cell_keys=cells,leak_flag=False,
        fields=dict(label='"Undrawn Fee"',band="'Debt Schedule'!F20:G20",representative="'Debt Schedule'!F20",
        calculation_kind='copied-column calculation',steps='Source calculation'+old.copy_lock_note(gold,cells)))
    payload=dict(records=[rec],agent_records=old.agent_records([rec]),audit={'untouched':True})
    stage=tmp_path/'original';save(stage/'tests/disclosure.json',payload)
    original_text='# Original task\n\n## Output\nReturn original results.\n'
    (stage/'instruction.md').write_text(original_text);shutil.copyfile(goldfile,stage/'delivered.xlsx')
    revised,changes=repair.corrected_payload(old,renderer,gold,payload)
    out=tmp_path/'candidate';shutil.copytree(stage,out);save(out/'tests/disclosure.json',revised)
    text=original_text.replace(renderer.OUTPUT_HEADING,'\n'+renderer.render_section(renderer.agent_records(revised['records']))+'\n## Output\n',1)
    (out/'instruction.md').write_text(text);bands=save(tmp_path/'bands.json',{'bands':[]})
    proof=dict(status='ISOLATED_SOURCE_DISCLOSURE_CANDIDATE',qualified=False,models=0,original_task_mutations=0,
        original_task=str(stage),candidate=str(out),repair=record(repair.__file__),retained_bands=bands,
        original_tree=D.tree(stage),candidate_tree=D.tree(out),candidate_instruction=record(out/'instruction.md'),
        candidate_disclosure=record(out/'tests/disclosure.json'),source=record(goldfile),changes=changes,
        changed_record_indices=[0],faith=dict(passed=True,faults=[]),verify=dict(passed=True,faults=[]))
    prior=[save(tmp_path/'prior.json',{'retained':'real rejected history is checked by Q'})];build={'original':'build'}
    p=dict(schema_version=D.SCHEMA,implementation=record(D.__file__),build=build,prior_rejections=prior,
        source=record(goldfile),original_renderer=record(oldpath),renderer=record(renderer.__file__),repair=record(repair.__file__),
        proof=save(tmp_path/'proof.json',proof),bands=bands)
    archive=tmp_path/'original.snapshot.md';archive.write_bytes((stage/'instruction.md').read_bytes());p['original_instruction']=record(archive)
    selection=NS(content=NS(stage=stage,bindings={'original_raw_source':record(goldfile),'source_file':record(goldfile)}))
    return tmp_path,p,NS(root=root),selection,prior,build


def check(data):
    home,p,w,s,prior,build=data
    return D.validate(save(home/'selection.json',p),worker=w,selection=s,build=build,prior_rejections=prior)


def test_complete_source_derivative_preserves_original_stage(candidate):
    _,p,_,s,_,_=candidate;before=D.tree(s.content.stage)
    _,proof,path=check(candidate)
    assert D.tree(s.content.stage)==before and path!=s.content.stage
    assert proof['changed_record_indices']==[0]


@pytest.mark.parametrize('defect',['source','build','history','original','candidate','mechanics','renderer'])
def test_derivative_rejects_foreign_or_changed_evidence(candidate,defect):
    home,p,w,s,prior,build=candidate
    if defect=='source':p['source']=save(home/'other.json',{})
    elif defect=='build':p['build']={'other':'build'}
    elif defect=='history':p['prior_rejections']=[]
    elif defect=='original':(s.content.stage/'instruction.md').write_text('changed')
    elif defect=='candidate':(home/'candidate/delivered.xlsx').write_bytes(b'changed')
    elif defect=='mechanics':
        proof=json.loads(Path(p['proof']['path']).read_text());proof['faith']['passed']=False;p['proof']=save(home/'failed-proof.json',proof)
    elif defect=='renderer':
        path=home/'foreign.py';path.write_text(Path(p['renderer']['path']).read_text().replace('the same name%s throughout the band','a fixed answer%s'))
        p['renderer']=record(path)
    with pytest.raises((ValueError,FileNotFoundError)):check(candidate)


def test_renderer_scope_restores_cache_and_paths_on_exception():
    before=(renderer.REPO_ROOT,renderer.REGISTRY_PATH,renderer._REGISTRY_CACHE)
    with pytest.raises(RuntimeError):
        with D.renderer_context(renderer,Path('/unused')):raise RuntimeError('rejected')
    assert (renderer.REPO_ROOT,renderer.REGISTRY_PATH,renderer._REGISTRY_CACHE)==before


def test_adoption_archives_original_bytes_instead_of_aliasing_candidate(candidate):
    home,p,w,s,prior,build=candidate;binding=save(home/'selection.json',p)
    authority=save(home/'authority.json',dict(source_derivative=binding))
    original=record(s.content.stage/'instruction.md')
    plan=dict(review_round=authority,instruction=record(home/'candidate/instruction.md'))
    assert A.instruction_archive(plan,original)==p['original_instruction']
    assert A.instruction_archive(plan,original)!=plan['instruction']


def test_owning_adoption_copies_only_reviewed_candidate_to_new_stage(candidate,monkeypatch):
    from dataclasses import dataclass
    from scripts.batch import harbor_instruction_worker as W
    home,p,w,s,prior,build=candidate;binding=save(home/'selection.json',p)
    authority=save(home/'authority.json',dict(source_derivative=binding,build=build,prior_rejections=prior))
    before=D.tree(s.content.stage);_,proof,path=check(candidate)
    @dataclass(frozen=True)
    class Content:
        stage:Path
    @dataclass(frozen=True)
    class Selection:
        content:Content
    selected=Selection(Content(s.content.stage));calls=[]
    def checked(*args):calls.append(args);return p,proof,path
    monkeypatch.setattr(Q,'reviewed_content',checked)
    monkeypatch.setattr(W,'refreshed',lambda value:value)
    directory=home/'owning-step';directory.mkdir()
    result=A.adopted_derivative(w,selected,dict(source_derivative=binding,aggregate={'actual':'passed aggregate'}),
        dict(review_round=authority),directory)
    assert len(calls)==1 and result.content.stage==directory/'source-derivative-stage'
    assert D.tree(result.content.stage)==proof['candidate_tree'] and D.tree(s.content.stage)==before
    assert result.content.stage!=path
    with pytest.raises(ValueError,match='already exists'):
        A.adopted_derivative(w,selected,dict(source_derivative=binding),dict(review_round=authority),directory)


def test_prior_partition_directory_cannot_be_reset_or_extended(tmp_path):
    from scripts.fleet import disclosure_review_prepare as P
    w=NS(run=tmp_path/'worker');(w.run/'disclosure-review-rounds/0002').mkdir(parents=True)
    first=save(tmp_path/'first.json',{'status':'http_rejected'})
    second=save(tmp_path/'second.json',{'schema_version':'harbor-completed-disclosure-partition-rejection/v1'})
    P.unused_round_directory(w,3,[first,second])
    (w.run/'disclosure-review-rounds/0003').mkdir()
    with pytest.raises(ValueError):P.unused_round_directory(w,3,[first,second])
    with pytest.raises(ValueError):P.unused_round_directory(w,1,[])


def derivative_request_case(candidate):
    from scripts.batch import harbor_disclosure_partition_review as R
    home,d,w,s,prior,build=candidate
    binding=save(home/'selection.json',d);_,proof,path=check(candidate)
    inputs=dict(instruction=record(path/'instruction.md'),disclosure=record(path/'tests/disclosure.json'),
                golden=proof['source'],delivered=record(path/'delivered.xlsx'),original_source=d['source'])
    authority=save(home/'request-authority.json',dict(inputs=inputs,source_derivative=binding,build=build,prior_rejections=prior))
    p=dict(inputs,review_round=authority)
    task=NS(stage=s.content.stage,stage_tree=D.tree(s.content.stage))
    before=dict(bindings=s.content.bindings)
    return R,task,p,before,home


def test_first_request_accepts_proved_candidate_without_mutating_live_stage(candidate):
    R,task,p,before,home=derivative_request_case(candidate)
    original={x:x.read_bytes()for x in task.stage.rglob('*')if x.is_file()}
    assert record(task.stage/'instruction.md')['sha256']!=p['instruction']['sha256']
    R.task_content(task,p,before)
    assert all(x.read_bytes()==b for x,b in original.items())


@pytest.mark.parametrize('fault',['original','candidate','financial_artifact','foreign_source','foreign_inputs','original_archive'])
def test_derivative_request_rejects_changed_live_or_reviewed_content(candidate,fault):
    R,task,p,before,home=derivative_request_case(candidate)
    if fault=='original':(task.stage/'instruction.md').write_text('changed')
    elif fault=='candidate':(home/'candidate/instruction.md').write_text('changed')
    elif fault=='financial_artifact':(home/'candidate/delivered.xlsx').write_bytes(b'changed')
    elif fault=='foreign_source':before['bindings']=dict(before['bindings'],original_raw_source=save(home/'foreign.json',{}))
    elif fault=='foreign_inputs':p['instruction']=record(task.stage/'instruction.md')
    elif fault=='original_archive':(home/'original.snapshot.md').write_text('changed')
    with pytest.raises(ValueError):R.task_content(task,p,before)


def test_original_request_content_rule_is_unchanged(candidate):
    R,task,p,before,home=derivative_request_case(candidate)
    p.pop('review_round')
    with pytest.raises(ValueError,match='partition task content changed'):R.task_content(task,p,before)
    p.update(instruction=record(task.stage/'instruction.md'),disclosure=record(task.stage/'tests/disclosure.json'))
    R.task_content(task,p,before)
