import copy
from pathlib import Path
from types import SimpleNamespace
import pytest
import openpyxl
from scripts.batch import harbor_disclosure_partition_evidence as E
from scripts.batch import harbor_semantic_readiness as S
from scripts.batch.harbor_admission import record

ROOT=Path(__file__).resolve().parents[1]

def body_fixture(tmp_path,monkeypatch,count=3,formula='=SUM(A1:A3)'):
    paths=[]
    for name in ('golden','delivered'):
        p=tmp_path/(name+'.xlsx');w=openpyxl.Workbook();ws=w.active;ws.title='Model';ws['A1']=5;ws['A3']=7;ws['B1']='Revenue';ws['B2']='USD millions';ws['C1']=formula;w.save(p);w.close();paths.append(p)
    rows=[dict(band=f'b{i:03d}',cells=[f'b{i:03d}'],cell_keys=['Model!C1'],source_cell='Model!A1') for i in range(count)]
    renderer=SimpleNamespace(compact_cells=lambda cells:cells[0],agent_records=lambda rows:copy.deepcopy(rows),render_section=lambda rows:'Disclosure\n'+'\n'.join('- '+r['band']+': calculation' for r in sorted(rows,key=lambda x:x['band']))+'\n')
    monkeypatch.setattr(S,'skill_module',lambda *a:renderer)
    ids=['record-'+S._hash(S._canonical(r))[:24] for r in rows];sentences=[]
    for i,row in enumerate(rows):
        s=dict(order=i,text='- '+row['band']+': calculation',record_ids=[ids[i]]);sentences.append({'sentence_id':'sentence-'+S._hash(S._canonical(s))[:24],**s})
    manifest=dict(schema_version='harbor-disclosure-sentence-manifest/v1',input_hashes={'renderer':record(ROOT/'.cursor/skills/task-disclosure/scripts/disclose.py')['sha256'],'instruction':E.hashlib.sha256(renderer.render_section(rows).encode()).hexdigest(),'disclosure':'b'*64,'golden':record(paths[0])['sha256'],'delivered':record(paths[1])['sha256']},agent_record_ids=ids,sentences=sentences)
    body=dict(manifest=manifest,manifest_sha256='a'*64,instruction=renderer.render_section(rows),disclosure=dict(records=rows,agent_records=copy.deepcopy(rows),method_assessments=[]),workbooks=dict(golden=S.workbook_evidence(record(paths[0])),delivered=S.workbook_evidence(record(paths[1]))))
    return body,paths

def test_source_ast_blank_closure_labels_and_complete_union(tmp_path,monkeypatch):
    body,paths=body_fixture(tmp_path,monkeypatch)
    plan=E.plan(body,*paths,root=ROOT,max_sentences=2)
    assert [len(x['manifest']['sentences']) for x in plan['batches']]==[2,1]
    shard=plan['batches'][0]['workbooks']['golden'];decoded=E.decode_workbook(shard)
    assert 'Model!A2' in decoded['blank_refs']
    assert {'Model!A1','Model!A3','Model!B1','Model!B2','Model!C1'}<=decoded['present'].keys()
    assert ['Model!C1','Model!A2'] in decoded['dependency_edges']
    assert E.validate_plan(plan,body,*paths,root=ROOT)==plan


def test_redacted_delivered_formula_keeps_its_own_dependency_evidence(tmp_path,monkeypatch):
    body,paths=body_fixture(tmp_path,monkeypatch,1,formula='=SUM(D1:D3)')
    for name,path in zip(('golden','delivered'),paths):
        w=openpyxl.load_workbook(path);ws=w.active
        ws['D1']=10 if name=='golden'else 12
        ws['D3']=30 if name=='golden'else None
        if name=='delivered':ws['C1']=None
        w.save(path);w.close();body['workbooks'][name]=S.workbook_evidence(record(path));body['manifest']['input_hashes'][name]=record(path)['sha256']
    p=E.plan(body,*paths,root=ROOT);pair={k:E.decode_workbook(v)for k,v in p['batches'][0]['workbooks'].items()}
    assert pair['golden']['selected_refs']==pair['delivered']['selected_refs']
    assert pair['golden']['present']['Model!D1']['value']==10
    assert pair['delivered']['present']['Model!D1']['value']==12
    assert {'Model!C1','Model!D2','Model!D3'}<=set(pair['delivered']['blank_refs'])
    assert pair['golden']['present']['Model!D3']['value']==30
    assert not any(edge[0]=='Model!C1'for edge in pair['delivered']['dependency_edges'])
    assert 'Model!Z999'not in pair['delivered']['selected_refs']and pair['delivered']['unknown_elsewhere']
    assert E.validate_plan(p,body,*paths,root=ROOT)==p
    with pytest.raises(E.ScopeDiagnosisRequired):E.plan(body,*paths,root=ROOT,max_bytes=100)


def test_paired_closure_reaches_dependencies_alternating_between_sources(tmp_path,monkeypatch):
    body,paths=body_fixture(tmp_path,monkeypatch,1,formula='=D1')
    for name,path in zip(('golden','delivered'),paths):
        w=openpyxl.load_workbook(path);ws=w.active
        if name=='golden':ws['D1']=5;ws['E1']='=F1';ws['F1']=17
        else:ws['C1']=None;ws['D1']='=E1';ws['E1']=9
        w.save(path);w.close();body['workbooks'][name]=S.workbook_evidence(record(path));body['manifest']['input_hashes'][name]=record(path)['sha256']
    p=E.plan(body,*paths,root=ROOT);pair={k:E.decode_workbook(v)for k,v in p['batches'][0]['workbooks'].items()}
    assert pair['golden']['selected_refs']==pair['delivered']['selected_refs']
    assert pair['golden']['present']['Model!F1']['value']==17
    assert 'Model!F1'in pair['delivered']['blank_refs']
    assert ['Model!D1','Model!E1']in pair['delivered']['dependency_edges']
    assert ['Model!E1','Model!F1']in pair['golden']['dependency_edges']

@pytest.mark.parametrize('mutation',['blank','present','edge','sentence','coverage','extra'])
def test_projection_tampering_rejected(tmp_path,monkeypatch,mutation):
    body,paths=body_fixture(tmp_path,monkeypatch);p=E.plan(body,*paths,root=ROOT,max_sentences=2)
    shard=p['batches'][0]['workbooks']['golden']
    if mutation=='blank':shard['blank_refs']=E.encode_refs([])
    elif mutation=='present':shard['sheets'][0]['present_rows'][0][1]=88
    elif mutation=='edge':shard['dependency_edges']=[]
    elif mutation=='sentence':p['batches'][0]['manifest']['sentences'].pop()
    elif mutation=='coverage':p['coverage']['record_ids'].pop()
    else:p['forged']=True
    with pytest.raises(ValueError):E.validate_plan(p,body,*paths,root=ROOT)

def test_source_mutation_rejected(tmp_path,monkeypatch):
    body,paths=body_fixture(tmp_path,monkeypatch);p=E.plan(body,*paths,root=ROOT)
    w=openpyxl.load_workbook(paths[0]);w.active['A1']=99;w.save(paths[0]);w.close()
    with pytest.raises(ValueError,match='full workbook evidence'):E.validate_plan(p,body,*paths,root=ROOT)

def test_dynamic_scope_retains_full_snapshots_or_fails_size(tmp_path,monkeypatch):
    body,paths=body_fixture(tmp_path,monkeypatch,1,'=INDIRECT("A1")')
    p=E.plan(body,*paths,root=ROOT)
    assert all(s['closure']=='full_snapshot_fallback' for s in p['batches'][0]['workbooks'].values())
    with pytest.raises(E.ScopeDiagnosisRequired):E.plan(body,*paths,root=ROOT,max_bytes=100)

def test_duplicate_manifest_and_decoder_extra_fields_rejected(tmp_path,monkeypatch):
    body,paths=body_fixture(tmp_path,monkeypatch);p=E.plan(body,*paths,root=ROOT)
    shard=p['batches'][0]['workbooks']['golden'];shard['extra']=1
    with pytest.raises(ValueError):E.decode_workbook(shard)
    body['manifest']['sentences'].append(body['manifest']['sentences'][0])
    with pytest.raises(ValueError,match='manifest'):E.plan(body,*paths,root=ROOT)

def test_cyclic_dependencies_terminate_and_are_retained(tmp_path,monkeypatch):
    body,paths=body_fixture(tmp_path,monkeypatch,1,'=D1')
    for name,p in zip(('golden','delivered'),paths):
        w=openpyxl.load_workbook(p);w.active['D1']='=C1';w.save(p);w.close();body['workbooks'][name]=S.workbook_evidence(record(p));body['manifest']['input_hashes'][name]=record(p)['sha256']
    plan=E.plan(body,*paths,root=ROOT)
    edges=E.decode_workbook(plan['batches'][0]['workbooks']['golden'])['dependency_edges']
    assert ['Model!C1','Model!D1'] in edges and ['Model!D1','Model!C1'] in edges


def test_static_named_reference_retains_exact_definition_and_complete_area(tmp_path,monkeypatch):
    from openpyxl.workbook.defined_name import DefinedName
    body,paths=body_fixture(tmp_path,monkeypatch,1,'=NamedInputs')
    for name,p in zip(('golden','delivered'),paths):
        w=openpyxl.load_workbook(p);w.defined_names.add(DefinedName('NamedInputs',attr_text='Model!$A$1:$A$3'));w.save(p);w.close();body['workbooks'][name]=S.workbook_evidence(record(p));body['manifest']['input_hashes'][name]=record(p)['sha256']
    plan=E.plan(body,*paths,root=ROOT)
    shard=plan['batches'][0]['workbooks']['golden']
    assert shard['closure']=='exact_ast_closure'
    assert 'Model!A2' in E.decode_workbook(shard)['blank_refs']
    assert shard['snapshot_metadata']['resolved_defined_names'][0]['definition']=='Model!$A$1:$A$3'
    assert E.validate_plan(plan,body,*paths,root=ROOT)==plan
    shard['snapshot_metadata']['resolved_defined_names'][0]['definition']='Model!$A$99'
    with pytest.raises(ValueError,match='differs'):E.validate_plan(plan,body,*paths,root=ROOT)


def test_cached_presence_and_unknown_addresses_are_not_conflated(tmp_path,monkeypatch):
    body,paths=body_fixture(tmp_path,monkeypatch,1)
    p=E.plan(body,*paths,root=ROOT);decoded=E.decode_workbook(p['batches'][0]['workbooks']['golden'])
    assert 'cached_value' in decoded['present']['Model!C1']
    assert 'cached_value' not in decoded['present']['Model!A1']
    assert 'Model!Z999' not in decoded['present'] and 'Model!Z999' not in decoded['blank_refs']
    assert decoded['unknown_elsewhere'] is True

def test_manifest_source_hash_and_reference_run_tampering_rejected(tmp_path,monkeypatch):
    body,paths=body_fixture(tmp_path,monkeypatch)
    body['manifest']['input_hashes']['golden']='0'*64
    with pytest.raises(ValueError,match='manifest inputs'):E.plan(body,*paths,root=ROOT)
    run=E.encode_refs(['Model!A1','Model!A2','Model!A3'])
    assert run['runs']==[['Model',1,1,3]] and E.decode_refs(run)==['Model!A1','Model!A2','Model!A3']
    run['runs'].append(['Model',1,2,1])
    with pytest.raises(ValueError):E.decode_refs(run)


def test_instruction_projection_preserves_context_and_assessments_scope(tmp_path,monkeypatch):
    body,paths=body_fixture(tmp_path,monkeypatch)
    body['instruction']='Original task context\n'+body['instruction']+'\nOriginal output instructions'
    body['manifest']['input_hashes']['instruction']=E.hashlib.sha256(body['instruction'].encode()).hexdigest()
    body['disclosure']['method_assessments']=[{'cell_keys':['Model!C1'],'status':'selected'},{'cell_keys':['Model!A3'],'status':'upstream_only'},{'status':'unscoped'}]
    p=E.plan(body,*paths,root=ROOT,max_sentences=1)
    first=p['batches'][0]
    assert 'Original task context' in first['instruction'] and 'Original output instructions' in first['instruction']
    assert '- b000' in first['instruction'] and '- b001' not in first['instruction']
    assert [x['status'] for x in first['disclosure']['method_assessments']]==['selected','unscoped']


def test_oversized_literal_range_never_reaches_expanding_resolver(tmp_path,monkeypatch):
    from xl_ast_graph import AstGraph
    body,paths=body_fixture(tmp_path,monkeypatch,1,'=SUM(A1:XFD1048576)')
    original=AstGraph.resolve
    def bounded(self,reference,sheet):
        if 'XFD1048576' in reference:raise AssertionError('oversized range was expanded')
        return original(self,reference,sheet)
    monkeypatch.setattr(AstGraph,'resolve',bounded)
    result=E.plan(body,*paths,root=ROOT)
    assert all(s['closure']=='full_snapshot_fallback' for s in result['batches'][0]['workbooks'].values())
