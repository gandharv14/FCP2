import hashlib
import openpyxl
import pytest
from openpyxl.workbook.defined_name import DefinedName
from scripts.batch import harbor_disclosure_partition_evidence as P
from scripts.batch import harbor_semantic_readiness as S
from scripts.batch.harbor_admission import record


def source(tmp_path, selector=1, definition="'Control'!$B$1", formula='=+INDEX(A:A,ROW(A50)+(CASE))'):
    path=tmp_path/'index.xlsx';w=openpyxl.Workbook();s=w.active;s.title='Model'
    s['N2']=formula;s['A50']='COGS';s['A51']='Case 1';s['A52']='Case 2';s['Z99']=987
    control=w.create_sheet('Control');control['B1']=selector
    w.defined_names.add(DefinedName('CASE',attr_text=definition));w.save(path);w.close()
    return path


def test_literal_named_selector_retains_definition_formula_and_dependencies(tmp_path):
    path=source(tmp_path);original=record(path);evidence=S.workbook_evidence(original)
    book=P._Workbook(path,evidence)
    try:
        refs,bad=book.dependencies('Model!N2')
        assert not bad and refs=={'Model!A50','Model!A51','Control!B1'}
        shard=book.shard({'Model!N2'});v=P.decode_workbook(shard)
        assert shard['closure']=='exact_ast_closure'
        assert v['present']['Model!N2']['value']=='=+INDEX(A:A,ROW(A50)+(CASE))'
        assert v['present']['Control!B1']['value']==1
        assert 'Model!Z99' not in v['present']
        assert shard['snapshot_metadata']['resolved_defined_names'][0]['definition']=="'Control'!$B$1"
        assert ['Model!N2','Control!B1'] in v['dependency_edges']
    finally:book.close()
    assert record(path)==original


@pytest.mark.parametrize('selector,definition,formula',[
    (None,"'Control'!$B$1",'=INDEX(A:A,ROW(A50)+CASE)'),
    ('1',"'Control'!$B$1",'=INDEX(A:A,ROW(A50)+CASE)'),
    (True,"'Control'!$B$1",'=INDEX(A:A,ROW(A50)+CASE)'),
    ('=1',"'Control'!$B$1",'=INDEX(A:A,ROW(A50)+CASE)'),
    (1.5,"'Control'!$B$1",'=INDEX(A:A,ROW(A50)+CASE)'),
    (-50,"'Control'!$B$1",'=INDEX(A:A,ROW(A50)+CASE)'),
    (1048576,"'Control'!$B$1",'=INDEX(A:A,ROW(A50)+CASE)'),
    (1,"'Control'!$B$1:$B$2",'=INDEX(A:A,ROW(A50)+CASE)'),
    (1,'OFFSET(Control!$B$1,0,0)','=INDEX(A:A,ROW(A50)+CASE)'),
    (1,"'Control'!$B$1",'=INDEX(A:B,ROW(A50)+CASE)'),
    (1,"'Control'!$B$1",'=INDEX(A:A,ROW(A50)*CASE)'),
])
def test_unknown_or_non_scalar_selectors_keep_full_snapshot(tmp_path,selector,definition,formula):
    path=source(tmp_path,selector,definition,formula);book=P._Workbook(path,S.workbook_evidence(record(path)))
    try:
        shard=book.shard({'Model!N2'})
        assert shard['closure']=='full_snapshot_fallback'
        assert 'Model!Z99' in P.decode_workbook(shard)['present']
    finally:book.close()


def test_new_source_selector_is_rederived_and_stale_evidence_rejected(tmp_path):
    path=source(tmp_path);old=S.workbook_evidence(record(path))
    w=openpyxl.load_workbook(path);w['Control']['B1']=2;w.save(path);w.close()
    with pytest.raises(ValueError,match='full workbook evidence'):
        P._Workbook(path,old)
    book=P._Workbook(path,S.workbook_evidence(record(path)))
    try:assert book.dependencies('Model!N2')==({'Model!A50','Model!A52','Control!B1'},False)
    finally:book.close()
