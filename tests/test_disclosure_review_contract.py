"""Empty disclosures still receive a real, source-complete independent review."""
import copy,json
from pathlib import Path
from types import SimpleNamespace
import openpyxl,pytest
from scripts.batch import harbor_disclosure_review_contract as C
from scripts.batch import harbor_disclosure_partition_evidence as E
from scripts.batch import harbor_disclosure_partition_review as R
from scripts.batch import harbor_disclosure_review_round as Q
from scripts.batch import harbor_semantic_readiness as S
from scripts.batch import harbor_transition_codec as codec,harbor_instruction_worker as W
from scripts.batch.harbor_admission import record
from scripts.batch.harbor_model_review_provenance import Policy
from test_disclosure_partition_evidence import body_fixture,ROOT


def write(path,value):
    path.write_text(json.dumps(value));return record(path)


def fixture(tmp_path,monkeypatch):
    body,paths=body_fixture(tmp_path,monkeypatch,count=0)
    renderer=S.skill_module(None);renderer.render_section=lambda rows:''
    body['instruction']='Reconstruct graded outputs from visible inputs.';body['manifest']['input_hashes']['instruction']=E.hashlib.sha256(body['instruction'].encode()).hexdigest()
    for name,path in zip(('golden','delivered'),paths):
        book=openpyxl.load_workbook(path);sheet=book.active
        if name=='delivered':sheet['C1']=None
        sheet['Z9']='=A1*2' # Nongraded rule outside the target dependency closure.
        book.save(path);book.close();body['workbooks'][name]=S.workbook_evidence(record(path));body['manifest']['input_hashes'][name]=record(path)['sha256']
    body['disclosure']['method_assessments']=[{'cell_keys':['Model!Z9'],'status':'nongraded'}]
    stage=tmp_path/'stage';(stage/'tests').mkdir(parents=True)
    (stage/'instruction.md').write_text(body['instruction']);write(stage/'tests/disclosure.json',body['disclosure']);write(stage/'tests/answer_key.json',{'targets':{'Model!C1':{'answer':'PRIVATE-GRADED-ANSWER-NEVER-SENT'}}})
    inventory={name:{'algorithm':'sha256','path':name,'sha256':record(stage/name)['sha256'],'size_bytes':(stage/name).stat().st_size}for name in ['instruction.md','tests/disclosure.json','tests/answer_key.json']}
    content=SimpleNamespace(stage=stage,stage_sha256=W.static.publication.canonical_hash(inventory))
    monkeypatch.setattr(codec,'decode',lambda x:SimpleNamespace(content=content))
    auth=tmp_path/'authorization';auth.write_text('Explicit additional independent review');monkeypatch.setattr(C,'AUTH',record(auth)['sha256'])
    p=dict(schema_version=C.CONTRACT,implementation=record(C.__file__),authorization=record(auth),build_result=write(tmp_path/'build.json',{'selection':{},'status':'STEP_COMPLETE','next_step':'disclosure-review'}),instruction=record(stage/'instruction.md'),disclosure=record(stage/'tests/disclosure.json'),answer_key=record(stage/'tests/answer_key.json'),stage_manifest=write(tmp_path/'stage.json',inventory),golden=record(paths[0]),delivered=record(paths[1]),targets=['Model!C1'])
    binding=write(tmp_path/'contract.json',p);original=Q.bound_module
    monkeypatch.setattr(Q,'bound_module',lambda row:C if row==record(C.__file__)else original(row))
    return body,paths,p,binding


def test_empty_manifest_requires_explicit_contract(tmp_path,monkeypatch):
    body,paths=body_fixture(tmp_path,monkeypatch,count=0)
    renderer=S.skill_module(None);renderer.render_section=lambda rows:''
    body['instruction']='Reconstruct graded outputs from visible inputs.';body['manifest']['input_hashes']['instruction']=E.hashlib.sha256(body['instruction'].encode()).hexdigest()
    with pytest.raises(ValueError,match='explicit full-source'):E.plan(body,*paths,root=ROOT)


def test_one_real_batch_keeps_full_source_and_excludes_answer_values(tmp_path,monkeypatch):
    body,paths,p,binding=fixture(tmp_path,monkeypatch)
    projection=E.plan(body,*paths,root=ROOT,empty_review_contract=binding)
    assert len(projection['batches'])==1
    batch=projection['batches'][0]
    assert batch['manifest']['sentences']==[] and batch['disclosure']==body['disclosure']
    for name,shard in batch['workbooks'].items():
        assert shard['closure']=='full_snapshot_fallback'
        decoded=E.decode_workbook(shard)
        assert decoded['present']['Model!Z9']['value']=='=A1*2'
        assert decoded['present']['Model!A1']['value']==5
        if name=='delivered':assert 'Model!C1'in decoded['blank_refs']
    request=R.payload({},projection,0,Policy('fixture-project'))
    assert 'PRIVATE-GRADED-ANSWER-NEVER-SENT'not in json.dumps(request)
    assert 'Reject missing nongraded source rules' in request['messages'][0]['content']
    assert E.validate_plan(projection,body,*paths,root=ROOT)==projection
    with pytest.raises(E.ScopeDiagnosisRequired):E.plan(body,*paths,root=ROOT,empty_review_contract=binding,max_bytes=100)


@pytest.mark.parametrize('change',['authorization','targets','build','inventory','instruction','answer_key','delivered','source','extra','nonempty'])
def test_changed_contract_or_source_fails_before_any_call(tmp_path,monkeypatch,change):
    body,paths,p,binding=fixture(tmp_path,monkeypatch)
    if change=='authorization':p['authorization']=p['instruction']
    elif change=='targets':p['targets']=['Model!Z9']
    elif change=='build':p['build_result']=write(tmp_path/'bad-build.json',{'selection':{},'status':'PUBLICATION_READY','next_step':'enqueue'})
    elif change=='inventory':p['stage_manifest']=write(tmp_path/'bad-stage.json',{})
    elif change in ('instruction','answer_key'):Path(p[change]['path']).write_bytes(b'changed')
    elif change=='source':
        book=openpyxl.load_workbook(paths[0]);book.active['A1']=99;book.save(paths[0]);book.close()
    elif change=='delivered':
        book=openpyxl.load_workbook(paths[1]);book.active['C1']='=A1';book.save(paths[1]);book.close();p['delivered']=record(paths[1]);body['workbooks']['delivered']=S.workbook_evidence(p['delivered']);body['manifest']['input_hashes']['delivered']=p['delivered']['sha256']
    elif change=='extra':p['waive_coverage']=True
    else:body['disclosure']['records']=[{'invented':'record'}]
    binding=write(tmp_path/'contract-2.json',p)
    with pytest.raises((ValueError,KeyError)):E.plan(body,*paths,root=ROOT,empty_review_contract=binding)


def test_empty_response_still_rejects_missing_nongraded_coverage(tmp_path,monkeypatch):
    body,paths,p,binding=fixture(tmp_path,monkeypatch);projection=E.plan(body,*paths,root=ROOT,empty_review_contract=binding);batch=projection['batches'][0]
    review=dict(schema_version='harbor-disclosure-faithfulness/v1',model='gpt-5.6-sol-high',input_hashes=batch['manifest']['input_hashes'],manifest_sha256=batch['manifest_sha256'],sentences=[],blocking_findings=['Missing nongraded source rule'],non_blocking_flags=[],passed=False)
    with pytest.raises(ValueError):S.validate_disclosure_review(review,batch['manifest'],batch['manifest_sha256'])
    review.update(passed=True,blocking_findings=[])
    S.validate_disclosure_review(review,batch['manifest'],batch['manifest_sha256'])


from dataclasses import asdict,replace
from test_fleet_model_gate import case as gate_case
from test_harbor_model_executor import Response,response_bytes
from test_disclosure_review_round import policy as allowance_policy
from scripts.batch import harbor_model_executor as executor,harbor_continuation_allowance as allowance


def rejected_call(gate_case,monkeypatch,tmp_path,*,passed=False):
    coordinator,campaign,prepared=gate_case
    manifest=dict(input_hashes={'source':'a'*64},agent_record_ids=[],sentences=[])
    body=dict(manifest=manifest,manifest_sha256='b'*64)
    review=dict(schema_version='harbor-disclosure-faithfulness/v1',model='gpt-5.6-sol-high',input_hashes=manifest['input_hashes'],manifest_sha256=body['manifest_sha256'],sentences=[],blocking_findings=[]if passed else['Missing nongraded rule'],non_blocking_flags=[],passed=passed)
    request=json.loads(Path(prepared.request['path']).read_text());request['messages'][1]['content']=json.dumps(body)
    bound=write(Path(prepared.request['path']),request)
    prepared=replace(prepared,request=bound,scope=executor.SemanticScope('0001','a'*64,'harbor-disclosure-writer','disclosure-review'))
    monkeypatch.setattr(executor,'_open',lambda *args:Response(response_bytes(prepared.policy,review)))
    actual=executor.execute(prepared,campaign,credential='fixture-only');assert actual.status=='completed'
    root=executor.location(prepared,campaign)
    row=write(tmp_path/'semantic-rejection.json',dict(schema_version=C.SCHEMA,implementation=record(C.__file__),execution=record(root/'plan.json'),outcome=record(root/'outcome.json')))
    return campaign,prepared,row,root


def test_real_rejected_call_reserves_one_new_round_preserving_original(gate_case,monkeypatch,tmp_path):
    campaign,prepared,row,root=rejected_call(gate_case,monkeypatch,tmp_path)
    original={p.name:p.read_bytes()for p in root.iterdir()};scope=asdict(prepared.scope)
    assert C.closed_rejection(row,scope,1)
    policy=allowance_policy(tmp_path,'a'*64);ledger=Path(policy['path']).parent/'review-reservations'
    selected=allowance.reserve_review(ledger,allowance=policy,semantic_scope=scope,prior_records=[row],original_limit=1,expected_ordinal=2,cumulative_limit=2,prove_closed=Q.closed_rejection)
    assert Q.validate_reservation(selected,scope,2,2)['prior_records']==[row]
    Q._primary_history(campaign,scope,2,[row])
    assert original=={p.name:p.read_bytes()for p in root.iterdir()}
    with pytest.raises((ValueError,FileExistsError)):allowance.reserve_review(ledger,allowance=policy,semantic_scope=scope,prior_records=[row],original_limit=1,expected_ordinal=2,cumulative_limit=2,prove_closed=Q.closed_rejection)


@pytest.mark.parametrize('fault',['passed','missing_response','scope','ordinal','implementation','outcome','request'])
def test_unproven_or_foreign_rejection_never_authorizes_more_calls(gate_case,monkeypatch,tmp_path,fault):
    campaign,prepared,row,root=rejected_call(gate_case,monkeypatch,tmp_path,passed=fault=='passed');scope=asdict(prepared.scope);ordinal=1
    if fault=='missing_response':(root/'response.raw.json').unlink()
    elif fault=='scope':scope['workbook_id']='0002'
    elif fault=='ordinal':ordinal=2
    elif fault=='request':(root/'request.json').write_text('{}')
    elif fault in ('implementation','outcome'):
        data=json.loads(Path(row['path']).read_text());data[fault]=record(root/'plan.json');row=write(tmp_path/'changed.json',data)
    with pytest.raises((ValueError,FileNotFoundError)):C.closed_rejection(row,scope,ordinal)
