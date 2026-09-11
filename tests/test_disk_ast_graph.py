"""Full graph equivalence and bounded temporary storage; no private corpus."""
import json
from pathlib import Path
import random
import tracemalloc

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula,DataTableFormula
import pytest

from xl_ast_graph import AstGraph,main,pick_default_focus,write_csv,write_json,write_graphml
from xl_graph_store import GraphStore


def workbook(path,case):
    w=openpyxl.Workbook();s=w.active;s.title="3 Statement's Data"
    s['A1']=3;s['A2']=4;s['B1']='Revenue';s['B2']=.1
    other=w.create_sheet('Results 2026')
    cases={
        'normal':['=A1*2+SUM(A1:A2)','=SQRT(A2)+SLN(10,2,4)',"='Results 2026'!A1+1"],
        'cycles':['=D2+1','=D1-1','=D3'],
        'range':['=SUM(A1:A200)','=SUM(INDEX(A1:A2,1):INDEX(A1:A2,2))','=SUM(A:A)'],
        'dynamic':['=INDIRECT("A"&"1")','=OFFSET(A1,1,0)','=IF(A1>0,A2,B2)'],
        'failed':['=UNKNOWN_FN(A1)','=A1+','=UnknownName+1'],
        'external':["='[other.xlsx]Data'!A1+2",'=A1/0','=TRUE'],
        'array':['=COLUMN(A1:B1)','=A1+2','=A2+3'],
        'data_table':['=A1+2','=A2+3','=A1*A2'],
    }
    for row,f in enumerate(cases[case],1):s.cell(row,4,f)
    other['A1']=2;other['B1']="='3 Statement''s Data'!D1"
    if case=='array':s['F1']=ArrayFormula(ref='F1:G1',text='=COLUMN(A1:B1)')
    if case=='data_table':s['F1']=DataTableFormula(ref='F1:G2',r1='A1')
    w.save(path);w.close()


@pytest.mark.parametrize('case',['normal','cycles','range','dynamic','failed','external','array','data_table'])
@pytest.mark.parametrize('values',[True,False])
def test_persisted_graph_and_scc_exactly_match_memory(tmp_path,case,values):
    path=tmp_path/'case.xlsx';workbook(path,case)
    with AstGraph(path,graph_storage='memory',verbose=False,read_values=values,max_range_expand=2) as memory, \
         AstGraph(path,graph_storage='disk',verbose=False,read_values=values,max_range_expand=2,storage_dir=tmp_path) as disk:
        memory.build();disk.build()
        assert {key:dict(value) for key,value in disk.nodes.items()}==memory.nodes
        assert [dict(value) for value in disk.edges]==memory.edges
        assert [list(c) for c in disk.cycles]==memory.cycles
        assert set(disk.cyclic_nodes)==memory.cyclic_nodes
        assert disk.stats(disk.nodes,disk.edges)==memory.stats(memory.nodes,memory.edges)
        assert pick_default_focus(disk.nodes,disk.edges)==pick_default_focus(memory.nodes,memory.edges)
        for name,graph in [('memory',memory),('disk',disk)]:
            directory=tmp_path/name;directory.mkdir()
            write_csv(directory,graph.nodes,graph.edges)
            write_json(directory,graph.nodes,graph.edges,{'same':'statistics','nested':{'x':[1,2]}})
            write_graphml(directory,graph.nodes,graph.edges)
        for filename in ('nodes.csv','edges.csv','graph.json','graph.graphml'):
            assert (tmp_path/'memory'/filename).read_bytes()==(tmp_path/'disk'/filename).read_bytes()
        root=disk._storage.root
    assert not root.exists()


@pytest.mark.parametrize('seed',range(20))
def test_scc_order_and_cycle_flags_match_iterative_tarjan(tmp_path,seed):
    rng=random.Random(seed);size=25
    memory=AstGraph('synthetic.xlsx',verbose=False,graph_storage='memory')
    with GraphStore(tmp_path,cache_records=3) as store:
        for i in range(size):
            node={'id':str(i),'kind':'formula','sheet':'S'}
            memory.nodes[str(i)]=dict(node);store.nodes[str(i)]=node
        for i in range(90):
            source,target=str(rng.randrange(size)),str(rng.randrange(size))
            memory._add_edge(source,target,'value',0,'+','', '',source,'S')
            edge=dict(source=source,target=target,role='value',arg_index=0,op='+',ref='',via_range='',cell=source,cross_sheet=False)
            store.add_edge(edge);store.add_edge(edge) # Duplicates remain suppressed.
        memory._postprocess();cycles,cyclic=store.postprocess()
        assert [list(group) for group in cycles]==memory.cycles
        assert set(cyclic)==memory.cyclic_nodes
        assert [dict(row) for row in store.edges]==memory.edges
        assert {key:dict(value) for key,value in store.nodes.items()}==memory.nodes
        assert len(store.cache)<=3 and len(store.ids)<=3


def test_evicted_live_record_proxy_preserves_later_updates(tmp_path):
    with GraphStore(tmp_path,cache_records=1) as store:
        store.nodes['A']={'id':'A','kind':'input','value':1}
        first=store.nodes['A']
        for i in range(20):store.nodes[str(i)]={'id':str(i),'kind':'label','value':'row'}
        first.update(value=7,extra='preserved')
        store.flush();store.cache.clear()
        assert dict(store.nodes['A'])=={'id':'A','kind':'input','value':7,'extra':'preserved'}
        del first['extra'];store.flush();store.cache.clear()
        assert 'extra' not in store.nodes['A']


def test_nonrecursive_scc_deep_chain_and_large_cycle(tmp_path):
    with GraphStore(tmp_path,cache_records=8) as store:
        for i in range(5000):store.nodes[str(i)]={'id':str(i),'kind':'op'}
        for i in range(4999):
            store.add_edge(dict(source=str(i),target=str(i+1),role='value',arg_index=0,op='+',ref='',via_range='',cell='',cross_sheet=False))
        store.add_edge(dict(source='4999',target='0',role='value',arg_index=0,op='+',ref='',via_range='',cell='',cross_sheet=False))
        cycles,cyclic=store.postprocess()
        assert len(cycles)==1 and len(cycles[0])==5000 and len(cyclic)==5000
        assert cycles[0][0]=='4999' and cycles[0][-1]=='0'
        assert set(store.root.iterdir())=={store.root/name for name in ('graph.sqlite','in-degree','out-degree','cyclic')}
        assert len(store.cache)<=8 and len(store.ids)<=8


def test_storage_python_heap_bounded_by_cache_not_node_count(tmp_path):
    tracemalloc.start()
    try:
        with GraphStore(tmp_path,cache_records=16) as store:
            for i in range(20_000):
                store.nodes[str(i)]={'id':str(i),'kind':'formula','formula':'=A1+'+str(i),'value':i}
            store.postprocess()
            _,peak=tracemalloc.get_traced_memory()
            assert peak<8*1024**2
            assert len(store.cache)<=16 and len(store.ids)<=16 and not store.dirty
    finally:tracemalloc.stop()


def test_empty_graph_and_single_self_loop(tmp_path):
    with GraphStore(tmp_path) as empty:
        cycles,nodes=empty.postprocess()
        assert len(cycles)==len(nodes)==0
    with GraphStore(tmp_path) as one:
        one.nodes['a']={'id':'a','kind':'formula'}
        one.add_edge(dict(source='a',target='a',role='value',arg_index=0,op='',ref='',via_range='',cell='a',cross_sheet=False))
        cycles,nodes=one.postprocess()
        assert [list(c) for c in cycles]==[['a']] and set(nodes)=={'a'}


def test_build_failure_closes_sqlite_mmaps_and_temp_directory(tmp_path,monkeypatch):
    path=tmp_path/'case.xlsx';workbook(path,'normal');roots=[]
    def failed(self):
        roots.append(self._storage.root)
        self._storage.array('partial',20,'Q')
        raise ValueError('stop before producing graph')
    monkeypatch.setattr(AstGraph,'_walk',failed)
    graph=AstGraph(path,graph_storage='disk',storage_dir=tmp_path,verbose=False)
    with pytest.raises(ValueError,match='stop before'):graph.build()
    assert graph._storage.closed and all(not root.exists() for root in roots)


def test_auto_policy_uses_formula_count_and_explicit_memory_override(tmp_path,monkeypatch):
    import xl_graph_store
    path=tmp_path/'case.xlsx';workbook(path,'normal')
    monkeypatch.setattr(xl_graph_store,'AUTO_FORMULA_THRESHOLD',1)
    with AstGraph(path,graph_storage='auto',verbose=False) as automatic:
        automatic.build();assert automatic._storage is not None
    with AstGraph(path,graph_storage='memory',verbose=False) as memory:
        memory.build();assert memory._storage is None


def test_production_cli_emits_only_csv_and_cleans_temporary_storage(tmp_path):
    path=tmp_path/'case.xlsx';workbook(path,'normal');temporary=tmp_path/'temporary';temporary.mkdir()
    main([str(path),'--production','--graph-storage','disk','--graph-storage-dir',str(temporary),'-o',str(tmp_path/'out'),'-q'])
    assert {p.name for p in (tmp_path/'out'/'case').iterdir()}=={'nodes.csv','edges.csv'}
    assert list(temporary.iterdir())==[]


@pytest.mark.parametrize('nodes,edges',[({},[]),({'A':{'id':'A','value':{'nested':[1,2]}}},[]),
                                      ({'A':{'id':'A'}},[{'source':'A','ref':'quote " newline\n'}])])
def test_streamed_json_is_byte_identical_to_previous_encoder(tmp_path,nodes,edges):
    stats={'empty':{},'nested':[1,{'x':None}]}
    write_json(tmp_path,nodes,edges,stats)
    expected=json.dumps({'stats':stats,'nodes':list(nodes.values()),'edges':edges},indent=1,default=str)
    assert (tmp_path/'graph.json').read_text()==expected


@pytest.mark.parametrize('method',['values','items'])
@pytest.mark.parametrize('cache_size',[1,4])
def test_mutating_future_node_during_iteration_never_loses_change(tmp_path,method,cache_size):
    with GraphStore(tmp_path,cache_records=cache_size) as store:
        for key in ('a','b','c','d'):store.nodes[key]={'id':key,'value':0}
        observed={}
        for value in getattr(store.nodes,method)():
            node=value[1] if method=='items' else value
            if node['id']=='a':store.nodes['d']['value']=42
            observed[node['id']]=node['value']
        store.flush();store.cache.clear()
        assert observed['d']==42 and store.nodes['d']['value']==42


def test_mutating_future_edge_during_iteration_preserves_change(tmp_path):
    with GraphStore(tmp_path,cache_records=2) as store:
        for key in ('a','b','c'):store.nodes[key]={'id':key,'kind':'formula'}
        for source,target in [('a','b'),('b','c')]:
            store.add_edge(dict(source=source,target=target,role='value',arg_index=0,op='+',ref='',via_range='',cell=source,cross_sheet=False))
        seen=[]
        for index,edge in enumerate(store.edges):
            if index==0:store.edges[1]['op']='-'
            seen.append(edge['op'])
        assert seen==['+','-'] and store.edges[1]['op']=='-'


def test_cycle_preview_keeps_exact_old_stable_order_and_members():
    from xl_ast_graph import cycle_previews
    cycles=[['z','x','y'],['c','b','a'],['q'],['f','e','d','c','b','a'],['foreign']]
    nodes={x for group in cycles[:-1] for x in group}
    old=sorted([g for g in cycles if any(k in nodes for k in g)],key=len,reverse=True)[:3]
    assert list(cycle_previews(cycles,nodes))==[(len(g),sorted(g)[:4]) for g in old]


def test_large_cycle_preview_uses_bounded_heap_not_full_member_sort():
    from xl_ast_graph import cycle_previews
    class LargeComponent:
        def __len__(self):return 100_000
        def __iter__(self):
            for i in range(100_000,0,-1):yield 'node-'+str(i).zfill(8)
    class AllNodes:
        def __contains__(self,key):return True
    tracemalloc.start()
    try:
        preview=list(cycle_previews([LargeComponent()],AllNodes()))
        _,peak=tracemalloc.get_traced_memory()
        assert peak<2*1024**2
    finally:tracemalloc.stop()
    assert preview==[(100_000,['node-00000001','node-00000002','node-00000003','node-00000004'])]


def test_invalid_edge_is_rejected_instead_of_ignored(tmp_path):
    import sqlite3
    with GraphStore(tmp_path) as store:
        store.nodes['a']={'id':'a'}
        with pytest.raises(sqlite3.IntegrityError):
            store.add_edge(dict(source='a',target='a',role=None,arg_index=0,op='',ref='',via_range='',cell='',cross_sheet=False))
        assert len(store.edges)==0

@pytest.mark.parametrize('case',['normal','cycles','range','dynamic','failed','array'])
def test_auto_spill_after_completed_formula_preserves_full_graph(tmp_path,monkeypatch,case):
    import xl_ast_graph as module
    monkeypatch.setattr(module,'AUTO_EDGE_THRESHOLD',2)
    path=tmp_path/'spill.xlsx';workbook(path,case)
    with AstGraph(path,graph_storage='memory',verbose=False,max_range_expand=2) as expected, \
         AstGraph(path,graph_storage='auto',verbose=False,max_range_expand=2,storage_dir=tmp_path) as actual:
        expected.build();actual.build()
        assert actual._storage is not None
        assert {k:dict(v) for k,v in actual.nodes.items()}==expected.nodes
        assert [dict(e) for e in actual.edges]==expected.edges
        assert [list(c) for c in actual.cycles]==expected.cycles
        assert actual.stats(actual.nodes,actual.edges)==expected.stats(expected.nodes,expected.edges)
        (tmp_path/'expected').mkdir(); (tmp_path/'actual').mkdir()
        write_csv(tmp_path/'expected',expected.nodes,expected.edges)
        write_csv(tmp_path/'actual',actual.nodes,actual.edges)
        for name in ('nodes.csv','edges.csv'):
            assert (tmp_path/'actual'/name).read_bytes()==(tmp_path/'expected'/name).read_bytes()
