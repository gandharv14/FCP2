from pathlib import Path
import zipfile
import openpyxl,pytest
from xl_volatile_formula_remediation import _defined_name_plan,build_plan,apply_plan,verify_manifest


def xml(names):
    return ('<workbook><definedNames>'+''.join('<definedName name="'+n+'">'+v+'</definedName>'for n,v in names)+'</definedNames></workbook>').encode()


def test_only_unreferenced_error_print_settings_are_selected():
    names=[('_xlnm.Print_Titles','#N/A'),('_xlnm.Print_Area','#REF!'),('_xlnm.Print_Area','Model!$A$1:$B$2'),('Other','#N/A')]
    actions,_=_defined_name_plan(xml(names),[])
    assert [(a['index'],a['rule_id'])for a in actions]==[(0,'unreferenced-error-print-setting/v1'),(1,'unreferenced-error-print-setting/v1')]

@pytest.mark.parametrize('reference',['_xlnm.Print_Titles','IF(TRUE,_xlnm.Print_Titles,0)','SUM(_XLNM.PRINT_TITLES)'])
def test_worksheet_or_other_package_formula_reference_preserves_error_name(reference):
    actions,_=_defined_name_plan(xml([('_xlnm.Print_Titles','#N/A')]),[reference])
    assert not actions


def test_transitive_defined_name_reference_preserves_print_error():
    actions,_=_defined_name_plan(xml([('_xlnm.Print_Titles','#N/A'),('Dependent','_xlnm.Print_Titles')]),[])
    assert not actions

@pytest.mark.parametrize('value',['Model!$1:$2','OFFSET(Model!$A$1,0,0)','[1]Other!$A$1','$A$1','', '#VALUE!'])
def test_other_print_setting_values_are_not_deleted_by_error_rule(value):
    actions,_=_defined_name_plan(xml([('_xlnm.Print_Titles',value)]),[])
    assert not any(a['rule_id']=='unreferenced-error-print-setting/v1'for a in actions)


def test_original_loader_failure_is_repaired_without_changing_cells_or_valid_print_area(tmp_path):
    original=tmp_path/'source.xlsx';candidate=tmp_path/'candidate.xlsx';book=openpyxl.Workbook();sheet=book.active;sheet.title='Model';sheet['A1']=3;sheet['A2']='=A1+4';sheet.print_area='A1:A2';book.save(original)
    with zipfile.ZipFile(original)as z:parts={n:z.read(n)for n in z.namelist()}
    parts['xl/workbook.xml']=parts['xl/workbook.xml'].replace(b'</definedNames>',b'<definedName name="_xlnm.Print_Titles" localSheetId="0">#N/A</definedName></definedNames>')
    with zipfile.ZipFile(original,'w')as z:
        for n,b in parts.items():z.writestr(n,b)
    with pytest.raises(ValueError,match='Unable to read workbook'):openpyxl.load_workbook(original,read_only=True)
    before=original.read_bytes();plan=build_plan(original);assert plan['status']=='eligible'
    manifest=apply_plan(original,plan,candidate);assert verify_manifest(original,candidate,plan,manifest)['status']=='verified'
    assert original.read_bytes()==before
    with zipfile.ZipFile(candidate)as z:
        assert z.read('xl/worksheets/sheet1.xml')==parts['xl/worksheets/sheet1.xml']
        assert b'_xlnm.Print_Area'in z.read('xl/workbook.xml')
        assert b'_xlnm.Print_Titles'not in z.read('xl/workbook.xml')
    result=openpyxl.load_workbook(candidate,read_only=True);assert result['Model']['A2'].value=='=A1+4';result.close()
