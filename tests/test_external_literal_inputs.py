"""External assumptions require original dependency bytes, never cached answers."""
import copy
import io
import json
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
import xl_external_literal_inputs as x


def rewrite(path, edit):
    with zipfile.ZipFile(path) as archive:
        infos=archive.infolist(); parts={i.filename:archive.read(i) for i in infos}
    edit(parts)
    with zipfile.ZipFile(path,'w') as archive:
        for info in infos:
            if info.filename in parts: archive.writestr(info,parts.pop(info.filename))
        for name,data in parts.items(): archive.writestr(name,data)


def setup(tmp_path,value=12,sheet="5yr O'Brien",formula=None,repeat=False):
    dep=tmp_path/'assumptions.xlsx'; wb=Workbook(); ws=wb.active; ws.title=sheet;ws['A1']=value
    if isinstance(value,datetime):ws['A1'].number_format='yyyy-mm-dd'
    if isinstance(value,str):ws['A1'].data_type='s'
    wb.save(dep)
    source=tmp_path/'source.xlsx'; wb=Workbook();ws=wb.active;ws.title='Task'
    ws['A1']=formula or "'[1]"+sheet.replace("'","''")+"'!$A$1"
    ws['A1']='='+ws['A1'].value
    if repeat:ws['A2']=ws['A1'].value
    ws['C1']=3;ws['C2']='=C1*2'
    if isinstance(value,datetime):ws['A1'].number_format='yyyy-mm-dd'
    wb.save(source)
    relpart='xl/externalLinks/_rels/externalLink1.xml.rels'
    target='file:///original/independent-assumptions.xlsx'
    def add(parts):
        book=ET.fromstring(parts['xl/workbook.xml'])
        refs=ET.SubElement(book,'{'+x.MAIN+'}externalReferences')
        ET.SubElement(refs,'{'+x.MAIN+'}externalReference',{'{'+x.DOCREL+'}id':'rIdExternal'})
        parts['xl/workbook.xml']=ET.tostring(book)
        rels=ET.fromstring(parts['xl/_rels/workbook.xml.rels'])
        ET.SubElement(rels,'{'+x.REL+'}Relationship',{'Id':'rIdExternal','Type':x.DOCREL+'/externalLink','Target':'externalLinks/externalLink1.xml'})
        parts['xl/_rels/workbook.xml.rels']=ET.tostring(rels)
        link=ET.Element('{'+x.MAIN+'}externalLink');external=ET.SubElement(link,'{'+x.MAIN+'}externalBook',{'{'+x.DOCREL+'}id':'rId1'})
        names=ET.SubElement(external,'{'+x.MAIN+'}sheetNames');ET.SubElement(names,'{'+x.MAIN+'}sheetName',{'val':sheet})
        # A poisonous cached link value must never become the supplied assumption.
        dataset=ET.SubElement(external,'{'+x.MAIN+'}sheetDataSet');data=ET.SubElement(dataset,'{'+x.MAIN+'}sheetData',{'sheetId':'0'})
        row=ET.SubElement(data,'{'+x.MAIN+'}row',{'r':'1'});cell=ET.SubElement(row,'{'+x.MAIN+'}cell',{'r':'A1'});ET.SubElement(cell,'{'+x.MAIN+'}v').text='999999'
        parts['xl/externalLinks/externalLink1.xml']=ET.tostring(link)
        rels=ET.Element('{'+x.REL+'}Relationships');ET.SubElement(rels,'{'+x.REL+'}Relationship',{'Id':'rId1','Type':x.DOCREL+'/externalLinkPath','TargetMode':'External','Target':target})
        parts[relpart]=ET.tostring(rels)
        content=ET.fromstring(parts['[Content_Types].xml']);ET.SubElement(content,'{http://schemas.openxmlformats.org/package/2006/content-types}Override',{'PartName':'/xl/externalLinks/externalLink1.xml','ContentType':'application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml'})
        parts['[Content_Types].xml']=ET.tostring(content)
    rewrite(source,add)
    with zipfile.ZipFile(source) as z: relhash=x.sha(z.read(relpart))
    evidence={'schema_version':x.EVIDENCE_SCHEMA,'source_sha256':x.sha(source.read_bytes()),'authority':'independent fixture approval',
      'dependencies':[{'link_index':1,'workbook_relationship_id':'rIdExternal','relationship_part':relpart,'relationship_id':'rId1',
       'relationship_part_sha256':relhash,'target_identity_sha256':x.sha(target.encode()),'path':str(dep),'sha256':x.sha(dep.read_bytes()),
       'literal_cells':[{'sheet':sheet,'coordinate':'A1','role':'external_assumption','provenance':'original assumption cell approved independently'}]}]}
    ep=tmp_path/'approved.json';ep.write_bytes(x.canonical(evidence))
    return source,dep,ep


def pin(ep):return x.sha(ep.read_bytes())


def repin(ep,fn):
    data=json.loads(ep.read_bytes());fn(data);ep.write_bytes(x.canonical(data))


@pytest.mark.parametrize('value',[12,1.25,True,False,'=this remains literal text','  spaced text  ',datetime(2024,3,1)])
def test_accept_typed_literal_with_exact_replay_and_preservation(tmp_path,value):
    source,dep,ep=setup(tmp_path,value)
    source_hash=x.sha(source.read_bytes());dep_hash=x.sha(dep.read_bytes())
    first=x.apply(source,ep,pin(ep),tmp_path/'first')
    second=x.apply(source,ep,pin(ep),tmp_path/'second')
    assert first==second==x.verify(source,ep,pin(ep),tmp_path/'first')
    assert first['qualification']=='unqualified_candidate'
    assert x.sha(source.read_bytes())==source_hash and x.sha(dep.read_bytes())==dep_hash
    wb=load_workbook(tmp_path/'first/candidate.xlsx',data_only=False)
    assert wb['Task']['A1'].value==value
    if isinstance(value,str):assert wb['Task']['A1'].data_type!='f'
    assert wb['Task']['C2'].value=='=C1*2'
    with zipfile.ZipFile(source) as a,zipfile.ZipFile(tmp_path/'first/candidate.xlsx') as b:
        assert not any('externalLink' in n for n in b.namelist())
        for member in first['plan']['members']:
            if member['disposition']=='unchanged': assert a.read(member['part'])==b.read(member['part'])


@pytest.mark.parametrize('sheet',['5yr','Space Name',"O'Brien",'Ordinary'])
def test_sheet_identity_and_repeated_occurrences(tmp_path,sheet):
    source,dep,ep=setup(tmp_path,sheet=sheet,repeat=True)
    p=x.plan(source,ep,pin(ep));assert len(p['dispositions'])==2
    assert {d['coordinate'] for d in p['dispositions']}=={'A1','A2'}


@pytest.mark.parametrize('formula',["SUM('[1]5yr O''Brien'!A1)","'[1]5yr O''Brien'!A1:A2","'[1]5yr O''Brien'!A1+1",'INDIRECT("[1]S!A1")',"'[1]5yr O''Brien'!NamedInput"])
def test_reject_non_direct_or_non_cell_external_expression(tmp_path,formula):
    source,dep,ep=setup(tmp_path,formula=formula)
    with pytest.raises(x.ExternalLiteralError):x.plan(source,ep,pin(ep))


@pytest.mark.parametrize('change',['missing_dependency','changed_dependency','wrong_identity','missing_identity','duplicate_identity','unknown_field','missing_provenance','unused_assumption','wrong_coordinate'])
def test_reject_incomplete_or_altered_bound_evidence(tmp_path,change):
    source,dep,ep=setup(tmp_path)
    if change=='missing_dependency':dep.unlink()
    elif change=='changed_dependency':dep.write_bytes(dep.read_bytes()+b'altered')
    elif change=='wrong_identity':repin(ep,lambda e:e['dependencies'][0].update(target_identity_sha256='0'*64))
    elif change=='missing_identity':repin(ep,lambda e:e.update(dependencies=[]))
    elif change=='duplicate_identity':repin(ep,lambda e:e['dependencies'].append(copy.deepcopy(e['dependencies'][0])))
    elif change=='unknown_field':repin(ep,lambda e:e.update(unsafe_override=True))
    elif change=='missing_provenance':repin(ep,lambda e:e['dependencies'][0]['literal_cells'][0].update(provenance=''))
    elif change=='unused_assumption':repin(ep,lambda e:e['dependencies'][0]['literal_cells'].append({'sheet':"5yr O'Brien",'coordinate':'C3','role':'external_assumption','provenance':'unreferenced'}))
    elif change=='wrong_coordinate':repin(ep,lambda e:e['dependencies'][0]['literal_cells'][0].update(coordinate='A2'))
    with pytest.raises((x.ExternalLiteralError,OSError)):x.plan(source,ep,pin(ep))


def test_approved_hash_is_independent_trust_anchor(tmp_path):
    source,dep,ep=setup(tmp_path);old=pin(ep)
    repin(ep,lambda e:e.update(authority='changed'))
    with pytest.raises(x.ExternalLiteralError,match='approved evidence hash'):x.plan(source,ep,old)


@pytest.mark.parametrize('encoding',['ordinary','shared','array','query'])
def test_never_materialize_calculated_or_query_backed_dependency(tmp_path,encoding):
    source,dep,ep=setup(tmp_path)
    def alter(parts):
        if encoding=='query':parts['xl/connections.xml']=b'<connections/>';return
        tree=ET.fromstring(parts['xl/worksheets/sheet1.xml']);cell=next(tree.iter('{'+x.MAIN+'}c'))
        for el in list(cell):cell.remove(el)
        attrs={} if encoding=='ordinary' else {'t':encoding,'ref':'A1:A2','si':'0'}
        ET.SubElement(cell,'{'+x.MAIN+'}f',attrs).text='6+6'
        ET.SubElement(cell,'{'+x.MAIN+'}v').text='12'
        parts['xl/worksheets/sheet1.xml']=ET.tostring(tree)
    rewrite(dep,alter);repin(ep,lambda e:e['dependencies'][0].update(sha256=x.sha(dep.read_bytes())))
    with pytest.raises(x.ExternalLiteralError):x.plan(source,ep,pin(ep))


def test_unaccounted_named_consumer_rejects_partial_link_removal(tmp_path):
    source,dep,ep=setup(tmp_path)
    def alter(parts):
        tree=ET.fromstring(parts['xl/workbook.xml']);names=ET.SubElement(tree,'{'+x.MAIN+'}definedNames');ET.SubElement(names,'{'+x.MAIN+'}definedName',{'name':'Hidden'}).text="'[1]S'!A1"
        parts['xl/workbook.xml']=ET.tostring(tree)
    rewrite(source,alter);repin(ep,lambda e:e.update(source_sha256=x.sha(source.read_bytes())))
    with pytest.raises(x.ExternalLiteralError,match='opaque consumer'):x.plan(source,ep,pin(ep))


def test_date_system_mismatch_rejected(tmp_path):
    source,dep,ep=setup(tmp_path,datetime(2024,3,1));wb=load_workbook(dep);wb.epoch=CALENDAR_MAC_1904;wb.save(dep)
    repin(ep,lambda e:e['dependencies'][0].update(sha256=x.sha(dep.read_bytes())))
    with pytest.raises(x.ExternalLiteralError,match='date system'):x.plan(source,ep,pin(ep))


def test_interruption_and_publication_conflict_preserve_evidence(tmp_path,monkeypatch):
    source,dep,ep=setup(tmp_path);dest=tmp_path/'result';real=x.os.rename
    monkeypatch.setattr(x.os,'rename',lambda *args: (_ for _ in ()).throw(OSError('interrupt')))
    with pytest.raises(OSError):x.apply(source,ep,pin(ep),dest)
    assert not dest.exists() and not list(tmp_path.glob('.external-literal-*'))
    monkeypatch.setattr(x.os,'rename',real);x.apply(source,ep,pin(ep),dest)
    before=(dest/'manifest.json').read_bytes()
    with pytest.raises(x.ExternalLiteralError,match='exists'):x.apply(source,ep,pin(ep),dest)
    assert (dest/'manifest.json').read_bytes()==before


@pytest.mark.parametrize('target',['candidate.xlsx','manifest.json'])
def test_altered_published_artifacts_fail_replay(tmp_path,target):
    source,dep,ep=setup(tmp_path);dest=tmp_path/'result';x.apply(source,ep,pin(ep),dest)
    if target=='candidate.xlsx':(dest/target).write_bytes((dest/target).read_bytes()+b'x')
    else:
        data=json.loads((dest/target).read_bytes());data['plan']['dispositions'][0]['literal_type']='answer';(dest/target).write_bytes(x.canonical(data))
    with pytest.raises(x.ExternalLiteralError):x.verify(source,ep,pin(ep),dest)


@pytest.mark.parametrize('consumer',['root_relation','duplicate_formula','dangling_content_type'])
def test_unaccounted_external_consumer_is_never_discarded(tmp_path,consumer):
    source,dep,ep=setup(tmp_path)
    def alter(parts):
        if consumer=='root_relation':
            part='_rels/.rels';tree=ET.fromstring(parts[part]);ET.SubElement(tree,'{'+x.REL+'}Relationship',{'Id':'unaccounted','Type':x.DOCREL+'/externalLink','Target':'xl/externalLinks/externalLink1.xml'})
        elif consumer=='dangling_content_type':
            part='[Content_Types].xml';tree=ET.fromstring(parts[part]);ET.SubElement(tree,'{http://schemas.openxmlformats.org/package/2006/content-types}Override',{'PartName':'/xl/externalLinks/missing.xml','ContentType':'application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml'})
        else:
            part='xl/worksheets/sheet1.xml';tree=ET.fromstring(parts[part]);cell=next(tree.iter('{'+x.MAIN+'}c'));ET.SubElement(cell,'{'+x.MAIN+'}f').text="'[1]5yr O''Brien'!A1"
        parts[part]=ET.tostring(tree)
    rewrite(source,alter);repin(ep,lambda e:e.update(source_sha256=x.sha(source.read_bytes())))
    with pytest.raises(x.ExternalLiteralError):x.plan(source,ep,pin(ep))


def test_duplicate_json_keys_rejected_even_with_matching_approved_hash(tmp_path):
    source,dep,ep=setup(tmp_path);raw=ep.read_bytes().decode();ep.write_text(raw.replace('"authority":','"authority":"duplicate","authority":',1))
    with pytest.raises(x.ExternalLiteralError,match='duplicate evidence key'):x.plan(source,ep,pin(ep))


def test_replay_rejects_changed_original_dependency_bytes(tmp_path):
    source,dep,ep=setup(tmp_path);dest=tmp_path/'result';x.apply(source,ep,pin(ep),dest)
    dep.write_bytes(dep.read_bytes()+b'changed')
    with pytest.raises(x.ExternalLiteralError,match='dependency source hash'):x.verify(source,ep,pin(ep),dest)


def test_cli_plan_exposes_only_unqualified_recovery_contract(tmp_path,capsys):
    source,dep,ep=setup(tmp_path)
    x.main(['plan','--source',str(source),'--evidence',str(ep),'--approved-evidence-sha256',pin(ep)])
    plan=json.loads(capsys.readouterr().out)
    assert plan['schema_version']==x.PLAN_SCHEMA
    assert 'disclosure_and_leak_review' in plan['required_downstream_gates']
    assert 'independent_reconstruction_without_answer_seeding' in plan['required_downstream_gates']
