from pathlib import Path
import importlib.util
import pytest
import openpyxl
from xl_seg import evaluate
from xl_ast_graph import AstGraph,write_csv
from xl_seg.model import load
from xl_seg.project import build
P=Path(__file__).resolve().parents[1]/'diagnostics/authoring-queue-extension-20260911T185027Z/four-resume-diagnostics/0306-semantic-diagnosis/first_error_probe.py'
s=importlib.util.spec_from_file_location('probe_under_test',P);probe=importlib.util.module_from_spec(s);s.loader.exec_module(probe)
@pytest.mark.parametrize('formula,handled',[('=MATCH(99,A2:A3,0)',False),('=IFERROR(MATCH(99,A2:A3,0),1)',True)])
def test_first_error_is_observation_even_when_enclosing_formula_handles_it(tmp_path,formula,handled):
 wb=openpyxl.Workbook();ws=wb.active;ws.title='Sheet';ws['A1']=formula;ws['A2']=1;ws['A3']=2;path=tmp_path/'probe.xlsx';wb.save(path);ast=AstGraph(path,verbose=False).build();out=tmp_path/'ast';out.mkdir();write_csv(out,ast.nodes,ast.edges);g=load(out,'probe');cg=build(g)
 plain=evaluate.Evaluator(g,cg,strict_proof=True,run_probes=False).run({'Sheet!A2','Sheet!A3'})
 assert (plain.values['Sheet!A1']==1) is handled
 cls=probe.observer(evaluate,{'Sheet!A1'});e=cls(g,cg,strict_proof=True,run_probes=False);e.frames=[]
 with pytest.raises(probe.FirstError) as caught:e.run({'Sheet!A2','Sheet!A3'})
 r=caught.value.evidence;assert r['operator']=='MATCH' and r['owner']=='Sheet!A1' and r['formula']==formula
 assert 'not proven causal' in r['interpretation'] and e.frames==[]
 if handled:assert 'IFERROR' in r['formula']  # discovery can probe before entering the wrapper

def test_other_owner_not_observed(tmp_path):
 wb=openpyxl.Workbook();wb.active['A1']='=MATCH(99,A2:A3,0)';wb.active['A2']=1;wb.active['A3']=2;path=tmp_path/'x.xlsx';wb.save(path);a=AstGraph(path,verbose=False).build();out=tmp_path/'ast';out.mkdir();write_csv(out,a.nodes,a.edges);g=load(out,'x');cg=build(g);cls=probe.observer(evaluate,{'Other!A1'});e=cls(g,cg,strict_proof=True,run_probes=False);e.frames=[]
 assert repr(e.run({'Sheet!A2','Sheet!A3'}).values['Sheet!A1'])=='#N/A'
