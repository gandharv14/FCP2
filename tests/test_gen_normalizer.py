from __future__ import annotations

import json
import os
import tempfile
import unittest
from pathlib import Path

import gen_normalizer as gn
import openpyxl
import plain_eligibility
import pytest
from xl_artifact_paths import resolve_workbook_artifact

ROOT = Path(__file__).resolve().parents[1]
# Representative sample, not the full nine: parenthesized sheets, ComEd
# fragments, a plus-in-name sheet, and a genuine zero-variable plain task.
SAMPLE_DRAFTS = ("0175", "0289", "0480", "0564")


class ExtractRefsTests(unittest.TestCase):
    def test_quoted_identity_and_absolute_addresses_remain_separate(self):
        name = "2026 O'Brien! $ Plan"
        text = "'2026 O''Brien! $ Plan'!$B$2:$C$3"
        sheets = {name.upper(): name}
        refs = gn.extract_refs(text, sheets)
        self.assertEqual(refs, [name + '!$B$2:$C$3'])
        self.assertEqual(gn.expand(refs[0], sheets), (name, ['B2', 'B3', 'C2', 'C3']))

    def test_parenthesised_sheet(self):
        refs = gn.extract_refs("Financial Summary (Model)!D9:F9 = 1, 2, 3")
        self.assertEqual(refs, ["Financial Summary (Model)!D9:F9"])

    def test_plus_in_sheet(self):
        refs = gn.extract_refs("2021 (Actual + Proj)!C6:N6 — 12 values")
        self.assertEqual(refs, ["2021 (Actual + Proj)!C6:N6"])

    def test_comed_fragment_is_full_sheet(self):
        text = (
            "`IL (ComEd) 6kW 10kWh - Non LMI!C40` = 7139.59<br>"
            "`IL (ComEd) 6kW 10kWh 40% ITC !C40` = 7139.59"
        )
        refs = gn.extract_refs(text)
        self.assertEqual(refs, [
            "IL (ComEd) 6kW 10kWh - Non LMI!C40",
            "IL (ComEd) 6kW 10kWh 40% ITC!C40",
        ])
        self.assertFalse(any(r.startswith("ComEd)") for r in refs))

    def test_quoted_sheet(self):
        refs = gn.extract_refs("'Model (Original)'!F32:J32 = [1]")
        self.assertEqual(refs, ["Model (Original)!F32:J32"])

    def test_plain_sheet_still_matches(self):
        refs = gn.extract_refs("Blueprint Standalone!F17:G17 — [9627.69647]")
        self.assertEqual(refs, ["Blueprint Standalone!F17:G17"])


@pytest.mark.private_corpus
@pytest.mark.parametrize("wb", SAMPLE_DRAFTS)
def test_draft_rows_have_live_sheet_refs(wb):
    roots = [Path(p) for p in os.environ.get("FCP_CORPUS_ROOTS", "").split(os.pathsep) if p]
    if not roots:
        pytest.skip("set FCP_CORPUS_ROOTS to run private workbook integration tests")
    draft_path = ROOT / "runs" / f"{wb}-variable-sources" / "draft.json"
    if not draft_path.is_file():
        pytest.skip(f"retained private audit is unavailable: {draft_path}")
    matches = [resolve_workbook_artifact(root, wb) for root in roots]
    matches = [path for path in matches if path.is_file()]
    assert len(matches) == 1, f"ambiguous or absent workbook {wb}: {matches}"
    draft = json.loads(draft_path.read_text(encoding="utf-8"))
    book = openpyxl.load_workbook(matches[0], read_only=True, data_only=True)
    sheets = {s.strip().upper(): s for s in book.sheetnames}
    book.close()
    skipped_no_cell = 0
    for row in draft["rows"]:
        refs = gn.extract_refs(row["value_text"])
        if not refs:
            skipped_no_cell += 1
            continue
        resolved = [gn.expand(ref, sheets)[0] for ref in refs]
        assert any(resolved), f"{wb} {row['draft_id']}: {refs} do not resolve"
    assert skipped_no_cell <= 6


class ExistingMcpSpecTests(unittest.TestCase):
    def test_0438_stays_mcp(self):
        run = ROOT / "runs" / "0438-variable-sources"
        if not (run / "normalized.json").is_file():
            self.skipTest("0438 normalized spec not present")
        spec = json.loads((run / "normalized.json").read_text(encoding="utf-8"))
        self.assertGreater(len(spec.get("variables") or []), 0)
        if (run / "normalization_report.json").is_file():
            result = plain_eligibility.evaluate(run)
            self.assertEqual(result["mode"], "mcp")


class PlainEligibilityTests(unittest.TestCase):
    def test_extraction_defect_is_not_plain(self):
        with tempfile.TemporaryDirectory() as temp:
            run = Path(temp)
            (run / "draft.json").write_text(json.dumps({
                "rows": [
                    {"draft_id": f"r{i}", "variable_text": "x", "value_text": "A!B1"}
                    for i in range(5)
                ],
            }), encoding="utf-8")
            (run / "normalized.json").write_text(json.dumps({
                "variables": [],
            }), encoding="utf-8")
            (run / "normalization_report.json").write_text(json.dumps({
                "workbook": "0001",
                "atomic_variables": 0,
                "first_atomic_variables": 0,
                "dispositions": [
                    {"draft_id": f"r{i}", "status": "excluded",
                     "reason": "bad", "code": "unparsable_reference"}
                    for i in range(5)
                ],
                "exclusion_reason_codes": {"unparsable_reference": 5},
            }), encoding="utf-8")
            (run / "0001-inputs-variable-sources.metadata.json").write_text(
                json.dumps({"inventory_rows": 20}), encoding="utf-8")
            result = plain_eligibility.evaluate(run)
            self.assertEqual(result["mode"], "fail")
            self.assertIn("extraction defect", result["reason"])

    def test_first_nonempty_cannot_downgrade(self):
        with tempfile.TemporaryDirectory() as temp:
            run = Path(temp)
            (run / "draft.json").write_text(json.dumps({
                "rows": [{"draft_id": "a", "variable_text": "x", "value_text": "S!A1"}] * 5,
            }), encoding="utf-8")
            (run / "normalized.json").write_text(json.dumps({"variables": []}), encoding="utf-8")
            (run / "normalization_report.json").write_text(json.dumps({
                "atomic_variables": 0,
                "first_atomic_variables": 4,
                "dispositions": [
                    {"draft_id": "a", "status": "excluded", "reason": "leak",
                     "code": "duplicate_value_leak"}
                ] * 5,
            }), encoding="utf-8")
            (run / "first_normalization.json").write_text(json.dumps({
                "atomic_variables": 4,
            }), encoding="utf-8")
            result = plain_eligibility.evaluate(run)
            self.assertEqual(result["mode"], "fail")
            self.assertIn("downgraded", result["reason"])

    def test_clean_all_excluded_is_plain(self):
        with tempfile.TemporaryDirectory() as temp:
            run = Path(temp)
            rows = [
                {"draft_id": f"r{i}", "variable_text": "rate", "value_text": "S!A1"}
                for i in range(5)
            ]
            (run / "draft.json").write_text(json.dumps({"rows": rows}), encoding="utf-8")
            (run / "normalized.json").write_text(json.dumps({"variables": []}), encoding="utf-8")
            (run / "normalization_report.json").write_text(json.dumps({
                "atomic_variables": 0,
                "first_atomic_variables": 0,
                "dispositions": [
                    {"draft_id": f"r{i}", "status": "excluded", "reason": "leak",
                     "code": "duplicate_value_leak"}
                    for i in range(5)
                ],
                "exclusion_reason_codes": {"duplicate_value_leak": 5},
            }), encoding="utf-8")
            (run / "0001-inputs-variable-sources.metadata.json").write_text(
                json.dumps({"inventory_rows": 10}), encoding="utf-8")
            result = plain_eligibility.evaluate(run)
            self.assertEqual(result["mode"], "plain")
            self.assertTrue(result["eligible"])

    def test_forced_exclusions_block_plain(self):
        with tempfile.TemporaryDirectory() as temp:
            run = Path(temp)
            (run / "draft.json").write_text(json.dumps({
                "rows": [
                    {"draft_id": f"r{i}", "variable_text": "x", "value_text": "S!A1"}
                    for i in range(5)
                ],
            }), encoding="utf-8")
            (run / "normalized.json").write_text(json.dumps({"variables": []}), encoding="utf-8")
            (run / "normalization_report.json").write_text(json.dumps({
                "atomic_variables": 0,
                "first_atomic_variables": 0,
                "dispositions": [
                    {"draft_id": f"r{i}", "status": "excluded", "reason": "forced",
                     "code": "forced_exclusion"}
                    for i in range(5)
                ],
            }), encoding="utf-8")
            (run / "oracle_forced_exclusions.json").write_text(
                json.dumps(["r0", "r1", "r2", "r3", "r4"]), encoding="utf-8")
            (run / "0001-inputs-variable-sources.metadata.json").write_text(
                json.dumps({"inventory_rows": 10}), encoding="utf-8")
            result = plain_eligibility.evaluate(run)
            self.assertEqual(result["mode"], "fail")
            self.assertIn("oracle_forced_exclusions", result["reason"])

    def test_plain_environment_rejects_stray_files(self):
        with tempfile.TemporaryDirectory() as temp:
            bundle = Path(temp)
            env = bundle / "environment"
            env.mkdir()
            (env / "0001-inputs.xlsx").write_bytes(b"x")
            (env / "Dockerfile").write_text("FROM scratch\n", encoding="utf-8")
            (env / "eval").mkdir()
            (env / "eval" / "tasks.jsonl").write_text("{}\n", encoding="utf-8")
            report = plain_eligibility.check_plain_environment(bundle, "0001")
            self.assertFalse(report["valid"])
            self.assertIn("eval/tasks.jsonl", report["unknown_files"])

    def test_plain_environment_allows_dialogue_notes_after_apply(self):
        with tempfile.TemporaryDirectory() as temp:
            bundle = Path(temp)
            env = bundle / "environment"
            env.mkdir()
            (env / "0001-inputs.xlsx").write_bytes(b"x")
            (env / "Dockerfile").write_text(
                "FROM scratch\nCOPY additional-assumptions.md "
                "/app/additional-assumptions.md\n",
                encoding="utf-8",
            )
            (env / "additional-assumptions.md").write_text("notes\n", encoding="utf-8")
            (bundle / "tests").mkdir()
            (bundle / "tests" / "dialogue-applied.json").write_text(
                '{"applied": true, "review_passed": true, '
                '"draft_passed": true}\n',
                encoding="utf-8",
            )
            report = plain_eligibility.check_plain_environment(bundle, "0001")
            self.assertTrue(report["valid"], report)

    def test_plain_environment_rejects_stray_dialogue_notes(self):
        with tempfile.TemporaryDirectory() as temp:
            bundle = Path(temp)
            env = bundle / "environment"
            env.mkdir()
            (env / "0001-inputs.xlsx").write_bytes(b"x")
            (env / "Dockerfile").write_text("FROM scratch\n", encoding="utf-8")
            (env / "additional-assumptions.md").write_text("notes\n", encoding="utf-8")
            report = plain_eligibility.check_plain_environment(bundle, "0001")
            self.assertFalse(report["valid"])
            self.assertIn("additional-assumptions.md", report["unknown_files"])

    def test_plain_environment_accepts_xlsm_inputs(self):
        with tempfile.TemporaryDirectory() as temp:
            bundle = Path(temp)
            env = bundle / "environment"
            env.mkdir()
            (env / "0624-inputs.xlsm").write_bytes(b"m")
            (env / "Dockerfile").write_text("FROM scratch\n", encoding="utf-8")
            report = plain_eligibility.check_plain_environment(bundle, "0624")
            self.assertTrue(report["valid"], report)
            self.assertEqual(report["workbooks"], ["0624-inputs.xlsm"])

    def test_plain_environment_rejects_both_workbook_suffixes(self):
        with tempfile.TemporaryDirectory() as temp:
            bundle = Path(temp)
            env = bundle / "environment"
            env.mkdir()
            (env / "0001-inputs.xlsx").write_bytes(b"x")
            (env / "0001-inputs.xlsm").write_bytes(b"m")
            (env / "Dockerfile").write_text("FROM scratch\n", encoding="utf-8")
            report = plain_eligibility.check_plain_environment(bundle, "0001")
            self.assertFalse(report["valid"])
            self.assertIn("exactly one workbook at environment root",
                          report["missing_files"])

    def test_plain_environment_missing_workbook_error_unchanged(self):
        with tempfile.TemporaryDirectory() as temp:
            bundle = Path(temp)
            env = bundle / "environment"
            env.mkdir()
            (env / "Dockerfile").write_text("FROM scratch\n", encoding="utf-8")
            report = plain_eligibility.check_plain_environment(bundle, "0001")
            self.assertFalse(report["valid"])
            self.assertIn("0001-inputs.xlsx", report["missing_files"])

    def test_plain_environment_rejects_notes_without_dockerfile_copy(self):
        with tempfile.TemporaryDirectory() as temp:
            bundle = Path(temp)
            env = bundle / "environment"
            env.mkdir()
            (env / "0001-inputs.xlsx").write_bytes(b"x")
            (env / "Dockerfile").write_text("FROM scratch\n", encoding="utf-8")
            (env / "additional-assumptions.md").write_text("notes\n", encoding="utf-8")
            (bundle / "tests").mkdir()
            (bundle / "tests" / "dialogue-applied.json").write_text(
                '{"applied": true, "review_passed": true, '
                '"draft_passed": true}\n',
                encoding="utf-8",
            )
            report = plain_eligibility.check_plain_environment(bundle, "0001")
            self.assertFalse(report["valid"])
            self.assertIn("additional-assumptions.md", report["unknown_files"])


if __name__ == "__main__":
    unittest.main()

class AtomicNormalizationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.run = self.root / 'runs' / '0001-variable-sources'
        self.run.mkdir(parents=True)
        self.inputs = self.root / 'inputs'
        self.inputs.mkdir()
        self.source = self.root / '0001.xlsx'
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = 'Model'
        for cell, value in {'A1': 47, 'A2': 91, 'A3': 113, 'B1': 47}.items():
            sheet[cell] = value
        book.save(self.source)
        book.save(self.inputs / '0001-inputs.xlsx')
        book.close()
        self.draft = {'rows': [{'draft_id': 'historical',
            'variable_text': 'Historical revenue', 'value_text': 'Model!A1:A3',
            'source_notes': 'Supplied ledger; exact public provenance unverified.',
            'source_urls': ['https://example.com/ledger'], 'normalization_status': 'needs_review'}]}
        (self.run / 'draft.json').write_text(json.dumps(self.draft))

    def tearDown(self):
        self.temp.cleanup()

    def analyse(self):
        return gn.analyse('0001', self.source, gn._sha256(self.source),
                          runs_root=self.root / 'runs', inputs_root=self.inputs)

    def review(self, rows, rejected=()):
        payload = {'schema_version': 'normalization-review/v1',
                   **rows[0]['evidence'], 'reviewer': 'independent-test-review',
                   'cells': {r['reference']: {'eligible': r['reference'] not in rejected,
                       'reason': 'Historical ledger evidence' if r['reference'] not in rejected else 'Internal projection',
                       'evidence_refs': [r['reference'], 'historical'],
                       'entity': 'Company A', 'period': 'FY2020',
                       'scenario': 'Actual', 'basis': 'Supplied revenue ledger'}
                       for r in rows if r['clean']}}
        gn.atomic_json(self.run / 'normalization-review.json', payload)

    def emit(self, draft, rows):
        gn.emit('0001', draft, rows, runs_root=self.root / 'runs')
        script = self.run / 'normalize_0001.py'
        namespace = {'__name__': '__main__', '__file__': str(script)}
        exec(compile(script.read_text(), str(script), 'exec'), namespace)
        return namespace

    def test_reviewed_atomic_name_replaces_overbroad_audit_label(self):
        self.draft['rows'][0]['variable_text'] = 'Revenue forecasts'
        (self.run / 'draft.json').write_text(json.dumps(self.draft))
        draft, rows = self.analyse()
        self.review(rows)
        path = self.run / 'normalization-review.json'
        review = json.loads(path.read_text())
        for cell in review['cells'].values():
            cell['variable_name'] = 'Historical revenue'
        gn.atomic_json(path, review)
        self.emit(draft, rows)
        spec = json.loads((self.run / 'normalized.json').read_text())
        for variable in spec['variables']:
            self.assertIn('Historical revenue', variable['name'])
            self.assertNotIn('forecasts', variable['name'])
            self.assertNotIn('forecasts', variable['question'])
            self.assertEqual(variable['scenario'], 'Actual')
        review['cells']['Model!A2']['variable_name'] = ''
        gn.atomic_json(path, review)
        with self.assertRaisesRegex(SystemExit, 'invalid reviewed variable_name'):
            self.emit(*self.analyse())

    def test_grouped_duplicate_does_not_drop_independent_siblings(self):
        draft, rows = self.analyse()
        self.assertEqual([r['reference'] for r in rows if r['clean']], ['Model!A2', 'Model!A3'])
        self.review(rows)
        self.emit(draft, rows)
        spec = json.loads((self.run / 'normalized.json').read_text())
        self.assertEqual(len(spec['variables']), 2)
        self.assertEqual(spec['variables'][0]['period'], 'FY2020')
        self.assertEqual(spec['variables'][0]['sources'][0]['url'], 'https://example.com/ledger')
        self.assertIn('unverified', spec['variables'][0]['source_notes'])
        report = json.loads((self.run / 'normalization_report.json').read_text())
        self.assertEqual(report['draft_rows'], 1)
        self.assertEqual(len(report['dispositions']), 1)
        self.assertEqual(report['cell_exclusions'][0]['reference'], 'Model!A1')
        hashes = {n: gn._sha256(self.run / n) for n in ('normalized.json', 'exclusions.json', 'normalization_report.json')}
        self.emit(*self.analyse())
        self.assertEqual(hashes, {n: gn._sha256(self.run / n) for n in hashes})

    def test_missing_and_altered_semantic_evidence_fail_closed(self):
        draft, rows = self.analyse()
        with self.assertRaisesRegex(SystemExit, 'require normalization-review'):
            self.emit(draft, rows)
        self.review(rows)
        path = self.run / 'normalization-review.json'
        payload = json.loads(path.read_text()); payload['source_sha256'] = '0' * 64
        gn.atomic_json(path, payload)
        with self.assertRaisesRegex(SystemExit, 'stale'):
            self.emit(draft, rows)

    def test_internal_projection_is_excluded_and_duplicate_closure_reruns(self):
        book = openpyxl.load_workbook(self.source)
        book['Model']['A3'] = 91
        book.save(self.source); book.save(self.inputs / '0001-inputs.xlsx'); book.close()
        draft, rows = self.analyse()
        self.review(rows, rejected={'Model!A3'})
        self.emit(draft, rows)
        report = json.loads((self.run / 'normalization_report.json').read_text())
        self.assertEqual(report['atomic_variables'], 0)
        self.assertEqual({r['reference']: r['code'] for r in report['cell_exclusions']},
                         {'Model!A1': 'duplicate_value_leak', 'Model!A2': 'duplicate_value_leak',
                          'Model!A3': 'not_research_input'})

    def test_generated_script_revalidates_review_and_input_evidence(self):
        draft, rows = self.analyse(); self.review(rows)
        namespace = self.emit(draft, rows)
        with (self.run / 'normalization-review.json').open('a') as stream:
            stream.write(' ')
        with self.assertRaisesRegex(SystemExit, 'evidence changed'):
            namespace['main']()
        self.review(rows)
        self.source.write_bytes(b'changed source')
        with self.assertRaisesRegex(SystemExit, 'evidence changed'):
            namespace['main']()

    def test_formula_cell_does_not_discard_typed_siblings(self):
        book = openpyxl.load_workbook(self.source); book['Model']['A3'] = '=A2*2'
        book.save(self.source); book.close()
        draft, rows = self.analyse()
        self.assertEqual([r['reference'] for r in rows if r['clean']], ['Model!A2'])
        self.assertEqual(rows[-1]['code'], 'formula_cell')
