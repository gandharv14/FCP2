"""Shared range products equal the historical expanded graph without its size."""
import random,tracemalloc
import pytest
from xl_seg.group_adjacency import RangeProducts

def legacy(events):
 adj={};radj={}
 for members,dst in events:
  for src in members:
   if src!=dst:
    adj.setdefault(src,set()).add(dst);radj.setdefault(dst,set()).add(src)
 return adj,radj

def test_overlaps_duplicates_self_edges_and_order_match_expanded_oracle():
 rng=random.Random(41)
 for _ in range(50):
  groups={str(i):tuple(rng.sample(list('abcdefgh'),rng.randrange(1,9))) for i in range(5)}
  p=RangeProducts();events=[]
  for _ in range(80):
   target=rng.choice('abcdefgh')
   if rng.randrange(3):
    identity=rng.choice(list(groups));members=groups[identity];p.add_range(identity,members,target)
   else:
    members=(rng.choice('abcdefgh'),);p.add_edge(members[0],target)
   events.append((members,target))
  adj,radj=p.seal();a,r=legacy(events)
  assert list(adj)==list(a) and list(radj)==list(r)
  assert dict(adj)==a and dict(radj)==r
  of={k:'band'+str(i//2) for i,k in enumerate('abcdefg')}
  expected={(of[src],of[dst]) for src,targets in a.items() for dst in targets if src in of and dst in of and of[src]!=of[dst]}
  assert set(p.band_pairs(of))==expected
  with pytest.raises(TypeError):p.add_edge('x','y')
  assert all(not hasattr(neighbors,'add') for neighbors in adj.values())

def test_hundred_million_edges_stay_grouped_and_band_collapse_never_expands(monkeypatch):
 members=tuple('s'+str(i) for i in range(9587));targets=tuple('t'+str(i) for i in range(11466))
 tracemalloc.start();p=RangeProducts()
 for target in targets:p.add_range('shared',members,target)
 adj,radj=p.seal();_,peak=tracemalloc.get_traced_memory();tracemalloc.stop()
 assert peak<20*1024**2
 assert len(adj)==9587 and len(radj)==11466
 assert len(adj[members[0]])==11466 and len(radj[targets[0]])==9587
 assert set(adj[members[0]])==set(targets)
 from xl_seg.group_adjacency import Adjacency
 monkeypatch.setattr(Adjacency,'items',lambda self:pytest.fail('global adjacency expansion'))
 of={**dict.fromkeys(members,'source-band'),**dict.fromkeys(targets,'target-band')}
 assert list(p.band_pairs(of))==[('source-band','target-band')]

def test_empty_ranges_and_only_self_edges_remain_absent():
 p=RangeProducts();p.add_range('empty',[],'a');p.add_range('self',['a'],'a')
 a,r=p.seal();assert dict(a)==dict(r)=={}

def test_actual_projection_bands_components_and_lineage_match_expanded_graph(tmp_path):
 import openpyxl
 from xl_ast_graph import AstGraph,write_csv
 from xl_seg import model,project,bands,condense,frontier,lineage
 w=openpyxl.Workbook();s=w.active;s.title='Data'
 for r in range(1,9):
  s.cell(r,1,r);s.cell(r,2,r*2)
  s.cell(r,4,'=SUM($A$1:$B$8)+A'+str(r))
 s['F1']='=D1';s['G1']='=F1+D2';s['G2']='=G1+D3'
 book=tmp_path/'source.xlsx';w.save(book);w.close()
 with AstGraph(book,max_range_expand=2,graph_storage='memory',verbose=False) as ast:
  ast.build();write_csv(tmp_path,ast.nodes,ast.edges)
 with model.load(tmp_path,'fixture',storage='disk') as graph:
  current=project.build(graph)
  # Independent historical product expansion from the source AST edges.
  cellids=set(current.info);ranges={n.id:[c for c in model.range_members(n.sheet,n.coordinate) if c in cellids] for n in graph.iter_nodes('range')}
  events=[]
  for edge in graph.edges:
   destination=graph.owner_of(edge.target)
   if destination is not None:
    source=graph.owner_of(edge.source)
    members=ranges[edge.source] if edge.source in ranges else ([source] if source else [])
    events.append((members,destination))
  a,r=legacy(events);old=project.CellGraph(graph,current.info,a,r)
  assert dict(current.adj)==a and dict(current.radj)==r
  oldbands,newbands=bands.build(old),bands.build(current)
  assert oldbands==newbands
  oldcomponents,newcomponents=condense.build(oldbands),condense.build(newbands)
  assert oldcomponents==newcomponents
  assert frontier.score_outputs(oldbands,oldcomponents)==frontier.score_outputs(newbands,newcomponents)
  inputs={n.id for n in graph.iter_nodes('input')}
  assert lineage.cell_trace(current,{},'Data!G2',inputs,100)==lineage.cell_trace(old,{},'Data!G2',inputs,100)

def test_runtime_proof_fallback_keeps_unvisited_large_groups_lazy():
 from types import SimpleNamespace
 import xl_segment
 members=tuple('s'+str(i) for i in range(9587));targets=tuple('t'+str(i) for i in range(11466));p=RangeProducts()
 for t in targets:p.add_range('shared',members,t)
 adj,radj=p.seal();info={t:SimpleNamespace(node=SimpleNamespace(kind='formula')) for t in targets}
 cg=SimpleNamespace(info=info,radj=radj);runtime={targets[0]:{members[0]}}
 tracemalloc.start();proof=xl_segment._runtime_proof_radj(cg,SimpleNamespace(runtime_radj=runtime));_,peak=tracemalloc.get_traced_memory();tracemalloc.stop()
 assert peak<8*1024**2
 assert proof[targets[0]]=={members[0]} and proof[targets[1]]==set(members)
 runtime[targets[0]].add('later-change');assert proof[targets[0]]=={members[0]}
 assert xl_segment._proof_cone([targets[0]],proof)=={targets[0],members[0]}

def test_proof_materialization_budget_counts_only_cone_and_accepts_boundary(monkeypatch):
 import xl_segment as x
 assert x.MAX_MATERIALIZED_PROOF_EDGES==4_000_000
 assert x.PROOF_EDGE_ESTIMATED_BYTES==256
 monkeypatch.setattr(x,'MAX_MATERIALIZED_PROOF_EDGES',2)
 graph={'t':frozenset(('s','u','outside')),'outside':frozenset(('s','u','t'))}
 cone={'s','u','t'}
 assert x._check_proof_edge_budget(graph,cone)==2
 assert x._proof_edge_signature(graph,cone)==(('s','t'),('u','t'))
 assert x._active_cycle_groups(cone,graph)==[]

def test_proof_materialization_budget_stops_lazy_scan_before_allocation(monkeypatch):
 import xl_segment as x
 monkeypatch.setattr(x,'MAX_MATERIALIZED_PROOF_EDGES',2)
 class Bounded:
  def __iter__(self):
   yield 'a';yield 'b';yield 'c'
   pytest.fail('budget overflow must stop the lazy scan')
 graph={'t':Bounded()};cone={'t','a','b','c'}
 for operation in (x._proof_edge_signature,lambda graph,cone:x._active_cycle_groups(cone,graph)):
  with pytest.raises(x.ProofResourceExhaustion) as caught:operation(graph,cone)
  assert caught.value.report=={'reason_code':'resource_exhaustion','resource':'materialized_proof_edges','observed_edge_count_lower_bound':3,'edge_limit':2,'estimated_allocation_bytes_lower_bound':768,'proof_truncated':False,'recovery':'streaming_proof_representation_required'}

def test_divergence_roots_lazy_union_retains_runtime_self_reference(monkeypatch):
 from types import SimpleNamespace
 import xl_segment as x
 import xl_seg.group_adjacency as grouped
 seen=[];original=grouped.Neighbors
 def recording(key,direct,groups):
  view=original(key,direct,groups);seen.append((key,groups,view));return view
 monkeypatch.setattr(grouped,'Neighbors',recording)
 static=frozenset(('input',))
 cg=SimpleNamespace(radj={'bad':static},info={'bad':SimpleNamespace(node=SimpleNamespace(kind='formula',value=1),is_literal=False)})
 result=SimpleNamespace(runtime_radj={'bad':{'bad'}},values={'bad':2})
 assert x._divergence_roots(cg,result,['bad'])=={'bad':'bad'}
 assert seen[0][0] is None and seen[0][1][0] is static
 assert set(seen[0][2])=={'input','bad'}
