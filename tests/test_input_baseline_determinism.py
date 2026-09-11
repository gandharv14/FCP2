"""Unchanged supplied assumptions produce the same hash across rebuild times."""
import shutil
import zipfile

from openpyxl import Workbook, load_workbook

from xl_input_mask import ASSUMPTIONS_PART, inject_assumptions


def test_embedded_assumptions_preserve_bytes_across_wall_clock_changes(tmp_path, monkeypatch):
    original = tmp_path / 'original.xlsx'
    workbook = Workbook()
    workbook.active['A1'] = 'Retained source'
    workbook.save(original)
    workbook.close()
    rows = [{'host': 'Sheet!B2', 'sheet': 'Sheet', 'label': 'Factor', 'value': 3.5}]
    results = []
    for index, clock in enumerate(((2024, 1, 2, 3, 4, 6, 0, 0, 0),
                                   (2026, 9, 9, 8, 7, 6, 0, 0, 0))):
        output = tmp_path / f'baseline-{index}.xlsx'
        shutil.copyfile(original, output)
        monkeypatch.setattr(zipfile.time, 'localtime', lambda *a, value=clock: value)
        sheet = inject_assumptions(output, rows, {'Sheet'})
        results.append(output.read_bytes())
        with zipfile.ZipFile(output) as archive:
            assert archive.getinfo(ASSUMPTIONS_PART).date_time == (1980, 1, 1, 0, 0, 0)
        delivered = load_workbook(output, data_only=True)
        assert delivered['Sheet']['A1'].value == 'Retained source'
        assert delivered[sheet]['D2'].value == 3.5
        delivered.close()
    assert results[0] == results[1]
