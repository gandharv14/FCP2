"""The delivery check must reject missing or altered reconstruction inputs."""
import datetime

import openpyxl
import pytest

from xl_input_mask import verify


def _verify(masked, source, **kwargs):
    return verify(masked, source, keep={}, frontier={}, formula_coords={},
                  deny=kwargs.get("deny", {}))


def test_missing_source_worksheet_is_reported_as_lost(tmp_path):
    original = openpyxl.Workbook()
    original.active.title = "Model"
    original.create_sheet("Assumptions")["A1"] = 123
    source = tmp_path / "source.xlsx"
    original.save(source)
    del original["Assumptions"]
    masked = tmp_path / "masked.xlsx"
    original.save(masked)
    faults = _verify(masked, source)
    assert any("Assumptions" in fault for fault in faults["lost"])


@pytest.mark.parametrize("replacement", [
    datetime.datetime(2030, 1, 1), "wrong date", 7,
])
def test_changed_date_input_is_rejected(tmp_path, replacement):
    original = openpyxl.Workbook()
    original.active["A1"] = datetime.datetime(2026, 1, 1)
    source = tmp_path / "source.xlsx"
    original.save(source)
    original.active["A1"] = replacement
    masked = tmp_path / "masked.xlsx"
    original.save(masked)
    assert "Sheet!A1" in _verify(masked, source)["changed"]


@pytest.mark.parametrize("value", ["confidential assumption", "=literal label", 123, True])
def test_denied_cell_cannot_survive_in_any_type(tmp_path, value):
    book = openpyxl.Workbook()
    book.active["A1"] = value
    if isinstance(value, str):
        book.active["A1"].data_type = "s"
    source = tmp_path / "source.xlsx"
    book.save(source)
    faults = _verify(source, source, deny={"Sheet": {(1, 1)}})
    assert any("Sheet!A1" in fault for fault in faults["leaked"])


@pytest.mark.parametrize("iso_dates", [False, True])
def test_unchanged_date_inputs_and_redacted_text_pass(tmp_path, iso_dates):
    original = openpyxl.Workbook(iso_dates=iso_dates)
    original.active["A1"] = datetime.datetime(2026, 1, 1, 12, 30)
    original.active["A2"] = datetime.date(2026, 2, 1)
    original.active["A3"] = "confidential assumption"
    source = tmp_path / "source.xlsx"
    original.save(source)
    original.active["A3"] = None
    masked = tmp_path / "masked.xlsx"
    original.save(masked)
    assert not any(_verify(masked, source, deny={"Sheet": {(3, 1)}}).values())
