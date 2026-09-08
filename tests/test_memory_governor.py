from __future__ import annotations

import json
import os
import threading
import time
import tracemalloc
import zipfile
from pathlib import Path

import openpyxl
import pytest

from scripts.batch import memory_governor
from scripts.batch.memory_governor import (
    LANE_CAP,
    BudgetTimeout,
    Calibration,
    MemoryBudget,
    MemoryGovernorError,
    WorkbookProfile,
    calibration_from_observations,
    estimate_peak_bytes,
    lanes_for_disk,
    load_calibration,
    plan_fleet,
    probe_workbook,
    record_observation,
    score_rows,
    seed_peak_bytes,
)

GIB = memory_governor.GIB


def _workbook(path: Path, *, rows: int = 20, sheets: int = 1) -> Path:
    workbook = openpyxl.Workbook()
    for index in range(sheets):
        sheet = (
            workbook.active if index == 0 else workbook.create_sheet(f"S{index}")
        )
        for row in range(1, rows + 1):
            sheet.cell(row=row, column=1, value=row)
            sheet.cell(row=row, column=2, value=f"=A{row}*2")
    workbook.save(path)
    workbook.close()
    return path


def _profile(**overrides: int) -> WorkbookProfile:
    base = {
        "file_bytes": 1000,
        "package_members": 8,
        "package_expanded_bytes": 5000,
        "sheet_count": 1,
        "shared_string_count": 0,
        "formula_cells": 100,
        "value_cells": 100,
    }
    base.update(overrides)
    return WorkbookProfile(**base)


def _budget(tmp_path: Path, *, total: int = 100 * GIB, reserve: int = 10 * GIB):
    return MemoryBudget(
        tmp_path / "budget",
        host_id="testvm",
        total_bytes=total,
        reserve_bytes=reserve,
        meminfo=tmp_path / "absent-meminfo",
        poll_seconds=0.01,
    )


# -- probe -----------------------------------------------------------------


def test_probe_counts_formula_and_value_cells(tmp_path: Path) -> None:
    profile = probe_workbook(_workbook(tmp_path / "wb.xlsx", rows=20, sheets=2))

    assert profile.sheet_count == 2
    assert profile.formula_cells == 40
    assert profile.value_cells == 40
    assert profile.package_expanded_bytes > 0
    assert profile.file_bytes == (tmp_path / "wb.xlsx").stat().st_size


def test_probe_scales_with_sheet_size(tmp_path: Path) -> None:
    small = probe_workbook(_workbook(tmp_path / "small.xlsx", rows=10))
    large = probe_workbook(_workbook(tmp_path / "large.xlsx", rows=500))

    assert large.formula_cells == 500
    assert large.formula_cells > small.formula_cells
    assert seed_peak_bytes(large) > seed_peak_bytes(small)


def test_probe_reads_shared_string_count(tmp_path: Path) -> None:
    # openpyxl writes inline strings, so a shared-string part is synthesized to
    # cover the branch that real Excel output takes.
    source = _workbook(tmp_path / "base.xlsx", rows=2)
    target = tmp_path / "shared.xlsx"
    with zipfile.ZipFile(source) as original:
        with zipfile.ZipFile(target, "w") as rebuilt:
            for item in original.infolist():
                rebuilt.writestr(item, original.read(item.filename))
            rebuilt.writestr(
                "xl/sharedStrings.xml",
                '<?xml version="1.0"?>'
                '<sst xmlns="http://schemas.openxmlformats.org/'
                'spreadsheetml/2006/main" count="7" uniqueCount="5">'
                "<si><t>a</t></si></sst>",
            )

    assert probe_workbook(target).shared_string_count == 5


def test_probe_rejects_a_non_workbook(tmp_path: Path) -> None:
    broken = tmp_path / "broken.xlsx"
    broken.write_text("not a workbook", encoding="utf-8")

    with pytest.raises(MemoryGovernorError):
        probe_workbook(broken)


def test_probe_rejects_an_oversized_package(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(memory_governor, "MAX_EXPANDED_BYTES", 16)

    with pytest.raises(MemoryGovernorError, match="expands to"):
        probe_workbook(_workbook(tmp_path / "wb.xlsx"))


def test_probe_refuses_a_symlink(tmp_path: Path) -> None:
    real = _workbook(tmp_path / "real.xlsx")
    link = tmp_path / "link.xlsx"
    link.symlink_to(real)

    with pytest.raises(MemoryGovernorError):
        probe_workbook(link)


# -- estimator -------------------------------------------------------------


def test_estimate_is_monotonic_in_each_feature() -> None:
    base = seed_peak_bytes(_profile())

    assert seed_peak_bytes(_profile(formula_cells=200)) > base
    assert seed_peak_bytes(_profile(value_cells=200)) > base
    assert seed_peak_bytes(_profile(package_expanded_bytes=50_000)) > base


def test_seed_covers_segmentation_on_the_largest_fleet_workbook() -> None:
    # 0568: 1.41M formula cells. xl_segment passed 101 GB on the old graph
    # code and 117.8 GB (killed at the whole-budget cap) after slots +
    # interning; the seed must say so itself instead of leaning on a
    # calibration scale capped at 4.0.
    big = _profile(
        formula_cells=1_410_000,
        value_cells=5_600_000,
        package_expanded_bytes=200 * 1024 * 1024,
    )
    seed = seed_peak_bytes(big)

    assert seed >= 90 * GIB
    assert seed >= memory_governor.SEGMENTATION_BYTES_PER_FORMULA * 1_410_000
    # Stages are sequential: the AST build term is not added on top.
    assert seed < (
        memory_governor.SEGMENTATION_BYTES_PER_FORMULA
        + memory_governor.AST_BUILD_BYTES_PER_FORMULA
    ) * 1_410_000
    # A 33k-formula workbook (0573 measured 0.6 GB) lands near the reserve floor.
    small = seed_peak_bytes(_profile(formula_cells=32_661, value_cells=22_058))
    assert 3 * GIB < small < 5 * GIB


def test_estimate_never_falls_below_the_floor() -> None:
    tiny = _profile(formula_cells=0, value_cells=0, package_expanded_bytes=0)

    assert estimate_peak_bytes(tiny) >= memory_governor.MIN_RESERVE_BYTES


def test_estimate_is_clamped_to_the_budget() -> None:
    huge = _profile(formula_cells=50_000_000)

    assert estimate_peak_bytes(huge, budget_total=90 * GIB) == 90 * GIB


def test_calibration_scales_the_estimate() -> None:
    profile = _profile(formula_cells=1_000_000)
    plain = estimate_peak_bytes(profile, Calibration(scale=1.0, sample_count=9))
    scaled = estimate_peak_bytes(profile, Calibration(scale=2.0, sample_count=9))

    assert scaled > plain


# -- calibration -----------------------------------------------------------


def _observation(ratio: float, status: str = "completed") -> dict[str, object]:
    profile = _profile()
    seed = seed_peak_bytes(profile)
    return {
        "status": status,
        "seed_bytes": seed,
        "observed_peak_bytes": int(seed * ratio),
        "profile": profile.to_record(),
    }


def test_calibration_is_the_measured_prior_below_the_sample_floor(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    # Too few completed rows to fit: the scale is the prior, not 1.0. The
    # fleet's first rows all peaked at 2.0x the seed model, so 1.0 meant every
    # one of them ran throttled until calibration caught up.
    monkeypatch.delenv(memory_governor.SEED_SCALE_ENV, raising=False)
    records = [_observation(3.0) for _ in range(memory_governor.CALIBRATION_MIN_SAMPLES - 1)]

    assert calibration_from_observations(records).scale == memory_governor.DEFAULT_SEED_SCALE
    assert memory_governor.DEFAULT_SEED_SCALE == 2.0


def test_the_seed_scale_is_applied_to_the_estimate_not_the_seed(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    # Folding the prior into seed_bytes would halve every recorded ratio and
    # make calibration converge on the square of the truth.
    monkeypatch.setenv(memory_governor.SEED_SCALE_ENV, "2.5")
    profile = _profile(formula_cells=1_000_000)

    raw = seed_peak_bytes(profile)
    unit = estimate_peak_bytes(profile, Calibration(scale=1.0, sample_count=9))
    prior = estimate_peak_bytes(profile)

    assert prior == pytest.approx(unit * 2.5, rel=0.01)
    assert memory_governor.seed_peak_bytes(profile) == raw
    monkeypatch.setenv(memory_governor.SEED_SCALE_ENV, "nope")
    with pytest.raises(memory_governor.MemoryGovernorError):
        memory_governor.seed_scale()
    monkeypatch.setenv(memory_governor.SEED_SCALE_ENV, "9")
    with pytest.raises(memory_governor.MemoryGovernorError):
        memory_governor.seed_scale()


def test_calibration_pools_sibling_logs_without_double_counting(tmp_path: Path) -> None:
    # A lane writes each observation to its own log and to the host-shared
    # one. Pooling must count the shared copy once and the sibling's rows
    # once, so three lanes reach the sample floor three times sooner.
    lane = tmp_path / "lane-a" / "observations.jsonl"
    shared = tmp_path / "shared" / "observations.jsonl"
    profile = _profile()
    for index in range(2):
        for target in (lane, shared):
            record_observation(
                target,
                profile,
                status="completed",
                observed_peak_bytes=int(seed_peak_bytes(profile) * 1.5),
                reserved_bytes=8 * GIB,
                workbook_id=f"000{index}",
                lane_id="lane-a",
                extra={"recorded_at": f"2026-01-01T00:00:0{index}Z"},
            )
    record_observation(
        shared,
        profile,
        status="completed",
        observed_peak_bytes=int(seed_peak_bytes(profile) * 1.5),
        reserved_bytes=8 * GIB,
        workbook_id="0009",
        lane_id="lane-b",
    )

    alone = load_calibration(lane)
    pooled = load_calibration(lane, shared)

    assert alone.sample_count == 2
    assert pooled.sample_count == 3
    assert pooled.scale == pytest.approx(1.5, rel=0.01)


def test_calibration_uses_the_high_quantile_of_completed_runs() -> None:
    records = [_observation(2.0) for _ in range(10)]

    result = calibration_from_observations(records)

    assert result.sample_count == 10
    assert result.scale == pytest.approx(2.0)


def test_calibration_is_clamped() -> None:
    high = calibration_from_observations([_observation(50.0) for _ in range(6)])
    low = calibration_from_observations([_observation(0.01) for _ in range(6)])

    assert high.scale == memory_governor.CALIBRATION_MAX_SCALE
    assert low.scale == memory_governor.CALIBRATION_MIN_SCALE


def test_a_killed_run_never_lowers_the_scale() -> None:
    # A killed attempt reports a truncated peak. Letting it into the quantile
    # would make the next estimate more optimistic, which is backwards.
    completed = [_observation(2.0) for _ in range(10)]
    truncated = [_observation(0.2, status="timed_out") for _ in range(20)]

    baseline = calibration_from_observations(completed).scale
    with_truncated = calibration_from_observations(completed + truncated)

    assert with_truncated.scale == pytest.approx(baseline)
    assert with_truncated.censored_count == 20


def test_a_killed_run_can_raise_the_scale() -> None:
    records = [_observation(1.0) for _ in range(6)]
    records.append(_observation(3.0, status="failed"))

    assert calibration_from_observations(records).scale == pytest.approx(3.0)


def test_observations_round_trip_through_the_log(tmp_path: Path) -> None:
    path = tmp_path / "memory" / "observations.jsonl"
    profile = _profile()
    for _ in range(6):
        record_observation(
            path,
            profile,
            status="completed",
            observed_peak_bytes=int(seed_peak_bytes(profile) * 1.5),
            reserved_bytes=8 * GIB,
            workbook_id="0001",
            lane_id="lane-a",
        )

    result = load_calibration(path)

    assert result.sample_count == 6
    assert result.scale == pytest.approx(1.5, rel=0.01)


def test_unparsable_observation_lines_are_skipped(tmp_path: Path) -> None:
    path = tmp_path / "observations.jsonl"
    path.write_text("not json\n" + json.dumps(_observation(2.0)) + "\n")

    assert load_calibration(path).censored_count == 0


# -- budget gate -----------------------------------------------------------


def test_budget_reserves_headroom_for_the_host(tmp_path: Path) -> None:
    budget = _budget(tmp_path, total=128 * GIB, reserve=16 * GIB)

    assert budget.total_bytes == 112 * GIB
    assert budget.host_bytes == 128 * GIB


def test_default_reserve_is_the_larger_of_floor_and_fraction() -> None:
    assert memory_governor.reserve_bytes_for(16 * GIB) == 8 * GIB
    assert memory_governor.reserve_bytes_for(128 * GIB) == int(128 * GIB * 0.12)


def test_two_lanes_inside_the_budget_both_run(tmp_path: Path) -> None:
    budget = _budget(tmp_path)

    first = budget.acquire("lane-a", 40 * GIB, workbook_id="0001")
    second = budget.acquire("lane-b", 40 * GIB, workbook_id="0002")

    assert first.bytes == 40 * GIB
    assert second.bytes == 40 * GIB
    assert budget.snapshot()["used_bytes"] == 80 * GIB


def test_an_over_budget_second_lane_waits(tmp_path: Path) -> None:
    budget = _budget(tmp_path)
    budget.acquire("lane-a", 60 * GIB, workbook_id="0001")

    with pytest.raises(BudgetTimeout):
        budget.acquire("lane-b", 60 * GIB, workbook_id="0002", timeout_seconds=0.1)


def test_a_waiting_lane_is_granted_after_a_release(tmp_path: Path) -> None:
    budget = _budget(tmp_path)
    held = budget.acquire("lane-a", 60 * GIB, workbook_id="0001")
    granted: list[int] = []

    def waiter() -> None:
        lease = budget.acquire(
            "lane-b", 60 * GIB, workbook_id="0002", timeout_seconds=5
        )
        granted.append(lease.bytes)

    thread = threading.Thread(target=waiter)
    thread.start()
    time.sleep(0.1)
    assert granted == []
    held.release()
    thread.join(timeout=5)

    assert granted == [60 * GIB]


def test_shrinking_never_blocks_and_frees_the_sibling(tmp_path: Path) -> None:
    budget = _budget(tmp_path)
    held = budget.acquire("lane-a", 85 * GIB, workbook_id="0001")

    held.resize(4 * GIB)
    sibling = budget.acquire(
        "lane-b", 60 * GIB, workbook_id="0002", timeout_seconds=1
    )

    assert held.bytes == 4 * GIB
    assert sibling.bytes == 60 * GIB


def test_growing_waits_for_room(tmp_path: Path) -> None:
    budget = _budget(tmp_path)
    budget.acquire("lane-a", 60 * GIB, workbook_id="0001")
    small = budget.acquire("lane-b", 4 * GIB, workbook_id="0002")

    with pytest.raises(BudgetTimeout):
        small.resize(60 * GIB, timeout_seconds=0.1)


def test_a_request_larger_than_the_budget_is_clamped(tmp_path: Path) -> None:
    budget = _budget(tmp_path)

    lease = budget.acquire("lane-a", 500 * GIB, workbook_id="0001")

    assert lease.bytes == budget.total_bytes


def test_a_dead_lane_reservation_is_reclaimed(tmp_path: Path) -> None:
    budget = _budget(tmp_path)
    budget.acquire("lane-a", 80 * GIB, workbook_id="0001")
    state = json.loads(budget.path.read_text(encoding="utf-8"))
    state["reservations"]["lane-a"]["pid"] = 2**30
    budget.path.write_text(json.dumps(state), encoding="utf-8")

    lease = budget.acquire(
        "lane-b", 80 * GIB, workbook_id="0002", timeout_seconds=1
    )

    assert lease.bytes == 80 * GIB
    assert "lane-a" not in budget.snapshot()["reservations"]


def test_an_older_heavy_waiter_is_served_before_a_newer_light_one(
    tmp_path: Path,
) -> None:
    budget = _budget(tmp_path)
    blocker = budget.acquire("lane-a", 85 * GIB, workbook_id="0001")
    order: list[str] = []

    def heavy() -> None:
        budget.acquire("lane-b", 80 * GIB, workbook_id="0002", timeout_seconds=5)
        order.append("heavy")

    def light() -> None:
        budget.acquire("lane-c", 1 * GIB, workbook_id="0003", timeout_seconds=5)
        order.append("light")

    heavy_thread = threading.Thread(target=heavy)
    heavy_thread.start()
    time.sleep(0.2)
    light_thread = threading.Thread(target=light)
    light_thread.start()
    time.sleep(0.2)

    assert order == []
    blocker.release()
    heavy_thread.join(timeout=5)
    light_thread.join(timeout=5)

    assert order[0] == "heavy"


def test_a_timed_out_waiter_leaves_no_ticket_behind(tmp_path: Path) -> None:
    budget = _budget(tmp_path)
    budget.acquire("lane-a", 85 * GIB, workbook_id="0001")

    with pytest.raises(BudgetTimeout):
        budget.acquire("lane-b", 80 * GIB, timeout_seconds=0.1)

    # A stale ticket would block every later request behind a lane that left.
    assert budget.snapshot()["waiters"] == []
    assert budget.acquire("lane-c", 1 * GIB, timeout_seconds=1).bytes == GIB


def test_low_host_availability_blocks_admission(tmp_path: Path) -> None:
    meminfo = tmp_path / "meminfo"
    meminfo.write_text(
        "MemTotal:       131072000 kB\nMemAvailable:      102400 kB\n",
        encoding="utf-8",
    )
    budget = MemoryBudget(
        tmp_path / "budget",
        host_id="testvm",
        total_bytes=100 * GIB,
        reserve_bytes=10 * GIB,
        meminfo=meminfo,
        poll_seconds=0.01,
    )

    # The budget file is empty, so only the live MemAvailable check can refuse.
    with pytest.raises(BudgetTimeout):
        budget.acquire("lane-a", 1 * GIB, timeout_seconds=0.1)


def test_release_is_idempotent(tmp_path: Path) -> None:
    budget = _budget(tmp_path)
    lease = budget.acquire("lane-a", 8 * GIB, workbook_id="0001")

    lease.release()
    lease.release()

    assert budget.snapshot()["used_bytes"] == 0


def test_the_lease_context_manager_always_releases(tmp_path: Path) -> None:
    budget = _budget(tmp_path)

    with pytest.raises(RuntimeError):
        with budget.lease("lane-a", 8 * GIB, workbook_id="0001"):
            raise RuntimeError("stage failed")

    assert budget.snapshot()["used_bytes"] == 0


def test_a_corrupt_budget_file_does_not_wedge_the_host(tmp_path: Path) -> None:
    budget = _budget(tmp_path)
    budget.directory.mkdir(parents=True, exist_ok=True)
    budget.path.write_text("{not json", encoding="utf-8")

    assert budget.acquire("lane-a", 8 * GIB, timeout_seconds=1).bytes == 8 * GIB


def test_budget_requires_known_host_memory(tmp_path: Path) -> None:
    with pytest.raises(MemoryGovernorError):
        MemoryBudget(tmp_path / "budget", meminfo=tmp_path / "absent")


def test_reserve_must_leave_something_to_schedule(tmp_path: Path) -> None:
    with pytest.raises(MemoryGovernorError):
        MemoryBudget(
            tmp_path / "budget",
            host_id="testvm",
            total_bytes=8 * GIB,
            reserve_bytes=8 * GIB,
            meminfo=tmp_path / "absent",
        )


def test_lane_ids_are_validated(tmp_path: Path) -> None:
    budget = _budget(tmp_path)

    with pytest.raises(MemoryGovernorError):
        budget.acquire("../escape", 1 * GIB)


# -- planner ---------------------------------------------------------------


def _scored(batch: str, workbook: str, estimate: int) -> dict[str, object]:
    return {
        "batch_id": batch,
        "workbook_id": workbook,
        "estimate_bytes": estimate,
    }


def test_planner_spreads_heavy_rows_across_vms() -> None:
    rows = [
        _scored("batch002", "0001", 80 * GIB),
        _scored("batch002", "0002", 78 * GIB),
        _scored("batch002", "0003", 76 * GIB),
        _scored("batch002", "0004", 2 * GIB),
        _scored("batch002", "0005", 2 * GIB),
        _scored("batch002", "0006", 2 * GIB),
    ]

    plan = plan_fleet(rows, vm_count=3, budget_total=110 * GIB)

    heaviest_per_vm = [vm["max_estimate_bytes"] for vm in plan["vms"]]
    assert sorted(heaviest_per_vm) == sorted(
        [80 * GIB, 78 * GIB, 76 * GIB]
    )


def test_planner_keeps_batches_in_separate_queues() -> None:
    rows = [
        _scored("batch002", "0001", 10 * GIB),
        _scored("batch003", "0002", 10 * GIB),
    ]

    plan = plan_fleet(rows, vm_count=1, budget_total=110 * GIB)

    queues = plan["vms"][0]["queues"]
    assert set(queues) == {"batch002", "batch003"}
    assert [row["workbook_id"] for row in queues["batch002"]] == ["0001"]
    assert [row["workbook_id"] for row in queues["batch003"]] == ["0002"]


def test_a_heavy_queue_starts_with_one_lane() -> None:
    rows = [_scored("batch002", f"000{index}", 80 * GIB) for index in range(1, 5)]

    plan = plan_fleet(rows, vm_count=1, budget_total=110 * GIB)

    assert plan["vms"][0]["lanes"] == 1
    assert plan["vms"][0]["heaviest_pair_fits"] is False


def test_a_light_queue_starts_with_the_lane_cap() -> None:
    rows = [_scored("batch002", f"000{index}", 5 * GIB) for index in range(1, 7)]

    plan = plan_fleet(rows, vm_count=1, budget_total=110 * GIB)

    assert plan["vms"][0]["lanes"] == LANE_CAP == 4
    assert plan["vms"][0]["memory_lanes"] == 4
    assert plan["vms"][0]["heaviest_pair_fits"] is True

    capped = plan_fleet(rows, vm_count=1, budget_total=110 * GIB, lane_cap=3)
    assert capped["vms"][0]["lanes"] == 3


def test_lanes_follow_the_heaviest_n_fit_not_the_median() -> None:
    # Heaviest three: 60 + 40 + 30 = 130 > 110, so three lanes could collide.
    # Heaviest two: 60 + 40 = 100 <= 110, so two lanes are safe. The median
    # (2 GiB) says nothing about that.
    rows = [
        _scored("batch002", "0001", 60 * GIB),
        _scored("batch002", "0002", 40 * GIB),
        _scored("batch002", "0003", 30 * GIB),
    ]
    rows += [_scored("batch002", f"001{index}", 2 * GIB) for index in range(9)]

    plan = plan_fleet(rows, vm_count=1, budget_total=110 * GIB)

    assert plan["vms"][0]["lanes"] == 2
    assert plan["vms"][0]["memory_lanes"] == 2
    assert plan["vms"][0]["heaviest_pair_fits"] is True


def test_one_monster_among_light_rows_still_allows_the_lane_cap() -> None:
    # 100 + 2 + 2 + 2 fits, so the light rows keep their concurrency; the
    # runtime gate is what serializes the single monster when it actually runs.
    rows = [_scored("batch002", "0001", 100 * GIB)]
    rows += [_scored("batch002", f"001{index}", 2 * GIB) for index in range(9)]

    plan = plan_fleet(rows, vm_count=1, budget_total=110 * GIB)

    assert plan["vms"][0]["lanes"] == LANE_CAP == 4
    assert plan["vms"][0]["max_estimate_bytes"] == 100 * GIB


def test_lanes_never_exceed_the_row_count() -> None:
    rows = [_scored("batch002", "0001", 2 * GIB), _scored("batch002", "0002", 2 * GIB)]

    plan = plan_fleet(rows, vm_count=1, budget_total=110 * GIB)

    assert plan["vms"][0]["lanes"] == 2


def test_disk_free_caps_lanes_per_vm() -> None:
    # VM1 has a 290 GB disk (~270 GiB) -> (270 - 120) // 60 = 2 lanes.
    # VM2/VM3 have 495 GB (~461 GiB) -> (461 - 120) // 60 = 5, capped at 4.
    rows = [_scored("batch002", f"00{index:02d}", 5 * GIB) for index in range(12)]

    plan = plan_fleet(
        rows,
        vm_count=3,
        budget_total=110 * GIB,
        disk_free_bytes=[270 * GIB, 461 * GIB, 461 * GIB],
    )

    lanes = {vm["vm_index"]: vm["lanes"] for vm in plan["vms"]}
    assert lanes == {0: 2, 1: 4, 2: 4}
    assert plan["vms"][0]["disk_lanes"] == 2
    assert plan["vms"][0]["disk_free_bytes"] == 270 * GIB
    assert plan["vms"][1]["disk_free_bytes"] == 461 * GIB


def test_lanes_for_disk_formula() -> None:
    assert lanes_for_disk(None) == LANE_CAP
    assert lanes_for_disk(100 * GIB) == 1  # below the floor still gets one
    assert lanes_for_disk(180 * GIB) == 1
    assert lanes_for_disk(240 * GIB) == 2
    assert lanes_for_disk(300 * GIB) == 3
    assert lanes_for_disk(10_000 * GIB) == LANE_CAP
    assert lanes_for_disk(300 * GIB, lane_cap=2) == 2


def test_disk_free_must_match_vm_count() -> None:
    with pytest.raises(MemoryGovernorError, match="one entry per VM"):
        plan_fleet([], vm_count=2, budget_total=GIB, disk_free_bytes=[GIB])


def test_plan_cli_accepts_disk_free_bytes(tmp_path: Path, capsys) -> None:
    queue = tmp_path / "queue.json"
    queue.write_text("[]", encoding="utf-8")
    code = memory_governor.main(
        [
            "--budget-total",
            str(110 * GIB),
            "plan",
            "--queues",
            str(queue),
            "--vms",
            "2",
            "--disk-free-bytes",
            str(270 * GIB),
            str(461 * GIB),
        ]
    )
    assert code == 0
    plan = json.loads(capsys.readouterr().out)
    assert [vm["disk_lanes"] for vm in plan["vms"]] == [2, 4]
    assert plan["lane_cap"] == 4


def test_planner_rejects_nonsense_arguments() -> None:
    with pytest.raises(MemoryGovernorError):
        plan_fleet([], vm_count=0, budget_total=GIB)
    with pytest.raises(MemoryGovernorError):
        plan_fleet([], vm_count=1, budget_total=0)


def test_score_rows_flags_an_unprobeable_row_as_whole_budget(
    tmp_path: Path,
) -> None:
    _workbook(tmp_path / "good.xlsx")
    (tmp_path / "bad.xlsx").write_text("nope", encoding="utf-8")
    rows = [
        {"batch_id": "batch002", "workbook_id": "0001", "run_source_path": "good.xlsx"},
        {"batch_id": "batch002", "workbook_id": "0002", "run_source_path": "bad.xlsx"},
    ]

    scored = score_rows(rows, budget_total=90 * GIB, queue_dir=tmp_path)

    assert scored[0]["unprobed"] is False
    assert scored[1]["unprobed"] is True
    assert scored[1]["estimate_bytes"] == 90 * GIB


# -- CLI -------------------------------------------------------------------


def test_cli_probe_prints_a_profile(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    workbook = _workbook(tmp_path / "wb.xlsx")

    assert memory_governor.main(["probe", str(workbook)]) == 0

    payload = json.loads(capsys.readouterr().out)
    assert payload["schema_version"] == memory_governor.PROFILE_SCHEMA
    assert payload["formula_cells"] == 20


def test_cli_estimate_reports_the_budget(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    workbook = _workbook(tmp_path / "wb.xlsx")

    exit_code = memory_governor.main(
        ["--host-bytes", str(128 * GIB), "estimate", str(workbook)]
    )

    assert exit_code == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["budget_total_bytes"] == 128 * GIB - int(128 * GIB * 0.12)
    assert payload["estimate_bytes"] >= memory_governor.MIN_RESERVE_BYTES


def test_cli_plan_reads_a_queue_file(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    _workbook(tmp_path / "wb1.xlsx")
    queue = tmp_path / "queue.json"
    queue.write_text(
        json.dumps(
            [
                {
                    "batch_id": "batch002",
                    "workbook_id": "0001",
                    "run_source_path": "wb1.xlsx",
                }
            ]
        ),
        encoding="utf-8",
    )

    exit_code = memory_governor.main(
        [
            "--host-bytes",
            str(128 * GIB),
            "plan",
            "--queues",
            str(queue),
            "--vms",
            "3",
        ]
    )

    assert exit_code == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["row_count"] == 1
    assert len(payload["vms"]) == 3


def test_cli_report_prints_budget_state(
    tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    exit_code = memory_governor.main(
        [
            "--budget-dir",
            str(tmp_path / "budget"),
            "--host-id",
            "testvm",
            "--host-bytes",
            str(64 * GIB),
            "report",
        ]
    )

    assert exit_code == 0
    payload = json.loads(capsys.readouterr().out)
    assert payload["host_id"] == "testvm"
    assert payload["used_bytes"] == 0


def test_two_lanes_each_needing_the_whole_budget_do_not_deadlock(
    tmp_path: Path,
) -> None:
    # Both lanes open with a small reservation, then both try to grow past what
    # the other is holding. Without handing the held bytes back while waiting,
    # neither can ever be granted.
    budget = _budget(tmp_path)
    first = budget.acquire("lane-a", 4 * GIB, workbook_id="0001")
    second = budget.acquire("lane-b", 4 * GIB, workbook_id="0002")
    outcomes: dict[str, int] = {}

    def grow(name: str, lease) -> None:
        lease.resize(budget.total_bytes, timeout_seconds=10)
        outcomes[name] = lease.bytes
        lease.release()

    threads = [
        threading.Thread(target=grow, args=("lane-a", first)),
        threading.Thread(target=grow, args=("lane-b", second)),
    ]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join(timeout=15)

    assert not any(thread.is_alive() for thread in threads)
    assert outcomes == {"lane-a": budget.total_bytes, "lane-b": budget.total_bytes}
    assert budget.snapshot()["used_bytes"] == 0


def test_a_failed_grow_restores_the_prior_reservation(tmp_path: Path) -> None:
    # Surrendering the held bytes is an internal step of the grow. If a failed
    # grow left them surrendered, the lane would run its stage completely
    # unreserved and a sibling would admit on top of it.
    budget = _budget(tmp_path)
    blocker = budget.acquire("lane-a", 85 * GIB, workbook_id="0001")
    waiter = budget.acquire("lane-b", 4 * GIB, workbook_id="0002")

    with pytest.raises(BudgetTimeout):
        waiter.resize(80 * GIB, timeout_seconds=0.2)

    reservations = budget.snapshot()["reservations"]
    assert blocker.bytes == 85 * GIB
    assert int(reservations["lane-b"]["bytes"]) == 4 * GIB


def test_a_lane_surrenders_its_bytes_while_actually_waiting(
    tmp_path: Path,
) -> None:
    budget = _budget(tmp_path)
    blocker = budget.acquire("lane-a", 85 * GIB, workbook_id="0001")
    waiter = budget.acquire("lane-b", 4 * GIB, workbook_id="0002")
    seen: list[int] = []

    def grow() -> None:
        waiter.resize(80 * GIB, timeout_seconds=5)

    thread = threading.Thread(target=grow)
    thread.start()
    time.sleep(0.2)
    seen.append(int(budget.snapshot()["used_bytes"]))
    blocker.release()
    thread.join(timeout=5)

    # While blocked the lane is not using memory, so it must not be counted.
    assert seen == [85 * GIB]
    assert waiter.bytes == 80 * GIB


def _big_sheet_workbook(path: Path, rows: int) -> Path:
    """Write a workbook with one large sheet, bypassing openpyxl's own cost."""

    body = "".join(
        f'<row r="{r}"><c r="A{r}"><v>{r}</v></c>'
        f'<c r="B{r}"><f>A{r}*2</f></c></row>'
        for r in range(1, rows + 1)
    )
    sheet = (
        '<?xml version="1.0"?><worksheet xmlns="http://schemas.openxmlformats.org'
        f'/spreadsheetml/2006/main"><sheetData>{body}</sheetData></worksheet>'
    )
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr(
            "xl/workbook.xml",
            '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/'
            '2006/main" xmlns:r="http://schemas.openxmlformats.org/'
            'officeDocument/2006/relationships"><sheets>'
            '<sheet name="Big" sheetId="1" r:id="rId1"/></sheets></workbook>',
        )
        archive.writestr(
            "xl/_rels/workbook.xml.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/'
            '2006/relationships"><Relationship Id="rId1" Target="worksheets/'
            'sheet1.xml" Type="x"/></Relationships>',
        )
        archive.writestr("xl/worksheets/sheet1.xml", sheet)
    return path


def test_probe_memory_stays_flat_on_a_large_sheet(tmp_path: Path) -> None:
    # The probe runs immediately before the agent on a swapless host, so its
    # own footprint must not scale with the sheet. Clearing a row is not
    # enough: the parser appends a whole read buffer to sheetData before
    # yielding that buffer's first end event, so the rows have to be dropped
    # as a group or they pile up.
    small = _big_sheet_workbook(tmp_path / "small.xlsx", 2_000)
    large = _big_sheet_workbook(tmp_path / "large.xlsx", 200_000)

    tracemalloc.start()
    tracemalloc.reset_peak()
    assert probe_workbook(small).formula_cells == 2_000
    _, small_peak = tracemalloc.get_traced_memory()
    tracemalloc.reset_peak()
    large_profile = probe_workbook(large)
    _, large_peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()

    # Counts must still be exact even though rows are dropped while their end
    # events are still queued.
    assert large_profile.formula_cells == 200_000
    assert large_profile.value_cells == 200_000
    # 100x the rows must not mean meaningfully more retained memory.
    assert large_peak < small_peak + 4 * 1024 * 1024
    assert large_peak < 8 * 1024 * 1024


def test_an_orchestrator_error_is_not_memory_evidence() -> None:
    # A launcher crash says nothing about how much memory the agent needed.
    records = [_observation(1.0) for _ in range(6)]
    records.append(_observation(4.0, status="orchestrator_error"))

    result = calibration_from_observations(records)

    assert result.scale == pytest.approx(1.0)
    assert result.censored_count == 0


def test_an_old_censored_spike_stops_pinning_the_scale() -> None:
    # The log is append-only. An unwindowed maximum would let one bad sample
    # hold the scale at its ceiling forever, stranding the host at one lane.
    records = [_observation(3.5, status="timed_out")]
    records += [_observation(1.0) for _ in range(10)]

    pinned = calibration_from_observations(records)
    records += [
        _observation(1.0, status="timed_out")
        for _ in range(memory_governor.CENSORED_WINDOW)
    ]
    recovered = calibration_from_observations(records)

    assert pinned.scale == pytest.approx(3.5)
    assert recovered.scale == pytest.approx(1.0)


def test_a_small_host_still_gets_a_working_budget() -> None:
    # Degrading to no gate at all is worst on the hosts with least room.
    assert memory_governor.reserve_bytes_for(8 * GIB) == 4 * GIB
    assert memory_governor.reserve_bytes_for(4 * GIB) == 2 * GIB


def test_report_does_not_write_the_budget_file(tmp_path: Path) -> None:
    # An operator running the report under sudo must not re-own the budget
    # file and lock the non-root lanes out of their own gate.
    budget = _budget(tmp_path)
    budget.acquire("lane-a", 8 * GIB, workbook_id="0001")
    before = budget.path.stat().st_mtime_ns

    snapshot = budget.snapshot()

    assert snapshot["used_bytes"] == 8 * GIB
    assert budget.path.stat().st_mtime_ns == before


def test_report_on_a_missing_budget_file_creates_nothing(tmp_path: Path) -> None:
    budget = _budget(tmp_path)

    snapshot = budget.snapshot()

    assert snapshot["used_bytes"] == 0
    assert not budget.path.exists()


def test_completed_observations_are_windowed() -> None:
    # A structural change to how much memory the pipeline uses -- disabling the
    # segment memory profiler, or overlapping two builds -- shifts real peaks.
    # An unwindowed quantile would keep weighting pre-change observations
    # forever, leaving every estimate permanently over-reserved.
    stale = [_observation(3.0) for _ in range(memory_governor.COMPLETED_WINDOW)]
    fresh = [_observation(1.0) for _ in range(memory_governor.COMPLETED_WINDOW)]

    assert calibration_from_observations(stale).scale == pytest.approx(3.0)
    assert calibration_from_observations(stale + fresh).scale == pytest.approx(1.0)


def test_the_window_still_needs_the_minimum_sample_count(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.delenv(memory_governor.SEED_SCALE_ENV, raising=False)
    few = [_observation(3.0) for _ in range(memory_governor.CALIBRATION_MIN_SAMPLES - 1)]

    assert calibration_from_observations(few).scale == memory_governor.DEFAULT_SEED_SCALE
