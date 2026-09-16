"""Defined logical constants retain their type and original reference provenance."""
import openpyxl
from openpyxl.workbook.defined_name import DefinedName
from xl_ast_graph import AstGraph, write_csv
from xl_seg.model import load
from xl_seg import evaluate, project


def test_defined_boolean_aliases_and_sheet_scope(tmp_path):
    path = tmp_path/'source.xlsx'
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = 'Budget'
    book.defined_names.add(DefinedName('Verify', attr_text='TRUE'))
    book.defined_names.add(DefinedName('Alias', attr_text='Verify'))
    sheet['A1'] = '=10*Verify'
    sheet['A2'] = '=IF(Alias,10,20)'
    other = book.create_sheet('Other')
    other.defined_names.add(DefinedName('Verify', attr_text='= false '))
    other['A1'] = '=10*Verify'
    other['A2'] = '=MissingName+1'
    book.save(path)
    graph = AstGraph(path, verbose=False).build()
    assert graph.resolve('Verify', 'Budget') == [('logical', 'TRUE')]
    assert graph.resolve('Alias', 'Budget') == [('logical', 'TRUE')]
    assert graph.resolve('Verify', 'Other') == [('logical', 'FALSE')]
    for owner, expected in [('Budget!A1', True), ('Budget!A2', True), ('Other!A1', False)]:
        nodes = [n for n in graph.nodes.values() if n['kind'] == 'const' and n['owner'] == owner and n['op'] == 'logical']
        assert len(nodes) == 1 and nodes[0]['value'] is expected
    assert 'name:TRUE' not in graph.nodes and 'name:FALSE' not in graph.nodes
    assert 'name:MissingName' in graph.nodes
    directory = tmp_path/'graph'
    directory.mkdir()
    write_csv(directory, graph.nodes, graph.edges)
    model = load(directory, 'source')
    for node in model.cells():
        if node.kind == 'formula':
            node.value, node.value_type = '99999', 'int'
    result = evaluate.Evaluator(model, project.build(model), strict_proof=True).run(set())
    assert result.values['Budget!A1'] == 10
    assert result.values['Budget!A2'] == 10
    assert result.values['Other!A1'] == 0


def test_only_exact_logical_definitions_resolve_without_assumptions(tmp_path):
    path = tmp_path/'source.xlsx'
    book = openpyxl.Workbook()
    for name, definition in [('TextFlag', '"TRUE"'), ('FormulaFlag', '1=1'), ('BrokenFlag', '#REF!'), ('CycleA', 'CycleB'), ('CycleB', 'CycleA')]:
        book.defined_names.add(DefinedName(name, attr_text=definition))
    book.active['A1'] = '=TextFlag+FormulaFlag+BrokenFlag+CycleA'
    book.save(path)
    graph = AstGraph(path, verbose=False).build()
    for name in ['TextFlag', 'FormulaFlag', 'BrokenFlag', 'CycleA']:
        assert all(item[0] != 'logical' for item in graph.resolve(name, book.active.title))
