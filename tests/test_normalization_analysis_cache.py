import json
from pathlib import Path
import openpyxl,pytest
import gen_normalizer as G
from scripts.fleet import normalization_analysis_cache as C

@pytest.fixture
def analysis(tmp_path):
    source=tmp_path/'0001.xlsx';inputs=tmp_path/'inputs';inputs.mkdir();baseline=inputs/'0001-inputs.xlsx'
    book=openpyxl.Workbook();book.active.title='Model';book.active['A1']=0.121;book.save(source);book.create_sheet('Embedded Assumptions')['D19']=0.121;book.save(baseline);book.close()
    run=tmp_path/'runs/0001-variable-sources';run.mkdir(parents=True);draft=run/'draft.json';draft.write_text(json.dumps(dict(rows=[dict(draft_id='v',variable_text='Historical value',value_text='Model!A1')])))
    runtime=tmp_path/'runtime.json';runtime.write_text('{}')
    selected=dict(workbook='0001',source=C.binding(source),baseline=C.binding(baseline),draft=C.binding(draft),normalizer=C.binding(G.__file__),dependencies=[C.binding(runtime)])
    args=('0001',source,C.binding(source)['sha256']);kwargs=dict(runs_root=run.parent,inputs_root=inputs)
    return selected,args,kwargs,runtime

def test_actual_analysis_once_exact_result_and_mutation_isolation(analysis):
    selected,args,kwargs,_=analysis;original=G.analyse;expected=original(*args,**kwargs)
    with C.installed(G,**selected)as metrics:
        first=G.analyse(*args,**kwargs);first[1][0]['clean']=True;first[0]['rows'].clear()
        assert G.analyse(*args,**kwargs)==expected
        assert metrics['original_calls']==1 and metrics['cache_hits']==1
    assert G.analyse is original
    assert expected[1][0]['code']=='duplicate_value_leak'

@pytest.mark.parametrize('key',['source','baseline','draft','runtime'])
def test_changed_bytes_reject_even_after_successful_analysis(analysis,key):
    selected,args,kwargs,runtime=analysis
    with C.installed(G,**selected):
        G.analyse(*args,**kwargs)
        path=runtime if key=='runtime'else Path(selected[key]['path']);path.write_bytes(path.read_bytes()+b' ')
        with pytest.raises(ValueError,match='dependency changed'):G.analyse(*args,**kwargs)

def test_changed_loaded_function_rejects_and_restores(analysis,monkeypatch):
    selected,args,kwargs,_=analysis;original=G.analyse
    with C.installed(G,**selected):
        G.analyse(*args,**kwargs)
        monkeypatch.setattr(G,'_analyse',lambda *a,**k:({},[]))
        with pytest.raises(ValueError,match='loaded dependency'):G.analyse(*args,**kwargs)
    assert G.analyse is original

def test_changed_arguments_use_original_failure_not_cached_pass(analysis):
    selected,args,kwargs,_=analysis
    with C.installed(G,**selected)as metrics:
        G.analyse(*args,**kwargs)
        with pytest.raises(SystemExit,match='source hash mismatch'):G.analyse(*args[:2],'0'*64,**kwargs)
        assert metrics['cache_hits']==0


def test_original_analysis_exception_restores_all_aliases(analysis):
    selected,args,kwargs,_=analysis;draft=Path(selected['draft']['path']);draft.write_text('{"rows":[{}]}');selected['draft']=C.binding(draft);original=G.analyse
    with pytest.raises(KeyError):
        with C.installed(G,**selected):G.analyse(*args,**kwargs)
    assert G.analyse is original


def test_nested_install_is_rejected_without_disturbing_outer(analysis):
    selected,args,kwargs,_=analysis;original=G.analyse
    with C.installed(G,**selected)as metrics:
        wrapper=G.analyse
        with pytest.raises(ValueError,match='original normalizer'):
            with C.installed(G,**selected):pass
        assert G.analyse is wrapper
        G.analyse(*args,**kwargs);G.analyse(*args,**kwargs)
        assert metrics['cache_hits']==1
    assert G.analyse is original


def test_in_place_function_code_assignment_is_rejected(analysis,monkeypatch):
    selected,args,kwargs,_=analysis
    with C.installed(G,**selected):
        G.analyse(*args,**kwargs)
        monkeypatch.setattr(G._analyse,'__code__',(lambda *a,**k:({},[])).__code__)
        with pytest.raises(ValueError,match='loaded dependency'):G.analyse(*args,**kwargs)


@pytest.fixture
def body_case(tmp_path):
    from test_harbor_review_context import body_fixture
    from scripts.batch import harbor_review_context as review
    body,args=body_fixture(tmp_path)
    records=[]
    for name in ('request','response','config','dependency'):
        path=tmp_path/(name+'.json');path.write_text('{}');records.append(C.binding(path))
    return review,body,args,dict(code=C.binding(review.__file__),dependencies=records)


def test_original_body_once_and_current_candidate_changes_reject(body_case):
    review,body,args,selected=body_case;original=review.normalization_body
    with C.body_installed(review,**selected)as metrics:
        assert review.normalization_body(body,**args)
        assert review.normalization_body(body,**args)
        assert metrics['original_calls']==1 and metrics['cache_hits']==1
        body['context']['Sheet'][0]['value']='Invented'
        with pytest.raises(ValueError,match='observations'):review.normalization_body(body,**args)
    assert review.normalization_body is original


@pytest.mark.parametrize('name',['source','baseline','draft','audit','request','response','config','dependency'])
def test_body_cache_rechecks_all_input_and_external_custody(body_case,name):
    review,body,args,selected=body_case
    with C.body_installed(review,**selected):
        review.normalization_body(body,**args)
        row=body[name]if name in body else next(r for r in selected['dependencies']if Path(r['path']).stem==name)
        path=Path(row['path']);path.write_bytes(path.read_bytes()+b' ')
        with pytest.raises(ValueError,match='input or dependency'):review.normalization_body(body,**args)


def test_body_exception_and_nested_install_restore_original(body_case):
    review,body,args,selected=body_case;original=review.normalization_body
    with pytest.raises(ValueError,match='observations'):
        with C.body_installed(review,**selected):
            wrapper=review.normalization_body
            with pytest.raises(ValueError,match='original normalization body'):
                with C.body_installed(review,**selected):pass
            assert review.normalization_body is wrapper
            body['context']['Sheet'][0]['value']='Wrong';review.normalization_body(body,**args)
    assert review.normalization_body is original


def test_body_keys_preserve_excel_date_types_and_reject_opaque_data():
    from datetime import datetime,date,time,timedelta
    for value in (datetime(2026,9,15,12,30),date(2026,9,15),time(12,30),timedelta(days=2)):
        assert C.body_key(value)==C.body_key(value)
        assert C.body_key(value)!=C.body_key(str(value))
    assert C.body_key([1])!=C.body_key((1,))
    assert C.body_key(True)!=C.body_key(1)
    with pytest.raises(TypeError):C.body_key(object())


def test_shared_body_key_matches_equal_copies_and_detects_mutation():
    import copy
    shared={'rate':[0.1,0.2]};rows=[shared]*100
    assert C.body_key(rows)==C.body_key([copy.deepcopy(shared)for _ in rows])
    before=C.body_key(rows);shared['rate'][1]=0.3
    assert C.body_key(rows)!=before


def test_body_dependency_rejects_symlinked_parent(body_case,tmp_path):
    review,body,args,selected=body_case;folder=tmp_path/'nested';folder.mkdir();file=folder/'dependency.json';file.write_text('{}');selected['dependencies'].append(C.binding(file))
    with C.body_installed(review,**selected):
        review.normalization_body(body,**args)
        folder.rename(tmp_path/'moved');folder.symlink_to(tmp_path/'moved',target_is_directory=True)
        with pytest.raises(ValueError,match='canonical cache'):review.normalization_body(body,**args)


def test_original_body_uses_two_parses_but_always_checks_current_claims(body_case):
    review,body,args,selected=body_case;loader=review.load_workbook;original=review.normalization_body
    with C.workbooks_installed(review,code=selected['code'],source=args['source'],baseline=args['baseline'],dependencies=selected['dependencies'])as metrics:
        assert review.normalization_body(body,**args)
        assert review.normalization_body(body,**args)
        assert review.normalization_body is original and metrics['workbook_loads']==2 and metrics['cache_hits']==3
        body['context']['Sheet'][0]['value']='Wrong'
        with pytest.raises(ValueError,match='observations'):review.normalization_body(body,**args)
        assert metrics['workbook_loads']==2
    assert review.load_workbook is loader


def test_workbook_view_falls_back_for_unseen_cells_and_other_callers(body_case):
    review,body,args,selected=body_case
    with C.workbooks_installed(review,code=selected['code'],source=args['source'],baseline=args['baseline'])as metrics:
        assert review.normalization_body(body,**args)
        body['context']['Sheet'].append(dict(reference='Sheet!C1',value=None,data_type='n',number_format='General'))
        assert review.normalization_body(body,**args)
        assert metrics['workbook_loads']==3
        book=review.load_workbook(args['source']['path'],data_only=False)
        assert isinstance(book,openpyxl.Workbook);book.close()


def test_workbook_cache_rejects_changed_input_and_nested_scope(body_case):
    review,body,args,selected=body_case;loader=review.load_workbook
    opts=dict(code=selected['code'],source=args['source'],baseline=args['baseline'])
    with C.workbooks_installed(review,**opts):
        assert review.normalization_body(body,**args)
        with pytest.raises(ValueError,match='original workbook parser'):
            with C.workbooks_installed(review,**opts):pass
        path=Path(args['baseline']['path']);path.write_bytes(path.read_bytes()+b' ')
        with pytest.raises(ValueError):review.normalization_body(body,**args)
    assert review.load_workbook is loader


@pytest.mark.parametrize('value',[None,False,0,0.01,'Text','=1+2','#REF!'])
def test_workbook_cache_preserves_exact_cell_types(body_case,value):
    review,body,args,selected=body_case
    for name in ('source','baseline'):
        path=Path(args[name]['path']);book=openpyxl.load_workbook(path);book.active['A1']=value;book.save(path)
        if name=='source':body['context']['Sheet'][0]=dict(reference='Sheet!A1',value=review.typed_cell_value(book.active['A1']),data_type=book.active['A1'].data_type,number_format=book.active['A1'].number_format)
        book.close();args[name]=C.binding(path);body[name]=args[name];body[name+'_sha256']=args[name]['sha256']
    with C.workbooks_installed(review,code=selected['code'],source=args['source'],baseline=args['baseline'])as metrics:
        assert review.normalization_body(body,**args)
        assert review.normalization_body(body,**args)
        assert metrics['workbook_loads']==2


def test_empty_body_still_runs_original_workbook_parser(body_case):
    from zipfile import BadZipFile
    review,body,args,selected=body_case
    path=Path(args['baseline']['path']);path.write_bytes(b'invalid workbook');args['baseline']=C.binding(path);body['baseline']=args['baseline'];body['baseline_sha256']=args['baseline']['sha256']
    args['rows']=[];body['candidates']=[];body['context']={};body['candidate_primitive_checks']=[]
    with C.workbooks_installed(review,code=selected['code'],source=args['source'],baseline=args['baseline']):
        with pytest.raises(BadZipFile):review.normalization_body(body,**args)
