import copy
from datetime import datetime
from pathlib import Path
import pytest
import openpyxl
from scripts.batch import harbor_normalization_compact as K
from scripts.batch import harbor_normalization_producer as N
from scripts.batch.harbor_admission import record


def evidence(tmp_path):
    p=tmp_path/'source.xlsx';w=openpyxl.Workbook();s=w.active;s.title="O'Brien 1"
    for row in range(1,24):
        s.cell(row,1,'Historical row '+str(row));s.cell(row,2,row*10)
        if row%3==0:s.cell(row,3,'period '+str(row))
        s.cell(row,4,datetime(2020,1,1));s.cell(row,5,'=B'+str(row)+'*2')
    s.cell(9,2,datetime(2019,1,1))
    w.save(p);refs=["O'Brien 1!B"+str(r) for r in (10,12,15,22)]
    binding=record(p);rows=[]
    for ref in refs:
        address=ref.rsplit('!',1)[1]
        rows.append(dict(reference=ref,cells=[[s.title,address,s[address].value]],duplicate_refs=[ref],
            clean=True,evidence=dict(source_sha256=binding['sha256']),row={'draft_id':'r1','source_notes':'exact row'},
            codes=[],extra=[],multi=False,evidence_paths={'source':binding}))
    ctx,checks=N.workbook_context(p,p,rows)
    return dict(candidates=rows,context=ctx,candidate_primitive_checks=checks,
        source=binding,baseline=binding,source_sha256=binding['sha256'],baseline_sha256=binding['sha256']),p


def test_exact_roundtrip_unknown_metadata_typed_values_and_order(tmp_path):
    body,p=evidence(tmp_path);body['metadata']={'arbitrary':[None,False,1,1.0,'001',{'x':'\r\n'}]}
    compact=K.compact(body)
    assert K.expand(compact)==body
    assert K.encoded(K.compact(body))==K.encoded(compact)
    assert len(compact['shared_candidate_groups'])==1


def test_subset_context_is_exact_original_producer_rule(tmp_path):
    body,p=evidence(tmp_path)
    for refs in ([body['candidates'][0]['reference']], [r['reference'] for r in body['candidates'][1:3]]):
        local=K.local_body(body,refs);rows=[r for r in body['candidates'] if r['reference'] in refs]
        context,checks=N.workbook_context(p,p,rows)
        assert (local['context'],local['candidate_primitive_checks'])==(context,checks)


def test_all_partitions_retain_exact_original_context_union(tmp_path):
    body,p=evidence(tmp_path)
    parts=K.partitions(body,lambda compact,refs:dict(evidence=compact,refs=refs),max_references=1)
    assert len(parts)==4
    assert [r for refs,_ in parts for r in refs]==[r['reference'] for r in body['candidates']]


@pytest.mark.parametrize('change',['negative','boolean','unknown','shadow','dropped_cell','duplicate'])
def test_compact_identity_and_coverage_cannot_be_forged(tmp_path,change):
    body,p=evidence(tmp_path);value=K.compact(body)
    if change=='negative':value['candidates'][0]['group']=-1
    elif change=='boolean':value['candidates'][0]['group']=True
    elif change=='unknown':value['candidates'][0]['group']=100
    elif change=='shadow':value['shared_candidate_groups'][0]['reference']='Fake!A1'
    elif change=='dropped_cell':
        sheet=next(iter(body['context']));body['context'][sheet]=[e for e in body['context'][sheet] if e['reference']!=body['candidates'][0]['reference']]
        with pytest.raises((ValueError,KeyError)):K.partitions(body,lambda c,r:c,max_references=1)
        return
    else:
        body['candidates'].append(body['candidates'][0])
        with pytest.raises(ValueError):K.partitions(body,lambda c,r:c)
        return
    with pytest.raises(ValueError):K.expand(value)


def test_single_candidate_never_drops_evidence_to_fit(tmp_path,monkeypatch):
    body,p=evidence(tmp_path);monkeypatch.setattr(K,'MAX_REQUEST_BYTES',100)
    with pytest.raises(ValueError,match='one complete candidate'):K.partitions(body,lambda c,r:c)


def test_byte_bounded_split_matches_greedy_reference_without_per_cell_rebuild(tmp_path,monkeypatch):
    body,p=evidence(tmp_path);refs=[r['reference'] for r in body['candidates']]
    sizes=[len(K.encoded(K.compact(K.local_body(body,refs[:n])))) for n in range(1,len(refs)+1)]
    assert sizes==sorted(sizes)
    monkeypatch.setattr(K,'MAX_REQUEST_BYTES',sizes[1])
    observed=[]
    parts=K.partitions(body,lambda c,r:(observed.append(tuple(r)) or c))
    assert all(len(K.encoded(req))<=sizes[1] for _,req in parts)
    assert [r for selected,_ in parts for r in selected]==refs
    for selected,req in parts:assert K.expand(req)==K.local_body(body,selected)
    assert parts[0][0]==refs[:2]


def test_normal_size_partition_uses_one_request_build_per_partition(tmp_path):
    body,p=evidence(tmp_path);seen=[]
    parts=K.partitions(body,lambda c,r:(seen.append(r) or c),max_references=2)
    assert len(parts)==len(seen)==2
