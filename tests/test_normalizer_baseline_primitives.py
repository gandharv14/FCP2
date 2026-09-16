import json
import openpyxl
import pytest
import gen_normalizer as G


@pytest.mark.parametrize('source_value,baseline_value,excluded',[
    (0.121,0.121,True),(1.5,1.5,True),('Acme','Acme',True),
    (1.5,2.5,False),(1,True,False),(1.5,'=1.5',False),
])
def test_new_baseline_primitive_is_considered_before_review(tmp_path,source_value,baseline_value,excluded):
    source=tmp_path/'0001.xlsx';inputs=tmp_path/'inputs';inputs.mkdir()
    book=openpyxl.Workbook();book.active.title='Model';book.active['A1']=source_value;book.save(source)
    book.create_sheet('Embedded Assumptions')['D3']=baseline_value;book.save(inputs/'0001-inputs.xlsx');book.close()
    run=tmp_path/'runs/0001-variable-sources';run.mkdir(parents=True)
    (run/'draft.json').write_text(json.dumps(dict(rows=[dict(draft_id='value',variable_text='Historical value',value_text='Model!A1')])) )
    raw=source.read_bytes();draft,rows=G.analyse('0001',source,G._sha256(source),runs_root=run.parent,inputs_root=inputs)
    assert source.read_bytes()==raw and len(rows)==1
    assert rows[0]['clean'] is not excluded
    assert rows[0]['duplicate_refs'].count('Model!A1')==1
    if excluded:
        assert rows[0]['code']=='duplicate_value_leak'
        assert rows[0]['leaks']==['Embedded Assumptions!D3']
        assert rows[0]['cells']==[('Model','A1',source_value)]


def test_historical_analysis_is_separate_from_current_production(tmp_path):
    source=tmp_path/'0001.xlsx';inputs=tmp_path/'inputs';inputs.mkdir()
    book=openpyxl.Workbook();book.active.title='Model';book.active['A1']=0.121;book.save(source)
    book.create_sheet('Embedded Assumptions')['D19']=0.121;book.save(inputs/'0001-inputs.xlsx');book.close()
    run=tmp_path/'runs/0001-variable-sources';run.mkdir(parents=True)
    (run/'draft.json').write_text(json.dumps(dict(rows=[dict(draft_id='value',variable_text='Historical value',value_text='Model!A1')])))
    args=('0001',source,G._sha256(source));kw=dict(runs_root=run.parent,inputs_root=inputs)
    old_draft,old=G.historical_source_analysis(*args,**kw)
    new_draft,current=G.analyse(*args,**kw)
    assert old_draft==new_draft
    assert old[0]['clean'] is True and old[0]['duplicate_refs']==['Model!A1']
    assert current[0]['clean'] is False and current[0]['leaks']==['Embedded Assumptions!D19']
    assert G.analyse(*args,**kw)==(new_draft,current)
    with pytest.raises(SystemExit,match='source hash mismatch'):
        G.historical_source_analysis('0001',source,'0'*64,**kw)
