"""Retained zero-offset endpoints require exact events and full-span proofs."""
import copy
import re
import pytest
from openpyxl import Workbook
from xl_ast_graph import AstGraph, write_csv, parse_formula
from xl_seg import model, project, restriction_cone as rc
from xl_segment import stabilize_runtime_proof
from test_runtime_proof_closure import CALC, _restricted_certificate_source, _passing_certificate_verification


def fixture(tmp_path, formula='=SUM(A1:OFFSET(C1,0,0))', middle=3):
    wb=Workbook(); ws=wb.active; ws.title='Sheet'
    for cell,value in {'A1':2,'B1':middle,'C1':5,'D1':7,'E1':11}.items(): ws[cell]=value
    ws['G1']=formula
    path=tmp_path/'book.xlsx'; wb.save(path)
    ast=AstGraph(path,verbose=False).build(); write_csv(tmp_path,ast.nodes,ast.edges)
    graph=model.load(tmp_path,'book'); graph.nodes['Sheet!G1'].value='999999'
    result, inputs, _, proof=stabilize_runtime_proof(graph,project.build(graph),
        {'Sheet!A1','Sheet!C1','Sheet!D1','Sheet!E1'}, {'Sheet!G1'}, calculation=CALC)
    events=[{'event':'volatile_function','function':'OFFSET','location':'Sheet!G1',
             'scope':'worksheet','token_offset':match.start()}
            for match in re.finditer('OFFSET',formula[1:])]
    return graph,result,inputs,proof,events


def cert(tmp_path, monkeypatch, graph, proof, events):
    source=_restricted_certificate_source(tmp_path,monkeypatch,events)
    verify=_passing_certificate_verification(proof,['Sheet!G1'])
    return rc.build_certificate(source_generation_dir=source,graph=graph,proof=proof,
        verification=verify,ordered_outputs=['Sheet!G1'],segmentation_fingerprints=verify['fingerprints'])


@pytest.mark.parametrize('formula,expected,calls',[
    ('=SUM(A1:OFFSET(C1,0,0))',10,1),
    ('=SUM($A$1:OFFSET($C$1,0.0,0E0))',10,1),
    ('=SUM(C1:OFFSET(A1,0,0))',10,1),
    ('=SUM(A1:OFFSET(C1,0,0))+SUM(A1:OFFSET(C1,0,0))',20,2),
    ('=SUM(A1:OFFSET(C1,0,0))+OFFSET(E1,0,0)',21,2),
    ('=SUM(A1:OFFSET(C1,0,0))+SUM(C1:OFFSET(E1,0,0))',33,2),
])
def test_exact_runtime_and_event_coverage(tmp_path,monkeypatch,formula,expected,calls):
    graph,result,inputs,proof,events=fixture(tmp_path,formula)
    assert result.values['Sheet!G1']==expected
    assert not result.unresolved
    assert 'Sheet!B1' in inputs
    c=cert(tmp_path,monkeypatch,graph,proof,events)
    assert c['event_count']==calls
    assert len({e['ast_operation_node'] for e in c['events']})==calls
    for event in c['events']:
        for span in event['enclosing_static_ranges']:
            assert set(span['static_targets'])==set(span['runtime_targets'])
            assert set(span['static_targets']) <= set(c['segmentation']['runtime_cone_cells'])


@pytest.mark.parametrize('tamper',['missing_event','duplicate_event','wrong_offset','missing_span','wrong_span','omitted_middle','altered_endpoint'])
def test_forged_or_incomplete_evidence_rejected(tmp_path,monkeypatch,tamper):
    graph,result,inputs,proof,events=fixture(tmp_path)
    colon=next(n for n in graph.nodes.values() if n.op==':')
    if tamper=='missing_event': events=[]
    elif tamper=='duplicate_event': events.append(copy.deepcopy(events[0]))
    elif tamper=='wrong_offset': events[0]['token_offset']+=1
    elif tamper=='missing_span': proof['resolved_operation_targets'].pop(colon.id)
    elif tamper=='wrong_span': proof['resolved_operation_targets'][colon.id]=['Sheet!A1','Sheet!C1']
    elif tamper=='omitted_middle': proof['runtime_radj']['Sheet!G1'].remove('Sheet!B1')
    elif tamper=='altered_endpoint': graph.nodes['Sheet!A1'].col=2
    with pytest.raises(rc.RestrictionConeError): cert(tmp_path,monkeypatch,graph,proof,events)


@pytest.mark.parametrize('formula',[
 '=SUM(A1:OFFSET(C1,D1,0))',
 '=SUM(A1:OFFSET(C1,0,0,2,1))','=SUM(A1:OFFSET(C1,TRUE,0))',
 "=SUM(A1:OFFSET('Other'!C1,0,0))",'=SUM(A1:OFFSET(C1:C2,0,0))',
])
def test_unproved_endpoint_stays_unsupported(formula):
    child=parse_formula(formula).args[0]
    assert child.kind=='func' and child.name.endswith(':OFFSET')


def test_interior_formula_is_reconstructed_from_its_inputs(tmp_path,monkeypatch):
    graph,result,inputs,proof,events=fixture(tmp_path,middle="=D1+E1")
    assert result.values['Sheet!G1']==25
    assert result.values['Sheet!B1']==18
    assert 'Sheet!B1' not in inputs
    assert {'Sheet!D1','Sheet!E1'} <= inputs
    certificate=cert(tmp_path,monkeypatch,graph,proof,events)
    assert 'Sheet!B1' in certificate['events'][0]['enclosing_static_ranges'][0]['static_targets']


def test_missing_interior_formula_ast_is_rejected(tmp_path,monkeypatch):
    graph,result,inputs,proof,events=fixture(tmp_path,middle="=D1+E1")
    graph.nodes['Sheet!B1'].parse_status='error'
    with pytest.raises(rc.RestrictionConeError,match='formula target is omitted'):
        cert(tmp_path,monkeypatch,graph,proof,events)
