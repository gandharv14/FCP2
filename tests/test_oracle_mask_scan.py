"""Read-only mask checks retain values and malformed-dimension behavior."""
import re
import zipfile

import openpyxl
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
import pytest

import xl_mcp_oracle as O


def _audit(refs):
    return [dict(variable_id='hidden', refs=refs, cell_value=12345,
                 mcp_value=12345, unit='')]


def _legacy_mask_result(path, refs):
    book = openpyxl.load_workbook(path, data_only=False, read_only=True)
    try:
        result = []
        for ref in sorted(O.expand_refs(refs)):
            sheet, _, coordinate = ref.rpartition('!')
            if sheet not in book.sheetnames:
                result.append(dict(cell=O.display_ref(ref), value='<unknown sheet>'))
            else:
                value = book[sheet][coordinate].value
                if value is not None:
                    result.append(dict(cell=O.display_ref(ref), value=repr(value)))
        return result
    finally:
        book.close()


def test_one_xml_scan_preserves_blank_formula_error_and_typed_mask_values(tmp_path, monkeypatch):
    path = tmp_path/'masked.xlsx'; book = openpyxl.Workbook(); ws = book.active; ws.title = 'Model'
    for row, value in enumerate([None, False, 0, '=1+2', '#N/A', '12345', None], 1):
        ws.cell(row, 1, value)
    ws['B7'] = 12345
    book.save(path)
    refs = ['Model!A1:A7', 'Missing!C4']
    expected = _legacy_mask_result(path, refs)
    def no_coordinate_reparse(self, key):
        raise AssertionError('covered mask triggered another worksheet XML parse')
    monkeypatch.setattr(ReadOnlyWorksheet, '__getitem__', no_coordinate_reparse)
    report = O.check_workbook(path, refs, _audit(refs), {('Model!B7', 'hidden'): 'retained duplicate'})
    assert report['masked_cells_not_blank'] == expected
    assert report['mask_cells'] == 8
    assert report['manifest_errors'] == report['unapproved_value_leaks'] == report['unused_allowlist_entries'] == []
    assert report['approved_value_leaks'] == [dict(cell='Model!B7', variable_id='hidden', value='12345', rendering='raw_number', reason='retained duplicate')]
    assert report['valid'] is False


@pytest.mark.parametrize('dimension', ['A1:A1', 'A1:B8'])
def test_underreported_dimensions_keep_original_outside_extent_lookup(tmp_path, monkeypatch, dimension):
    path = tmp_path/'dimensions.xlsx'; book = openpyxl.Workbook(); ws = book.active; ws.title = 'Model'
    ws['A1'] = 'label'; ws['D8'] = 12345; book.save(path)
    with zipfile.ZipFile(path) as z:
        items = {i.filename: z.read(i) for i in z.infolist()}
    items['xl/worksheets/sheet1.xml'] = re.sub(rb'<dimension ref="[^"]+"', ('<dimension ref="'+dimension+'"').encode(), items['xl/worksheets/sheet1.xml'])
    with zipfile.ZipFile(path, 'w', compression=zipfile.ZIP_DEFLATED) as z:
        for name, data in items.items(): z.writestr(name, data)
    refs = ['Model!D8', 'Model!Z99']; expected = _legacy_mask_result(path, refs)
    original = ReadOnlyWorksheet.__getitem__; queried = []
    def counted(self, key):
        queried.append(key); return original(self, key)
    monkeypatch.setattr(ReadOnlyWorksheet, '__getitem__', counted)
    actual = O.check_workbook(path, refs, _audit(refs), {})
    assert actual['masked_cells_not_blank'] == expected == [dict(cell='Model!D8', value='12345')]
    assert queried == ['D8', 'Z99']
    assert not actual['valid']
