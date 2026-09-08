"""Output-cone pruning: attribute verifier failures to outputs, drop only those.

Covers the attribution helper, the ``curation.toml`` rewrite, the retention
record and gate vocabulary, and one end-to-end segmentation where a single
output disagrees with the workbook and the rest must survive.
"""

from __future__ import annotations

import json
import re
import zipfile
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

import xl_ast_graph
import xl_segment
from xl_seg import emit, retention


# -- attribution ------------------------------------------------------------


def _bands():
    # Two single-cell outputs and one two-cell band.
    return [("S!O1", ["S!O1"]), ("S!O2", ["S!O2"]), ("S!O3:P3", ["S!O3", "S!P3"])]


def _radj():
    # O1 <- M1 <- I1 ; O2 <- M1, M2 ; O3 <- M3 ; P3 <- M3
    return {
        "S!O1": ["S!M1"],
        "S!O2": ["S!M1", "S!M2"],
        "S!O3": ["S!M3"],
        "S!P3": ["S!M3"],
        "S!M1": ["S!I1"],
        "S!M2": [],
        "S!M3": [],
    }


def test_cone_masks_label_shared_cells_with_every_output() -> None:
    mask = xl_segment._output_cone_masks(_radj(), _bands())
    assert mask["S!M1"] == 0b011  # feeds O1 and O2
    assert mask["S!M2"] == 0b010
    assert mask["S!M3"] == 0b100
    assert mask["S!I1"] == 0b011
    assert mask["S!P3"] == 0b100


def test_output_cell_failure_blames_only_its_band() -> None:
    attribution = {
        "output_failures": [{"cell": "S!O2", "verdict": "mismatch"}],
        "leaked_cells": [],
        "oracle_leak_cells": [],
        "unknown_op_cells": {},
        "blocking_cycles": [],
        "cone_wide": [],
    }
    implicated, cone_wide, unattributed = xl_segment._attribute_to_outputs(
        attribution, None, _radj(), _bands(), ["output_mismatch"]
    )
    assert implicated == {"S!O2": ["output_mismatch"]}
    assert cone_wide == [] and unattributed == []


def test_shared_leaked_cell_blames_both_cones_and_spares_the_third() -> None:
    attribution = {
        "output_failures": [],
        "leaked_cells": ["S!M1"],
        "oracle_leak_cells": [],
        "unknown_op_cells": {"S!M3": ["TABLE"]},
        "blocking_cycles": [],
        "cone_wide": [],
    }
    implicated, cone_wide, unattributed = xl_segment._attribute_to_outputs(
        attribution,
        None,
        _radj(),
        _bands(),
        ["seeded_inside_output_cone", "unknown_operation"],
    )
    assert implicated == {
        "S!O1": ["seeded_inside_output_cone"],
        "S!O2": ["seeded_inside_output_cone"],
        "S!O3:P3": ["unknown_operation"],
    }
    assert cone_wide == [] and unattributed == []


def test_cone_wide_reason_and_unexplained_code_stop_pruning() -> None:
    attribution = {
        "output_failures": [{"cell": "S!O1", "verdict": "unresolved"}],
        "leaked_cells": [],
        "oracle_leak_cells": [],
        "unknown_op_cells": {},
        "blocking_cycles": [],
        "cone_wide": ["runtime_dependency_closure_not_stabilized"],
    }
    implicated, cone_wide, _ = xl_segment._attribute_to_outputs(
        attribution,
        None,
        _radj(),
        _bands(),
        ["output_unresolved", "runtime_dependency_closure_not_stabilized", "ast_integrity_error"],
    )
    assert implicated == {"S!O1": ["output_unresolved"]}
    # Both the declared cone-wide code and the blocking code no cell explained.
    assert cone_wide == [
        "runtime_dependency_closure_not_stabilized",
        "ast_integrity_error",
    ]


def test_blocking_cycle_members_and_certificate_locations_are_blamed() -> None:
    from xl_seg.restriction_cone import RestrictionConeError

    attribution = {
        "output_failures": [],
        "leaked_cells": [],
        "oracle_leak_cells": [],
        "unknown_op_cells": {},
        "blocking_cycles": [
            {"cycle_id": 1, "reasons": ["active_cycle_uncertified"], "cells": ["S!M2"]}
        ],
        "cone_wide": [],
    }
    error = RestrictionConeError("OFFSET targets differ", locations=["S!M3"])
    implicated, cone_wide, unattributed = xl_segment._attribute_to_outputs(
        attribution, error, _radj(), _bands(), ["active_cycle_uncertified"]
    )
    assert implicated == {
        "S!O2": ["active_cycle_uncertified"],
        "S!O3:P3": ["restriction_certificate"],
    }
    assert cone_wide == [] and unattributed == []
    # A structural certificate error belongs to no cell: cone-wide.
    structural = RestrictionConeError("restriction event ledger hash is invalid")
    _, cone_wide, _ = xl_segment._attribute_to_outputs(
        attribution, structural, _radj(), _bands(), []
    )
    assert "restriction_certificate" in cone_wide
    # The generic "not fully passing" message adds nothing the verifier did not.
    generic = RestrictionConeError(
        "strict segmentation verification is not fully passing"
    )
    _, cone_wide, _ = xl_segment._attribute_to_outputs(
        attribution, generic, _radj(), _bands(), ["active_cycle_uncertified"]
    )
    assert "restriction_certificate" not in cone_wide


def test_cell_outside_every_cone_is_reported_not_guessed() -> None:
    attribution = {
        "output_failures": [],
        "leaked_cells": ["S!Z9"],
        "oracle_leak_cells": [],
        "unknown_op_cells": {},
        "blocking_cycles": [],
        "cone_wide": [],
    }
    implicated, _, unattributed = xl_segment._attribute_to_outputs(
        attribution, None, _radj(), _bands(), ["seeded_inside_output_cone"]
    )
    assert implicated == {}
    assert unattributed == [{"cell": "S!Z9", "reason": "seeded_inside_output_cone"}]


# -- curation rewrite and retention record ----------------------------------


CURATION = """# Output curation for workbook t.
# Auto-included at score >= 6.0; top 40 candidates listed.

[[output]]
band = "S!A1"
score = 9.0
include = true
name = "IRR"

[[output]]
band = "S!A2"
score = 8.0
include = true  # fallback: top-4 auto-include
name = "NPV"

[[output]]
band = "S!A3"
score = 7.0
include = false  # heuristic: true
name = "ROI"

[[output]]
band = "S!A4"
score = 6.5
include = false
name = "Bare"

[[output]]
band = "S!A5"
score = 6.2
include = false  # excluded: non_financial bands.csv kind=label
name = "Year"

[[output]]
band = "S!A6"
score = 2.0
include = false
name = "Low score"
"""


def test_apply_pruning_marks_only_included_bands_and_is_idempotent(tmp_path: Path) -> None:
    path = tmp_path / "curation.toml"
    path.write_text(CURATION, encoding="utf-8")
    changed = emit.apply_pruning(
        path, {"S!A1": ["output_mismatch", "unknown_operation"], "S!A4": ["x"]}
    )
    assert changed == 1  # A4 was already excluded and is left alone
    text = path.read_text(encoding="utf-8")
    assert "include = false  # pruned: output_mismatch,unknown_operation" in text
    assert 'band = "S!A4"\nscore = 6.5\ninclude = false\n' in text
    again = path.read_bytes()
    assert emit.apply_pruning(path, {"S!A1": ["output_mismatch"]}) == 0
    assert path.read_bytes() == again
    entries = {e["band"]: e["include"] for e in emit.read_curation(path)}
    assert entries == {
        "S!A1": False, "S!A2": True, "S!A3": False,
        "S!A4": False, "S!A5": False, "S!A6": False,
    }


def test_output_retention_classifies_every_marker(tmp_path: Path) -> None:
    path = tmp_path / "curation.toml"
    path.write_text(CURATION, encoding="utf-8")
    emit.apply_pruning(path, {"S!A1": ["output_mismatch"]})
    record = retention.output_retention(path)
    assert record["threshold"] == 6.0
    assert record["candidates_listed"] == 6
    assert record["candidates_auto_included"] == 5
    assert record["retained"] == 1
    assert record["pruned"] == 1
    assert record["manual_excluded"] == 1
    assert record["adjudicated_excluded"] == 1
    assert record["unexplained_excluded"] == 1
    assert record["reason_codes"] == ["non_financial", "output_mismatch"]
    assert record["bands"]["unexplained_excluded"] == ["S!A4"]
    assert record["bands"]["manual_excluded"][0]["note"] == "bands.csv kind=label"


def test_gate_findings_vocabulary_cap_and_resource_evidence() -> None:
    def record(manual, unexplained=(), auto=10):
        return {
            "candidates_auto_included": auto,
            "bands": {
                "unexplained_excluded": list(unexplained),
                "manual_excluded": [
                    {"band": f"S!B{i}", "code": code, "base": code.split(":")[0]}
                    for i, code in enumerate(manual)
                ],
            },
        }

    assert retention.gate_findings(record([])) == []
    assert retention.gate_findings(record(["non_financial", "check_cell"])) == []
    assert any("without a marker" in f for f in retention.gate_findings(record([], ["S!A4"])))
    assert any("not in the vocabulary" in f for f in retention.gate_findings(record(["because"])))
    assert any("no matching attempt evidence" in f for f in retention.gate_findings(record(["memory"])))
    assert retention.gate_findings(
        record(["memory"]), resource_evidence={"memory": True}
    ) == []
    assert any("no matching attempt evidence" in f for f in retention.gate_findings(
        record(["budget"]), resource_evidence={"memory": True}
    ))
    # Four hand exclusions out of ten auto-included exceed max(3, 25%) = 3.
    over = retention.gate_findings(record(["non_financial"] * 4))
    assert any("exceed the cap" in f for f in over)
    assert retention.gate_findings(record(["non_financial"] * 3)) == []
    assert retention.gate_findings(
        record(["non_financial"] * 3 + ["cone_wide:ast_integrity_error"])
    ) == []
    assert retention.manual_cap(40) == 10
    assert retention.manual_cap(4) == 3


# -- end to end -------------------------------------------------------------


def _inject_cached_value(path: Path, coord: str, value: str) -> None:
    with zipfile.ZipFile(path) as src:
        parts = {item.filename: src.read(item.filename) for item in src.infolist()}
    sheet = "xl/worksheets/sheet1.xml"
    xml = parts[sheet]
    attr = xml.find(f'r="{coord}"'.encode())
    assert attr != -1
    start = xml.rfind(b"<c", 0, attr)
    close = xml.find(b"</c>", attr)
    end = close + len(b"</c>")
    cell = xml[start:end]
    new_v = f"<v>{value}</v>".encode()
    if re.search(rb"<v\b", cell):
        cell = re.sub(rb"<v\b[^/]*/>|<v>[^<]*</v>", new_v, cell, count=1)
    else:
        cell = cell.replace(b"</c>", new_v + b"</c>", 1)
    parts[sheet] = xml[:start] + cell + xml[end:]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for name, data in parts.items():
            out.writestr(name, data)


def _fixture(root: Path, wb: str) -> None:
    """Two sink outputs; ``Net Income`` carries a wrong cached value."""
    (root / "src").mkdir(parents=True)
    path = root / "src" / f"{wb}.xlsx"
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "Summary"
    ws["A1"], ws["B1"] = "Revenue", 100
    ws["A2"], ws["B2"] = "Growth", 0.1
    ws["A3"], ws["B3"] = "Cost", 40
    ws["A4"], ws["B4"] = "IRR", "=B1*(1+B2)"
    ws["A5"], ws["B5"] = "EBITDA", "=B1-B3"
    ws["A6"], ws["B6"] = "Net Income", "=B5*(1-B2)"
    book.save(path)
    _inject_cached_value(path, "B4", "110")
    _inject_cached_value(path, "B5", "60")
    _inject_cached_value(path, "B6", "999")  # model says 54
    xl_ast_graph.main([str(path), "-o", str(root / "ast"), "-q"])
    (root / "empty.env").write_text("", encoding="utf-8")


def _args(root: Path, **overrides):
    base = dict(
        ast_dir=str(root / "ast"),
        source=str(root / "src"),
        out=str(root / "seg"),
        threshold=6.0,
        top=40,
        sample=400,
        lineage_max=4000,
        recurate=False,
        llm=False,
        model="unused",
        env_file=str(root / "empty.env"),
        no_verify=False,
        allow_legacy_ast=True,
        prune_failing_outputs=True,
        prune_rounds=1,
    )
    base.update(overrides)
    return SimpleNamespace(**base)


def test_end_to_end_prunes_the_mismatching_output_and_publishes_the_rest(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.delenv(xl_segment.PRUNE_TIME_BUDGET_ENV, raising=False)
    monkeypatch.delenv(xl_segment.PRUNE_DEADLINE_ENV, raising=False)
    wb = "prun01"
    _fixture(tmp_path, wb)

    # Without pruning the run fails on the one bad output.
    plain = xl_segment.segment(wb, _args(tmp_path, prune_failing_outputs=False, out=str(tmp_path / "seg_plain")))
    assert plain["verification"]["status"] == "fail"
    last = json.loads((tmp_path / "seg_plain" / wb / "last-verification.json").read_text())
    assert last["passed"] is False
    blamed = {row["band"]: row for row in last["by_output"]}
    assert blamed["Summary!B6"]["implicated"] is True
    assert blamed["Summary!B6"]["reasons"] == ["output_mismatch"]
    assert blamed["Summary!B4"]["implicated"] is False
    assert not (tmp_path / "seg_plain" / wb / "pruning-log.json").exists()

    payload = xl_segment.segment(wb, _args(tmp_path))
    assert payload["verification"]["status"] == "pass"
    seg = tmp_path / "seg" / wb
    curation = (seg / "curation.toml").read_text(encoding="utf-8")
    assert 'band = "Summary!B6"' in curation
    assert curation.count("# pruned: output_mismatch") == 1
    entries = {e["band"]: e["include"] for e in emit.read_curation(seg / "curation.toml")}
    assert entries["Summary!B4"] is True and entries["Summary!B6"] is False

    log = json.loads((seg / "pruning-log.json").read_text())
    assert log["final_status"] == "verified"
    assert [(r["round"], len(r["pruned"]), len(r["retained"])) for r in log["rounds"]] == [(0, 1, 1)]
    assert log["rounds"][0]["pruned"][0]["band"] == "Summary!B6"

    generations = sorted((seg / "generations").iterdir())
    assert len(generations) == 1
    manifest = json.loads((generations[0] / "generation-manifest.json").read_text())
    record = manifest["output_retention"]
    assert record["candidates_auto_included"] == 2
    assert record["retained"] == 1
    assert record["pruned"] == 1
    assert record["unexplained_excluded"] == 0
    assert record["bands"]["pruned"] == [{"band": "Summary!B6", "reasons": ["output_mismatch"]}]
    # The uncapped attribution never enters the published record.
    segments = json.loads((generations[0] / "segments.json").read_text())
    assert "attribution" not in segments["verification"]

    # Rerunning on the pruned curation republishes the same generation.
    again = xl_segment.segment(wb, _args(tmp_path))
    assert again["verification"]["status"] == "pass"
    assert sorted((seg / "generations").iterdir()) == generations
    assert json.loads((seg / "pruning-log.json").read_text())["final_status"] == "verified"


def test_time_budget_exhausted_exits_3_with_the_pruned_curation_on_disk(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    wb = "prun02"
    _fixture(tmp_path, wb)
    monkeypatch.setenv(xl_segment.PRUNE_TIME_BUDGET_ENV, "0")
    with pytest.raises(SystemExit) as exit_info:
        xl_segment.segment(wb, _args(tmp_path))
    assert exit_info.value.code == xl_segment.PRUNE_EXIT_CODE
    seg = tmp_path / "seg" / wb
    assert "# pruned: output_mismatch" in (seg / "curation.toml").read_text()
    log = json.loads((seg / "pruning-log.json").read_text())
    assert log["final_status"] == "pruned_pending_rerun:time_budget"
    assert not (seg / "generations").exists()

    # The rerun (budget lifted) picks up the pruned cone and passes.
    monkeypatch.delenv(xl_segment.PRUNE_TIME_BUDGET_ENV)
    payload = xl_segment.segment(wb, _args(tmp_path))
    assert payload["verification"]["status"] == "pass"
    log = json.loads((seg / "pruning-log.json").read_text())
    assert log["final_status"] == "verified" and log["verified_on_rerun"] is True


def test_round_cap_of_zero_exits_3_and_deadline_env_is_honoured(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    wb = "prun03"
    _fixture(tmp_path, wb)
    monkeypatch.delenv(xl_segment.PRUNE_TIME_BUDGET_ENV, raising=False)
    with pytest.raises(SystemExit) as exit_info:
        xl_segment.segment(wb, _args(tmp_path, prune_rounds=0))
    assert exit_info.value.code == xl_segment.PRUNE_EXIT_CODE
    log = json.loads((tmp_path / "seg" / wb / "pruning-log.json").read_text())
    assert log["final_status"] == "pruned_pending_rerun:rounds_exhausted"

    wb2 = "prun04"
    root2 = tmp_path / "two"
    _fixture(root2, wb2)
    monkeypatch.setenv(xl_segment.PRUNE_DEADLINE_ENV, "1")  # long past
    with pytest.raises(SystemExit) as exit_info:
        xl_segment.segment(wb2, _args(root2))
    assert exit_info.value.code == xl_segment.PRUNE_EXIT_CODE
    log = json.loads((root2 / "seg" / wb2 / "pruning-log.json").read_text())
    assert log["final_status"] == "pruned_pending_rerun:time_budget"
    assert log["deadline_epoch"] == 1.0
