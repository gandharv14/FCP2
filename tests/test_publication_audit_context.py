"""Current generation applicability is distinct from historical audit integrity."""
import copy
import hashlib
import json
from pathlib import Path
from unittest.mock import Mock

import openpyxl
import pytest

import xl_variable_source_audit as audit
from grader.tests.test_variable_source_audit import GOOD_MODEL_TABLE, write_two_variable_fixture
from scripts import qualified_publication as publication
from xl_seg.publication import write_inputs_sidecar


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def binding(path):
    return {'path': str(path), 'sha256': sha(path)}


def segmentation_metadata(context):
    b = context['bindings']
    return {'generation_id': b['segmentation_generation_id'],
            'manifest_sha256': b['segmentation_manifest_sha256'],
            'source_generation_id': b['source_generation_id'],
            'inventory_approval_sha256': b['inventory_approval_sha256'],
            'recalc_signals_sha256': b['recalc_signals_sha256']}


@pytest.fixture
def case(tmp_path, monkeypatch):
    artifact, original_segments = write_two_variable_fixture(tmp_path)
    source_dir = tmp_path / 'sources/0001/generations' / ('a' * 64)
    seg_dir = tmp_path / 'segments/0001/generations' / ('b' * 64)
    source_dir.mkdir(parents=True); seg_dir.mkdir(parents=True)
    (source_dir / 'generation-manifest.json').write_text('{"fixture":"source"}')
    (seg_dir / 'generation-manifest.json').write_text('{"fixture":"segmentation"}')
    (seg_dir / 'segments.json').write_bytes((original_segments / 'segments.json').read_bytes())
    source = source_dir / '0001.xlsx'; source.write_bytes(artifact.read_bytes())
    current = {'source_generation_id': source_dir.name, 'segmentation_generation_id': seg_dir.name,
        'source_manifest_sha256': sha(source_dir / 'generation-manifest.json'),
        'segmentation_manifest_sha256': sha(seg_dir / 'generation-manifest.json'),
        'inventory_approval_sha256': None, 'recalc_signals_sha256': None,
        'source_health_sha256': 'c' * 64, 'restriction_evidence_sha256': None,
        'release_id': None, 'task_generation_id': None}
    context = {'source_dir': str(source_dir), 'source_path': str(source),
               'source_sha256': sha(source), 'segmentation_dir': str(seg_dir), 'bindings': current}
    context['ast_dir'] = str(source_dir / 'ast')
    context['source_manifest'] = {'ast_provenance': {'builder': {'code': [
        {'name': p.name, 'sha256': sha(p), 'size_bytes': p.stat().st_size}
        for p in publication.source_publication._default_builder_paths()]}}}
    from xl_seg import publication as seg_publication
    live_resolver = Mock(side_effect=lambda *a, **k: (Path(context['segmentation_dir']), {}))
    monkeypatch.setattr(seg_publication, 'resolve_generation_by_id', live_resolver)
    resolver = Mock(side_effect=lambda *a, **k: context)
    monkeypatch.setattr(publication.release, 'resolve_build_context', resolver)
    sidecar = write_inputs_sidecar(artifact, seg_dir, {'generation_id': seg_dir.name}, pipeline_bindings=current)
    monkeypatch.setattr(audit, '_call_model', lambda *a, **k: (GOOD_MODEL_TABLE, 'stop'))
    output = tmp_path / 'audit'
    audit.generate_audit('0001', str(artifact), seg_dir, output, 'fixture-key', segmentation=segmentation_metadata(context))
    bindings = {key: binding(output / ('0001-inputs-variable-sources' + suffix)) for key, suffix in [
        ('variable_source_metadata', '.metadata.json'), ('variable_source_inventory', '.inventory.json'),
        ('variable_source_audit', '.md')]}
    bindings.update(source_generation={'path': str(source_dir), 'id': source_dir.name, 'manifest_sha256': current['source_manifest_sha256']},
                    segmentation_generation={'path': str(seg_dir), 'id': seg_dir.name, 'manifest_sha256': current['segmentation_manifest_sha256']},
                    source_file=binding(source), baseline_inputs=binding(artifact), baseline_segmentation=binding(sidecar))
    return {'bindings': bindings, 'context': context, 'resolver': resolver, 'artifact': artifact,
            'seg_dir': seg_dir, 'source_dir': source_dir, 'sidecar': sidecar, 'output': output, 'live_resolver': live_resolver}


def assert_history_valid(case):
    b = case['bindings']
    return audit.validate_saved_audit_coverage(*(b[k]['path'] for k in
        ('variable_source_metadata', 'variable_source_inventory', 'variable_source_audit')))


def rebind_baseline(case):
    b, context = case['bindings'], case['context']
    case['sidecar'] = write_inputs_sidecar(case['artifact'], context['segmentation_dir'],
        {'generation_id': context['bindings']['segmentation_generation_id']}, pipeline_bindings=context['bindings'])
    b['baseline_inputs'] = binding(case['artifact'])
    b['baseline_segmentation'] = binding(case['sidecar'])


def test_exact_current_context_binds_source_pair_and_baseline_without_mutation(case):
    before = {p: p.read_bytes() for p in case['output'].rglob('*') if p.is_file()}
    records = publication.validate_audit_bindings(case['bindings'])
    assert len(records) == 9
    assert all(sha(p) == digest for p, digest in records.items())
    assert case['live_resolver'].call_args.kwargs['validate_live_evidence'] is True
    assert case['live_resolver'].call_args.kwargs['require_pass'] is True
    assert {p: p.read_bytes() for p in before} == before
    case['resolver'].assert_called_once_with('0001', source_root=case['source_dir'].parent.parent,
        source_generation_id=case['source_dir'].name, segmentation_root=case['seg_dir'].parent.parent,
        segmentation_generation_id=case['seg_dir'].name)


@pytest.mark.parametrize('key', ['source_generation', 'segmentation_generation', 'source_file', 'baseline_inputs', 'baseline_segmentation'])
def test_missing_or_changed_current_evidence_fails(case, key):
    item = case['bindings'].pop(key)
    with pytest.raises(ValueError, match='missing qualified current audit binding'):
        publication.validate_audit_bindings(case['bindings'])
    case['bindings'][key] = {**item, ('manifest_sha256' if 'generation' in key and key != 'baseline_segmentation' else 'sha256'): '0' * 64}
    with pytest.raises(ValueError, match='changed qualified current audit evidence'):
        publication.validate_audit_bindings(case['bindings'])


@pytest.mark.parametrize('reason', ['current verifier fingerprint changed', 'source health changed', 'forged trusted receipt'])
def test_current_production_proof_failure_is_refresh_not_historical_acceptance(case, reason):
    assert assert_history_valid(case)['passed']
    case['resolver'].side_effect = ValueError(reason)
    with pytest.raises(audit.AuditRevalidationRequired, match=reason):
        publication.validate_audit_bindings(case['bindings'])


@pytest.mark.parametrize('kind', ['source', 'segmentation'])
def test_identical_financial_rows_cannot_rebind_history_to_new_generation(case, kind):
    assert assert_history_valid(case)['passed']
    context, b = case['context'], case['bindings']
    key = kind + '_generation'
    old = Path(b[key]['path']); new = old.with_name('d' * 64)
    import shutil
    shutil.copytree(old, new)
    b[key] = {**b[key], 'id': new.name, 'path': str(new)}
    context['bindings'][key + '_id'] = new.name
    context['source_dir' if kind == 'source' else 'segmentation_dir'] = str(new)
    if kind == 'source':
        context['source_path'] = str(new / '0001.xlsx'); b['source_file'] = binding(new / '0001.xlsx')
    rebind_baseline(case)
    with pytest.raises(audit.AuditRevalidationRequired, match='inventory/current source/segmentation context'):
        publication.validate_audit_bindings(b)


@pytest.mark.parametrize('field', ['inventory_approval_sha256', 'recalc_signals_sha256'])
def test_changed_approval_or_recalculation_context_cannot_reuse_saved_prompt(case, field):
    case['context']['bindings'][field] = 'e' * 64
    rebind_baseline(case)
    assert assert_history_valid(case)['passed']
    with pytest.raises(audit.AuditRevalidationRequired, match='context requires refresh'):
        publication.validate_audit_bindings(case['bindings'])


@pytest.mark.parametrize('field', ['value', 'document_properties', 'label', 'selection'])
def test_current_inventory_is_reconstructed_not_accepted_from_saved_hash(case, field):
    if field in ('value', 'document_properties'):
        wb = openpyxl.load_workbook(case['artifact'])
        if field == 'value': wb['Inputs']['B1'] = 0.3
        else: wb.properties.title = 'Changed workbook context'
        wb.save(case['artifact']); rebind_baseline(case)
    else:
        path = case['seg_dir'] / 'segments.json'; data = json.loads(path.read_text())
        if field == 'label': data['inputs'][1]['label'] = 'Changed financial basis'
        else: data['inputs'].pop()
        path.write_text(json.dumps(data))
    assert assert_history_valid(case)['passed']
    with pytest.raises(audit.AuditRevalidationRequired, match='context requires refresh'):
        publication.validate_audit_bindings(case['bindings'])


@pytest.mark.parametrize('change', ['pipeline', 'file', 'path'])
def test_rehashed_invalid_baseline_sidecar_still_fails(case, change):
    path = case['sidecar']; data = json.loads(path.read_text())
    if change == 'pipeline': data['pipeline_bindings']['source_health_sha256'] = 'wrong'
    elif change == 'file': data['inputs_sha256'] = 'wrong'
    else: path = path.with_name('wrong.segmentation.json')
    path.write_text(json.dumps(data)); case['bindings']['baseline_segmentation'] = binding(path)
    with pytest.raises(ValueError, match='sidecar'):
        publication.validate_audit_bindings(case['bindings'])


def test_rehashed_metadata_with_wrong_artifact_is_not_current(case):
    b = case['bindings']; path = Path(b['variable_source_metadata']['path']); meta = json.loads(path.read_text())
    meta['artifact'] = str(case['artifact'].with_name('other.xlsx'))
    path.write_text(json.dumps(meta)); b['variable_source_metadata'] = binding(path)
    assert assert_history_valid(case)['passed']
    with pytest.raises(audit.AuditRevalidationRequired, match='baseline artifact context'):
        publication.validate_audit_bindings(b)


def test_rehashed_metadata_with_wrong_workbook_is_not_current(case):
    b = case['bindings']; path = Path(b['variable_source_metadata']['path']); meta = json.loads(path.read_text())
    meta['workbook'] = 'other'; path.write_text(json.dumps(meta)); b['variable_source_metadata'] = binding(path)
    assert assert_history_valid(case)['passed']
    with pytest.raises(ValueError, match='workbook identity'):
        publication.validate_audit_bindings(b)


def test_source_binding_cannot_point_to_identical_unrelated_file(case):
    p = case['artifact'].with_name('unrelated.xlsx'); p.write_bytes(Path(case['context']['source_path']).read_bytes())
    case['bindings']['source_file'] = binding(p)
    with pytest.raises(ValueError, match='current source/segmentation binding mismatch'):
        publication.validate_audit_bindings(case['bindings'])


def test_real_inactive_source_segmentation_and_current_verifier_gate(tmp_path, monkeypatch):
    """Exercise actual production resolvers, proof, sidecar and auditor together."""
    import zipfile
    import xl_segment
    import xl_source_recalc
    from xl_seg import publication as seg_publication
    raw = tmp_path / '0001.xlsx'
    wb = openpyxl.Workbook(); ws = wb.active
    ws['A1'] = 3; ws['B1'] = '=A1*2'
    wb.calculation.fullCalcOnLoad = False; wb.calculation.calcMode = 'auto'; wb.save(raw)
    with zipfile.ZipFile(raw) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    parts['xl/worksheets/sheet1.xml'] = parts['xl/worksheets/sheet1.xml'].replace(b'<f>A1*2</f><v></v>', b'<f>A1*2</f><v>6</v>').replace(b'<f>A1*2</f><v />', b'<f>A1*2</f><v>6</v>')
    with zipfile.ZipFile(raw, 'w') as z:
        for n, data in parts.items(): z.writestr(n, data)
    srcroot = tmp_path / 'sources/0001'
    src, sm = xl_source_recalc.prepare_source_generation(raw, '0001', srcroot)
    assert xl_segment.main(['0001', '--source-generation-path', str(src), '-o', str(tmp_path / 'segments'), '--threshold', '0', '--no-auto-adjudication']) == 0
    seg = next((tmp_path / 'segments/0001/generations').iterdir())
    context = publication.release.resolve_build_context('0001', source_root=srcroot, source_generation_id=src.name,
        segmentation_root=seg.parent.parent, segmentation_generation_id=seg.name)
    baseline = tmp_path / '0001-inputs.xlsx'; baseline.write_bytes(raw.read_bytes())
    sidecar = write_inputs_sidecar(baseline, seg, {'generation_id': seg.name}, pipeline_bindings=context['bindings'])
    # Match the auditor's independently resolved metadata, not a test-local copy.
    _, segmeta = audit.resolve_segmentation_directory(tmp_path / 'segments', '0001', segmentation_generation_id=seg.name,
        source_generation_root=tmp_path / 'sources', source_generation_id=src.name)
    assert segmeta == segmentation_metadata(context)
    inventory = audit.build_inventory('0001', baseline, seg, segmentation=segmeta)
    assert inventory['variables']
    table = '| Variable | Inventory row IDs | Plausible external source(s) |\n| --- | --- | --- |\n' + '\n'.join(
        f"| Supplied assumption {i} | {row['row_id']} | Supplied workbook input |" for i, row in enumerate(inventory['variables'], 1))
    monkeypatch.setattr(audit, '_call_model', lambda *a, **k: (table, 'stop'))
    output = tmp_path / 'audit'; audit.generate_audit('0001', str(baseline), seg, output, 'fixture-key', segmentation=segmeta)
    bindings = {key: binding(output / ('0001-inputs-variable-sources' + suffix)) for key, suffix in [
        ('variable_source_metadata', '.metadata.json'), ('variable_source_inventory', '.inventory.json'), ('variable_source_audit', '.md')]}
    bindings.update(source_generation={'path': str(src), 'id': src.name, 'manifest_sha256': sha(src / 'generation-manifest.json')},
        segmentation_generation={'path': str(seg), 'id': seg.name, 'manifest_sha256': sha(seg / 'generation-manifest.json')},
        source_file=binding(context['source_path']), baseline_inputs=binding(baseline), baseline_segmentation=binding(sidecar))
    assert len(publication.validate_audit_bindings(bindings)) == 9
    assert not (srcroot / 'current.json').exists()
    audit.validate_saved_audit_coverage(*(bindings[k]['path'] for k in ['variable_source_metadata', 'variable_source_inventory', 'variable_source_audit']))
    changed = tmp_path / 'changed-verifier.py'; changed.write_text('# changed verifier\n')
    original = seg_publication.verifier_paths()
    monkeypatch.setattr(seg_publication, 'verifier_paths', lambda: (*original, changed))
    with pytest.raises(audit.AuditRevalidationRequired, match='proof requires refresh'):
        publication.validate_audit_bindings(bindings)


def test_current_builder_content_differs_from_retained_provenance(case):
    case['context']['source_manifest']['ast_provenance']['builder']['code'][0]['sha256'] = '0' * 64
    with pytest.raises(audit.AuditRevalidationRequired, match='current production AST builder changed'):
        publication.validate_audit_bindings(case['bindings'])


def test_live_segmentation_failure_not_hidden_by_pinned_context(case):
    case['live_resolver'].side_effect = ValueError('changed live verification evidence')
    with pytest.raises(audit.AuditRevalidationRequired, match='changed live verification evidence'):
        publication.validate_audit_bindings(case['bindings'])
