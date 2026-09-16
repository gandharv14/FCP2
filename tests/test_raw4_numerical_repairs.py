"""Prospective regression proofs; never rewrite or replay corpus attempts.

PERCENTILE semantics: https://support.microsoft.com/en-us/excel/functions/percentile-function
Empty arrays: https://learn.microsoft.com/en-us/office/vba/api/excel.worksheetfunction.percentile_inc
"""
import math

import openpyxl
import pytest

from xl_ast_graph import AstGraph, SheetIndex, write_csv
from xl_seg.evaluate import Evaluator, ExcelError, RangeValues, Unresolved
from xl_seg.model import load
from xl_seg.project import build as project
from test_evaluate import evaluator, make_node


def apply_percentile(args, name='PERCENTILE'):
    node = make_node('Sheet!Z1#0:' + name, 'op', owner='Sheet!Z1',
                     op=name, op_kind='func', arity=len(args))
    ev = evaluator([node])
    ev._args = lambda unused: args
    result = ev._apply(node)
    assert not ev.unknown_ops
    return result


@pytest.mark.parametrize('name', ['PERCENTILE', 'PERCENTILE.INC', '_XLFN.PERCENTILE.INC'])
@pytest.mark.parametrize('values,k,expected', [
    ([1, 3, 2, 4], .3, 1.9),
    ([40, 10, 20, 30], 0, 10),
    ([40, 10, 20, 30], 1, 40),
    ([7], .83, 7),
    ([2, 2, 2, 8], .75, 3.5),
    ([-9, -5, -1], .25, -7),
    ([1, None, '99', False, 3], .5, 2),
    ([[1, 2], [3, 4]], .5, 2.5),
    ([-1e308, 1e308], .5, 0),
])
def test_percentile_inclusive_interpolation(name, values, k, expected):
    assert apply_percentile([values, k], name) == pytest.approx(expected)


@pytest.mark.parametrize('values,k,code', [
    ([], .5, '#NUM!'),
    ([None, True, '7'], .5, '#NUM!'),
    ([1, 2], -.01, '#NUM!'),
    ([1, 2], 1.01, '#NUM!'),
    ([1, 2], math.inf, '#NUM!'),
    ([1, 2], 'invalid', '#VALUE!'),
    ([1, 2], '2026-01-01', '#VALUE!'),
    ([1, 2], math.nan, '#NUM!'),
    ([1, 2], RangeValues([.2, .8], 2, 1), '#VALUE!'),
    ([1, math.inf], .5, '#NUM!'),
])
def test_percentile_invalid_inputs_are_not_silently_coerced(values, k, code):
    result = apply_percentile([values, k])
    assert isinstance(result, ExcelError) and result.code == code


def test_percentile_scalar_reference_error_and_unknown_propagation():
    assert apply_percentile([[10, 30], RangeValues([.25], 1, 1)]) == 15
    assert apply_percentile([[10, 30], '0.25']) == 15
    error = ExcelError('#N/A')
    assert apply_percentile([[10, error, 30], .5]) is error
    unknown = Unresolved('missing-source')
    assert apply_percentile([[10, unknown, 30], .5]) is unknown
    assert isinstance(apply_percentile([[1, 2]]), Unresolved)


@pytest.mark.parametrize('sheet,formula', [
    ('Cost_Engine', "=SUM(Cost_Engine!B5:'Cost_Engine'!B8)+Cost_Engine!B12"),
    ('Cost_Engine', "=SUM('Cost_Engine'!$B$5:Cost_Engine!$B$8)+Cost_Engine!B12"),
    ('Cost_Engine', "=SUM(cost_engine!B5:'COST_ENGINE'!B8)+Cost_Engine!B12"),
    ("O'Brien", "=SUM('O''Brien'!B5:'O''Brien'!B8)+'O''Brien'!B12"),
    ('Bang!Sheet', "=SUM('Bang!Sheet'!B5:'Bang!Sheet'!B8)+'Bang!Sheet'!B12"),
    ('Cost Engine', "=SUM('Cost Engine'!B8:'Cost Engine'!B5)+'Cost Engine'!B12"),
])
@pytest.mark.parametrize('storage', ['memory', 'disk'])
def test_repeated_sheet_endpoint_range_retains_all_dependencies(tmp_path, sheet, formula, storage):
    path = tmp_path / 'source.xlsx'
    book = openpyxl.Workbook(); inputs = book.active; inputs.title = sheet
    for row, value in enumerate([10, 20, 30, 40], 5):
        inputs.cell(row, 2, value)
    inputs['B12'] = 5
    output = book.create_sheet('Output'); output['A1'] = formula
    book.save(path); book.close(); original = path.read_bytes()
    ast_dir = tmp_path / 'ast'; ast_dir.mkdir()
    with AstGraph(path, verbose=False, graph_storage=storage, storage_dir=tmp_path) as ast:
        ast.build(); write_csv(ast_dir, ast.nodes, ast.edges)
        dependencies = {e['source'] for e in ast.edges
                        if e['cell'] == 'Output!A1' and ast.nodes[e['source']]['kind'] == 'input'}
        assert dependencies == {f'{sheet}!B{r}' for r in (5, 6, 7, 8, 12)}
        assert not any(n['kind'] in ('name', 'external') for n in ast.nodes.values())
    graph = load(ast_dir, 'test', storage='memory')
    ev = Evaluator(graph, project(graph), strict_proof=True, run_probes=False)
    ev.run(dependencies)
    assert ev._eval_node('Output!A1') == 105
    assert path.read_bytes() == original


def test_cross_sheet_or_unknown_repeated_endpoints_fail_closed():
    graph = AstGraph('not-read.xlsx', verbose=False)
    graph.index = {name: SheetIndex() for name in ('First', 'Last')}
    for name in graph.index:
        graph.index[name].add(1, 1); graph.index[name].finalise()
    graph.sheet_names = ['First', 'Last']
    for token in ('First!A1:Last!A1', "Missing!A1:'Missing'!A2",
                  "First!XFE1:'First'!XFE2", "First!A1048577:'First'!A1048578"):
        assert graph.resolve(token, 'First') == [('name', token)]
    token = "'[other.xlsx]First'!A1:'[other.xlsx]First'!A2"
    assert graph.resolve(token, 'First') == [('ext', token)]
    # Genuine 3-D sheet-span syntax retains its established interpretation.
    assert graph.resolve('First:Last!A1', 'First') == [('cell', 'First', 1, 1), ('cell', 'Last', 1, 1)]
