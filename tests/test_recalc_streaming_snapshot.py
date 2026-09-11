"""Exact semantic equivalence and bounded XML retention for large worksheets."""
from pathlib import Path
import gc
import tracemalloc
import zipfile
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook
import xl_source_recalc as recalc
from test_table_recalc_observation import fixture as table_fixture


def dom_records(archive, part):
    root = ET.fromstring(archive.read(part))
    for cell in root.findall('.//{*}c'):
        yield 'cell', cell
    for merge in root.findall('.//{*}mergeCell'):
        yield 'merge', merge


@pytest.mark.parametrize('tables', [False, True])
def test_complete_snapshot_identical_to_retained_dom_semantics(tmp_path, monkeypatch, tables):
    path = tmp_path / 'fixture.xlsx'
    if tables:
        table_fixture(path)
    else:
        wb = Workbook(); ws = wb.active; ws.title = "1 Owner's model"
        ws.append(['Revenue', 123.5, '=B1*2', True, '#N/A'])
        ws['A3'] = 'Rich & literal <text>'; ws.merge_cells('A3:B3')
        ws['C4'] = '=SUM($B$1,C1)'
        other = wb.create_sheet('Cash flow'); other['A1'] = "='1 Owner''s model'!$C$4"
        wb.save(path)
    actual = recalc.semantic_workbook_snapshot(path)
    with monkeypatch.context() as context:
        context.setattr(recalc, '_worksheet_records', dom_records)
        expected = recalc.semantic_workbook_snapshot(path)
    assert actual == expected
    assert recalc.semantic_workbook_diff(actual, expected)['semantic_changes'] == []
    if tables:
        assert any(v.get('type') == 'dataTableResult' for s in actual['formula_caches'].values() for v in s.values())


def test_large_worksheet_parser_memory_does_not_retain_completed_rows(tmp_path):
    path = tmp_path / 'sheet.zip'
    rows = ''.join(f'<row r="{r}"><c r="A{r}"><f>1+2</f><v>3</v></c></row>' for r in range(1, 15001))
    xml = '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>' + rows + '</sheetData></worksheet>'
    with zipfile.ZipFile(path, 'w') as archive: archive.writestr('sheet.xml', xml)
    del xml, rows
    def measure(reader):
        gc.collect(); tracemalloc.start()
        with zipfile.ZipFile(path) as archive:
            count = sum(1 for kind, cell in reader(archive, 'sheet.xml') if kind == 'cell' and cell.find('{*}f').text == '1+2')
        _, peak = tracemalloc.get_traced_memory(); tracemalloc.stop()
        assert count == 15000
        return peak
    dom = measure(dom_records); streamed = measure(recalc._worksheet_records)
    assert streamed < dom / 4, (streamed, dom)


@pytest.mark.parametrize('xml', [
    '<worksheet><sheetData><row><c r="A1"><v>1</v></c></row></sheetData></broken>',
    '<worksheet><c r="A1"><c r="B1"><v>1</v></c></c></worksheet>',
])
def test_malformed_tail_and_nested_cells_fail_closed(tmp_path, xml):
    path = tmp_path / 'sheet.zip'
    with zipfile.ZipFile(path, 'w') as archive: archive.writestr('sheet.xml', xml)
    with zipfile.ZipFile(path) as archive, pytest.raises((ET.ParseError, recalc.RecalculationError)):
        list(recalc._worksheet_records(archive, 'sheet.xml'))
