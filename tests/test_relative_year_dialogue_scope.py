"""Relative project-year grouping proves every original source link and header."""
import copy
import hashlib
import importlib
import json
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.utils import get_column_letter


@pytest.fixture
def relative_case(tmp_path, monkeypatch):
    scripts = Path.cwd()/'.cursor/skills/additional-assumptions-dialogue/scripts'
    monkeypatch.syspath_prepend(str(scripts))
    scope = importlib.import_module('claim_scope')
    extract = importlib.import_module('extract_claims')
    book = Workbook()
    book.active.title = 'Summary'
    book.create_sheet('Operations')
    claims, rows = [], []
    for year in range(11):
        target = f'Summary!{get_column_letter(year+3)}24'
        origin = f'Operations!{get_column_letter(year+2)}139'
        th = f'Summary!{get_column_letter(year+3)}16'
        sh = f'Operations!{get_column_letter(year+2)}3'
        formula = '=' + origin.replace('139', '$139')
        book['Summary'][target.split('!')[1]] = formula
        book['Summary'][th.split('!')[1]] = f'Year {year}'
        book['Operations'][sh.split('!')[1]] = f'Year {year}'
        claim = extract.build_claim(dict(entry='source_selection', value='source', cells=[target],
            fields=dict(label='"Income Before Fixed Charges"', ingredient=f'cell {origin} on the row labelled "Income Before Fixed Charges"')), year+1)
        claims.append(claim)
        rows.append(dict(claim_id=claim['record_id'], target_reference=target, source_reference=origin,
            target_period_header=th, source_period_header=sh, target_period_axis_start='Summary!C16',
            source_period_axis_start='Operations!B3', period_kind='relative_year', period=year,
            original_formula=formula))
    unrelated = dict(record_id='unrelated-method', sheet='Summary', row_label='Expense',
        spoken='Multiply the input by two.', must_say=['multiply', 'two'],
        reviewer_only=dict(entry='method_operating_expense', cells=['Summary!P40']))
    claims.insert(3, unrelated)
    source = tmp_path/'original.xlsx'
    disclosure = tmp_path/'disclosure.json'
    disclosure.write_text('{}')
    evidence = tmp_path/'scope.json'
    payload = dict(schema_version='dialogue-source-link-scope/v3', claim_ids=[c['record_id'] for c in claims], rows=rows)
    def prove(*, rebind=True):
        book.save(source)
        if rebind:
            payload.update(source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
                           disclosure_sha256=hashlib.sha256(disclosure.read_bytes()).hexdigest())
        evidence.write_text(json.dumps(payload))
        chosen = copy.deepcopy(claims)
        result = scope.prove_groups(chosen, disclosure, source, evidence)
        return result, chosen
    return book, claims, payload, prove, disclosure


def test_all_eleven_source_links_share_one_exchange_without_losing_claims(relative_case):
    book, claims, payload, prove, _ = relative_case
    result, chosen = prove()
    assert [len(group) for group in result['groups']] == [11, 1]
    assert sorted(c for g in result['groups'] for c in g) == sorted(payload['claim_ids'])
    assert result['proven_claim_ids'] == [r['claim_id'] for r in payload['rows']]
    assert chosen[3] == claims[3]
    for old, current in zip(claims, chosen):
        assert old['reviewer_only'] == current['reviewer_only']
        if old['record_id'] != 'unrelated-method':
            assert current['must_say'] == old['must_say'] + ['matching period']
    assert not any('matching period' in c['must_say'] for c in claims)


@pytest.mark.parametrize('bad', ['1', 1, True, 'Price 1', 'Year 01', 'Year -1', 'Year 1 USD',
                               '= "Year 1"', 'Year one', 'Year 2', None])
def test_relative_year_needs_typed_original_consecutive_labels(relative_case, bad):
    book, _, _, prove, _ = relative_case
    book['Operations']['C3'] = bad
    with pytest.raises(ValueError):
        prove()


@pytest.mark.parametrize('fault', ['different_period', 'bool_period', 'foreign_axis', 'wrong_axis_row',
    'axis_after_header', 'wrong_header_column', 'header_below_value', 'wrong_link', 'arithmetic_link',
    'literal_formula', 'unknown_claim', 'duplicate_claim', 'missing_claim', 'target_substitution',
    'v2_schema', 'unbound_hash', 'disclosure_hash'])
def test_reject_ambiguous_or_changed_scope_evidence(relative_case, fault):
    book, claims, payload, prove, disclosure = relative_case
    row = payload['rows'][0]
    if fault == 'different_period': row['period'] = 1
    elif fault == 'bool_period': row['period'] = False
    elif fault == 'foreign_axis': row['source_period_axis_start'] = 'Summary!B3'
    elif fault == 'wrong_axis_row': row['source_period_axis_start'] = 'Operations!B4'
    elif fault == 'axis_after_header': row['source_period_axis_start'] = 'Operations!C3'
    elif fault == 'wrong_header_column': row['source_period_header'] = 'Operations!C3'
    elif fault == 'header_below_value': row['target_period_header'] = 'Summary!C25'
    elif fault == 'wrong_link': book['Summary']['C24'] = '=Operations!C$139'
    elif fault == 'arithmetic_link':
        book['Summary']['C24'] = row['original_formula'] = '=Operations!B$139+0'
    elif fault == 'literal_formula': book['Summary']['C24'].data_type = 's'
    elif fault == 'unknown_claim': row['claim_id'] = 'unknown'
    elif fault == 'duplicate_claim': payload['rows'].append(copy.deepcopy(row))
    elif fault == 'missing_claim': payload['claim_ids'].pop()
    elif fault == 'target_substitution': row['target_reference'] = 'Summary!D24'
    elif fault == 'v2_schema': payload['schema_version'] = 'dialogue-source-link-scope/v2'
    elif fault == 'unbound_hash':
        prove()
        book['Summary']['Z99'] = 'changed source'
    else:
        prove()
        disclosure.write_text('{"changed":true}')
    with pytest.raises((ValueError, KeyError)):
        prove(rebind=fault not in {'unbound_hash', 'disclosure_hash'})


def test_partial_scope_keeps_unproved_source_claim_individual(relative_case):
    _, claims, payload, prove, _ = relative_case
    omitted = payload['rows'].pop(5)['claim_id']
    result, chosen = prove()
    assert [omitted] in result['groups']
    assert len(result['proven_claim_ids']) == 10
    assert sorted(c for g in result['groups'] for c in g) == sorted(payload['claim_ids'])
    assert next(c for c in chosen if c['record_id'] == omitted) == next(c for c in claims if c['record_id'] == omitted)
