"""Decimal precision and multiple rounding must retain different meanings."""
import importlib.util
from pathlib import Path

import openpyxl
import pytest
from test_disclosure_copied_scope import disclose
from xl_ast_graph import parse_formula


@pytest.fixture
def spoken():
    p = Path(__file__).resolve().parents[1] / '.cursor/skills/additional-assumptions-dialogue/scripts/spoken_formula.py'
    spec = importlib.util.spec_from_file_location('rounding_spoken', p)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.mark.parametrize('digits', [0, 2, -2])
def test_round_precision_survives_ast_and_spoken_rendering(spoken, digits):
    steps = disclose.describe_ast(parse_formula(f'=ROUND(MAX(13.25,9),{digits})'), None, 'S', '')
    text = spoken.flatten_spoken(steps)
    assert 'decimal places' in text and 'nearest multiple' not in text
    assert 'greater' in text and '13.25' in text and '9' in text
    assert f'{digits} decimal places' in text


def test_mround_keeps_existing_multiple_semantics(spoken):
    steps = disclose.describe_ast(parse_formula('=MROUND(13.25,5)'), None, 'S', '')
    assert spoken.flatten_spoken(steps) == 'round 13.25 to the nearest multiple of 5'


@pytest.mark.parametrize('text', [
    'round applied to take the greater of (13.25) and (9), 0',
    'round (13.25) to an unspecified scale',
    'round (13.25) to the nearest multiple of',
])
def test_unrecognized_rounding_prose_cannot_invent_a_multiple(spoken, text):
    assert spoken.flatten_spoken(text) == text


def test_actual_affected_formula_keeps_every_reference_and_decimal_precision(spoken, tmp_path):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Matrix'
    for row in range(1, 33):
        ws.cell(row, 3, f'Category {row}')
    for row, label in [(10, '60+'), (11, 'FCL'), (22, 'Loss Mitigation'), (31, 'Loss Mitigation')]:
        ws.cell(row, 3, label)
    ws['I31'] = '=ROUND(MAX(G31,(I10+I11)/I22),0)'
    source = tmp_path/'source.xlsx'; wb.save(source)
    book = disclose.Book(source)
    steps = disclose.describe_ast(parse_formula(ws['I31'].value), book, 'Matrix', 'Loss Mitigation')
    assert all(f'Matrix!{c}' in steps for c in ('G31', 'I10', 'I11', 'I22'))
    text = spoken.speak_steps(steps, 'Matrix!I31', ws['I31'].value, 'Matrix', 'Loss Mitigation', 'only for this cell')
    assert '0 decimal places' in text and 'nearest multiple' not in text
    assert 'last period' in text and 'this period' in text and 'greater' in text
    assert not spoken.CELL_PHRASE_RE.search(text)


def test_nested_round_and_mround_retain_both_operators(spoken):
    steps = disclose.describe_ast(parse_formula('=MROUND(ROUND(13.25,1),5)'), None, 'S', '')
    text = spoken.flatten_spoken(steps)
    assert text == 'round round 13.25 to 1 decimal places to the nearest multiple of 5'
