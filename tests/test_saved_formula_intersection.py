"""Saved formula intersection uses syntax and context, never cached answers."""
import pytest
from openpyxl.worksheet.formula import ArrayFormula
from test_blank_formula_results import graph_for, run
from xl_seg import evaluate as candidate

def test_saved_normal_average_scalar_intersects_same_column(tmp_path):
    graph=graph_for(tmp_path,{'B3':'=AVERAGE(A1:C1/-A2)'}, {'A1':2,'B1':10,'C1':30,'A2':-2})
    _,result=run(candidate,graph)
    assert result.values['S!B3']==5
    assert result.runtime_radj['S!B3']=={'S!B1','S!A2'}


def test_saved_normal_average_scalar_intersects_same_row(tmp_path):
    graph=graph_for(tmp_path,{'C2':'=AVERAGE(A1:A3/2)'}, {'A1':2,'A2':10,'A3':30})
    _,result=run(candidate,graph)
    assert result.values['S!C2']==5
    assert result.runtime_radj['S!C2']=={'S!A2'}


def test_unselected_inputs_do_not_poison_value_or_dependency_proof(tmp_path):
    graph=graph_for(tmp_path,{'B3':'=AVERAGE(A1:C1/2)'}, {'A1':2,'B1':10,'C1':30})
    ev,result=run(candidate,graph,omit={'S!A1','S!C1'})
    assert result.values['S!B3']==5
    assert result.runtime_radj['S!B3']=={'S!B1'}
    assert set(ev.read_attempts.get('S!B3',())) <= {'S!B1'}


def test_missing_selected_input_remains_unresolved(tmp_path):
    graph=graph_for(tmp_path,{'B3':'=AVERAGE(A1:C1/2)'}, {'A1':2,'B1':10,'C1':30})
    _,result=run(candidate,graph,omit={'S!B1'})
    assert isinstance(result.values['S!B3'],candidate.Unresolved)


def test_no_intersection_is_value_error_not_first_or_average(tmp_path):
    graph=graph_for(tmp_path,{'D3':'=AVERAGE(A1:C1/2)'}, {'A1':2,'B1':10,'C1':30})
    _,result=run(candidate,graph)
    assert result.values['S!D3'].code=='#VALUE!'


def test_unproved_two_dimensional_reference_fails_closed(tmp_path):
    graph=graph_for(tmp_path,{'D3':'=AVERAGE(A1:B2/2)'}, {'A1':2,'B1':10,'A2':30,'B2':40})
    _,result=run(candidate,graph)
    assert isinstance(result.values['S!D3'],candidate.Unresolved)
    assert result.values['S!D3'].reason=='implicit-intersection-shape-unproved'


@pytest.mark.parametrize('formula',[
    ArrayFormula(ref='B3',text='=AVERAGE(A1:C1/2)'),
    ArrayFormula(ref='B3:B4',text='=AVERAGE(A1:C1/2)'),
])
def test_saved_explicit_array_including_modern_single_output_retains_ae(tmp_path,formula):
    graph=graph_for(tmp_path,{'B3':formula},{'A1':2,'B1':10,'C1':30})
    _,result=run(candidate,graph)
    assert result.values['S!B3']==7
    assert result.runtime_radj['S!B3']=={'S!A1','S!B1','S!C1'}
    assert graph.nodes['S!B3'].array_formula


@pytest.mark.parametrize('formula,expected',[
    ('=SUMPRODUCT((A1:C1>5)*A1:C1)',40),
    ('=AVERAGE(SUMPRODUCT((A1:C1>5)*A1:C1)/2)',20),
    ('=AVERAGE(A1:C1)',14),
    ('=SUM(A1:C1)',42),
])
def test_intrinsic_array_functions_and_direct_range_aggregate_unchanged(tmp_path,formula,expected):
    graph=graph_for(tmp_path,{'B3':formula},{'A1':2,'B1':10,'C1':30})
    _,result=run(candidate,graph)
    assert result.values['S!B3']==expected
    assert result.runtime_radj['S!B3']=={'S!A1','S!B1','S!C1'}


def test_explicit_spill_arithmetic_unchanged(tmp_path):
    graph=graph_for(tmp_path,{'A3':ArrayFormula(ref='A3:C3',text='=A1:C1/2')}, {'A1':2,'B1':10,'C1':30})
    _,result=run(candidate,graph)
    assert [result.values[f'S!{c}3'] for c in 'ABC']==[1,5,15]


def test_scalar_error_and_blank_use_existing_arithmetic_rules(tmp_path):
    graph=graph_for(tmp_path,{'B3':'=AVERAGE(A1:C1/0)','B4':'=AVERAGE(A1:C1/2)'},{'A1':2,'C1':30})
    _,result=run(candidate,graph)
    assert result.values['S!B3'].code=='#DIV/0!'
    assert result.values['S!B4']==0


@pytest.mark.parametrize('formula',[
    '=AVERAGE(A1:C1/{1,2,3})',
    '=AVERAGE(A1:C1/TRANSPOSE({1;2;3}))',
])
def test_array_constants_and_other_function_contexts_do_not_opt_in(tmp_path,formula):
    from xl_seg import evaluate
    graph=graph_for(tmp_path,{'B3':formula},{'A1':2,'B1':10,'C1':30})
    ev,result=run(candidate,graph)
    _,existing=run(evaluate,graph)
    assert result.values['S!B3']==existing.values['S!B3']
    assert not any(ev._legacy_average_arithmetic(n) for n in graph.nodes.values() if n.kind=='op')


def test_trusted_roundtrip_gate_rejects_same_text_changed_to_single_cell_array(tmp_path):
    import openpyxl
    from scripts.recalc.excel_runner import _replace_candidate_from_staging, RunnerError
    formula='=AVERAGE(A1:C1/2)'
    original=tmp_path/'original.xlsx'; saved=tmp_path/'saved.xlsx'
    for path,value in [(original,formula),(saved,ArrayFormula(ref='B3',text=formula))]:
        book=openpyxl.Workbook();book.active['B3']=value;book.save(path)
    before=original.read_bytes()
    with pytest.raises(RunnerError,match='changed formula encoding'):
        _replace_candidate_from_staging(saved,original)
    assert original.read_bytes()==before


def test_trusted_roundtrip_gate_accepts_unchanged_saved_normal_formula(tmp_path):
    import openpyxl
    from scripts.recalc.excel_runner import _replace_candidate_from_staging
    formula='=AVERAGE(A1:C1/2)'
    original=tmp_path/'original.xlsx'; saved=tmp_path/'saved.xlsx'
    for path in (original,saved):
        book=openpyxl.Workbook();book.active['B3']=formula;book.save(path)
    _replace_candidate_from_staging(saved,original)
    book=openpyxl.load_workbook(original,data_only=False)
    assert book.active['B3'].value==formula
    assert book.active['B3'].data_type=='f'
