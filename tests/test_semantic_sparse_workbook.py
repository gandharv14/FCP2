"""Retain exact review evidence without allocating the empty sheet rectangle."""
import ast,copy,datetime,json,zipfile
from pathlib import Path
import pytest
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from scripts.batch import harbor_semantic_readiness as S
from scripts.batch.harbor_admission import record


def original_reader():
 source=Path(S.__file__).read_text()
 start=source.index('            # iter_rows()');end=source.index("        return _json",start)
 old="""            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is None:continue
                    value=typed_cell_value(cell)
                    item={'address':cell.coordinate,'value':value,'number_format':cell.number_format,'data_type':cell.data_type}
                    if cell.data_type=='f':item['cached_value']=cached[ws.title][cell.coordinate].value
                    cells.append(item)
            sheets.append({'sheet':ws.title,'state':ws.sheet_state,'dimensions':ws.calculate_dimension(),'merged_ranges':[str(x) for x in ws.merged_cells.ranges],'cells':cells})
"""
 tree=ast.parse(source[:start]+old+source[end:]);fn=next(n for n in tree.body if isinstance(n,ast.FunctionDef)and n.name=='workbook_evidence');ns=dict(S.__dict__);exec(compile(ast.Module(body=[fn],type_ignores=[]),'retained_rectangular_reader','exec'),ns);return ns['workbook_evidence']


def workbook(tmp_path,extent='J20'):
 path=tmp_path/'evidence.xlsx';w=Workbook();s=w.active;s.title='Values'
 # Deliberately insert out of row/column order.
 s['G7']='=B2*2';s['B2']=1.25;s['D4']=True;s['C3']='label';s['E5']=datetime.date(2025,1,2);s['F6']='#N/A';s['I9']=''
 s['G7'].number_format='0.00';s.merge_cells('C12:E12');s['C12']='merged';s[extent].number_format='0%'
 h=w.create_sheet('Hidden');h['D8']='off-origin';h.sheet_state='hidden';w.create_sheet('Empty');w.save(path);w.close()
 # Preserve and exercise an actual cached formula value.
 with zipfile.ZipFile(path)as z:files={n:z.read(n)for n in z.namelist()}
 files['xl/worksheets/sheet1.xml']=files['xl/worksheets/sheet1.xml'].replace(b'<f>B2*2</f><v></v>',b'<f>B2*2</f><v>2.5</v>').replace(b'<f>B2*2</f><v />',b'<f>B2*2</f><v>2.5</v>')
 with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED)as z:
  for name,data in files.items():z.writestr(name,data)
 return record(path)


def test_sparse_reader_has_exact_original_canonical_evidence(tmp_path):
 bound=workbook(tmp_path);before=original_reader()(bound);after=S.workbook_evidence(bound)
 assert S._canonical(after)==S._canonical(before)
 formula=next(c for c in after['sheets'][0]['cells']if c['address']=='G7')
 assert formula['value']=='=B2*2'and formula['cached_value']==2.5 and formula['number_format']=='0.00'
 assert after['sheets'][1]['dimensions']=='A1:D8'and after['sheets'][2]['dimensions']=='A1:A1'


def test_extreme_formatted_extent_does_not_expand_blank_cells(tmp_path,monkeypatch):
 bound=workbook(tmp_path,'XFD1048576');books=[];load=S.load_workbook
 def loaded(*a,**k):
  b=load(*a,**k);books.append(b);return b
 monkeypatch.setattr(S,'load_workbook',loaded)
 def forbidden(*a,**k):raise AssertionError('rectangular blank-cell allocation')
 monkeypatch.setattr(Worksheet,'iter_rows',forbidden)
 result=S.workbook_evidence(bound)
 assert result['sheets'][0]['dimensions']=='A1:XFD1048576'
 assert len(result['sheets'][0]['cells'])==7
 assert all(len(b['Values']._cells)<20 for b in books)
 assert result['sheets'][0]['cells'][0]['address']=='B2'
