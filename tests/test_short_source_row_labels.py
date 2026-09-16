"""Short source category names must survive without promoting unit columns."""
import importlib.util
from pathlib import Path
import openpyxl
import pytest

@pytest.fixture
def D():
    p=Path(__file__).resolve().parents[1]/'.cursor/skills/task-disclosure/scripts/disclose.py'
    s=importlib.util.spec_from_file_location('short_source_labels',p)
    m=importlib.util.module_from_spec(s);s.loader.exec_module(m);return m

def book(D,tmp_path,label,*,numeric=False,stamp=False,isolated=False,formula=False):
    w=openpyxl.Workbook();s=w.active;s.title='Employees'
    for r,name in enumerate(['Corp Training','Acct Mgmt','Surveillance','Reporting','Other','Support Functions'],1):
        s.cell(r,3,r if numeric else '%' if stamp else name)
        s.cell(r,9,10);s.cell(r,14,f'=I{r}/$I$68*$N$68')
    if isolated:
        for r in range(1,7):s.cell(r,3).value=None
    s['C7']='="'+label+'"' if formula else label
    s['I7']=2;s['N7']='=I7/$I$68*$N$68';s['I68']=279;s['N68']=149455
    p=tmp_path/'source.xlsx';w.save(p);return D.Book(p)

@pytest.mark.parametrize('label',['HR','QC','BK'])
def test_source_department_name_survives_complete_formula_rendering(D,tmp_path,label):
    g=book(D,tmp_path,label)
    assert D.is_unit_stamp(label) # Context-free behavior is deliberately retained.
    assert g.row_label('Employees!N7')==label
    profile=D.formula_profile(g,D.make_band(g,'Employees',7,[(14,'Employees!N7')]))
    fields=D.custom_sentence_fields(profile,g)
    assert fields['_coverage_complete'] and fields['label']==D.q(label)
    assert all(x in fields['steps']for x in ['cell Employees!I7','fixed cell Employees!I68','fixed cell Employees!N68'])

@pytest.mark.parametrize('label',['MM','BN','BC','NA','X','M','hr','qc','KW','MW','GW','TW'])
def test_units_and_unqualified_short_tokens_remain_excluded(D,tmp_path,label):
    assert book(D,tmp_path,label).row_label('Employees!N7')==''

@pytest.mark.parametrize('mode',['numeric','stamp','isolated','formula'])
def test_short_token_requires_independent_literal_label_column(D,tmp_path,mode):
    assert book(D,tmp_path,'HR',**{mode:True}).row_label('Employees!N7')==''
