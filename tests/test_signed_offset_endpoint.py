"""0117 reduced to synthetic cells; no private workbook or cached answers."""
import copy
import re
import pytest
from openpyxl import Workbook
from xl_ast_graph import AstGraph, write_csv, parse_formula
from xl_seg import model, project, restriction_cone as rc
from xl_segment import stabilize_runtime_proof
from test_runtime_proof_closure import CALC, _restricted_certificate_source, _passing_certificate_verification


def build(tmp_path, formula, *, target='G1', cells=None):
    tmp_path.mkdir(parents=True, exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.title = 'Sheet'
    cells = cells or {'A1': 2, 'B1': '=D1+E1', 'C1': 5, 'D1': 7, 'E1': 11,
                      'A2': 13, 'B2': 17, 'C2': 19}
    for cell, value in cells.items(): ws[cell] = value
    ws[target] = formula
    file = tmp_path/'book.xlsx'; wb.save(file)
    ast = AstGraph(file, verbose=False).build(); write_csv(tmp_path, ast.nodes, ast.edges)
    graph = model.load(tmp_path, 'book')
    graph.nodes['Sheet!'+target].value = '987654321'
    seeds = {'Sheet!'+cell for cell,value in cells.items() if type(value) in {int,float}}
    result, inputs, _, proof = stabilize_runtime_proof(graph, project.build(graph), seeds,
                                                      {'Sheet!'+target}, calculation=CALC)
    return graph, result, inputs, proof


def certificate(tmp_path, monkeypatch, graph, proof, *, target='G1'):
    formula = graph.nodes['Sheet!'+target].formula
    events = [{'event':'volatile_function','function':'OFFSET','location':'Sheet!'+target,
               'scope':'worksheet','token_offset':m.start()}
              for m in re.finditer('OFFSET',formula[1:])]
    source = _restricted_certificate_source(tmp_path, monkeypatch, events)
    verification = _passing_certificate_verification(proof, ['Sheet!'+target])
    return rc.build_certificate(source_generation_dir=source,graph=graph,proof=proof,
        verification=verification,ordered_outputs=['Sheet!'+target],
        segmentation_fingerprints=verification['fingerprints'])


@pytest.mark.parametrize('formula,direct',[
    ('=SUM(A1:OFFSET(C2,-1,0))','=SUM(A1:C1)'),
    ('=SUM(A1:OFFSET(B1,0,+1))','=SUM(A1:C1)'),
    ('=SUM(A1:OFFSET(C1,+1,-1))','=SUM(A1:B2)'),
    ('=SUM($A$1:OFFSET($C$2,-1.0,+0E0))','=SUM(A1:C1)'),
    ('=SUM(C2:OFFSET(B1,+1,-1))','=SUM(A2:C2)'),
])
def test_signed_literal_endpoints_equal_direct_ranges(tmp_path,monkeypatch,formula,direct):
    graph,result,inputs,proof=build(tmp_path/'offset',formula)
    _,direct_result,_,_=build(tmp_path/'direct',direct)
    assert not result.unresolved
    assert result.values['Sheet!G1'] == direct_result.values['Sheet!G1']
    assert 'Sheet!B1' not in inputs  # formula reconstructed independently
    c=certificate(tmp_path/'offset',monkeypatch,graph,proof)
    event=c['events'][0]
    assert event['enclosing_static_ranges']
    for span in event['enclosing_static_ranges']:
        assert set(span['static_targets']) == set(span['runtime_targets'])
        assert span['endpoint'] in event['static_targets']


def test_observed_0117_address_anchor_is_not_a_value_dependency(tmp_path,monkeypatch):
    cells={f'F{row}': row-60 for row in range(61,69)}
    cells['F60']=0.1
    graph,result,inputs,proof=build(tmp_path,'=(SUM(F61:OFFSET(F69,-1,0)))*(1+F$60)',
                                  target='F69',cells=cells)
    assert result.values['Sheet!F69'] == pytest.approx(39.6)
    assert 'Sheet!F69' not in inputs
    assert 'Sheet!F69' not in proof['runtime_radj']['Sheet!F69']
    c=certificate(tmp_path,monkeypatch,graph,proof,target='F69')
    span=c['events'][0]['enclosing_static_ranges'][0]
    assert span['base']=='Sheet!F69' and span['endpoint']=='Sheet!F68'
    assert span['displacement']==[-1,0]
    assert span['static_targets']==[f'Sheet!F{row}' for row in range(61,69)]


@pytest.mark.parametrize('formula',[
    '=SUM(A1:OFFSET(C2,D1,0))', '=SUM(A1:OFFSET(C2,UNKNOWN(),0))',
    '=SUM(A1:OFFSET(C2,1-2,0))', '=SUM(A1:OFFSET(C2,-0.5,0))',
    '=SUM(A1:OFFSET(C2,-TRUE,0))', '=SUM(A1:OFFSET(C2,"-1",0))',
    '=SUM(A1:OFFSET(C2,-1,0,1,1))', '=SUM(A1:OFFSET(C2:C3,-1,0))',
    "=SUM(A1:OFFSET('Other'!C2,-1,0))", '=SUM(A1:OFFSET([1]Other!C2,-1,0))',
    '=SUM(A1:OFFSET(C2,-2,0))', '=SUM(A1:OFFSET(C2,0,-3))',
    '=SUM(A1:OFFSET(C2,1048576,0))', '=SUM(A1:OFFSET(XFD2,0,+1))',
    '=SUM(A1:OFFSET(XFE2,-1,0))', '=SUM(A1:OFFSET(C2,--1,0))',
])
def test_unsupported_endpoint_remains_closed(formula):
    child=parse_formula(formula).args[0]
    assert child.kind=='func' and child.name.endswith(':OFFSET')


@pytest.mark.parametrize('tamper',['offset_target','span_target','literal','signed_literal','missing_middle',
                                  'dropped_unary','unknown_unary','base_sheet','base_coordinate'])
def test_independent_certificate_rejects_forged_endpoints(tmp_path,monkeypatch,tamper):
    graph,result,inputs,proof=build(tmp_path,'=SUM(A1:OFFSET(C2,-1,0))')
    offset=next(n for n in graph.nodes.values() if n.op=='OFFSET')
    colon=next(n for n in graph.nodes.values() if n.op==':')
    unary=next(n for n in graph.nodes.values() if n.op=='u-')
    if tamper=='offset_target': proof['resolved_operation_targets'][offset.id]=['Sheet!C2']
    elif tamper=='span_target': proof['resolved_operation_targets'][colon.id]=['Sheet!A1','Sheet!C1']
    elif tamper=='missing_middle': proof['runtime_radj']['Sheet!G1'].remove('Sheet!B1')
    elif tamper=='dropped_unary': graph.in_edges[unary.id]=[]
    elif tamper=='unknown_unary': unary.op='UNPROVEN'
    elif tamper=='signed_literal': unary.op='u+'
    elif tamper=='literal':
        graph.nodes[graph.in_edges[unary.id][0].source].value='2'
    elif tamper=='base_sheet': graph.nodes['Sheet!C2'].sheet='Other'
    elif tamper=='base_coordinate': graph.nodes['Sheet!C2'].row=3
    with pytest.raises(rc.RestrictionConeError): certificate(tmp_path,monkeypatch,graph,proof)
