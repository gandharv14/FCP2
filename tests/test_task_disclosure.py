from __future__ import annotations

import importlib.util
import json
from pathlib import Path
from types import SimpleNamespace


ROOT = Path(__file__).resolve().parents[1]
DISCLOSE_PATH = (
    ROOT / ".cursor" / "skills" / "task-disclosure" / "scripts" / "disclose.py"
)


def _load_disclose():
    spec = importlib.util.spec_from_file_location("task_disclose", DISCLOSE_PATH)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


def test_expand_method_band_passes_the_expanded_run_to_make_band(
    monkeypatch,
) -> None:
    disclose = _load_disclose()
    gold = SimpleNamespace(formula={
        "Sheet!A1": "=A2",
        "Sheet!B1": "=B2",
        "Sheet!C1": "=C2",
    })
    delivered = SimpleNamespace(has=lambda _key: False)
    captured = {}

    monkeypatch.setattr(disclose, "r1c1ish", lambda *_args: "same-pattern")

    def capture(_gold, sheet, row, run, full=None):
        captured.update(sheet=sheet, row=row, run=run, full=full)
        return {"expanded": True}

    monkeypatch.setattr(disclose, "make_band", capture)

    result = disclose.expand_method_band(
        gold,
        delivered,
        {
            "sheet": "Sheet",
            "row": 1,
            "cell_keys": ["Sheet!B1"],
            "pattern": "same-pattern",
        },
    )

    assert result == {"expanded": True}
    assert captured == {
        "sheet": "Sheet",
        "row": 1,
        "run": [
            (1, "Sheet!A1"),
            (2, "Sheet!B1"),
            (3, "Sheet!C1"),
        ],
        "full": None,
    }


def test_disclosure_quotes_preserve_visible_label_text() -> None:
    disclose = _load_disclose()
    labels = [
        "Revenue",
        "Shareholders' Equity",
        'Ad Cost Per Thousand Impressions ("CPM")',
        'Customer "A"\'s Revenue',
    ]

    for label in labels:
        assert disclose._unquote_label(disclose.q(label)) == label


def test_vertical_singletons_with_different_row_labels_stay_separate(
    monkeypatch,
) -> None:
    disclose = _load_disclose()
    labels = {
        "Analysis!D16": "Term loan",
        "Analysis!D17": "Senior unsecured note",
    }
    gold = SimpleNamespace(row_label=labels.__getitem__)
    bands = [
        {
            "band": "Analysis!D16",
            "cell_keys": ["Analysis!D16"],
            "stated_cell_keys": ["Analysis!D16"],
        },
        {
            "band": "Analysis!D17",
            "cell_keys": ["Analysis!D17"],
            "stated_cell_keys": ["Analysis!D17"],
        },
    ]
    monkeypatch.setattr(disclose, "copy_equivalent", lambda *_args: True)
    monkeypatch.setattr(
        disclose,
        "full_copied_scope_vertical",
        lambda _gold, keys, _targets: keys,
    )

    merged = disclose.merge_vertical_singletons(gold, bands)

    assert merged == bands


def test_method_band_splits_cells_with_different_row_labels() -> None:
    disclose = _load_disclose()
    labels = {
        "Analysis!D16": "Term loan",
        "Analysis!D17": "Senior unsecured note",
    }
    gold = SimpleNamespace(
        formula={
            "Analysis!D16": "=E16*$D$24*(H16/100)",
            "Analysis!D17": "=E17*$D$24*(H17/100)",
        },
        value={"Analysis!D16": 1, "Analysis!D17": 2},
        row_label=labels.__getitem__,
    )
    band = {
        "band": "Analysis!D16:D17",
        "sheet": "Analysis",
        "row": 16,
        "cell_keys": ["Analysis!D16"],
        "stated_cell_keys": ["Analysis!D16", "Analysis!D17"],
        "label": "Term loan",
    }

    split = disclose.method_subbands(gold, band)

    assert [item["band"] for item in split] == [
        "Analysis!D16",
        "Analysis!D17",
    ]
    assert [item["label"] for item in split] == [
        "Term loan",
        "Senior unsecured note",
    ]


def test_faithcheck_accepts_complete_heterogeneous_scope_partition() -> None:
    disclose = _load_disclose()
    labels = {
        "Analysis!D16": "Term loan",
        "Analysis!D17": "Senior unsecured note",
    }
    gold = SimpleNamespace(row_label=labels.__getitem__)
    records = [
        {
            "entry": "method_debt_movement",
            "cells": ["Analysis!D16"],
            "cell_keys": ["Analysis!D16"],
            "fields": {"label": '"Term loan"'},
        },
        {
            "entry": "method_debt_movement",
            "cells": ["Analysis!D17"],
            "cell_keys": ["Analysis!D17"],
            "fields": {"label": '"Senior unsecured note"'},
        },
    ]

    assert disclose._heterogeneous_scope_partitioned(
        gold,
        ["Analysis!D16", "Analysis!D17"],
        records,
        "method_debt_movement",
    )


def test_compensation_type_markers_are_not_row_names() -> None:
    disclose = _load_disclose()
    book = object.__new__(disclose.Book)

    assert not book.is_row_name("Cost Assumptions Matrix", 5, "Salary Only")
    assert not book.is_row_name("Cost Assumptions Matrix", 5, "Hourly Only")
    assert book.is_row_name("Cost Assumptions Matrix", 2, "District Managers")


def test_ungraded_row_label_year_is_licensed_from_target_collision(
    tmp_path: Path,
) -> None:
    disclose = _load_disclose()
    task = tmp_path / "0167-outputs"
    (task / "tests").mkdir(parents=True)
    (task / "tests" / "answer_key.json").write_text(
        json.dumps({"targets": {"Overview!L12": 2013}}),
        encoding="utf-8",
    )
    record = {
        "entry": "source_selection",
        "family": "source_selection",
        "value": "source",
        "disposition": "disclosed",
        "source": "convention_detector",
        "cell_keys": ["Overview!F55"],
        "fields": {
            "label": '"Adjusted EBITDA (FY 2013)"',
            "ingredient": 'cell LBO!X20 on the row labelled "Adjusted EBITDA"',
        },
        "leak_flag": False,
    }

    import openpyxl
    workbook = openpyxl.Workbook(); workbook.active.title = 'Overview'
    workbook.active['D55'] = 'Adjusted EBITDA (FY 2013)'
    workbook.active['F55'] = '=1+2'
    path = tmp_path / 'gold.xlsx'; workbook.save(path)
    gold = disclose.Book(path)
    allowed = disclose.trusted_row_label_literal_counts([record], task, gold)

    assert allowed[2013.0] == 1
    assert disclose.audit_text(
        disclose.render_sentence(record),
        task,
        records=[record],
        gold=gold,
    ) == []


def test_source_label_collision_needs_exact_occurrence_provenance(tmp_path):
    import openpyxl
    disclose = _load_disclose()
    task = tmp_path / '0001-outputs'; (task / 'tests').mkdir(parents=True)
    (task / 'tests/answer_key.json').write_text(json.dumps({'targets': {'Model!Z1': 6}}))
    wb = openpyxl.Workbook(); wb.active.title = 'Model'; ws = wb.create_sheet('Inputs')
    wb['Model']['B10'] = '=Inputs!B20'; wb['Model']['K10'] = '=Inputs!K20'
    ws['B20'] = 'Care home 6'; ws['K20'] = '=1+1'
    path = tmp_path / 'source.xlsx'; wb.save(path)
    gold = disclose.Book(path)
    # Establish the linked literal label even without an Excel display cache.
    gold.value['Model!B10'] = 'Care home 6'
    record = {'entry': 'source_selection', 'value': 'source', 'disposition': 'disclosed',
        'cell_keys': ['Model!K10'], 'source_cell': 'Inputs!K20',
        'fields': {'label': '"Care home 6"',
                   'ingredient': 'cell Inputs!K20 on the row labelled "Care home 6"'}}
    sentence = disclose.render_sentence(record)
    assert disclose.trusted_row_label_literal_counts([record], task, gold)[6] == 2
    assert disclose.audit_text(sentence, task, [record], gold=gold) == []
    assert disclose.audit_text(sentence + ' The answer is 6.', task, [record], gold=gold)
    assert disclose.audit_text(sentence, task, [record])
    changed = {**record, 'source_cell': 'Inputs!K21'}
    assert disclose.audit_text(sentence, task, [changed], gold=gold)
    gold.formula['Model!K10'] = '=-Inputs!K20'
    assert disclose.audit_text(sentence, task, [record], gold=gold)


def test_verify_uses_only_hash_bound_custom_root_gold_for_leak_audit(tmp_path, monkeypatch):
    import hashlib
    import openpyxl
    disclose = _load_disclose()
    task = tmp_path / 'custom' / '0001-outputs'
    (task / 'tests').mkdir(parents=True)
    path = tmp_path / 'custom-golden.xlsx'
    workbook = openpyxl.Workbook()
    workbook.active.title = 'Model'
    workbook.active['A1'] = 3
    workbook.save(path)
    environment = tmp_path / 'delivered.xlsx'
    workbook.save(environment)
    payload = {'golden': str(path), 'golden_sha256': hashlib.sha256(path.read_bytes()).hexdigest()}
    (task / 'tests/disclosure.json').write_text(json.dumps(payload))
    (task / 'tests/answer_key.json').write_text(json.dumps({'targets': {'Model!A1': 3}}))
    (task / 'instruction.md').write_text(disclose.SECTION_START + '\n')
    monkeypatch.setattr(disclose, 'find_environment', lambda *_: environment)
    monkeypatch.setattr(disclose, 'find_golden', lambda *_: (_ for _ in ()).throw(
        AssertionError('legacy lookup must not be used')))
    observed = []
    monkeypatch.setattr(disclose, 'audit_text', lambda *a, **kw: observed.append(kw['gold']) or [])
    result = disclose.verify_task(task)
    # This fixture intentionally has no segmentation proof: that separate gate
    # still fails, while leak auditing receives the independently bound book.
    assert not result['passed']
    assert observed[-1].value['Model!A1'] == 3
    assert not any('golden binding' in f for f in result['faults'])
    path.write_bytes(b'altered evidence')
    result = disclose.verify_task(task)
    assert observed[-1] is None
    assert 'disclosure golden binding is missing or changed' in result['faults']
