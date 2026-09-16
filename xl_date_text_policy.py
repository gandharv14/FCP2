"""Explicit source-bound date interpretation for primitive text only.

This is interpretation evidence, not permission to run, retry, recalculate or
publish a task. Plain strings retain existing behavior without this selection.
Microsoft DATEVALUE documents that date order depends on the host locale; a
workbook's formatting alone never selects our interpretation implicitly.
"""
from pathlib import Path
from datetime import datetime
import hashlib,json,re

SCHEMA='source-primitive-date-text/v1'
LITERAL_SCHEMA='source-explicit-date-text/v2'
AUTHORIZATION_SHA256='acffe75c8ed401f66484a5bb6c39a15a0a3e209fc976967da0bddd6125126fd0'


def require(ok,message):
    if not ok:raise ValueError(message)


def binding(path):
    p=Path(path);require(p.is_absolute() and p.is_file() and not any(x.is_symlink()for x in(p,*p.parents)),'regular date interpretation evidence required')
    return dict(path=str(p),sha256=hashlib.sha256(p.read_bytes()).hexdigest())


class DateText(str):
    """Preserve text identity; numeric coercion uses the declared date order."""
    def __new__(cls,text,order):
        require(type(text)is str and order in ('MDY','DMY'),'explicit date order required')
        match=re.fullmatch(r'([0-9]{1,2})/([0-9]{1,2})/([0-9]{4})',text)
        require(match is not None,'only complete slash dates with four-digit years are supported')
        a,b,year=map(int,match.groups());month,day=(a,b)if order=='MDY'else(b,a)
        require(1900<=year<=9999,'unsupported date year')
        # Gregorian validity: no two-digit-year pivot, overflow, time, timezone,
        # or fictitious 1900-02-29 interpretation is guessed.
        stamp=datetime(year,month,day)
        from xl_seg.evaluate import _datetime_to_serial
        instance=str.__new__(cls,text);instance._serial=_datetime_to_serial(stamp);return instance
    def __float__(self):return float(self._serial)


class DateSelection(dict):
    """Primitive text plus separately verified formula-literal selections."""
    def __init__(self, primitives, literals):
        super().__init__(primitives)
        self.formula_literals=literals


def graph_literals(selection,graph):
    """Bind selected source literals to the actual formula AST, never caches."""
    result={}
    for row in getattr(selection,'formula_literals',[]):
        owner=row['sheet']+'!'+row['coordinate'];formula=graph.nodes.get(owner)
        require(formula is not None and formula.kind=='formula' and formula.formula==row['formula'],'selected date formula differs from source AST')
        nodes=[n for n in graph.nodes.values()if n.owner==owner and n.kind=='const' and n.op=='text' and n.value==row['text']]
        require(bool(nodes),'selected date literal absent from source AST')
        for n in nodes:
            require(n.id not in result,'duplicate date literal AST selection')
            result[n.id]=DateText(row['text'],row['date_order'])
    return result


def load(row,source):
    require(type(row)is dict and set(row)=={'path','sha256'} and binding(row['path'])==row,'date policy changed')
    p=json.loads(Path(row['path']).read_bytes())
    extended=p.get('schema_version')==LITERAL_SCHEMA
    require(set(p)==({'schema_version','source','authorization','date_order','cells'}|({'formula_literals'}if extended else set())) and p['schema_version']in(SCHEMA,LITERAL_SCHEMA),'closed primitive date policy required')
    require(p['source']==binding(p['source']['path']) and p['source']['sha256']==binding(source)['sha256'],'date policy source differs')
    require(p['authorization']==binding(p['authorization']['path']) and p['authorization']['sha256']==AUTHORIZATION_SHA256,'date interpretation is not authorized')
    literals=p['formula_literals']if extended else []
    require(type(literals)is list and p['date_order']in('MDY','DMY') and type(p['cells'])is list and 0<len(p['cells'])+len(literals)<=100,'bounded explicit date interpretation required')
    selected={}
    for r in p['cells']:
        require(set(r)=={'sheet','coordinate','text','number_format'},'closed primitive cell selection required')
        key=r['sheet']+'!'+r['coordinate'];require(key not in selected,'duplicate primitive date cell');selected[key]=r
        DateText(r['text'],p['date_order'])
    selected_literals={}
    for r in literals:
        require(type(r)is dict and set(r)=={'sheet','coordinate','formula','text'},'closed formula literal selection required')
        key=(r['sheet']+'!'+r['coordinate'],r['text'])
        require(key not in selected_literals and type(r['formula'])is str and r['formula'].startswith('='),'duplicate or invalid formula literal selection')
        DateText(r['text'],p['date_order']);selected_literals[key]=r
    import openpyxl
    book=openpyxl.load_workbook(source,read_only=True,data_only=False)
    try:
        require(book.epoch==datetime(1899,12,30),'1904 date system requires separate interpretation')
        observed={};observed_literals=[]
        for sheet in {r['sheet']for r in p['cells']+literals}:
            require(sheet in book.sheetnames,'selected date sheet absent')
            for cells in book[sheet].iter_rows():
                for cell in cells:
                    key=sheet+'!'+getattr(cell,'coordinate','')
                    if key in selected:
                        r=selected[key];require(cell.data_type=='s' and type(cell.value)is str and cell.value==r['text'] and cell.number_format==r['number_format'],'source primitive date text or format changed')
                        observed[key]=DateText(cell.value,p['date_order'])
                    for (owner,text),r in selected_literals.items():
                        if owner!=key:continue
                        require(cell.data_type=='f' and type(cell.value)is str and cell.value==r['formula'],'source date formula changed')
                        from openpyxl.formula import Tokenizer
                        tokens=Tokenizer(cell.value).items
                        require(any(t.type=='OPERAND' and t.subtype=='TEXT' and t.value[1:-1].replace('""','"')==text for t in tokens),'selected date is not a literal in the source formula')
                        observed_literals.append({**r,'date_order':p['date_order']})
        require(set(observed)==set(selected),'source primitive date cell absent')
        require(len(observed_literals)==len(literals),'source date formula absent')
    finally:book.close()
    require(binding(source)['sha256']==p['source']['sha256'],'source changed during date interpretation')
    return DateSelection(observed,observed_literals)
