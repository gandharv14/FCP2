"""Proof-bound direct external literal assumptions; standalone, not source approval.

The independently approved evidence SHA-256 is a required caller trust anchor.
This tool never derives it from workbook caches or creates identity approvals.
Candidates still require source health, recalculation where required, independent
reconstruction, disclosure/leak review, and all publication gates.
"""
from __future__ import annotations

import argparse
import copy
import fcntl
import hashlib
import io
import json
import os
import posixpath
import re
import shutil
import tempfile
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.styles.numbers import is_date_format
from xl_source_health import validate_ooxml_package

POLICY = 'external-direct-literal-input/v1'
EVIDENCE_SCHEMA = 'external-literal-identity-evidence/v1'
PLAN_SCHEMA = 'external-literal-input-plan/v1'
MANIFEST_SCHEMA = 'external-literal-input-manifest/v1'
MAIN = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL = 'http://schemas.openxmlformats.org/package/2006/relationships'
DOCREL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
NS = {'s': MAIN, 'r': REL}
MAX_BYTES = 128 * 1024 * 1024
DIRECT = re.compile(r"^(?:'\[(?P<quoted_index>[1-9][0-9]*)\](?P<quoted_sheet>(?:[^']|'')+)'|\[(?P<index>[1-9][0-9]*)\](?P<sheet>[A-Za-z0-9_][A-Za-z0-9_.]*))!\$?(?P<col>[A-Z]{1,3})\$?(?P<row>[1-9][0-9]*)$")
EXTERNAL = re.compile(r'\[[^\]]+\][^!\r\n]*!')
CELL_SPAN = re.compile(rb'<(?:[A-Za-z_][\w.-]*:)?c\b[^>]*>.*?</(?:[A-Za-z_][\w.-]*:)?c>', re.S)
ALLOWED = re.compile(r'^(?:\[Content_Types\]\.xml|_rels/\.rels|docProps/(?:app|core)\.xml|xl/(?:workbook\.xml|styles\.xml|sharedStrings\.xml|_rels/workbook\.xml\.rels|theme/theme[0-9]+\.xml|worksheets/sheet[0-9]+\.xml|externalLinks/externalLink[0-9]+\.xml|externalLinks/_rels/externalLink[0-9]+\.xml\.rels))$')


class ExternalLiteralError(ValueError):
    pass


def canonical(value):
    return (json.dumps(value, sort_keys=True, separators=(',', ':'), ensure_ascii=False) + '\n').encode()


def sha(data):
    return hashlib.sha256(data).hexdigest()


def _object(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise ExternalLiteralError('duplicate evidence key')
        result[key] = value
    return result


def _json(path, max_bytes=1024 * 1024):
    raw = Path(path).read_bytes()
    if len(raw) > max_bytes:
        raise ExternalLiteralError('evidence exceeds size bound')
    try:
        return raw, json.loads(raw, object_pairs_hook=_object)
    except (UnicodeError, json.JSONDecodeError) as exc:
        raise ExternalLiteralError('invalid evidence JSON') from exc


def _keys(value, expected, description):
    if not isinstance(value, dict) or set(value) != set(expected):
        raise ExternalLiteralError(f'{description} has missing or unknown fields')


def _package(path):
    path = Path(path)
    if path.suffix.lower() != '.xlsx' or path.is_symlink() or not path.is_file():
        raise ExternalLiteralError('only regular original .xlsx files are supported')
    validate_ooxml_package(path)
    raw = path.read_bytes()
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        infos = archive.infolist()
        if sum(item.file_size for item in infos) > MAX_BYTES:
            raise ExternalLiteralError('package exceeds bounded transformer size')
        parts = {item.filename: archive.read(item) for item in infos}
    for name, data in parts.items():
        if not ALLOWED.fullmatch(name):
            raise ExternalLiteralError('unsupported package consumer: ' + name)
        if b'<!DOCTYPE' in data.upper() or b'<!ENTITY' in data.upper():
            raise ExternalLiteralError('XML entities are unsupported')
        try:
            ET.fromstring(data)
        except ET.ParseError as exc:
            raise ExternalLiteralError('invalid XML package member') from exc
    required = {'[Content_Types].xml', '_rels/.rels', 'xl/workbook.xml', 'xl/_rels/workbook.xml.rels'}
    if not required <= set(parts):
        raise ExternalLiteralError('required workbook package identity is missing')
    _validate_consumers(parts)
    return raw, infos, parts


def _validate_consumers(parts):
    """Conservative closed route: reject consumers this producer cannot prove."""
    forbidden = {'definedName', 'dataValidation', 'conditionalFormatting', 'extLst',
                 'oleObject', 'control', 'drawing', 'tablePart', 'webPublishItem',
                 'pivotTable', 'legacyDrawing', 'smartTag', 'customSheetView', 'AlternateContent'}
    for name, data in parts.items():
        root = ET.fromstring(data)
        if any(el.tag.rsplit('}', 1)[-1] in forbidden for el in root.iter()):
            raise ExternalLiteralError('unsupported indirect or opaque consumer: ' + name)
        if any(key.rsplit('}', 1)[-1] == 'Ignorable' for el in root.iter() for key in el.attrib):
            raise ExternalLiteralError('namespace-dependent compatibility metadata is unsupported')
        if name.endswith('.rels'):
            seen = set()
            for relation in root:
                rid = relation.get('Id')
                if not rid or rid in seen:
                    raise ExternalLiteralError('ambiguous relationship identity')
                seen.add(rid)
                typ = relation.get('Type', '').rsplit('/', 1)[-1]
                if typ not in {'officeDocument', 'core-properties', 'extended-properties',
                               'worksheet', 'styles', 'theme', 'sharedStrings',
                               'externalLink', 'externalLinkPath'}:
                    raise ExternalLiteralError('unsupported relationship consumer')
                if (typ == 'externalLink' and name != 'xl/_rels/workbook.xml.rels'
                        or typ == 'externalLinkPath' and not name.startswith('xl/externalLinks/_rels/')):
                    raise ExternalLiteralError('unaccounted external relationship consumer')
                if relation.get('TargetMode') == 'External' and typ != 'externalLinkPath':
                    raise ExternalLiteralError('unsupported external relationship')
        for element in root.iter('{' + MAIN + '}f'):
            if element.attrib:
                raise ExternalLiteralError('shared, array or data-table formula encoding is unsupported')


def _resolve_part(base, target):
    if not isinstance(target, str) or '\\' in target or ':' in target:
        raise ExternalLiteralError('invalid internal relationship target')
    resolved = posixpath.normpath(target.lstrip('/') if target.startswith('/') else posixpath.join(base, target))
    if resolved.startswith('../') or not resolved.startswith('xl/'):
        raise ExternalLiteralError('internal relationship escapes workbook')
    return resolved


def _layout(parts):
    workbook = ET.fromstring(parts['xl/workbook.xml'])
    relations = ET.fromstring(parts['xl/_rels/workbook.xml.rels'])
    byid = {el.get('Id'): el for el in relations}
    sheets = {}
    for sheet in workbook.findall('s:sheets/s:sheet', NS):
        name = sheet.get('name')
        rel = byid.get(sheet.get('{' + DOCREL + '}id'))
        if name in sheets or rel is None or not rel.get('Type', '').endswith('/worksheet'):
            raise ExternalLiteralError('ambiguous worksheet identity')
        part = _resolve_part('xl', rel.get('Target'))
        if part not in parts:
            raise ExternalLiteralError('worksheet package is missing')
        if part in sheets.values():
            raise ExternalLiteralError('multiple sheets share one package part')
        sheets[name] = part
    references = workbook.findall('s:externalReferences/s:externalReference', NS)
    links = []
    for index, ref in enumerate(references, 1):
        rid = ref.get('{' + DOCREL + '}id'); relation = byid.get(rid)
        if relation is None or not relation.get('Type', '').endswith('/externalLink'):
            raise ExternalLiteralError('external reference has no exact relationship')
        part = _resolve_part('xl', relation.get('Target'))
        relpart = posixpath.join(posixpath.dirname(part), '_rels', posixpath.basename(part) + '.rels')
        if part not in parts or relpart not in parts:
            raise ExternalLiteralError('external dependency relationship is missing')
        root = ET.fromstring(parts[part]); book = root.find('s:externalBook', NS)
        if book is None or len(root) != 1:
            raise ExternalLiteralError('DDE, OLE or unsupported external book')
        source_rel_id = book.get('{' + DOCREL + '}id')
        rels = list(ET.fromstring(parts[relpart]))
        if (len(rels) != 1 or rels[0].get('Id') != source_rel_id
                or not rels[0].get('Type', '').endswith('/externalLinkPath')
                or rels[0].get('TargetMode') != 'External'):
            raise ExternalLiteralError('ambiguous external dependency relationship')
        # Cached sheetData may exist, but is never used to identify or read a value.
        if any(child.tag.rsplit('}', 1)[-1] not in {'sheetNames', 'sheetDataSet'} for child in book):
            raise ExternalLiteralError('external names or unsupported external consumers')
        names = [el.get('val') for el in book.findall('s:sheetNames/s:sheetName', NS)]
        if len(names) != len(set(names)):
            raise ExternalLiteralError('ambiguous external sheet identity')
        links.append({'link_index': index, 'workbook_relationship_id': rid,
                      'relationship_part': relpart, 'relationship_id': source_rel_id,
                      'relationship_part_sha256': sha(parts[relpart]),
                      'target_identity_sha256': sha(rels[0].get('Target', '').encode()),
                      'part': part, 'sheet_names': names})
    referenced = {item['part'] for item in links} | {item['relationship_part'] for item in links}
    actual = {name for name in parts if name.startswith('xl/externalLinks/')}
    relids = {el.get('Id') for el in relations if el.get('Type', '').endswith('/externalLink')}
    overrides = [el.get('PartName', '').lstrip('/') for el in ET.fromstring(parts['[Content_Types].xml'])
                 if el.get('ContentType', '').endswith('externalLink+xml')]
    if len(overrides) != len(set(overrides)) or set(overrides) != {item['part'] for item in links}:
        raise ExternalLiteralError('unaccounted external content type')
    if referenced != actual or relids != {item['workbook_relationship_id'] for item in links}:
        raise ExternalLiteralError('unaccounted external package or relationship')
    return sheets, links


def _coordinate(text):
    if not isinstance(text, str) or not re.fullmatch(r'[A-Z]{1,3}[1-9][0-9]*', text):
        raise ExternalLiteralError('single canonical A1 coordinate required')
    row, col = coordinate_to_tuple(text)
    if row > 1048576 or col > 16384:
        raise ExternalLiteralError('cell coordinate outside Excel bounds')
    return text


def _cells(parts, sheetpart):
    result = {}
    for cell in ET.fromstring(parts[sheetpart]).iter('{' + MAIN + '}c'):
        coordinate = _coordinate(cell.get('r'))
        if coordinate in result:
            raise ExternalLiteralError('duplicate cell identity')
        result[coordinate] = cell
    return result


def _literal(parts, cell, source_cell, destination_cell, dependency_epoch, source_epoch):
    if cell.find('s:f', NS) is not None:
        raise ExternalLiteralError('dependency is calculated, never a literal assumption')
    if set(cell.attrib) - {'r', 's', 't'}:
        raise ExternalLiteralError('dependency cell metadata is not supported')
    typ = cell.get('t', 'n')
    if typ in {'s', 'inlineStr'}:
        expected_child = 'v' if typ == 's' else 'is'
        if [child.tag.rsplit('}', 1)[-1] for child in cell] != [expected_child]:
            raise ExternalLiteralError('ambiguous literal text payload')
        if typ == 's':
            raw = cell.findtext('s:v', namespaces=NS)
            try:
                strings = list(ET.fromstring(parts['xl/sharedStrings.xml']))
                if raw is None or not re.fullmatch(r'[0-9]+', raw): raise ValueError()
                value_root = strings[int(raw)]
            except (KeyError, ValueError, IndexError) as exc:
                raise ExternalLiteralError('shared literal string is missing') from exc
        else:
            value_root = cell.find('s:is', NS)
        if value_root is None or len(value_root) != 1 or value_root[0].tag != '{' + MAIN + '}t':
            raise ExternalLiteralError('rich or ambiguous text is unsupported')
        value = value_root[0].text or ''
        payload = ET.Element('is'); text = ET.SubElement(payload, 't')
        text.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); text.text = value
        return 'text', 'inlineStr', [payload], value
    raw = cell.findtext('s:v', namespaces=NS)
    if raw is None or [child.tag.rsplit('}', 1)[-1] for child in cell] != ['v']:
        raise ExternalLiteralError('blank or unsupported literal cell')
    if typ == 'b':
        if raw not in {'0', '1'}: raise ExternalLiteralError('invalid boolean literal')
        value, kind = raw == '1', 'boolean'
    elif typ == 'n':
        try:
            if not Decimal(raw).is_finite(): raise InvalidOperation()
        except InvalidOperation as exc:
            raise ExternalLiteralError('invalid numeric literal') from exc
        value, kind = raw, 'number'
        if source_cell.is_date:
            if (dependency_epoch != source_epoch or not is_date_format(destination_cell.number_format)
                    or destination_cell.number_format != source_cell.number_format):
                raise ExternalLiteralError('date system or destination date format differs')
            kind = 'date_serial'
    else:
        raise ExternalLiteralError('only numeric, text, boolean and styled date literals are supported')
    element = ET.Element('v'); element.text = raw
    return kind, typ, [element], value


def _read_evidence(path, expected_sha256, source_sha256):
    raw, evidence = _json(path)
    if not re.fullmatch('[0-9a-f]{64}', expected_sha256 or '') or sha(raw) != expected_sha256:
        raise ExternalLiteralError('independently approved evidence hash differs')
    _keys(evidence, {'schema_version', 'source_sha256', 'authority', 'dependencies'}, 'identity evidence')
    if evidence['schema_version'] != EVIDENCE_SCHEMA or evidence['source_sha256'] != source_sha256:
        raise ExternalLiteralError('identity evidence is not bound to original source')
    if not isinstance(evidence['authority'], str) or not evidence['authority'].strip() or len(evidence['authority']) > 4096:
        raise ExternalLiteralError('independent identity evidence authority is missing')
    if not isinstance(evidence['dependencies'], list) or not evidence['dependencies']:
        raise ExternalLiteralError('complete external dependency evidence is required')
    return evidence


def _build(source, evidence_path, expected_evidence_sha256):
    raw, infos, parts = _package(source)
    source_sha = sha(raw); evidence = _read_evidence(evidence_path, expected_evidence_sha256, source_sha)
    sheets, links = _layout(parts)
    if not links: raise ExternalLiteralError('source has no external dependency links')
    fields = {'link_index', 'workbook_relationship_id', 'relationship_part', 'relationship_id',
              'relationship_part_sha256', 'target_identity_sha256', 'path', 'sha256', 'literal_cells'}
    if len(evidence['dependencies']) > 32:
        raise ExternalLiteralError('dependency count exceeds narrow transformer bound')
    supplied = {}
    for row in evidence['dependencies']:
        _keys(row, fields, 'dependency identity')
        index = row['link_index']
        if type(index) is not int or index in supplied: raise ExternalLiteralError('ambiguous dependency identity')
        supplied[index] = row
    if set(supplied) != {link['link_index'] for link in links}:
        raise ExternalLiteralError('missing or extra external dependency identity')
    wb = load_workbook(io.BytesIO(raw), data_only=False, keep_links=False)
    deps = {}; dependency_evidence = []; cell_cache = {}
    expanded_total = sum(len(data) for data in parts.values())
    try:
        for link in links:
            row = supplied[link['link_index']]
            for key in fields - {'path', 'sha256', 'literal_cells'}:
                if row[key] != link[key]: raise ExternalLiteralError('dependency identity does not match exact link')
            depraw, _, depparts = _package(row['path'])
            if sha(depraw) != row['sha256']: raise ExternalLiteralError('dependency source hash differs')
            expanded_total += sum(len(data) for data in depparts.values())
            if expanded_total > MAX_BYTES:
                raise ExternalLiteralError('total dependency package bound exceeded')
            depsheets, deplinks = _layout(depparts)
            if deplinks: raise ExternalLiteralError('dependency workbook has unresolved external links')
            if not isinstance(row['literal_cells'], list) or not row['literal_cells']:
                raise ExternalLiteralError('literal assumption provenance is missing')
            approved = {}
            for literal in row['literal_cells']:
                _keys(literal, {'sheet', 'coordinate', 'role', 'provenance'}, 'literal assumption')
                key = (literal['sheet'], _coordinate(literal['coordinate']))
                if (key in approved or literal['role'] != 'external_assumption'
                        or not isinstance(literal['provenance'], str) or not literal['provenance'].strip() or len(literal['provenance']) > 4096):
                    raise ExternalLiteralError('ambiguous or unjustified literal assumption')
                if literal['sheet'] not in depsheets: raise ExternalLiteralError('dependency sheet is missing')
                approved[key] = literal
            depwb = load_workbook(io.BytesIO(depraw), data_only=False, keep_links=False)
            deps[link['link_index']] = (depparts, depsheets, depwb, approved, set())
            dependency_evidence.append({'link_index': link['link_index'], 'path': str(Path(row['path']).resolve()),
                                        'sha256': row['sha256'], 'literal_cells': row['literal_cells']})
        updated = dict(parts); dispositions = []; used_links = set()
        for sheet, part in sheets.items():
            edits = {}
            for coordinate, cell in _cells(parts, part).items():
                formula = cell.find('s:f', NS)
                if formula is None: continue
                text = formula.text or ''
                if not EXTERNAL.search(text): continue
                match = DIRECT.fullmatch(text)
                if match is None: raise ExternalLiteralError('only direct external single-cell identities are supported')
                index = int(match.group('quoted_index') or match.group('index'))
                if index not in deps: raise ExternalLiteralError('formula external index has no dependency identity')
                dependency_sheet = (match.group('quoted_sheet') or match.group('sheet')).replace("''", "'")
                dependency_cell = _coordinate(match.group('col') + match.group('row'))
                depparts, depsheets, depwb, approved, used = deps[index]
                key = (dependency_sheet, dependency_cell)
                if dependency_sheet not in links[index - 1]['sheet_names'] or key not in approved:
                    raise ExternalLiteralError('referenced cell lacks exact literal assumption evidence')
                cache_key = (index, dependency_sheet)
                if cache_key not in cell_cache:
                    cell_cache[cache_key] = _cells(depparts, depsheets[dependency_sheet])
                depcell = cell_cache[cache_key].get(dependency_cell)
                if depcell is None: raise ExternalLiteralError('dependency literal cell is missing')
                if (set(cell.attrib) - {'r', 's', 't'}
                        or [child.tag.rsplit('}', 1)[-1] for child in cell] not in [['f'], ['f', 'v']]):
                    raise ExternalLiteralError('destination formula metadata is unsupported')
                kind, typ, payload, value = _literal(depparts, depcell, depwb[dependency_sheet][dependency_cell],
                    wb[sheet][coordinate], depwb.epoch, wb.epoch)
                replacement = ET.Element('c', {key: val for key, val in cell.attrib.items() if key != 't'})
                replacement.set('t', typ)
                for child in payload: replacement.append(child)
                edits[coordinate] = ET.tostring(replacement, encoding='utf-8')
                dispositions.append({'sheet': sheet, 'coordinate': coordinate, 'part': part,
                    'original_formula_sha256': sha(text.encode()), 'link_index': index,
                    'dependency_sha256': supplied[index]['sha256'], 'dependency_sheet': dependency_sheet,
                    'dependency_coordinate': dependency_cell, 'literal_type': kind,
                    'literal_evidence_sha256': sha(canonical({'type': kind, 'value': value})),
                    'assumption_provenance': approved[key]['provenance'], 'replacement_cell_sha256': sha(edits[coordinate])})
                if len(dispositions) > 10_000:
                    raise ExternalLiteralError('reference occurrence bound exceeded')
                used.add(key); used_links.add(index)
            if edits:
                replaced = set()
                def replace(match):
                    chunk = match.group(0)
                    # Namespace prefixes on cell elements are outside the first route.
                    if not chunk.startswith(b'<c '): raise ExternalLiteralError('prefixed worksheet cells unsupported')
                    ref = re.search(rb'\br="([A-Z]+[0-9]+)"', chunk)
                    coordinate = ref.group(1).decode() if ref else None
                    if coordinate not in edits: return chunk
                    if coordinate in replaced: raise ExternalLiteralError('duplicate destination cell span')
                    replaced.add(coordinate)
                    return edits[coordinate]
                updated[part] = CELL_SPAN.sub(replace, parts[part])
                if replaced != set(edits): raise ExternalLiteralError('formula cell XML spans are incomplete')
        if used_links != set(supplied): raise ExternalLiteralError('unused or unaccounted external link')
        for _, _, _, approved, used in deps.values():
            if used != set(approved): raise ExternalLiteralError('unused or incomplete assumption dispositions')
        book = ET.fromstring(parts['xl/workbook.xml'])
        book.remove(book.find('s:externalReferences', NS))
        updated['xl/workbook.xml'] = ET.tostring(book, encoding='utf-8', xml_declaration=True)
        rels = ET.fromstring(parts['xl/_rels/workbook.xml.rels'])
        for el in list(rels):
            if el.get('Type', '').endswith('/externalLink'): rels.remove(el)
        updated['xl/_rels/workbook.xml.rels'] = ET.tostring(rels, encoding='utf-8', xml_declaration=True)
        removed = sorted(name for name in parts if name.startswith('xl/externalLinks/'))
        for name in removed: updated.pop(name)
        types = ET.fromstring(parts['[Content_Types].xml'])
        for el in list(types):
            if el.get('PartName', '').lstrip('/') in removed: types.remove(el)
        updated['[Content_Types].xml'] = ET.tostring(types, encoding='utf-8', xml_declaration=True)
        for name, data in updated.items():
            # Scan executable fields in every remaining supported member. This
            # deliberately rejects residual external-like text rather than guessing.
            tree = ET.fromstring(data)
            for el in tree.iter():
                if EXTERNAL.search(el.text or '') or any(EXTERNAL.search(v) for v in el.attrib.values()):
                    raise ExternalLiteralError('unaccounted external reference remains')
        plan = {'schema_version': PLAN_SCHEMA, 'policy_version': POLICY,
                'generator_sha256': sha(Path(__file__).read_bytes()), 'source_sha256': source_sha,
                'identity_evidence_sha256': expected_evidence_sha256, 'authority': evidence['authority'],
                'dependencies': dependency_evidence, 'dispositions': dispositions,
                'members': [{'part': name, 'before_sha256': sha(parts[name]),
                             'after_sha256': sha(updated[name]) if name in updated else None,
                             'disposition': 'removed_external_link' if name not in updated else
                                            'changed' if updated[name] != parts[name] else 'unchanged'}
                            for name in sorted(parts)],
                'required_downstream_gates': ['fresh_source_health', 'trusted_recalculation_where_required',
                    'independent_reconstruction_without_answer_seeding', 'disclosure_and_leak_review',
                    'all_harbor_quality_and_publication_gates']}
        plan['plan_sha256'] = sha(canonical(plan))
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, 'w') as archive:
            for info in infos:
                if info.filename in updated: archive.writestr(copy.copy(info), updated[info.filename])
        return plan, stream.getvalue()
    finally:
        wb.close()
        for _, _, depwb, _, _ in deps.values(): depwb.close()


def plan(source, evidence, approved_evidence_sha256):
    return _build(source, evidence, approved_evidence_sha256)[0]


def apply(source, evidence, approved_evidence_sha256, destination):
    """Atomically install candidate and manifest together; existing output is immutable."""
    destination = Path(destination)
    if destination.exists(): raise ExternalLiteralError('destination already exists')
    result, candidate = _build(source, evidence, approved_evidence_sha256)
    manifest = {'schema_version': MANIFEST_SCHEMA, 'plan': result, 'candidate_sha256': sha(candidate),
                'candidate_name': 'candidate.xlsx', 'qualification': 'unqualified_candidate'}
    manifest['manifest_sha256'] = sha(canonical(manifest))
    destination.parent.mkdir(parents=True, exist_ok=True)
    temp = Path(tempfile.mkdtemp(prefix='.external-literal-', dir=destination.parent))
    try:
        for name, data in [('candidate.xlsx', candidate), ('manifest.json', canonical(manifest))]:
            with (temp / name).open('xb') as handle:
                handle.write(data); handle.flush(); os.fsync(handle.fileno())
        descriptor = os.open(temp, os.O_RDONLY)
        try: os.fsync(descriptor)
        finally: os.close(descriptor)
        # Serialize cooperating publishers; never replace an existing generation.
        with (destination.parent / '.external-input-publication.lock').open('a') as lock:
            fcntl.flock(lock, fcntl.LOCK_EX)
            if destination.exists(): raise ExternalLiteralError('publication conflict')
            os.rename(temp, destination)
        descriptor = os.open(destination.parent, os.O_RDONLY)
        try: os.fsync(descriptor)
        finally: os.close(descriptor)
    finally:
        if temp.exists(): shutil.rmtree(temp)
    return manifest


def verify(source, evidence, approved_evidence_sha256, destination):
    """Recompute from original dependency bytes, never trust manifest assertions alone."""
    directory = Path(destination)
    _, manifest = _json(directory / 'manifest.json', max_bytes=MAX_BYTES)
    result, candidate = _build(source, evidence, approved_evidence_sha256)
    expected = {'schema_version': MANIFEST_SCHEMA, 'plan': result, 'candidate_sha256': sha(candidate),
                'candidate_name': 'candidate.xlsx', 'qualification': 'unqualified_candidate'}
    expected['manifest_sha256'] = sha(canonical(expected))
    if (manifest != expected or (directory / 'candidate.xlsx').read_bytes() != candidate
            or {path.name for path in directory.iterdir()} != {'manifest.json', 'candidate.xlsx'}):
        raise ExternalLiteralError('candidate or manifest evidence differs from exact replay')
    return manifest


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('command', choices=['plan', 'apply', 'verify'])
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--evidence', type=Path, required=True)
    parser.add_argument('--approved-evidence-sha256', required=True)
    parser.add_argument('--destination', type=Path)
    args = parser.parse_args(argv)
    if args.command != 'plan' and args.destination is None:
        parser.error('--destination is required for apply/verify')
    try:
        result = (plan(args.source, args.evidence, args.approved_evidence_sha256) if args.command == 'plan'
                  else globals()[args.command](args.source, args.evidence, args.approved_evidence_sha256, args.destination))
    except (ExternalLiteralError, OSError, ValueError) as exc:
        parser.exit(1, f'external literal input FAIL: {exc}\n')
    print(json.dumps(result, sort_keys=True))


if __name__ == '__main__':
    main()
