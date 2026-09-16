"""Materialize sensitivity tables only after proving exclusion from targets.

The inspection generation is never a source approval. The candidate must still
pass source publication, independent reconstruction, and every packaging gate.
"""
from __future__ import annotations

from argparse import Namespace
import json
import math
from pathlib import Path
import shutil
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.utils.cell import range_boundaries, get_column_letter

from scripts.batch.hardened_runner import atomic_write_json
from xl_volatile_formula_remediation import (_package, _sheet_parts, _cell_blocks,
    FORMULA_TAG_RE, _attrs, _write_package, _value_ledger, sha256_file,
    canonical_bytes, sha256_bytes)

POLICY = "presentation-data-tables/v2"
COMPOSED_POLICY = "presentation-data-tables-remediation/v1"


def composable_source(health, source):
    """Select two existing proofs; this is never a source qualification."""
    reasons = set(health.get("reason_codes", []))
    if ("data_tables_present" not in reasons or len(reasons) < 2
            or not reasons <= {"data_tables_present", "external_links_present",
                               "unsupported_volatile_formulas_present"}):
        return False
    from xl_volatile_formula_remediation import build_plan
    plan = build_plan(source)
    return plan["status"] == "eligible" and plan["unresolved"] == []


def table_inventory(members):
    tables = []
    for sheet, part in _sheet_parts(members):
        root = ET.fromstring(members[part])
        for cell in root.findall(".//{*}sheetData/{*}row/{*}c"):
            formula = cell.find("{*}f")
            if formula is None or formula.get("t") != "dataTable":
                continue
            bounds = range_boundaries(formula.get("ref", ""))
            if None in bounds:
                raise ValueError("data table has no bounded result rectangle")
            c1, r1, c2, r2 = bounds
            if (c2-c1+1)*(r2-r1+1) > 100000:
                raise ValueError("data table exceeds inspection budget")
            cells = [f"{sheet}!{get_column_letter(c)}{r}"
                     for c in range(c1, c2+1) for r in range(r1, r2+1)]
            if f"{sheet}!{cell.get('r')}" not in cells:
                raise ValueError("data table anchor is outside result rectangle")
            tables.append(dict(sheet=sheet, part=part, anchor=cell.get("r"),
                               rectangle=formula.get("ref"), cells=cells))
    if not tables:
        raise ValueError("source contains no data tables")
    return tables


def check_exclusion(tables, payload):
    from xl_preflight import classify
    from xl_seg.restriction_cone import _proof_cone
    verification = payload.get("verification", {})
    if classify(verification, True)[0] != "healthy":
        raise ValueError("original workbook reconstruction is not healthy")
    outputs = verification.get("ordered_output_cells")
    radj = payload.get("proof", {}).get("runtime_radj")
    if not outputs or not isinstance(radj, dict):
        raise ValueError("complete output dependency proof is required")
    cone = _proof_cone(outputs, radj)
    table_cells = {cell for table in tables for cell in table["cells"]}
    if table_cells & cone:
        raise ValueError("data table intersects a graded dependency cone")
    return outputs, sorted(cone)


def scan_leaks(source, tables, outputs):
    book = openpyxl.load_workbook(source, data_only=True, read_only=True)
    try:
        def value(key):
            sheet, coordinate = key.rsplit("!", 1)
            return book[sheet][coordinate].value
        targets = [value(key) for key in outputs]
        for table in tables:
            for key in table["cells"]:
                observed = value(key)
                for target in targets:
                    if isinstance(target, (float, int)) and not isinstance(target, bool):
                        if isinstance(observed, (float, int)) and math.isclose(observed, target, rel_tol=1e-6, abs_tol=1e-9):
                            raise ValueError("presentation table discloses a graded numeric answer")
                    elif observed is not None and observed == target:
                        raise ValueError("presentation table discloses a graded answer")
    finally:
        book.close()


def transformed_members(members, tables):
    after = dict(members)
    for part in sorted({table["part"] for table in tables}):
        xml = members[part].decode("utf-8")
        blocks = _cell_blocks(xml)
        for table in (t for t in tables if t["part"] == part):
            block = blocks[table["anchor"]]["block"]
            matches = list(FORMULA_TAG_RE.finditer(block))
            if len(matches) != 1 or _attrs(matches[0].group("attrs") or matches[0].group("self_attrs") or "").get("t") != "dataTable":
                raise ValueError("data table XML does not match inventory")
            replacement = FORMULA_TAG_RE.sub("", block)
            if xml.count(block) != 1:
                raise ValueError("data table anchor XML is ambiguous")
            xml = xml.replace(block, replacement, 1)
        after[part] = xml.encode("utf-8")
    if _value_ledger(members) != _value_ledger(after):
        raise ValueError("materialization changed workbook values or caches")
    return after


def inspection_metadata(source):
    """Derive a presentation-only loader repair without changing any cell part."""
    from xl_volatile_formula_remediation import build_plan, _rewrite_defined_names
    infos, before = _package(Path(source))
    plan = build_plan(source)
    actions = [a for a in plan['actions'] if a.get('rule_id') == 'unreferenced-error-print-setting/v1']
    if not actions:
        return infos, before, None
    if plan['status'] != 'eligible' or plan['unresolved']:
        raise ValueError('inspection metadata requires an eligible original remediation plan')
    after = dict(before)
    after['xl/workbook.xml'] = _rewrite_defined_names(before['xl/workbook.xml'].decode('utf-8'), actions).encode('utf-8')
    edge = dict(schema_version='unreferenced-error-print-settings-inspection/v1',
        original_sha256=sha256_file(source), actions=actions,
        workbook_xml_before_sha256=sha256_bytes(before['xl/workbook.xml']),
        workbook_xml_after_sha256=sha256_bytes(after['xl/workbook.xml']))
    return infos, after, edge


def verify_inspection_source(source, inspected, edge):
    if edge is None:
        if sha256_file(source) != sha256_file(inspected):
            raise ValueError('inspection source is not the original workbook')
        return
    _, expected, actual_edge = inspection_metadata(source)
    if edge != actual_edge or _package(Path(inspected))[1] != expected:
        raise ValueError('inspection metadata edge or unchanged package content differs')


def materialize(source, workbook_id, curation, output_root, *, allow_initial_curation=False,
                remediate=False, redaction_policy=None, cache_policy=None):
    from xl_ast_graph import main as ast_main
    from xl_segment import segment
    from xl_seg.publication import validate_generation_directory
    from xl_source_health import inspect_workbook
    source, curation, root = Path(source), Path(curation), Path(output_root)
    if source.is_symlink() or not source.is_file() or (not curation.is_file() and not allow_initial_curation):
        raise ValueError("materialization requires a regular source and retained curation")
    if root.exists():
        raise ValueError("materialization output already exists")
    if type(remediate) is not bool:
        raise ValueError("explicit Boolean remediation selection required")
    if remediate and not composable_source(inspect_workbook(source), source):
        raise ValueError("source is not eligible for the bounded table/remediation chain")
    root.mkdir(parents=True)
    infos, members = _package(source)
    tables = table_inventory(members)
    if redaction_policy is not None:
        from xl_presentation_redaction import selected
        selected(redaction_policy, source)
    if cache_policy is not None:
        from xl_table_scoped_cache import policy
        policy(cache_policy, source)
        if not remediate:
            raise ValueError('table-cache proof requires the composed route')
    private = root / "inspection-source"
    private.mkdir()
    inspected = private / f"{workbook_id}.xlsx"
    metadata = None
    if remediate:
        inspection_infos, inspection_members, metadata = inspection_metadata(source)
    if metadata is None:
        shutil.copyfile(source, inspected)
    else:
        _write_package(inspection_infos, inspection_members, inspected)
    verify_inspection_source(source, inspected, metadata)
    ast_main([str(inspected), "-o", str(root / "ast"), "--production", "--quiet"])
    from xl_source_publication import create_ast_provenance
    ast_provenance = create_ast_provenance(inspected, root / "ast" / workbook_id,
                                         builder_args=["--production", "--quiet"])
    segmentation = root / "inspection-segments" / workbook_id
    segmentation.mkdir(parents=True)
    retained = curation.is_file()
    if retained:
        shutil.copyfile(curation, segmentation / "curation.toml")
    options = Namespace(ast_dir=str(root / "ast"), source=str(private), out=str(segmentation.parent),
        allow_legacy_ast=True, threshold=6.0, top=40, sample=400, lineage_max=4000,
        recurate=False, llm=False, no_auto_adjudication=True, model="", env_file="", no_verify=False)
    payload = segment(workbook_id, options)
    if not retained:
        curation = segmentation / "curation.toml"
    elif curation.read_bytes() != (segmentation / "curation.toml").read_bytes():
        raise ValueError("retained curation changed during inspection")
    generation = Path(payload["generation"]["directory"])
    validate_generation_directory(generation, require_pass=True)
    outputs, cone = check_exclusion(tables, payload)
    if redaction_policy is None:
        scan_leaks(inspected, tables, outputs)
    final_candidate = root / f"{workbook_id}.xlsx"
    candidate = (root / 'manual-projection' / final_candidate.name
                 if cache_policy is not None else final_candidate)
    candidate.parent.mkdir(parents=True, exist_ok=True)
    table_candidate = candidate
    if remediate:
        table_candidate = candidate.parent / "table-candidate" / candidate.name
        table_candidate.parent.mkdir()
    if redaction_policy is None:
        after = transformed_members(members, tables)
    else:
        from xl_presentation_redaction import transformed_members as redact_members
        after = redact_members(members, tables)
    _write_package(infos, after, table_candidate)
    remediation = None
    if remediate:
        from xl_volatile_formula_remediation import build_plan, apply_plan, verify_manifest
        plan = build_plan(table_candidate)
        if plan["status"] != "eligible" or plan["unresolved"]:
            raise ValueError("materialized source has unresolved remediation")
        manifest = apply_plan(table_candidate, plan, candidate)
        verify_manifest(table_candidate, candidate, plan, manifest)
        atomic_write_json(root / "remediation-plan.json", plan)
        atomic_write_json(root / "remediation-manifest.json", manifest)
        remediation = dict(intermediate=str(table_candidate),
            intermediate_sha256=sha256_file(table_candidate), plan=plan, manifest=manifest)
    health = inspect_workbook(candidate)
    atomic_write_json(root / "health.json", health)
    if cache_policy is None and health["route"] not in {"pass", "restricted_pass"}:
        raise ValueError("materialized candidate still requires source support")
    scan_leaks(candidate, tables, outputs)
    verify_inspection_source(source, inspected, metadata)
    receipt = dict(policy_version=COMPOSED_POLICY if remediate else POLICY, original_sha256=sha256_file(source),
        candidate_sha256=sha256_file(candidate), curation=str(curation), curation_sha256=sha256_file(curation),
        curation_origin="retained" if retained else "initial_selection_on_metadata_normalized_original" if metadata else "initial_selection_on_original",
        inspection_generation=str(generation), inspection_manifest_sha256=sha256_file(generation / "generation-manifest.json"),
        inspection_ast_provenance=ast_provenance,
        ordered_outputs=outputs, excluded_cone_sha256=sha256_bytes(canonical_bytes(cone)),
        tables=tables, leak_scan="passed", disposition="presentation_only_materialized")
    if metadata is not None:
        receipt["inspection_metadata"] = metadata
    if redaction_policy is not None:
        receipt["redaction_policy"] = redaction_policy
    if remediate:
        receipt["remediation"] = remediation
    receipt["manifest_sha256"] = sha256_bytes(canonical_bytes(receipt))
    if cache_policy is not None:
        from xl_table_scoped_cache import create
        from xl_source_cache_equivalence import binding
        projection_receipt = candidate.parent / 'manifest.json'
        atomic_write_json(projection_receipt, receipt)
        edge = create(source, candidate, binding(projection_receipt), cache_policy, final_candidate)
        receipt = {**receipt, 'candidate_sha256': sha256_file(final_candidate),
                   'cache_equivalence': edge}
        receipt['manifest_sha256'] = sha256_bytes(canonical_bytes(
            {k:v for k,v in receipt.items() if k != 'manifest_sha256'}))
        candidate = final_candidate
    atomic_write_json(root / "manifest.json", receipt)
    return candidate, receipt


def verify_materialization(source, candidate, curation, receipt):
    if 'cache_equivalence' in receipt:
        from xl_source_cache_equivalence import read
        from xl_table_scoped_cache import verify
        edge = receipt['cache_equivalence']
        projection = read(edge['scope']['projection_receipt'])
        expected = {**projection, 'candidate_sha256': sha256_file(candidate),
                    'cache_equivalence': edge}
        expected['manifest_sha256'] = sha256_bytes(canonical_bytes(
            {k:v for k,v in expected.items() if k != 'manifest_sha256'}))
        if receipt != expected or receipt['curation_sha256'] != sha256_file(curation):
            raise ValueError('table-cache materialization scope or receipt differs')
        verify(source, candidate, edge)
        return {'status': 'verified', 'manifest_sha256': receipt['manifest_sha256']}
    return _verify_materialization(source, candidate, curation, receipt)


def verify_projection(source, candidate, curation, receipt):
    """Authenticate an intermediate; this never admits a source generation."""
    if 'cache_equivalence' in receipt:
        raise ValueError('nested table-cache projection forbidden')
    return _verify_materialization(source, candidate, curation, receipt, manual=True)


def _verify_materialization(source, candidate, curation, receipt, *, manual=False):
    from xl_seg.publication import validate_generation_directory, resolve_generation_by_id
    from xl_source_health import inspect_workbook
    from xl_source_publication import validate_ast_provenance, _default_builder_paths
    unsigned = {k:v for k,v in receipt.items() if k != "manifest_sha256"}
    composed = receipt.get("policy_version") == COMPOSED_POLICY
    if (receipt.get("policy_version") not in {POLICY, COMPOSED_POLICY}
            or receipt.get("manifest_sha256") != sha256_bytes(canonical_bytes(unsigned))
            or receipt.get("original_sha256") != sha256_file(source)
            or receipt.get("candidate_sha256") != sha256_file(candidate)
            or receipt.get("curation_sha256") != sha256_file(curation)):
        raise ValueError("materialization evidence is changed or incomplete")
    generation = Path(receipt["inspection_generation"])
    validate_generation_directory(generation, require_pass=True,
        expected_manifest_sha256=receipt["inspection_manifest_sha256"])
    payload = json.loads((generation / "segments.json").read_text())
    _, before = _package(Path(source))
    tables = table_inventory(before)
    outputs, cone = check_exclusion(tables, payload)
    if (tables != receipt["tables"] or outputs != receipt["ordered_outputs"]
            or sha256_bytes(canonical_bytes(cone)) != receipt["excluded_cone_sha256"]):
        raise ValueError("materialization coverage changed")
    # The inspection source is retained beside its generation; bind it back
    # to the original, not just to the receipt's assertion.
    inspection_root = generation.parents[3]
    inspected = inspection_root / "inspection-source" / Path(candidate).name
    if "inspection_metadata" in receipt and not composed:
        raise ValueError("inspection metadata requires the explicit composed route")
    verify_inspection_source(source, inspected, receipt.get("inspection_metadata"))
    ast_dir = inspection_root / "ast" / Path(candidate).stem
    provenance = receipt.get("inspection_ast_provenance")
    if not isinstance(provenance, dict):
        raise ValueError("materialization lacks inspection AST provenance")
    validate_ast_provenance(inspected, ast_dir, provenance)
    # Immutable provenance may point to a still-valid historical code snapshot.
    # Require the running parser too; relocating identical code is harmless.
    recorded = provenance["builder"]["code"]
    current = _default_builder_paths()
    if sorted((r["name"], r["sha256"], r["size_bytes"]) for r in recorded) != sorted(
            (p.name, sha256_file(p), p.stat().st_size) for p in current):
        raise ValueError("materialization inspection AST builder changed")
    resolve_generation_by_id(generation.parents[1], generation.name,
        source_path=inspected, ast_dir=ast_dir, require_pass=True,
        validate_live_evidence=True)
    table_candidate = Path(candidate)
    if composed:
        from xl_volatile_formula_remediation import verify_manifest
        value = receipt.get("remediation")
        if not isinstance(value, dict) or set(value) != {"intermediate", "intermediate_sha256", "plan", "manifest"}:
            raise ValueError("closed remediation edge required")
        table_candidate = Path(value["intermediate"])
        if (table_candidate != Path(candidate).parent / "table-candidate" / Path(candidate).name
                or table_candidate.is_symlink()
                or sha256_file(table_candidate) != value["intermediate_sha256"]):
            raise ValueError("table intermediate binding differs")
        if not composable_source(inspect_workbook(source), source):
            raise ValueError("original composition source differs")
        verify_manifest(table_candidate, candidate, value["plan"], value["manifest"])
        health = inspect_workbook(candidate)
        admissible = (health['route'] == 'recalc_candidate'
                      and set(health['reason_codes']) == {'manual_calculation_mode'}) if manual else health['route'] in {'pass', 'restricted_pass'}
        if not admissible:
            raise ValueError("composed candidate is not source-admissible")
    elif "remediation" in receipt:
        raise ValueError("legacy table receipt cannot contain a remediation edge")
    elif manual:
        raise ValueError('manual projection requires authenticated composition')
    _, after = _package(table_candidate)
    if "redaction_policy" in receipt:
        from xl_presentation_redaction import selected, transformed_members as redact_members
        selected(receipt["redaction_policy"], source)
        expected = redact_members(before, tables)
    else:
        expected = transformed_members(before, tables)
    if expected != after:
        raise ValueError("candidate contains unapproved edits")
    scan_leaks(candidate if "redaction_policy" in receipt else inspected, tables, outputs)
    return {"status": "verified", "manifest_sha256": receipt["manifest_sha256"]}
